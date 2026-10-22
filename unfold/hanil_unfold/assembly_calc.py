"""조립품 전체 계산: 모든 부품을 전개해 두께·면적·무게를 모으고, 재질·두께별로 합친다.

철판 발주와 견적에 쓰려는 것이다. 펼칠 수 없는 부품(파이프 등)은 이유와 함께 따로 표시한다.
"""
from __future__ import annotations

from dataclasses import dataclass

from . import metrics as mt
from .unfolder import unfold


@dataclass
class PartResult:
    index: int                 # 0부터
    name: str
    result: object = None      # UnfoldResult 또는 None
    error: str = ""
    metrics: mt.Metrics | None = None

    @property
    def ok(self):
        return self.result is not None

    @property
    def thickness(self):
        return self.result.thickness if self.result else None


def compute_all(model, settings, mode="auto", thickness=None, progress=None):
    """progress(i, n, 이름) 가 False 를 돌려주면 멈춘다."""
    out = []
    n = len(model.solids)
    for i in range(n):
        name = model.names[i] if model.names else f"부품 {i + 1}"
        if progress is not None and progress(i, n, name) is False:
            break
        try:
            r = unfold(model.solids[i], thickness=thickness, settings=settings, mode=mode)
            out.append(PartResult(i, name, r, metrics=mt.compute(r, settings.mat.density)))
        except Exception as e:
            out.append(PartResult(i, name, error=str(e).splitlines()[0]))
    return out


@dataclass
class Group:
    material: str
    thickness: float
    count: int = 0
    blank_area: float = 0.0
    rect_area: float = 0.0
    weight_kg: float = 0.0
    rect_weight_kg: float = 0.0
    cut_length: float = 0.0


def summarize(parts, material_code):
    """(재질·두께별 묶음 목록, 전체 합계 Group)."""
    groups = {}
    total = Group(material_code, 0.0)
    for p in parts:
        if not p.ok:
            continue
        t = round(p.thickness, 2)
        g = groups.setdefault(t, Group(material_code, t))
        for target in (g, total):
            target.count += 1
            target.blank_area += p.metrics.blank_area
            target.rect_area += p.metrics.rect_area
            target.weight_kg += p.metrics.weight_kg
            target.rect_weight_kg += p.metrics.rect_weight_kg
            target.cut_length += p.metrics.cut_length
    return [groups[t] for t in sorted(groups)], total


def save_xlsx(path, model, parts, material_code):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    head_font = Font(name="맑은 고딕", bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="28405A")
    wb = Workbook()

    ws = wb.active
    ws.title = "두께별 합계"
    cols = ["재질", "두께(mm)", "부품 수", "블랭크 면적(m²)", "사각 원판 면적(m²)",
            "블랭크 무게(kg)", "사각 원판 무게(kg)", "절단 길이(m)"]
    for j, h in enumerate(cols, 1):
        c = ws.cell(1, j, h)
        c.font, c.fill = head_font, fill
        ws.column_dimensions[c.column_letter].width = 16
    groups, total = summarize(parts, material_code)
    for i, g in enumerate(groups + [total], 2):
        vals = [g.material, g.thickness if g is not total else "합계", g.count, g.blank_area / 1e6,
                g.rect_area / 1e6, g.weight_kg, g.rect_weight_kg, g.cut_length / 1000]
        for j, v in enumerate(vals, 1):
            c = ws.cell(i, j, round(v, 4) if isinstance(v, float) else v)
            if isinstance(v, float):
                c.number_format = "#,##0.000"
            if g is total:
                c.font = Font(bold=True)

    ws2 = wb.create_sheet("부품별")
    cols = ["번호", "부품명", "전개 방식", "두께(mm)", "가로(mm)", "세로(mm)", "블랭크 면적(mm²)",
            "블랭크 무게(kg)", "사각 원판 무게(kg)", "절단 길이(mm)", "절곡 수", "비고"]
    for j, h in enumerate(cols, 1):
        c = ws2.cell(1, j, h)
        c.font, c.fill = head_font, fill
        ws2.column_dimensions[c.column_letter].width = 14
    ws2.column_dimensions["B"].width = 36
    ws2.column_dimensions["L"].width = 50
    for i, p in enumerate(parts, 2):
        if p.ok:
            m, r = p.metrics, p.result
            vals = [p.index + 1, p.name, "성형품(근사)" if r.mode == "forming" else "판금", r.thickness,
                    m.width, m.height, m.blank_area, m.weight_kg, m.rect_weight_kg, m.cut_length,
                    len(r.bends), " / ".join(w for w in r.warnings if "근사 전개:" not in w)]
        else:
            vals = [p.index + 1, p.name, "펼칠 수 없음"] + [None] * 8 + [p.error]
        for j, v in enumerate(vals, 1):
            c = ws2.cell(i, j, round(v, 4) if isinstance(v, float) else v)
            if isinstance(v, float):
                c.number_format = "#,##0.000"
    ws2.freeze_panes = "C2"
    wb.save(path)
