"""공정별 성형 해석 (AutoForm 의 공정 체계처럼): OP10 드로잉 → OP20 트림·피어싱 → OP30 플랜지 → OP40 리스트라이크.

OP10 드로잉   — 다이 페이스 (제품 + 어덴덤 + 바인더, dieface) 위 원스텝: 드로잉 블랭크·끌려 들어감·제품 부분의 판정.
OP20 트림·피어싱 — 트림선 (제품 바깥 경계) 과 막고 해석한 피어싱 구멍의 둘레로 전단력 F = L·t·τ, 버리는 판.
OP30 플랜지   — 트림선 가까이 선 벽 (프레스 방향에 거의 나란한 면) 을 트림 뒤 꺾어 세우는 플랜지로 보고,
                가장자리 늘림 (늘림 플랜지 → 구멍 확장률 HER) · 줄임 (줄임 플랜지 → 주름) · 스프링백.
OP40 리스트라이크·캠 — 고장력강·스프링백이 크면 다시 눌러 맞춤, 플랜지가 프레스 방향과 비스듬하면 캠.
각 공정의 힘은 권장 프레스 (× 1.3) 로.
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field

import numpy as np

from . import diecalc as dc
from . import formsim as fs


@dataclass
class Op:
    code: str
    name: str
    verdict: str
    force_n: float
    lines: list = field(default_factory=list)


@dataclass
class Ops:
    ops: list
    dieface: object = None
    fa_die: object = None
    notes: list = field(default_factory=list)

    @property
    def press_tonf(self):
        return max((o.force_n for o in self.ops), default=0.0) * 1.3 / dc.TON

    def rows(self):
        return [[o.code, o.name, o.verdict, o.force_n / 1000, o.force_n * 1.3 / dc.TON, " · ".join(o.lines)]
                for o in self.ops]


HEAD = ["공정", "이름", "판정", "힘 kN", "프레스 tf (× 1.3)", "내용"]


def _loop_len(P):
    return float(np.linalg.norm(np.roll(P, -1, axis=0) - P, axis=1).sum())


def flange_mask(V, F, z, t, reach=80.0):
    """트림 뒤 꺾어 세우는 플랜지 (삼각형 표시): 트림선에서 시작해 선 벽 (프레스 방향에 거의 나란함) 을 따라 이어지고,
    트림선 쪽이 낮은 (안쪽 면에서 아래로 꺾여 내려가는 — 후드 둘레처럼) 띠. 컵 벽처럼 트림선이 위 끝인 벽은
    드로잉에서 만드는 벽이라 빼낸다. (mask, 트림선 중 플랜지 비, 플랜지 벽이 프레스 방향과 이루는 각)."""
    from .forming import boundary_loops
    loops = boundary_loops(F)
    m = np.zeros(len(F), bool)
    if not loops:
        return m, 0.0, None
    outer = max(loops, key=lambda lp: _loop_len(V[lp]))
    n = np.cross(V[F[:, 1]] - V[F[:, 0]], V[F[:, 2]] - V[F[:, 0]])
    n /= np.maximum(np.linalg.norm(n, axis=1), 1e-30)[:, None]
    steep = np.abs(n @ z) < 0.5
    h = V @ z
    hc = h[F].mean(1)
    # 꼭짓점 → 삼각형, 삼각형 이웃 (모서리 공유)
    vt = {}
    for ti, tri in enumerate(F):
        for k in tri:
            vt.setdefault(int(k), []).append(ti)
    edge = {}
    for ti, tri in enumerate(F):
        for k in range(3):
            key = tuple(sorted((int(tri[k]), int(tri[(k + 1) % 3]))))
            edge.setdefault(key, []).append(ti)
    nb = [[] for _ in range(len(F))]
    for ts in edge.values():
        if len(ts) == 2:
            nb[ts[0]].append(ts[1])
            nb[ts[1]].append(ts[0])
    C = V[F].mean(1)
    seeds = {ti for k in outer for ti in vt.get(int(k), ()) if steep[ti]}
    # 선 벽을 따라 넓혀 감 (트림선에서 reach 안)
    P = V[outer]
    from scipy.spatial import cKDTree
    tree = cKDTree(P)
    stack = list(seeds)
    seen = set(seeds)
    while stack:
        ti = stack.pop()
        for tj in nb[ti]:
            if tj not in seen and steep[tj] and tree.query(C[tj])[0] < reach:
                seen.add(tj)
                stack.append(tj)
    comp = np.zeros(len(F), bool)
    comp[list(seen)] = True
    # 트림선 쪽이 낮은 곳만 (벽이 트림선에서 위로 올라감 = 안쪽에서 아래로 꺾여 내려온 플랜지)
    flagv = np.zeros(len(outer), bool)
    for i, k in enumerate(outer):
        ts = [ti for ti in vt.get(int(k), ()) if comp[ti]]
        if not ts:
            continue
        near = np.nonzero(comp & (np.linalg.norm(C - P[i], axis=1) < reach))[0]
        flagv[i] = len(near) > 0 and hc[near].mean() > h[int(k)] + 2 * t
    keep_seed = {ti for i, k in enumerate(outer) if flagv[i] for ti in vt.get(int(k), ()) if comp[ti]}
    stack = list(keep_seed)
    seen = set(keep_seed)
    while stack:
        ti = stack.pop()
        for tj in nb[ti]:
            if tj not in seen and comp[tj]:
                seen.add(tj)
                stack.append(tj)
    m[list(seen)] = True
    seg = np.linalg.norm(np.roll(P, -1, axis=0) - P, axis=1)
    frac = float(seg[flagv].sum() / max(seg.sum(), 1e-9))
    if frac > 0.3:
        # 트림선 대부분이 플랜지면 (후드 둘레처럼) 귀처럼 덜 가파른 곳까지 트림선에 붙은 선 벽 전부를 플랜지로 —
        # 부분만 떼면 드로잉 모양의 가장자리가 들쭉날쭉해 다이 페이스를 못 붙인다
        m = comp
    tilt = float(np.degrees(np.arcsin(np.clip(np.median(np.abs(n[m] @ z)), 0, 1)))) if m.any() else None
    return m, frac, tilt


def _draw_shape(V, F, mask):
    """플랜지를 떼어 낸 드로잉 모양 (가장 큰 덩어리)."""
    from .forming import _keep_largest_component
    mask = mask.copy()
    for _ in range(20):          # 한 점에서만 맞닿는 꼭짓점 (경계가 두 번 지나감) 이 없어질 때까지 그 둘레를 더 뗌
        keep = F[~mask]
        cnt = {}
        for tri in keep:
            for k in range(3):
                e = (int(tri[k]), int(tri[(k + 1) % 3]))
                key = (min(e), max(e))
                cnt[key] = cnt.get(key, 0) + 1
        deg = {}
        for (a, b), c in cnt.items():
            if c == 1:
                deg[a] = deg.get(a, 0) + 1
                deg[b] = deg.get(b, 0) + 1
        bad = {v for v, d in deg.items() if d > 2}
        if not bad:
            break
        hit = np.isin(F, list(bad)).any(1)
        mask |= hit
    keep = F[~mask]
    if len(keep) < 10:
        return None
    used = np.unique(keep)
    remap = -np.ones(len(V), int)
    remap[used] = np.arange(len(used))
    V2, F2 = _keep_largest_component(V[used], remap[keep])
    return V2, F2


def build(fa, result, settings, dieface_on=True):
    """제품 성형 해석 fa (formsim.analyze) 에서 공정별 결과."""
    from . import dieface as dfm
    mat = settings.mat
    fi = fa.fi
    t = fa.thickness
    ops, notes = [], []
    z = z0 = None
    df = fa_die = None
    fmask, frac_f, tilt = np.zeros(0, bool), 0.0, None
    if fi is not None and fi.V is not None:
        z0, _ = fs.press_direction(fi.V, fi.F)
        if z0 is not None:
            from .forming import boundary_loops
            loops = boundary_loops(fi.F)
            if loops:
                outer = max(loops, key=lambda lp: _loop_len(fi.V[lp]))
                if (fi.V[outer] @ z0).mean() < (fi.V @ z0).mean():
                    z0 = -z0
            fmask, frac_f, tilt = flange_mask(fi.V, fi.F, z0, t)
            z = z0
    if fi is not None and fi.V is not None and dieface_on:
        try:
            rd = None
            if fa.ldr:
                g = fa.draw_geom or {}
                if g:
                    rd = dc.die_radius(g["L_blank"] / math.pi, g["L_prod"] / math.pi, t)
            Vd, Fd = fi.V, fi.F
            if fmask.any() and frac_f > 0.1:          # 트림 뒤 꺾는 플랜지는 떼고 드로잉 모양에 다이 페이스
                ds = _draw_shape(fi.V, fi.F, fmask)
                if ds is not None:
                    Vd, Fd = ds
                    notes.append(f"플랜지 ({frac_f:.0%} 의 트림선) 는 떼어 낸 드로잉 모양에 다이 페이스 — 플랜지는 OP30 에서")
            df = dfm.build(Vd, Fd, t, r_die=rd)
            if df is not None:
                dfm.solve(df, t, mat)
                fa_die = fs.analyze(dfm.as_result(df, t, mat), settings)
                z = df.z
            else:
                notes.append("다이 페이스를 자동으로 못 만듦 (트림선이 벽을 따라 내려가는 끝이 잘린 부품 등) — "
                             "OP10 은 제품 원스텝 결과")
        except Exception as e:                       # 다이 페이스는 덤 — 제품 결과로 계속
            notes.append(f"다이 페이스 해석 실패: {e}")
            df = fa_die = None

    # 트림선 전개 (플랜지를 드로잉 형면 위로 펼침) — 플랜지가 없으면 제품 가장자리가 트림선
    trim = None
    if fi is not None and fi.V is not None and z is not None:
        from . import formtrim
        try:
            z_t = z0 if z0 is not None else z
            trim = formtrim.develop(fi.V, fi.F, fi.U, fmask if frac_f > 0.1 else None, z_t,
                                    df=df if (df is not None and frac_f > 0.1) else None, t=t, mat=mat)
        except Exception as e:                       # 트림선은 덤
            notes.append(f"트림선 전개 실패: {e}")

    # OP10 드로잉
    src = fa_die or fa
    lines = []
    if df is not None and df.U is not None:
        lines += [f"{a} {b}" for a, b in dfm.summary(df, fa_die)]
        prod = df.region == dfm.PRODUCT
        zone = src.zone[prod]
        w = src.fi._weights()[prod]
        frac = [float(w[zone == i].sum() / max(w.sum(), 1e-30)) for i in range(len(fs.ZONES))]
        v = fs.NG if frac[0] > 0.001 else fs.WARN if (frac[1] > 0.001 or frac[2] > 0.001 or frac[6] > 0.02) else fs.OK
        wa = src.fi._weights()[~prod]
        za = src.zone[~prod]
        crack_a = float(wa[za == 0].sum() / max(wa.sum(), 1e-30)) if wa.sum() > 0 else 0.0
        if crack_a > 0.001:
            lines.append(f"어덴덤·바인더 파단 {crack_a:.1%} — 제품이 아니라 어덴덤 설계 문제: 다이 R 을 키우거나 벽 각을 눕히거나 "
                         "바인더 폭을 바꾸세요 (자동 다이 페이스는 첫 어림)")
        fa_die.verdict = v                            # 그림 제목은 제품 부분 판정으로
    else:
        frac = fa.zone_frac
        v = fa.verdict if fi is not None else fs.OK
        lines.append(f"제품 원스텝: 얇아짐 최대 {fa.thin_hi:.1%} · 한계 비 {fa.risk_top:.2f}")
    lines.append(f"제품 부분 파단 {frac[0]:.1%} · 위험 {frac[1]:.1%} · 주름 {frac[6]:.1%}")
    f_draw = sum(f for name, f, _ in src.loads if not name.startswith("굽힘")) or \
        sum(f for name, f, _ in fa.loads if not name.startswith("굽힘"))
    ops.append(Op("OP10", "드로잉 (다이 페이스)" if df is not None else "드로잉 (제품 원스텝)", v, f_draw, lines))
    # OP20 트림·피어싱
    tau = dc.shear_strength(mat)
    L_trim = 0.0
    if fi is not None and fi.V is not None:
        from .forming import boundary_loops
        loops = boundary_loops(fi.F)
        L_trim = max((_loop_len(fi.V[lp]) for lp in loops), default=0.0)
    holes = getattr(fi, "holes", []) if fi is not None else []
    L_pierce = sum(math.pi * h["d"] for h in holes)
    f_trim = (L_trim + L_pierce) * t * tau
    lines = [f"트림선 {L_trim:,.0f} mm", f"피어싱 {len(holes)} 개 · 둘레 {L_pierce:,.0f} mm",
             f"전단 저항 τ {tau:.0f} MPa"]
    if df is not None and df.U is not None:
        from .unfolder import _tri_areas2
        scrap = float(_tri_areas2(df.U, df.F).sum() - _tri_areas2(df.U, df.F[:df.n_product]).sum())
        lines.append(f"버리는 판 (어덴덤·바인더) {scrap:,.0f} mm² · {scrap * t * mat.density / 1e6:.2f} kg")
    v20 = fs.OK
    if trim is not None:
        lines += trim.summary()
        if trim.cam_frac > 0.02:
            v20 = fs.WARN
            lines.append("트림 각 30° 넘는 구간 — 캠 트림 (또는 트림 공정을 나눔)")
    ops.append(Op("OP20", "트림 · 피어싱", v20, f_trim, lines))
    # OP30 플랜지
    if fi is not None and fi.V is not None and z is not None:
        L_out = L_trim
        if frac_f > 0.1:
            e, where, her = fa.edge if fa.edge else (0.0, "", None)
            fl_lines = [f"트림선의 {frac_f:.0%} 가 아래로 꺾어 세우는 플랜지 (벽이 프레스 방향과 {tilt:.0f}°)"
                        if tilt is not None else f"트림선의 {frac_f:.0%} 가 꺾어 세우는 플랜지"]
            fl_lines.append(f"가장자리 늘림 최대 {e:.0%} ({where})" + (f" — HER {her:.0%}" if her else ""))
            v = fa.edge_verdict
            shrink = fa.zone_frac[6]
            if shrink > 0.02:
                fl_lines.append(f"줄임 플랜지 두꺼워짐 {shrink:.1%} — 주름 (패드 힘·플랜지 높이)")
                v = fs.WARN if v == fs.OK else v
            r_in = max(2 * t, 3.0)
            ang = 90.0
            if trim is not None and trim.bend:
                bb = np.concatenate(trim.bend)
                bb = bb[np.isfinite(bb) & (bb > 5)]
                if len(bb):
                    ang = float(np.median(bb))
            after, die, _ = dc.springback(mat, r_in, t, ang)
            fl_lines.append(f"스프링백 (안쪽 R {r_in:g}, 꺾는 각 {ang:.0f}° — 트림선 전개에서): {after:.1f}° 로 벌어짐 → "
                            f"금형 {die:.1f}°")
            f_fl = L_out * frac_f * t * dc.flow_stress(mat, 0.05) * 0.6
            ops.append(Op("OP30", "플랜지", v, f_fl, fl_lines))
            need = mat.tensile_mpa >= 590 or abs(ang - after) > 1.0
            cam = tilt is not None and tilt > 15
            l4 = []
            if need:
                l4.append(f"스프링백 {abs(ang - after):.1f}° · TS {mat.tensile_mpa:g} MPa — 다시 눌러 각도·R 맞춤")
            if cam:
                l4.append(f"플랜지가 프레스 방향과 {tilt:.0f}° 비스듬 — 캠 플랜지")
            if need or cam:
                ops.append(Op("OP40", "리스트라이크" + (" · 캠" if cam else ""), fs.WARN if need else fs.OK,
                              f_fl * 1.5, l4))
        else:
            notes.append("트림 뒤 꺾는 플랜지가 거의 없음 — 플랜지·리스트라이크 공정 없이 드로잉에서 다 만듦")
    o = Ops(ops, df, fa_die, notes)
    o.flange = fmask
    o.trim = trim
    return o


def add_sheet(wb, ops: Ops, name=""):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws = wb.create_sheet("공정별 해석")
    ws["A1"] = f"공정별 성형 해석 — {name}"
    ws["A1"].font = Font(bold=True, size=13)
    hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E8C")
    for j, h in enumerate(HEAD, 1):
        c = ws.cell(3, j, h)
        c.font, c.fill = hf, hfill
    for i, row in enumerate(ops.rows(), 4):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, round(v, 1) if isinstance(v, float) else v)
            if j == len(row):
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, wdt in zip("ABCDEF", (8, 22, 10, 10, 14, 110)):
        ws.column_dimensions[col].width = wdt
    r = 5 + len(ops.ops)
    for n in ops.notes:
        ws.cell(r, 1, "• " + n)
        r += 1
    return ws
