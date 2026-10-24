"""견적: 도면(3D) → 부품별 공정·필요 톤수·사용 기계·공정 시간, 재료비, 가공비(기계 임률), 부자재,
관리비·이윤 → 생산 견적서 초안, 그리고 금형 견적서 초안.

  · 자작 부품: 자동 공정 설계(procplan)로 금형 형식·공정을 정하고, costing 으로 공정마다 힘(톤) →
    기계 선택 → 시간 × 임률 (+ 준비 시간 ÷ 로트)
  · 재료비: BOM 원자재 소요량(사각 원판·환봉 kg) × 재질 단가 − 스크랩 × 고철 단가 (부품별 직접 입력 가능)
  · 부자재: BOM 의 표준품·구매품 + 직접 넣는 줄 (도장·포장 …)
  · 금형비: costing.die_quote (설계·강재·인서트·열처리·가공·표준 부품·조립·트라이)
  · 관리비·이윤: 제조원가에 % 로

단가·기계·임률은 모두 예시값이다 — 견적 창에서 회사 값으로 바꾸면 저장된다.
"""
from __future__ import annotations

import copy
import datetime as _dt
import json
import math
from dataclasses import dataclass, field

from . import bendcalc as bc
from . import costing as cs

# 원/kg (예시값, 2026 국내 유통가 어림 — 반드시 회사 단가로 바꿀 것)
MATERIAL_PRICE = {
    "SPCC": 1300, "SPHC": 1100, "SECC": 1500, "SGCC": 1450, "SS275": 1150, "SAPH440": 1250,
    "SGARC440": 1600, "SPFC590": 1700, "SPFC980": 2300, "SUS304": 4500, "SUS430": 3000,
    "AL5052": 6500, "AL1050": 5500, "AL6061": 7500, "C1100": 15000, "C2600": 12000,
}

DEFAULT_RATES = {
    "material": dict(MATERIAL_PRICE),
    "scrap_kg": 300,              # 고철 회수 원/kg
    "overhead_pct": 10.0,         # 일반관리비 %
    "profit_pct": 10.0,           # 이윤 %
    "vat_pct": 10.0,
    "supplier": {"대표": "", "주소": "", "전화": "", "사업자번호": "", "담당": ""},
    "costing": cs.default_costing(),
}


def load_rates(text: str | None):
    r = copy.deepcopy(DEFAULT_RATES)
    if text:
        try:
            saved = json.loads(text)
            for k, v in saved.items():
                if k == "costing" and isinstance(v, dict):
                    if isinstance(v.get("machines"), list) and v["machines"]:
                        r["costing"]["machines"] = v["machines"]
                    if isinstance(v.get("die"), dict):
                        r["costing"]["die"].update(v["die"])
                    if "safety" in v:
                        r["costing"]["safety"] = v["safety"]
                elif isinstance(v, dict) and isinstance(r.get(k), dict):
                    r[k].update(v)
                elif k in r:
                    r[k] = v
        except (ValueError, TypeError, AttributeError):
            pass
    return r


@dataclass
class PartCost:
    no: int
    name: str
    qty: int                 # 세트당
    category: str
    spec: str                # 재질·두께
    method: str              # 금형 형식 또는 레이저 + 절곡기
    opc: list                # costing.OpCost 목록
    weight: float            # 제품 kg (개당)
    gross_kg: float          # 원자재 소요 kg (개당)
    material_auto: float     # 원/개 (스크랩 뺀)
    die: object = None       # costing.DieQuote
    note: str = ""
    process_override: float | None = None
    material_override: float | None = None

    @property
    def ops(self):
        return [o.op for o in self.opc]

    @property
    def n_ops(self):
        return len(self.opc)

    @property
    def material(self):
        return self.material_auto if self.material_override is None else self.material_override

    @property
    def process(self):
        return sum(o.cost for o in self.opc)

    @property
    def process_cost(self):
        return self.process if self.process_override is None else self.process_override

    @property
    def seconds(self):
        return sum(o.seconds for o in self.opc)

    @property
    def need_tonf(self):
        """프레스에 걸리는 가장 큰 힘 (tf). 한 프레스 금형은 공정 합, 아니면 공정 중 최대."""
        press = [o for o in self.opc if o.kind == cs.PRESS]
        if press and ("순차" in self.method or "트랜스퍼" in self.method or "복합" in self.method):
            return sum(o.tonf for o in press)
        return max((o.tonf for o in self.opc), default=0.0)

    @property
    def machines(self):
        out = []
        for o in self.opc:
            if o.machine and o.machine not in out:
                out.append(o.machine)
        return out

    @property
    def die_cost(self):
        return self.die.cost if self.die is not None else 0.0

    @property
    def unit(self):
        return self.material + self.process_cost

    @property
    def amount(self):
        return self.unit * self.qty


@dataclass
class SubItem:
    name: str
    spec: str
    qty: float
    unit_price: float
    unit: str = "EA"
    note: str = ""

    @property
    def amount(self):
        return self.qty * self.unit_price


@dataclass
class Quote:
    title: str
    sets: int
    annual: int
    parts: list
    subs: list
    rates: dict
    amortize_die: bool = False
    customer: str = ""
    number: str = ""
    date: _dt.date = field(default_factory=_dt.date.today)
    valid_days: int = 30
    extra: float = 0.0       # 조립·용접·도장 등 세트당 직접 입력

    # ---- 합계 (세트 1 대당)
    @property
    def material(self):
        return sum(p.material * p.qty for p in self.parts)

    @property
    def process(self):
        return sum(p.process_cost * p.qty for p in self.parts) + self.extra

    @property
    def sub_total(self):
        return sum(s.amount for s in self.subs)

    @property
    def die_total(self):
        return sum(p.die_cost for p in self.parts)

    @property
    def die_per_set(self):
        return self.die_total / max(self.sets, 1) if self.amortize_die else 0.0

    @property
    def manufacturing(self):
        return self.material + self.process + self.sub_total + self.die_per_set

    @property
    def overhead(self):
        return self.manufacturing * self.rates["overhead_pct"] / 100

    @property
    def profit(self):
        return (self.manufacturing + self.overhead) * self.rates["profit_pct"] / 100

    @property
    def unit_price(self):
        """세트 단가 (원, 10 원 단위 올림)."""
        return math.ceil((self.manufacturing + self.overhead + self.profit) / 10) * 10

    @property
    def supply(self):
        return self.unit_price * self.sets + (0 if self.amortize_die else self.die_total)

    @property
    def vat(self):
        return round(self.supply * self.rates["vat_pct"] / 100)

    @property
    def total(self):
        return self.supply + self.vat

    @property
    def total_ops(self):
        return sum(p.n_ops for p in self.parts)

    @property
    def max_tonf(self):
        return max((p.need_tonf for p in self.parts), default=0.0)

    @property
    def machine_list(self):
        out = []
        for p in self.parts:
            for m in p.machines:
                if m not in out:
                    out.append(m)
        return out

    @property
    def dies(self):
        return [p.die for p in self.parts if p.die is not None]


# ---------------------------------------------------------------- 부품 하나

def part_cost(item, settings: bc.BendSettings, rates, annual, cache=None, lot=1):
    """BomItem (자작) → PartCost. cache: {(품번, 생산량): (형식, 공정, 절단 길이, 피어싱)}
    (단가·임률만 바꿀 때 공정 설계를 다시 안 함). lot: 이 부품 로트 수량 (준비 시간 나눔)."""
    from . import metrics as mt
    from . import procplan as pp
    mat = settings.mat
    price = rates["material"].get(mat.code, MATERIAL_PRICE.get(mat.code, 1300))
    rm = item.rm or {}
    gross = float(rm.get("kg", item.weight))
    net = float(rm.get("net_kg", item.weight))
    t = item.thickness or 0.0
    costing = rates["costing"]
    machined = item.result is None or "선반" in item.category or item.category.endswith("(가공)")
    if machined:
        material = gross * price - max(gross - net, 0) * rates["scrap_kg"]
        opc = cs.op_costs(item, settings, "기계 가공", [], 0, 0, costing, lot)
        return PartCost(item.no, item.name, item.qty, item.category, rm.get("spec", ""), "기계 가공", opc,
                        item.weight, gross, material, note="가공비는 무게 어림 — 확인 필요")
    material = gross * price - max(gross - net, 0) * rates["scrap_kg"]
    spec = f"{mat.code} t{t:g}"
    key = (item.no, annual)
    if cache is not None and key in cache:
        method, ops, cut_len, pierce = cache[key]
        ops = list(ops)
    else:
        try:
            rt = pp.route(item.result, settings, annual)
            method, ops = rt.die_type, [s.op for s in rt.steps]
        except Exception:
            method, ops = "금형 없이 레이저 절단 + 절곡기", []
        m = mt.compute(item.result, mat.density)
        cut_len, pierce = m.cut_length, m.pierce
        if cache is not None:
            cache[key] = (method, tuple(ops), cut_len, pierce)
    opc = cs.op_costs(item, settings, method, ops, cut_len, pierce, costing, lot)
    die = cs.die_quote(item, method, ops, costing, settings)
    return PartCost(item.no, item.name, item.qty, item.category, spec, method, opc, item.weight, gross,
                    material, die)


def build(model, parts, settings: bc.BendSettings, rates, sets=1, annual=100_000, excluded=(),
          title="", items=None, cache=None):
    from . import bom
    if items is None:
        items = bom.build(model, parts, settings.mat.density, excluded, settings.mat.code)
    pcs, subs = [], []
    for it in items:
        if it.excluded:
            continue
        if it.procure == "자작":
            pcs.append(part_cost(it, settings, rates, annual, cache, lot=sets * it.qty))
        else:
            subs.append(SubItem(it.name or f"품번 {it.no}", it.size_text, it.qty, 0.0, "EA",
                                it.category))
    return Quote(title, sets, annual, pcs, subs, rates)


# ---------------------------------------------------------------- 금액 글

_DIGITS = "영일이삼사오육칠팔구"
_SMALL = ["", "십", "백", "천"]
_BIG = ["", "만", "억", "조"]


def korean_amount(n: int) -> str:
    """1234500 → '일백이십삼만사천오백'."""
    n = int(round(n))
    if n == 0:
        return "영"
    out = []
    for gi, big in enumerate(_BIG):
        chunk = n % 10000
        n //= 10000
        if chunk:
            s = ""
            for i in range(3, -1, -1):
                d = chunk // (10 ** i) % 10
                if d:
                    s += _DIGITS[d] + _SMALL[i]
            out.append(s + big)
        if not n:
            break
    return "".join(reversed(out))


def short_method(m):
    return (m.replace("금형 없이 레이저 절단 + 절곡기", "레이저 + 절곡기").replace("(프로그레시브) ", "")
            .replace(" (블랭킹 + 피어싱 한 번에)", "").replace(" (공정별 단발 금형)", "")[:16])


def die_note(q):
    if not q.die_total:
        return "금형이 필요 없는 공정입니다 (레이저 절단 + 절곡기 / 기계 가공)."
    return "금형비는 " + ("세트 단가에 나눠 넣었습니다." if q.amortize_die else "별도이며 한 번만 청구합니다.")


def won(x):
    return f"{int(round(x)):,}"


# ---------------------------------------------------------------- 견적서 (엑셀)

def save_xlsx(q: Quote, path: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    from . import company
    from .report import _add_image

    wb = Workbook()
    ws = wb.active
    ws.title = "견적서"
    thin = Side(style="thin", color="888888")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="1F4E8C")
    hf = Font(bold=True, color="FFFFFF")
    widths = [6, 34, 22, 8, 7, 14, 16, 22]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.merge_cells("A1:H1")
    ws["A1"] = "견  적  서"
    ws["A1"].font = Font(bold=True, size=22)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[1].height = 40
    sup = q.rates.get("supplier", {})
    left = [("견적 번호", q.number or f"Q-{q.date:%Y%m%d}"), ("견적 일자", f"{q.date:%Y-%m-%d}"),
            ("수신", f"{q.customer} 귀하" if q.customer else "            귀하"),
            ("유효 기간", f"견적일로부터 {q.valid_days}일"), ("제품", q.title)]
    right = [("상호", company.NAME), ("대표", sup.get("대표", "")), ("주소", sup.get("주소", "")),
             ("전화", sup.get("전화", "")), ("사업자번호", sup.get("사업자번호", "")),
             ("홈페이지", company.HOMEPAGE)]
    for i, (k, v) in enumerate(left):
        ws.cell(3 + i, 1, k).font = Font(bold=True)
        ws.merge_cells(start_row=3 + i, start_column=2, end_row=3 + i, end_column=3)
        ws.cell(3 + i, 2, v)
    ws.cell(2, 6, "공급자").font = Font(bold=True)
    for i, (k, v) in enumerate(right):
        ws.cell(3 + i, 6, k).font = Font(bold=True)
        ws.merge_cells(start_row=3 + i, start_column=7, end_row=3 + i, end_column=8)
        ws.cell(3 + i, 7, v)
    tmp = []
    try:
        tmp.append(_add_image(ws, company.logo_image(60), "H2", 90))
    except Exception:          # 화면 없이 (명령줄) 저장할 때는 로고 없이
        pass
    r = 10
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    ws.cell(r, 1, "아래와 같이 견적합니다.")
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    c = ws.cell(r, 1, f"합계 금액: 일금 {korean_amount(q.total)}원정  (₩ {won(q.total)})  · 부가세 포함")
    c.font = Font(bold=True, size=13)
    c.fill = PatternFill("solid", fgColor="EAF0FA")
    r += 2
    heads = ["번호", "품명", "규격", "수량", "단위", "단가", "금액", "비고"]
    for j, h in enumerate(heads, 1):
        c = ws.cell(r, j, h)
        c.font, c.fill, c.border = hf, head, box
        c.alignment = Alignment(horizontal="center")
    lines = [[1, q.title or "제품", f"자작 {len(q.parts)}종 · 부자재 {len(q.subs)}종", q.sets, "SET",
              q.unit_price, q.unit_price * q.sets, f"공정 {q.total_ops} · 최대 {q.max_tonf:,.0f} tf"]]
    if q.die_total and not q.amortize_die:
        lines.append([2, "금형비", f"{sum(1 for p in q.parts if p.die_cost)} 부품", 1, "식", q.die_total,
                      q.die_total, "별도 (한 번)"])
    for line in lines:
        r += 1
        for j, v in enumerate(line, 1):
            c = ws.cell(r, j, v)
            c.border = box
            if j in (6, 7):
                c.number_format = "#,##0"
    for k, v in (("공급가액", q.supply), (f"부가세 ({q.rates['vat_pct']:g} %)", q.vat), ("합계", q.total)):
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        ws.cell(r, 1, k).alignment = Alignment(horizontal="right")
        ws.cell(r, 1).font = Font(bold=True)
        c = ws.cell(r, 7, v)
        c.number_format = "#,##0"
        c.font = Font(bold=True)
        c.border = box
    r += 2
    notes = ["1. 단가는 도면(3D) 전개 결과와 단가표로 계산한 초안입니다.",
             "2. " + die_note(q),
             "3. 도장·열처리·포장·운송은 따로 적지 않은 경우 포함되지 않습니다.",
             f"4. 연간 예상 {q.annual:,} 개 기준으로 금형 형식을 정했습니다.",
             f"5. 필요 톤수 최대 {q.max_tonf:,.1f} tf · 사용 기계: {', '.join(q.machine_list) or '-'}"]
    for n in notes:
        ws.cell(r, 1, n)
        r += 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # 산출 근거
    w2 = wb.create_sheet("산출 근거")
    h2 = ["품번", "품명", "세트당 수량", "구분", "재질·두께", "공정 방식", "공정수", "공정", "필요 톤수 (tf)",
          "사용 기계", "제품 kg", "소요 kg", "재료비/개", "가공비/개", "단가/개", "세트당 금액", "금형비", "비고"]
    for j, h in enumerate(h2, 1):
        c = w2.cell(1, j, h)
        c.font, c.fill = hf, head
        w2.column_dimensions[c.column_letter].width = max(9, len(h) * 2)
    w2.column_dimensions["B"].width = 26
    w2.column_dimensions["F"].width = 28
    w2.column_dimensions["H"].width = 50
    w2.column_dimensions["J"].width = 30
    for i, p in enumerate(q.parts, 2):
        vals = [p.no, p.name, p.qty, p.category, p.spec, p.method, p.n_ops, " → ".join(p.ops),
                round(p.need_tonf, 1), ", ".join(p.machines),
                round(p.weight, 3), round(p.gross_kg, 3), round(p.material), round(p.process_cost), round(p.unit),
                round(p.amount), round(p.die_cost), p.note]
        for j, v in enumerate(vals, 1):
            c = w2.cell(i, j, v)
            if 13 <= j <= 17:
                c.number_format = "#,##0"
    r2 = len(q.parts) + 3
    w2.cell(r2, 1, "부자재").font = Font(bold=True)
    for j, h in enumerate(["품명", "규격", "수량", "단위", "단가", "금액", "비고"], 2):
        c = w2.cell(r2 + 1, j, h)
        c.font, c.fill = hf, head
    for i, s in enumerate(q.subs, r2 + 2):
        for j, v in enumerate([s.name, s.spec, s.qty, s.unit, s.unit_price, s.amount, s.note], 2):
            c = w2.cell(i, j, v)
            if j in (6, 7):
                c.number_format = "#,##0"
    r3 = r2 + len(q.subs) + 3
    summary = [("재료비 (세트)", q.material), ("가공비 (세트, 조립 등 포함)", q.process), ("부자재비 (세트)", q.sub_total),
               ("금형비 상각 (세트)", q.die_per_set), ("제조원가 (세트)", q.manufacturing),
               (f"일반관리비 {q.rates['overhead_pct']:g} %", q.overhead), (f"이윤 {q.rates['profit_pct']:g} %", q.profit),
               ("세트 단가", q.unit_price), ("금형비 합계", q.die_total)]
    for i, (k, v) in enumerate(summary, r3):
        w2.cell(i, 2, k).font = Font(bold=True)
        w2.cell(i, 3, round(v)).number_format = "#,##0"
    # 공정별 (톤수·기계·시간·비용)
    w4 = wb.create_sheet("공정별")
    h4 = ["품번", "품명", "공정", "기계 종류", "힘 (tf)", "사용 기계", "시간 (초/개)", "가공비 (원/개)", "비고"]
    for j, h in enumerate(h4, 1):
        c = w4.cell(1, j, h)
        c.font, c.fill = hf, head
        w4.column_dimensions[c.column_letter].width = 14
    w4.column_dimensions["B"].width = 24
    w4.column_dimensions["F"].width = 22
    w4.column_dimensions["I"].width = 48
    r4 = 2
    for p in q.parts:
        for o in p.opc:
            for j, v in enumerate([p.no, p.name, o.op, o.kind, round(o.tonf, 2), o.machine, round(o.seconds, 1),
                                   round(o.cost, 1), o.note], 1):
                w4.cell(r4, j, v)
            r4 += 1
    # 기계·임률
    w5 = wb.create_sheet("기계·임률")
    for j, (_, h) in enumerate(cs.MACHINE_COLS, 1):
        c = w5.cell(1, j, h)
        c.font, c.fill = hf, head
        w5.column_dimensions[c.column_letter].width = 16
    used = set(q.machine_list)
    for i, mch in enumerate(q.rates["costing"]["machines"], 2):
        for j, (k, _) in enumerate(cs.MACHINE_COLS, 1):
            c = w5.cell(i, j, mch.get(k, ""))
            if mch["name"] in used:
                c.font = Font(bold=True)
    w3 = wb.create_sheet("단가표")
    w3.append(["항목", "값"])
    for k, v in q.rates["material"].items():
        w3.append([f"재료 {k} (원/kg)", v])
    for k in RATE_NAMES:
        w3.append([RATE_NAMES[k], q.rates[k]])
    w3.append(["기계 선택 여유 (배)", q.rates["costing"].get("safety", cs.SAFETY)])
    wb.save(path)
    import os
    for t in tmp:
        try:
            os.remove(t)
        except (OSError, TypeError):
            pass


RATE_NAMES = {
    "scrap_kg": "고철 회수 (원/kg)", "overhead_pct": "일반관리비 (%)", "profit_pct": "이윤 (%)",
    "vat_pct": "부가세 (%)",
}


# ---------------------------------------------------------------- 견적서 (PDF)

def save_pdf(q: Quote, path: str):
    from PySide6.QtCore import QMarginsF, QRectF, Qt
    from PySide6.QtGui import QColor, QPageLayout, QPageSize, QPainter, QPdfWriter, QPen

    from . import company
    from .report import _brand, _font, ensure_gui, version_text
    ensure_gui()
    pdf = QPdfWriter(path)
    pdf.setResolution(300)
    pdf.setPageLayout(QPageLayout(QPageSize(QPageSize.A4), QPageLayout.Portrait, QMarginsF(0, 0, 0, 0)))
    pdf.setTitle(f"견적서 - {q.title}")
    pdf.setCreator(f"한일 판금 전개 {version_text()}")
    k = 300 / 25.4
    W, H, m = 210.0, 297.0, 15.0

    def R(x, y, w, h):
        return QRectF(x * k, y * k, w * k, h * k)

    p = QPainter(pdf)
    p.setRenderHint(QPainter.Antialiasing)

    def text(x, y, w, h, s, size=9, bold=False, align=Qt.AlignLeft, color=QColor(20, 20, 20)):
        p.setFont(_font(size, bold))
        p.setPen(color)
        p.drawText(R(x, y, w, h), align | Qt.AlignVCenter, s)

    def table(x, y, cols, rows, size=8.5, rh=6.5):
        tw = sum(c[1] for c in cols)
        p.setPen(Qt.NoPen)
        p.setBrush(QColor(31, 78, 140))
        p.drawRect(R(x, y, tw, rh))
        cx = x
        for title, cw, al in cols:
            text(cx + 1.2, y, cw - 2.4, rh, title, size, True, Qt.AlignCenter, QColor(255, 255, 255))
            cx += cw
        y += rh
        for i, row in enumerate(rows):
            p.setPen(QPen(QColor(200, 200, 200), 0.15 * k))
            p.setBrush(QColor(245, 248, 252) if i % 2 else QColor(255, 255, 255))
            p.drawRect(R(x, y, tw, rh))
            cx = x
            for (title, cw, al), v in zip(cols, row):
                text(cx + 1.2, y, cw - 2.4, rh, str(v), size, False, al)
                cx += cw
            y += rh
        return y

    # ---- 1 쪽: 견적서
    text(m, m, W - 2 * m, 14, "견  적  서", 24, True, Qt.AlignCenter)
    sup = q.rates.get("supplier", {})
    y = m + 20
    for i, (kk, v) in enumerate([("견적 번호", q.number or f"Q-{q.date:%Y%m%d}"), ("견적 일자", f"{q.date:%Y-%m-%d}"),
                                 ("수신", f"{q.customer} 귀하" if q.customer else "귀하"),
                                 ("유효 기간", f"견적일로부터 {q.valid_days}일"), ("제품", q.title)]):
        text(m, y + i * 6.5, 22, 6, kk, 9, True)
        text(m + 22, y + i * 6.5, 70, 6, v, 9)
    bx = 112
    p.setPen(QPen(QColor(120, 120, 120), 0.25 * k))
    p.setBrush(Qt.NoBrush)
    p.drawRect(R(bx, y - 2, W - m - bx, 44))
    text(bx + 2, y - 1, 30, 6, "공급자", 9, True, color=QColor(31, 78, 140))
    _brand(p, R(bx + 2, y - 1, W - m - bx - 4, 5))
    for i, (kk, v) in enumerate([("상호", company.NAME), ("대표", sup.get("대표", "")), ("주소", sup.get("주소", "")),
                                 ("전화", sup.get("전화", "")), ("사업자번호", sup.get("사업자번호", "")),
                                 ("홈페이지", company.HOMEPAGE)]):
        text(bx + 2, y + 6 + i * 6, 20, 6, kk, 8.5, True)
        text(bx + 22, y + 6 + i * 6, W - m - bx - 24, 6, v, 8.5)
    y += 48
    text(m, y, W - 2 * m, 6, "아래와 같이 견적합니다.", 10)
    y += 8
    p.setPen(Qt.NoPen)
    p.setBrush(QColor(234, 240, 250))
    p.drawRect(R(m, y, W - 2 * m, 10))
    text(m + 3, y, W - 2 * m - 6, 10, f"합계 금액:  일금 {korean_amount(q.total)}원정   (₩ {won(q.total)},  부가세 포함)",
         11, True)
    y += 15
    cols = [("번호", 12, Qt.AlignCenter), ("품명", 52, Qt.AlignLeft), ("규격", 36, Qt.AlignLeft),
            ("수량", 14, Qt.AlignRight), ("단위", 12, Qt.AlignCenter), ("단가", 26, Qt.AlignRight),
            ("금액", 28, Qt.AlignRight)]
    rows = [[1, q.title or "제품", f"자작 {len(q.parts)}종 · 부자재 {len(q.subs)}종", q.sets, "SET",
             won(q.unit_price), won(q.unit_price * q.sets)]]
    if q.die_total and not q.amortize_die:
        rows.append([2, "금형비 (별도, 한 번)", f"{sum(1 for x in q.parts if x.die_cost)} 부품", 1, "식",
                     won(q.die_total), won(q.die_total)])
    y = table(m, y, cols, rows)
    for kk, v in (("공급가액", q.supply), (f"부가세 {q.rates['vat_pct']:g} %", q.vat), ("합계", q.total)):
        text(m, y, 152, 6.5, kk, 9.5, True, Qt.AlignRight)
        text(m + 152, y, 28, 6.5, won(v), 9.5, True, Qt.AlignRight)
        y += 6.5
    y += 6
    text(m, y, W - 2 * m, 6, "견적 요약 (세트 1 대)", 10, True, color=QColor(31, 78, 140))
    y += 7
    srows = [["재료비", won(q.material)], ["가공비 (조립 등 포함)", won(q.process)], ["부자재비", won(q.sub_total)],
             ["금형비 상각", won(q.die_per_set)], ["제조원가", won(q.manufacturing)],
             [f"일반관리비 {q.rates['overhead_pct']:g} %", won(q.overhead)], [f"이윤 {q.rates['profit_pct']:g} %", won(q.profit)],
             ["세트 단가", won(q.unit_price)]]
    y = table(m, y, [("항목", 60, Qt.AlignLeft), ("금액 (원)", 40, Qt.AlignRight)], srows)
    text(m + 108, y - 8 * 6.5, W - 2 * m - 108, 6.5, f"공정 합계 {q.total_ops} · 자작 {len(q.parts)}종", 9)
    text(m + 108, y - 7 * 6.5, W - 2 * m - 108, 6.5, f"연간 예상 {q.annual:,} 개 기준 금형 형식", 9)
    text(m + 108, y - 6 * 6.5, W - 2 * m - 108, 6.5, f"필요 톤수 최대 {q.max_tonf:,.1f} tf", 9, True)
    yy = y - 5 * 6.5
    for mname in q.machine_list[:5]:
        text(m + 108, yy, W - 2 * m - 108, 6, "· " + mname, 8.5)
        yy += 6
    y += 6
    for n in ["※ 도면(3D) 전개 결과와 단가표로 계산한 초안입니다. 확정 전 검토하세요.",
              "※ " + die_note(q),
              "※ 도장·열처리·포장·운송은 따로 적지 않은 경우 포함되지 않습니다."]:
        text(m, y, W - 2 * m, 5.5, n, 8.5, color=QColor(90, 90, 90))
        y += 5.5

    # ---- 2 쪽: 산출 근거
    pdf.newPage()
    text(m, m, W - 2 * m, 9, f"산출 근거 — {q.title}", 14, True)
    y = m + 12
    cols = [("품번", 8, Qt.AlignCenter), ("품명", 26, Qt.AlignLeft), ("수량", 8, Qt.AlignRight),
            ("재질", 17, Qt.AlignLeft), ("방식", 24, Qt.AlignLeft), ("공정", 8, Qt.AlignRight),
            ("톤 tf", 11, Qt.AlignRight), ("기계", 26, Qt.AlignLeft),
            ("재료/개", 16, Qt.AlignRight), ("가공/개", 16, Qt.AlignRight), ("세트 금액", 20, Qt.AlignRight)]
    rows = [[x.no, x.name[:14], x.qty, x.spec, short_method(x.method)[:12], x.n_ops, f"{x.need_tonf:,.1f}",
             ", ".join(x.machines)[:16], won(x.material), won(x.process_cost), won(x.amount)] for x in q.parts]
    y = table(m, y, cols, rows, 7, 5.6)
    y += 4
    for x in q.parts:
        if y > H - 40:
            break
        steps = " → ".join(f"{o.op}" + (f" ({o.machine.split()[0] if o.machine else ''} {o.tonf:.0f}tf)"
                                         if o.tonf >= 0.5 else "") for o in x.opc)
        text(m, y, W - 2 * m, 5, f"{x.no}. {x.name[:16]}: {steps}" + (f"   ({x.note})" if x.note else ""),
             6.8, color=QColor(60, 60, 60))
        y += 4.6
    if q.subs:
        y += 4
        text(m, y, W - 2 * m, 6, "부자재", 10, True, color=QColor(31, 78, 140))
        y += 7
        cols = [("품명", 60, Qt.AlignLeft), ("규격", 40, Qt.AlignLeft), ("수량", 16, Qt.AlignRight),
                ("단가", 30, Qt.AlignRight), ("금액", 34, Qt.AlignRight)]
        y = table(m, y, cols, [[s.name[:30], s.spec[:22], f"{s.qty:g}", won(s.unit_price), won(s.amount)]
                               for s in q.subs if y < H - 30], 7.5, 5.6)
    text(m, H - m - 5, W - 2 * m, 5, f"한일 판금 전개 {version_text()}  ·  {company.MAKER}", 7.5,
         color=QColor(130, 130, 130))
    p.end()


# ---------------------------------------------------------------- 금형 견적서

def die_total_with_vat(q: Quote):
    supply = sum(d.cost for d in q.dies)
    vat = round(supply * q.rates["vat_pct"] / 100)
    return supply, vat, supply + vat


def save_die_xlsx(q: Quote, path: str):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from . import company
    wb = Workbook()
    ws = wb.active
    ws.title = "금형 견적서"
    head = PatternFill("solid", fgColor="1F4E8C")
    hf = Font(bold=True, color="FFFFFF")
    for i, w in enumerate([8, 30, 26, 12, 46, 16], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.merge_cells("A1:F1")
    ws["A1"] = "금 형  견 적 서"
    ws["A1"].font = Font(bold=True, size=20)
    ws["A1"].alignment = Alignment(horizontal="center")
    supply, vat, total = die_total_with_vat(q)
    sup = q.rates.get("supplier", {})
    info = [("견적 일자", f"{q.date:%Y-%m-%d}"), ("수신", f"{q.customer} 귀하" if q.customer else "귀하"),
            ("제품", q.title), ("공급자", f"{company.NAME}  {sup.get('전화', '')}"),
            ("합계 금액", f"일금 {korean_amount(total)}원정 (₩ {won(total)}, 부가세 포함)")]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=6)
        ws.cell(i, 2, v)
    r = 9
    for j, h in enumerate(["번호", "금형", "형식", "공정수", "크기 (길이×폭×두께)", "금액"], 1):
        c = ws.cell(r, j, h)
        c.font, c.fill = hf, head
    for i, d in enumerate(q.dies, 1):
        r += 1
        for j, v in enumerate([i, f"{d.part_name} (품번 {d.part_no})", short_method(d.die_type), d.stations,
                               f"{d.size[0]:.0f} × {d.size[1]:.0f} × {d.size[2]:.0f}", round(d.cost)], 1):
            c = ws.cell(r, j, v)
            if j == 6:
                c.number_format = "#,##0"
    for k, v in (("공급가액", supply), (f"부가세 {q.rates['vat_pct']:g} %", vat), ("합계", total)):
        r += 1
        ws.cell(r, 5, k).font = Font(bold=True)
        c = ws.cell(r, 6, v)
        c.number_format = "#,##0"
        c.font = Font(bold=True)
    r += 2
    for n in ["1. 금형 구조·크기는 전개도와 공정 설계로 어림한 초안입니다 (금형 설계 후 확정).",
              "2. 트라이 후 제품 수정·설계 변경 비용은 포함되지 않습니다.",
              "3. 금형 보증 타수·보관·수리 조건은 따로 정합니다."]:
        ws.cell(r, 1, n)
        r += 1
    ws2 = wb.create_sheet("금형 산출 근거")
    r = 1
    for d in q.dies:
        ws2.cell(r, 1, f"{d.part_name} — {d.die_type}, {d.stations} 공정, "
                       f"{d.size[0]:.0f}×{d.size[1]:.0f}×{d.size[2]:.0f} mm").font = Font(bold=True)
        r += 1
        for name, formula, amt in d.lines:
            ws2.cell(r, 1, name)
            ws2.cell(r, 2, formula)
            ws2.cell(r, 3, round(amt)).number_format = "#,##0"
            r += 1
        ws2.cell(r, 1, "소계").font = Font(bold=True)
        ws2.cell(r, 3, round(d.cost)).number_format = "#,##0"
        r += 2
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 60
    ws2.column_dimensions["C"].width = 16
    ws3 = wb.create_sheet("금형 단가")
    for k, v in q.rates["costing"]["die"].items():
        ws3.append([cs.DIE_NAMES.get(k, k), v])
    ws3.column_dimensions["A"].width = 34
    wb.save(path)


def save_die_pdf(q: Quote, path: str):
    from PySide6.QtCore import QMarginsF, QRectF, Qt
    from PySide6.QtGui import QColor, QPageLayout, QPageSize, QPainter, QPdfWriter, QPen

    from . import company
    from .report import _brand, _font, ensure_gui, version_text
    ensure_gui()
    pdf = QPdfWriter(path)
    pdf.setResolution(300)
    pdf.setPageLayout(QPageLayout(QPageSize(QPageSize.A4), QPageLayout.Portrait, QMarginsF(0, 0, 0, 0)))
    pdf.setTitle(f"금형 견적서 - {q.title}")
    k = 300 / 25.4
    W, H, m = 210.0, 297.0, 15.0

    def R(x, y, w, h):
        return QRectF(x * k, y * k, w * k, h * k)

    p = QPainter(pdf)

    def text(x, y, w, h, s, size=9, bold=False, align=Qt.AlignLeft, color=QColor(20, 20, 20)):
        p.setFont(_font(size, bold))
        p.setPen(color)
        p.drawText(R(x, y, w, h), align | Qt.AlignVCenter, s)

    supply, vat, total = die_total_with_vat(q)
    text(m, m, W - 2 * m, 14, "금 형  견 적 서", 22, True, Qt.AlignCenter)
    _brand(p, R(m, m + 16, W - 2 * m, 6))
    y = m + 25
    for kk, v in [("견적 일자", f"{q.date:%Y-%m-%d}"), ("수신", f"{q.customer} 귀하" if q.customer else "귀하"),
                  ("제품", q.title), ("공급자", company.NAME)]:
        text(m, y, 25, 6, kk, 9.5, True)
        text(m + 25, y, 140, 6, v, 9.5)
        y += 6.5
    y += 3
    p.setPen(Qt.NoPen)
    p.setBrush(QColor(234, 240, 250))
    p.drawRect(R(m, y, W - 2 * m, 10))
    text(m + 3, y, W - 2 * m - 6, 10, f"합계 금액: 일금 {korean_amount(total)}원정 (₩ {won(total)}, 부가세 포함)",
         11, True)
    y += 15
    for d in q.dies:
        if y > H - 60:
            pdf.newPage()
            y = m
        text(m, y, W - 2 * m, 7, f"{d.part_name} — {short_method(d.die_type)}, {d.stations} 공정, "
                                 f"{d.size[0]:.0f}×{d.size[1]:.0f}×{d.size[2]:.0f} mm", 10.5, True,
             color=QColor(31, 78, 140))
        y += 8
        for name, formula, amt in d.lines:
            text(m + 2, y, 50, 5.6, name, 8.5)
            text(m + 52, y, 98, 5.6, formula, 8, color=QColor(90, 90, 90))
            text(m + 150, y, 30, 5.6, won(amt), 8.5, align=Qt.AlignRight)
            y += 5.6
        p.setPen(QPen(QColor(120, 120, 120), 0.2 * k))
        p.drawLine(R(m, y, W - 2 * m, 0).topLeft(), R(m, y, W - 2 * m, 0).topRight())
        text(m + 2, y + 0.5, 50, 6, "소계", 9, True)
        text(m + 150, y + 0.5, 30, 6, won(d.cost), 9, True, Qt.AlignRight)
        y += 10
    for kk, v in (("공급가액", supply), (f"부가세 {q.rates['vat_pct']:g} %", vat), ("합계", total)):
        text(m + 100, y, 50, 6.5, kk, 10, True, Qt.AlignRight)
        text(m + 150, y, 30, 6.5, won(v), 10, True, Qt.AlignRight)
        y += 6.5
    y += 4
    for n in ["※ 금형 구조·크기는 전개도와 공정 설계로 어림한 초안입니다 (금형 설계 후 확정).",
              "※ 트라이 후 제품 수정·설계 변경 비용은 포함되지 않습니다."]:
        text(m, y, W - 2 * m, 5.5, n, 8.5, color=QColor(90, 90, 90))
        y += 5.5
    text(m, H - m - 5, W - 2 * m, 5, f"한일 판금 전개 {version_text()}  ·  {company.MAKER}", 7.5,
         color=QColor(130, 130, 130))
    p.end()


# ---------------------------------------------------------------- 기록용

def record_values(q: Quote, die=False):
    """History.add_doc 에 넘길 값 (생산 견적 / 금형 견적)."""
    if die:
        supply, vat, total = die_total_with_vat(q)
        return dict(title=q.title, customer=q.customer, sets=q.sets, annual=q.annual, unit_price=None,
                    supply=supply, total=total, die_total=supply, max_tonf=q.max_tonf,
                    machines=", ".join(q.machine_list), total_ops=sum(d.stations for d in q.dies),
                    parts=len(q.dies),
                    summary=f"금형 {len(q.dies)} 벌 · 공급가 {won(supply)} 원 · "
                            + ", ".join(f"{d.part_name} {short_method(d.die_type)} {d.stations}공정" for d in q.dies))
    return dict(title=q.title, customer=q.customer, sets=q.sets, annual=q.annual, unit_price=q.unit_price,
                supply=q.supply, total=q.total, die_total=q.die_total, max_tonf=q.max_tonf,
                machines=", ".join(q.machine_list), total_ops=q.total_ops, parts=len(q.parts),
                summary=f"세트 단가 {won(q.unit_price)} × {q.sets:,} · 재료 {won(q.material)} · 가공 {won(q.process)}"
                        f" · 부자재 {won(q.sub_total)} · 금형 {won(q.die_total)}"
                        + (" (상각)" if q.amortize_die else ""))


def record_detail(q: Quote, die=False):
    if die:
        return [{"no": d.part_no, "name": d.part_name, "method": d.die_type, "n_ops": d.stations,
                 "size": list(d.size), "lines": [[a, b, round(c)] for a, b, c in d.lines], "die": round(d.cost)}
                for d in q.dies]
    return [{"no": p.no, "name": p.name, "qty": p.qty, "spec": p.spec, "method": p.method, "n_ops": p.n_ops,
             "ops": " → ".join(p.ops), "tonf": round(p.need_tonf, 1), "machines": ", ".join(p.machines),
             "material": round(p.material), "process": round(p.process_cost), "die": round(p.die_cost),
             "material_override": p.material_override is not None,
             "process_override": p.process_override is not None} for p in q.parts] + \
        [{"no": "", "name": s.name, "qty": s.qty, "spec": s.spec, "method": "부자재", "material": round(s.unit_price)}
         for s in q.subs]
