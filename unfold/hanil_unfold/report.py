"""전개 결과 보고서 — PDF 와 엑셀.

PDF  : A4 한 장(절곡이 많으면 두 장). 부품 정보, 재료·무게·절단 길이, 치수가 들어간 전개도,
       3D 그림, 절곡 표(연신·근거), 성형품이면 두께 변화 색 지도, 주의 사항.
엑셀 : [보고서] 시트(같은 정보 + 그림), [절곡표] 시트, 성형품이면 [두께 변화] 시트.
여러 파일을 한꺼번에 처리하면 [부품 목록] 엑셀도 만든다 (견적용).
"""
from __future__ import annotations

import datetime as _dt
import math
import os
import sys
import tempfile
from dataclasses import dataclass, field

import numpy as np

from . import __version__, version_text
from . import bendcalc as bc
from . import metrics as mt


@dataclass
class ReportData:
    part_name: str
    source_path: str
    result: object
    settings: bc.BendSettings
    solid: object = None
    solid_index: int = 1
    solid_count: int = 1
    part_title: str = ""         # 조립품 안의 부품 이름 (조립도 번호와 같은 순서)
    created: _dt.datetime = field(default_factory=_dt.datetime.now)
    user_dims: list = field(default_factory=list)   # 화면에서 찍은 치수 [(p1, p2)]

    @property
    def material(self) -> bc.Material:
        return self.settings.mat

    @property
    def metrics(self) -> mt.Metrics:
        if not hasattr(self, "_m"):
            self._m = mt.compute(self.result, self.material.density)
        return self._m

    @property
    def dims(self):
        if not hasattr(self, "_d"):
            from . import dims as dimsmod
            self._d = dimsmod.compute(self.result, self.material.density)
        return self._d

    @property
    def check_rows(self):
        """전개 검증표 (unfoldcheck) — 3D 솔리드가 있을 때."""
        if not hasattr(self, "_chk"):
            self._chk = []
            try:
                from . import unfoldcheck
                self._chk = unfoldcheck.check(self.solid, self.result, self.settings)
            except Exception:  # noqa: BLE001 — 검증은 덤
                self._chk = []
        return self._chk

    def dim_rows(self):
        """보고서 치수 표: 요약 표와 겹치는 면적·무게 줄은 빼고 위치·구멍만."""
        return [r for r in self.dims.rows()
                if not r[0].startswith(("두께", "전체 크기", "블랭크", "외곽", "사각", "절단"))]

    @property
    def title(self):
        return "성형품 블랭크 전개 보고서" if self.result.mode == "forming" else "판금 전개 보고서"

    @property
    def display_name(self):
        if self.part_title and self.solid_count > 1:
            return f"{self.part_name} — {self.solid_index}. {self.part_title}"
        if self.solid_count > 1:
            return f"{self.part_name} ({self.solid_index}/{self.solid_count})"
        return self.part_name


# ---------------------------------------------------------------- 표에 들어갈 값

def summary_rows(d: ReportData):
    r, m = d.result, d.metrics
    rows = [
        ("파일", os.path.basename(d.source_path)),
        ("부품", d.display_name),
        ("전개 방식", "성형품 근사 전개 (곡면)" if r.mode == "forming" else "판금 정확 전개"),
        ("재질", f"{d.material.code}  {d.material.name}"),
        ("두께", f"{r.thickness:.3f} mm"),
    ]
    if r.mode != "forming":
        rows.append(("계산 조건", r.settings_text))
    rows += [
        ("전개 크기 (가로 × 세로)", f"{m.width:.2f} × {m.height:.2f} mm"),
        ("블랭크 면적 (구멍 뺌)", f"{m.blank_area:,.1f} mm²  ({m.blank_area / 1e6:.4f} m²)"),
        ("외곽 면적 (구멍 포함)", f"{m.outer_area:,.1f} mm²"),
        ("사각 원판 면적", f"{m.rect_area:,.1f} mm²   수율 {m.yield_ratio:.1%}"),
        ("블랭크 무게", f"{m.weight_kg:.3f} kg   (사각 원판 {m.rect_weight_kg:.3f} kg, "
                     f"비중 {d.material.density:g})"),
        ("절단 길이", f"{m.cut_length:,.1f} mm   (피어싱 {m.pierce}곳, 구멍 {m.holes}개)"),
    ]
    if r.mode == "forming":
        fi = r.forming
        lo, hi = fi.thinning_range
        rows += [
            ("제품 면적 (두께 가운데)", f"{fi.product_area:,.1f} mm²"),
            ("블랭크 / 제품 면적", f"{fi.blank_area / fi.product_area:.3f}"),
            ("최대 얇아짐", f"{max(hi, 0):.1%}  (두께 약 {r.thickness * (1 - max(hi, 0)):.2f} mm)"),
            ("최대 두꺼워짐", f"{max(-lo, 0):.1%}"),
            ("전개 방법", fi.method),
        ]
        if fi.flc0 > 0:
            split, marginal, wrinkle, top = fi.flc_summary
            rows += [
                ("성형 한계 (FLC0)", f"진변형률 {fi.flc0:.3f}  ({fi.material}, 두께 {r.thickness:g} mm)"),
                ("성형 한계 판정", f"최대 {top:.0%} — 넘음 {split:.1%} · 80 % 이상 {marginal:.1%} · "
                                 f"주름 가능 {wrinkle:.1%}"),
            ]
    else:
        rows.append(("절곡 수", f"{len(r.bends)}곳"))
    for txt in r.rolled:
        rows.append(("말기", txt))
    return rows


BEND_HEADER = ["번호", "방향", "각도(°)", "모델 R", "실제 R", "K", "연신", "90° 연신",
               "펼친 길이", "절곡선 길이", "근거"]


def _f(x, n=3):
    return "-" if x is None or (isinstance(x, float) and math.isnan(x)) else f"{x:.{n}f}"


def bend_rows(d: ReportData):
    return [[f"B{b.number}", "UP" if b.direction == "UP" else "DOWN", f"{b.angle_deg:.2f}",
             f"{b.model_radius:.2f}", b.radius_text, f"{b.k:.3f}", _f(b.deduction), _f(b.deduction90),
             f"{b.allowance:.3f}", f"{b.length:.2f}", b.source] for b in d.result.bends]


def notes(d: ReportData):
    out = list(d.result.warnings) + list(getattr(d.result, "advice", []))
    rows = d.check_rows
    if rows:
        from . import unfoldcheck
        out.append(f"전개 검증 {unfoldcheck.verdict(rows)}: " + " · ".join(
            f"{r[0].split(' (')[0]} {r[3] or r[2]}" for r in rows))
    if d.result.mode == "forming":
        out.append("두께 변화는 블랭크→제품 면적 비로 어림한 값입니다 (체적 일정). "
                   "실제 성형성은 성형 해석으로 확인하세요.")
    else:
        out.append("연신 = 2·tan(θ/2)·(R+t) − θ·(R+K·t).  문헌값은 출발점이므로 시험 절곡으로 확인하세요.")
    out.append("방향 UP/DOWN 은 전개도를 앞에서 본 기준입니다.")
    from . import tolerances as tl
    w, h = d.result.size
    out.append(f"{tl.NOTE} — 전개 가로 {w:,.1f} {tl.text(w)}, 세로 {h:,.1f} {tl.text(h)}")
    return out


# ---------------------------------------------------------------- Qt 준비

def new_image(w, h, dpr=1.0):
    """흰 바탕 QImage 와 QPainter. dpr (화면 배율 1.25 · 1.5 · 2) 을 주면 실제 픽셀은 그만큼 크게 그리고 논리 크기는
    w × h 로 — 창에 넣었을 때 흐리지 않게 (고해상도 화면). 보고서 (PDF · 엑셀) 는 1."""
    from PySide6.QtGui import QColor, QImage, QPainter
    ensure_gui()
    dpr = max(float(dpr or 1.0), 1.0)
    img = QImage(int(round(w * dpr)), int(round(h * dpr)), QImage.Format_ARGB32)
    img.setDevicePixelRatio(dpr)
    img.fill(QColor(255, 255, 255))
    p = QPainter(img)
    p.setRenderHint(QPainter.Antialiasing)
    p.setRenderHint(QPainter.TextAntialiasing)
    p.setRenderHint(QPainter.SmoothPixmapTransform)
    return img, p


def render_fld(fi, w=900, h=700, dpr=1.0):
    """성형 한계도 (FLD): 삼각형마다 (부 변형률, 주 변형률) 점과 성형 한계 곡선."""
    import numpy as np
    from PySide6.QtCore import QPointF, QRectF, Qt
    from PySide6.QtGui import QColor, QImage, QPainter, QPen, QPolygonF

    from .forming import flc_limit

    img, p = new_image(w, h, dpr)
    keep = fi.reliable if fi.reliable is not None and fi.reliable.any() else np.ones(len(fi.F), bool)
    e1, e2 = fi.major[keep], fi.minor[keep]
    top = max(0.1, fi.flc0 * 1.6, float(np.percentile(e1, 99.5)) * 1.15 if len(e1) else 0)
    side = max(0.1, float(np.percentile(np.abs(e2), 99.5)) * 1.15 if len(e2) else 0, top * 0.6)
    bottom = -max(0.05, top * 0.15)
    L, T, R_, B = 80, 40, w - 30, h - 70

    def pt(x, y):
        return QPointF(L + (x + side) / (2 * side) * (R_ - L), B - (y - bottom) / (top - bottom) * (B - T))

    p.setFont(_font(13))
    p.setPen(QPen(QColor(200, 200, 200), 1))
    for i in range(-4, 5):
        x = side * i / 4
        p.setPen(QPen(QColor(210, 210, 210), 1))
        p.drawLine(pt(x, bottom), pt(x, top))
        p.setPen(QColor(90, 90, 90))
        p.drawText(QRectF(pt(x, bottom).x() - 40, B + 6, 80, 20), Qt.AlignCenter, f"{x:.2f}")
    for i in range(0, 6):
        y = bottom + (top - bottom) * i / 5
        p.setPen(QPen(QColor(210, 210, 210), 1))
        p.drawLine(pt(-side, y), pt(side, y))
        p.setPen(QColor(90, 90, 90))
        p.drawText(QRectF(0, pt(0, y).y() - 10, L - 8, 20), Qt.AlignRight | Qt.AlignVCenter, f"{y:.2f}")
    p.setPen(QPen(QColor(60, 60, 60), 1.5))
    p.drawLine(pt(-side, 0), pt(side, 0))
    p.drawLine(pt(0, bottom), pt(0, top))
    if fi.flc0 > 0:
        rk = fi.risk[keep]
        colors = [(rk >= 1, QColor(220, 30, 30)), ((rk >= 0.8) & (rk < 1), QColor(240, 170, 0)),
                  (rk < 0.8, QColor(40, 150, 70))]
    else:
        colors = [(np.ones(len(e1), bool), QColor(40, 110, 200))]
    wr = fi.wrinkle[keep]
    step = max(1, len(e1) // 6000)
    p.setClipRect(QRectF(L, T, R_ - L, B - T))
    p.setPen(Qt.NoPen)
    for sel, col in colors:
        idx = np.nonzero(sel)[0][::step]
        for i in idx:
            c = QColor(150, 80, 200) if wr[i] else col
            p.setBrush(c)
            p.drawEllipse(pt(float(e2[i]), float(e1[i])), 2.2, 2.2)
    if fi.flc0 > 0:
        xs = np.linspace(-side, side, 81)
        for scale, pen in ((1.0, QPen(QColor(200, 0, 0), 3)), (0.8, QPen(QColor(230, 150, 0), 2, Qt.DashLine))):
            ys = flc_limit(xs, fi.flc0, fi.material) * scale
            p.setPen(pen)
            p.setBrush(Qt.NoBrush)
            p.drawPolyline(QPolygonF([pt(float(a), float(b)) for a, b in zip(xs, ys) if b <= top]))
    p.setPen(QPen(QColor(150, 80, 200), 1.5, Qt.DotLine))       # 두께 일정 (ε1 = −ε2) 선
    p.drawLine(pt(-min(side, top), min(side, top)), pt(0, 0))
    p.setClipping(False)
    p.setPen(QColor(30, 30, 30))
    p.setFont(_font(15))
    p.drawText(QRectF(0, 6, w, 28), Qt.AlignCenter,
               f"성형 한계도 (FLD) — {fi.material}  FLC0 = {fi.flc0:.3f}"
               + (" · 실측 FLC" if fi.material in __import__("hanil_unfold.forming", fromlist=["x"]).MEASURED_FLC else "")
               if fi.flc0 > 0 else "변형률 분포")
    p.setFont(_font(13))
    p.drawText(QRectF(0, h - 34, w, 24), Qt.AlignCenter,
               "가로: 부 변형률 ε2   세로: 주 변형률 ε1 (진변형률)   빨강 실선 FLC · 주황 점선 80 % · "
               "보라 = 두꺼워짐(주름 가능)")
    p.end()
    return img


def ensure_gui():
    from PySide6.QtGui import QGuiApplication
    if QGuiApplication.instance() is None:
        if sys.platform.startswith("linux") and not os.environ.get("DISPLAY"):
            os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
        ensure_gui.app = QGuiApplication([])
    return QGuiApplication.instance()


def _brand(p, rect):
    """머리글 오른쪽: [로고] 제작사 : 한일특장(주). 쓴 너비(점)."""
    from . import company
    return company.paint_brand(p, rect, _font(8.5, True))


def _font(size_pt, bold=False):
    from PySide6.QtGui import QFont
    f = QFont()
    f.setFamilies(["Malgun Gothic", "맑은 고딕", "Noto Sans CJK KR", "NanumGothic", "Sans Serif"])
    f.setPointSizeF(size_pt)
    f.setBold(bold)
    return f


def render_flat_image(d: ReportData, w=1600, h=1000, dims=True):
    from PySide6.QtCore import QRectF
    from PySide6.QtGui import QColor, QImage, QPainter

    from .drawing import LIGHT, FlatTransform, paint_dimensions, paint_flat, paint_legend

    ensure_gui()
    img = QImage(w, h, QImage.Format_ARGB32)
    img.fill(LIGHT.bg)
    p = QPainter(img)
    p.setFont(_font(14))
    legend_h = 70 if d.result.forming is not None else 0
    rect = QRectF(20, 20, w - 110, h - 90 - legend_h)
    tf = FlatTransform.fit(d.result, rect)
    paint_flat(p, d.result, tf, LIGHT, line_scale=2.0)
    if dims:
        paint_dimensions(p, d.result, tf, LIGHT, d.dims, line_width=1.5)
    if d.user_dims:
        from .drawing import paint_user_dims
        paint_user_dims(p, d.user_dims, tf, LIGHT, line_width=1.8, color=QColor(200, 0, 120))
    if legend_h:
        paint_legend(p, QRectF(w * 0.25, h - legend_h - 10, w * 0.5, 22), LIGHT)
    p.end()
    return img


def render_model_image(d: ReportData, w=900, h=700):
    from PySide6.QtCore import QRectF
    from PySide6.QtGui import QImage, QPainter

    from .drawing import LIGHT, paint_model, solid_polylines

    ensure_gui()
    img = QImage(w, h, QImage.Format_ARGB32)
    img.fill(LIGHT.bg)
    if d.solid is not None:
        p = QPainter(img)
        paint_model(p, solid_polylines(d.solid), QRectF(10, 10, w - 20, h - 20), LIGHT, line_width=2.0)
        p.end()
    return img


# ---------------------------------------------------------------- PDF

def save_pdf(d: ReportData, path: str):
    from PySide6.QtCore import QMarginsF, QPointF, QRectF, Qt
    from PySide6.QtGui import QColor, QPageLayout, QPageSize, QPainter, QPdfWriter, QPen

    from .drawing import LIGHT, FlatTransform, paint_dimensions, paint_flat, paint_legend

    ensure_gui()
    pdf = QPdfWriter(path)
    pdf.setResolution(300)
    pdf.setPageLayout(QPageLayout(QPageSize(QPageSize.A4), QPageLayout.Portrait, QMarginsF(0, 0, 0, 0)))
    pdf.setTitle(f"{d.title} - {d.display_name}")
    pdf.setCreator(f"한일 판금 전개 {version_text()}")

    k = 300 / 25.4                      # mm → 점

    def R(x, y, w, h):
        return QRectF(x * k, y * k, w * k, h * k)

    p = QPainter(pdf)
    page_w, page_h, margin = 210.0, 297.0, 12.0
    inner_w = page_w - 2 * margin
    gray = QColor(110, 110, 110)
    line = QColor(180, 180, 180)

    def header(y):
        p.setPen(QColor(20, 20, 20))
        p.setFont(_font(16, True))
        p.drawText(R(margin, y, inner_w, 9), Qt.AlignLeft | Qt.AlignVCenter, d.title)
        p.setFont(_font(8))
        p.setPen(gray)
        bw = _brand(p, R(margin, y + 1.5, inner_w, 6)) / k
        p.setFont(_font(8))
        p.setPen(gray)
        p.drawText(R(margin, y, inner_w - bw - 4, 9), Qt.AlignRight | Qt.AlignVCenter,
                   f"{d.created:%Y-%m-%d %H:%M}  ·  한일 판금 전개 {version_text()}")
        p.setPen(QPen(QColor(20, 20, 20), 0.4 * k))
        p.drawLine(QPointF(margin * k, (y + 10) * k), QPointF((page_w - margin) * k, (y + 10) * k))
        return y + 13

    y = header(margin)

    # 요약 표
    rows = summary_rows(d)
    p.setFont(_font(8.5))
    row_h = 5.4
    key_w = 42.0
    for i, (key, val) in enumerate(rows):
        ry = y + i * row_h
        if i % 2 == 0:
            p.fillRect(R(margin, ry, inner_w, row_h), QColor(244, 246, 249))
        p.setPen(gray)
        p.drawText(R(margin + 2, ry, key_w, row_h), Qt.AlignLeft | Qt.AlignVCenter, key)
        p.setPen(QColor(20, 20, 20))
        p.drawText(R(margin + key_w, ry, inner_w - key_w - 2, row_h), Qt.AlignLeft | Qt.AlignVCenter, val)
    y += len(rows) * row_h + 4

    # 전개도 + 3D
    bends = bend_rows(d)
    dim_rows = d.dim_rows()
    table_h = (7 + 5 * min(len(bends), 12) if bends else 0) + (6 + 4.6 * min(len(dim_rows), 10))
    note_list = notes(d)
    notes_h = 4 + 4.2 * len(note_list)
    avail = page_h - margin - y - table_h - notes_h - 6
    draw_h = max(70.0, min(avail, 130.0))
    flat_w = inner_w * 0.68
    box = R(margin, y, flat_w, draw_h)
    p.setPen(QPen(line, 0.2 * k))
    p.drawRect(box)
    p.setFont(_font(7))
    legend_h = 9 if d.result.forming is not None else 0
    inner = R(margin + 3, y + 5, flat_w - 14, draw_h - 13 - legend_h)
    tf = FlatTransform.fit(d.result, inner, margin=0.78)
    paint_flat(p, d.result, tf, LIGHT, line_scale=0.25 * k)
    paint_dimensions(p, d.result, tf, LIGHT, d.dims, line_width=0.15 * k)
    if legend_h:
        paint_legend(p, R(margin + flat_w * 0.2, y + draw_h - legend_h - 2, flat_w * 0.6, 3), LIGHT)
    p.setPen(gray)
    p.drawText(R(margin + 2, y + 0.5, flat_w, 4), Qt.AlignLeft, "전개도 (mm)")

    mx = margin + flat_w + 3
    mw = inner_w - flat_w - 3
    p.setPen(QPen(line, 0.2 * k))
    p.drawRect(R(mx, y, mw, draw_h))
    if d.solid is not None:
        from .drawing import paint_model, solid_polylines
        paint_model(p, solid_polylines(d.solid), R(mx + 2, y + 6, mw - 4, draw_h - 8), LIGHT,
                    line_width=0.2 * k)
    p.setPen(gray)
    p.drawText(R(mx + 2, y + 0.5, mw, 4), Qt.AlignLeft, "3D 형상")
    y += draw_h + 5

    # 절곡 표
    if bends:
        y = _bend_table(p, R, bends, y, margin, inner_w, page_h, header)

    # 치수 표 (절곡선 위치·플랜지 길이·구멍) — 두 단
    if dim_rows:
        y = _dim_table(p, R, dim_rows, y, margin, inner_w, page_h, header)

    # 주의 사항
    p.setFont(_font(7.5))
    if y + notes_h > page_h - margin:
        pdf.newPage()
        y = header(margin)
    p.setPen(QColor(20, 20, 20))
    p.drawText(R(margin, y, inner_w, 4), Qt.AlignLeft, "참고")
    y += 4.5
    p.setPen(gray)
    for n in note_list:
        p.drawText(R(margin + 2, y, inner_w - 2, 4.2), Qt.AlignLeft | Qt.TextWordWrap, "· " + n)
        y += 4.2
    p.end()


def _dim_table(p, R, rows, y, margin, inner_w, page_h, header):
    from PySide6.QtCore import Qt
    from PySide6.QtGui import QColor

    row_h = 4.6
    p.setFont(_font(7.5, True))
    p.setPen(QColor(20, 20, 20))
    p.drawText(R(margin, y, inner_w, 5), Qt.AlignLeft | Qt.AlignVCenter,
               "치수 (mm, 원점 = 전개도 왼쪽 아래)")
    y += 5.5
    p.setFont(_font(7))
    col_w = (inner_w - 4) / 2
    half = (len(rows) + 1) // 2
    for i, (k, v) in enumerate(rows):
        col, r = divmod(i, half)
        ry = y + r * row_h
        if ry + row_h > page_h - 14:
            p.device().newPage()
            y = header(12) - r * row_h
            ry = y + r * row_h
        x = margin + col * (col_w + 4)
        if r % 2 == 0:
            p.fillRect(R(x, ry, col_w, row_h), QColor(244, 246, 249))
        p.setPen(QColor(110, 110, 110))
        p.drawText(R(x + 1.5, ry, 34, row_h), Qt.AlignLeft | Qt.AlignVCenter, k)
        p.setPen(QColor(20, 20, 20))
        p.drawText(R(x + 36, ry, col_w - 37, row_h), Qt.AlignLeft | Qt.AlignVCenter, v)
    return y + half * row_h + 4


def _bend_table(p, R, bends, y, margin, inner_w, page_h, header):
    from PySide6.QtCore import Qt
    from PySide6.QtGui import QColor

    widths = [9, 11, 12, 11, 15, 11, 12, 13, 14, 15]
    widths.append(inner_w - sum(widths))
    row_h = 5.0

    def head(yy):
        p.fillRect(R(margin, yy, inner_w, row_h + 1), QColor(40, 60, 90))
        p.setPen(QColor(255, 255, 255))
        p.setFont(_font(7, True))
        x = margin
        for h, w in zip(BEND_HEADER, widths):
            p.drawText(R(x, yy, w, row_h + 1), Qt.AlignCenter, h)
            x += w
        return yy + row_h + 1

    y = head(y)
    p.setFont(_font(7))
    for i, row in enumerate(bends):
        if y + row_h > page_h - 14:
            p.device().newPage()
            y = head(header(12))
            p.setFont(_font(7))
        if i % 2:
            p.fillRect(R(margin, y, inner_w, row_h), QColor(244, 246, 249))
        p.setPen(QColor(20, 20, 20))
        x = margin
        for j, (v, w) in enumerate(zip(row, widths)):
            align = Qt.AlignLeft | Qt.AlignVCenter if j == len(row) - 1 else Qt.AlignCenter
            p.drawText(R(x + (1 if j == len(row) - 1 else 0), y, w, row_h), align, v)
            x += w
        y += row_h
    return y + 4


# ---------------------------------------------------------------- 엑셀

def _xl_styles():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    thin = Side(style="thin", color="BBBBBB")
    return {
        "title": Font(name="맑은 고딕", size=16, bold=True),
        "head": Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF"),
        "head_fill": PatternFill("solid", fgColor="28405A"),
        "key": Font(name="맑은 고딕", size=10, color="555555"),
        "val": Font(name="맑은 고딕", size=10),
        "zebra": PatternFill("solid", fgColor="F4F6F9"),
        "border": Border(left=thin, right=thin, top=thin, bottom=thin),
        "center": Alignment(horizontal="center", vertical="center"),
        "left": Alignment(horizontal="left"),
        "wrap": Alignment(wrap_text=True, vertical="top"),
    }


def _xl_font(base, **kw):
    """openpyxl Font 를 바꿔 새로 (copy() 가 없는 판도 있다)."""
    from openpyxl.styles import Font
    return Font(name=base.name, size=base.sz, bold=kw.get("bold", base.b),
                color=kw.get("color", base.color))


def _add_image(ws, qimage, anchor, width_px):
    from openpyxl.drawing.image import Image as XLImage

    fd, tmp = tempfile.mkstemp(suffix=".png")
    os.close(fd)
    qimage.save(tmp, "PNG")
    img = XLImage(tmp)
    ratio = width_px / img.width
    img.width, img.height = int(img.width * ratio), int(img.height * ratio)
    ws.add_image(img, anchor)
    return tmp


def save_xlsx(d: ReportData, path: str):
    from openpyxl import Workbook

    st = _xl_styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "보고서"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 70
    ws["A1"] = d.title
    ws["A1"].font = st["title"]
    from . import company
    ws["A2"] = company.MAKER
    ws["A2"].font = _xl_font(st["val"], bold=True)
    ws["A2"].hyperlink = company.HOMEPAGE
    ws["A3"] = f"{d.created:%Y-%m-%d %H:%M}  ·  한일 판금 전개 {version_text()}"
    ws["A3"].font = st["key"]
    logo_tmp = _add_image(ws, company.logo_image(96), "B2", 90)
    ws.row_dimensions[2].height = 32
    r = 5
    for i, (k, v) in enumerate(summary_rows(d)):
        ws.cell(r, 1, k).font = st["key"]
        ws.cell(r, 2, v).font = st["val"]
        if i % 2 == 0:
            ws.cell(r, 1).fill = st["zebra"]
            ws.cell(r, 2).fill = st["zebra"]
        r += 1

    # 숫자로도 넣어 둔다 (계산·견적에 바로 쓰게)
    m = d.metrics
    r += 1
    ws.cell(r, 1, "숫자 값").font = st["val"]
    r += 1
    for k, v, unit in [("가로", m.width, "mm"), ("세로", m.height, "mm"), ("두께", d.result.thickness, "mm"),
                       ("블랭크 넓이", m.blank_area, "mm²"), ("사각 원판 넓이", m.rect_area, "mm²"),
                       ("수율", m.yield_ratio, ""), ("블랭크 무게", m.weight_kg, "kg"),
                       ("사각 원판 무게", m.rect_weight_kg, "kg"), ("절단 길이", m.cut_length, "mm"),
                       ("피어싱 수", m.pierce, "곳"), ("절곡 수", len(d.result.bends), "곳")]:
        ws.cell(r, 1, f"{k} ({unit})" if unit else k).font = st["key"]
        c = ws.cell(r, 2, round(float(v), 4))
        c.font = st["val"]
        c.number_format = "0.0%" if k == "수율" else "#,##0.000"
        c.alignment = st["left"]
        r += 1

    r += 1
    ws.cell(r, 1, "참고").font = st["val"]
    for n in notes(d):
        r += 1
        ws.cell(r, 2, n).alignment = st["wrap"]
        ws.cell(r, 2).font = st["key"]

    tmps = [
        _add_image(ws, render_flat_image(d), "D1", 760),
        _add_image(ws, render_model_image(d), "D27", 420),
    ]

    # 절곡표
    if d.result.bends:
        wsb = wb.create_sheet("절곡표")
        for j, h in enumerate(BEND_HEADER, 1):
            c = wsb.cell(1, j, h)
            c.font, c.fill, c.alignment, c.border = st["head"], st["head_fill"], st["center"], st["border"]
            wsb.column_dimensions[c.column_letter].width = 12
        wsb.column_dimensions[wsb.cell(1, len(BEND_HEADER)).column_letter].width = 44
        for i, b in enumerate(d.result.bends, 2):
            vals = [f"B{b.number}", b.direction, b.angle_deg, b.model_radius, b.inner_radius, b.k,
                    b.deduction, b.deduction90, b.allowance, b.length, b.source]
            for j, v in enumerate(vals, 1):
                if isinstance(v, float):
                    v = None if math.isnan(v) else round(v, 4)
                c = wsb.cell(i, j, v)
                c.border = st["border"]
                c.alignment = st["center"] if j < len(vals) else st["wrap"]
                if isinstance(v, float):
                    c.number_format = "0.000"
            if b.kind in ("원뿔", "곡선 단면"):
                wsb.cell(i, 5, b.radius_text)
        wsb.freeze_panes = "A2"

    rows = d.check_rows
    if rows:
        from openpyxl.styles import Font

        from . import unfoldcheck
        wsc = wb.create_sheet("전개 검증")
        for j, h in enumerate(unfoldcheck.HEAD, 1):
            c = wsc.cell(1, j, h)
            c.font = Font(bold=True)
        for i, r in enumerate(rows, 2):
            for j, v in enumerate(r, 1):
                wsc.cell(i, j, v)
        for col, wd in zip("ABCDEF", (40, 24, 24, 14, 22, 8)):
            wsc.column_dimensions[col].width = wd

    # 치수
    wsd = wb.create_sheet("치수")
    wsd.column_dimensions["A"].width = 28
    wsd.column_dimensions["B"].width = 60
    for j, h in enumerate(("항목", "값 (mm, 원점 = 전개도 왼쪽 아래)"), 1):
        c = wsd.cell(1, j, h)
        c.font, c.fill = st["head"], st["head_fill"]
    for i, (k, v) in enumerate(d.dims.rows(), 2):
        wsd.cell(i, 1, k)
        wsd.cell(i, 2, v)
    r0 = len(d.dims.rows()) + 3
    if d.dims.holes:
        for j, h in enumerate(("구멍", "지름", "가로", "세로", "중심 X", "중심 Y"), 1):
            c = wsd.cell(r0, j, h)
            c.font, c.fill = st["head"], st["head_fill"]
        for i, hole in enumerate(d.dims.holes, r0 + 1):
            vals = [hole.label, hole.diameter, hole.size[0] if hole.size else None,
                    hole.size[1] if hole.size else None, hole.center[0], hole.center[1]]
            for j, v in enumerate(vals, 1):
                c = wsd.cell(i, j, round(v, 3) if isinstance(v, float) else v)
                if isinstance(v, float):
                    c.number_format = "0.00"

    # 성형품 두께 변화 분포
    if d.result.forming is not None:
        tmps.append(_thinning_sheet(wb, d, st))
    # 제품 성형 해석 (성형성 존·절곡 균열·하중) — 전개 전용판에는 없음
    from . import edition
    if not edition.unfold_only():
        try:
            from . import formsim
            tmps += formsim.add_sheets(wb, formsim.analyze(d.result, d.settings), d.part_title or d.part_name)
        except Exception:
            import traceback
            traceback.print_exc()

    from .company import stamp_workbook
    stamp_workbook(wb)
    wb.save(path)
    for t in tmps + [logo_tmp]:
        try:
            os.remove(t)
        except OSError:
            pass


def _thinning_sheet(wb, d, st):
    fi = d.result.forming
    from .unfolder import _tri_areas2
    areas = _tri_areas2(fi.U, fi.F)
    th = fi.thinning
    ws = wb.create_sheet("두께 변화")
    ws["A1"] = "두께 변화 분포 (블랭크 면적 기준)"
    ws["A1"].font = st["val"]
    head = ["구간", "넓이 (mm²)", "비율", "판정"]
    for j, h in enumerate(head, 1):
        c = ws.cell(3, j, h)
        c.font, c.fill, c.alignment = st["head"], st["head_fill"], st["center"]
        ws.column_dimensions[c.column_letter].width = 22
    edges = [-1, -0.15, -0.05, 0.05, 0.15, 0.25, 1]
    names = ["15% 넘게 두꺼워짐", "5~15% 두꺼워짐", "±5% 이내", "5~15% 얇아짐", "15~25% 얇아짐", "25% 넘게 얇아짐"]
    verdict = ["주름 위험", "주의", "양호", "양호", "주의", "터짐 위험"]
    total = areas.sum()
    for i in range(len(names)):
        sel = (th >= edges[i]) & (th < edges[i + 1])
        a = float(areas[sel].sum())
        ws.cell(4 + i, 1, names[i])
        ws.cell(4 + i, 2, round(a, 1)).number_format = "#,##0.0"
        ws.cell(4 + i, 3, a / total if total else 0).number_format = "0.0%"
        ws.cell(4 + i, 4, verdict[i])
    r = 4 + len(names) + 1
    lo, hi = fi.thinning_range
    ws.cell(r, 1, "최대 얇아짐")
    ws.cell(r, 2, float(max(hi, 0))).number_format = "0.0%"
    ws.cell(r + 1, 1, "최대 두꺼워짐")
    ws.cell(r + 1, 2, float(max(-lo, 0))).number_format = "0.0%"
    ws.cell(r + 2, 1, "(메시 조각 삼각형과 넓이 기준 위아래 0.5 % 는 뺀 값)")
    ws.cell(r + 3, 1, "전개 방법")
    ws.cell(r + 3, 2, fi.method)
    if fi.flc0 > 0:
        split, marginal, wrinkle, top = fi.flc_summary
        for i, (k, v) in enumerate([("성형 한계 FLC0 (진변형률)", round(fi.flc0, 4)),
                                    ("한계 넘음 (파단 위험)", split), ("한계 80 % 이상", marginal),
                                    ("주름 가능 (5 % 넘게 두꺼워짐)", wrinkle), ("가장 큰 한계 비", top)]):
            ws.cell(r + 5 + i, 1, k)
            c = ws.cell(r + 5 + i, 2, v)
            c.number_format = "0.000" if i == 0 else "0.0%"
    return _add_image(ws, render_fld(fi, 900, 700), "F3", 600)


def save_parts_list(items, path: str):
    """여러 부품을 한 표로 (견적·재료 발주용). items: ReportData 목록."""
    from openpyxl import Workbook

    st = _xl_styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "부품 목록"
    head = ["파일", "부품", "전개 방식", "재질", "두께(mm)", "가로(mm)", "세로(mm)", "블랭크 넓이(mm²)",
            "블랭크 무게(kg)", "사각 원판 무게(kg)", "수율", "절단 길이(mm)", "피어싱", "절곡 수", "주의"]
    for j, h in enumerate(head, 1):
        c = ws.cell(1, j, h)
        c.font, c.fill, c.alignment, c.border = st["head"], st["head_fill"], st["center"], st["border"]
        ws.column_dimensions[c.column_letter].width = 14
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["O"].width = 50
    for i, d in enumerate(items, 2):
        m, r = d.metrics, d.result
        vals = [os.path.basename(d.source_path), d.display_name,
                "성형품(근사)" if r.mode == "forming" else "판금", d.material.code, r.thickness,
                m.width, m.height, m.blank_area, m.weight_kg, m.rect_weight_kg, m.yield_ratio,
                m.cut_length, m.pierce, len(r.bends),
                " / ".join(w for w in r.warnings if "근사 전개:" not in w)]
        for j, v in enumerate(vals, 1):
            if isinstance(v, (float, np.floating)):
                v = round(float(v), 4)
            c = ws.cell(i, j, v)
            c.border = st["border"]
            if j == 11:
                c.number_format = "0.0%"
            elif isinstance(v, float):
                c.number_format = "#,##0.000"
    n = len(items) + 1
    ws.cell(n + 1, 1, "합계").font = st["val"]
    for col in ("I", "J", "L"):
        c = ws[f"{col}{n + 1}"]
        c.value = f"=SUM({col}2:{col}{n})"
        c.number_format = "#,##0.000"
    ws.freeze_panes = "A2"
    from .company import stamp_workbook
    stamp_workbook(wb)
    wb.save(path)


# ---------------------------------------------------------------- 조립도 (번호 풍선 + 부품표)

def _bbox_size(solid):
    from OCP.Bnd import Bnd_Box
    from OCP.BRepBndLib import BRepBndLib
    b = Bnd_Box()
    BRepBndLib.Add_s(solid, b)
    lo, hi = b.CornerMin(), b.CornerMax()
    return (hi.X() - lo.X(), hi.Y() - lo.Y(), hi.Z() - lo.Z())


def save_assembly_pdf(model, path: str, parts=None, yaw=None, pitch=None, highlight=None):
    """조립품 전체 그림에 부품 번호 풍선을 달고, 다음 쪽부터 번호·부품명·크기 표를 붙인다.
    번호는 화면의 [부품] 목록, 전개 기록, 보고서의 번호와 같다."""
    from PySide6.QtCore import QMarginsF, QPointF, QRectF, Qt
    from PySide6.QtGui import QColor, QPageLayout, QPageSize, QPainter, QPdfWriter, QPen

    from .drawing import LIGHT, assembly_polylines, paint_assembly

    ensure_gui()
    parts = parts if parts is not None else assembly_polylines(model.raw_solids)
    kw = {}
    if yaw is not None:
        kw["yaw"] = yaw
    if pitch is not None:
        kw["pitch"] = pitch

    pdf = QPdfWriter(path)
    pdf.setResolution(300)
    pdf.setPageLayout(QPageLayout(QPageSize(QPageSize.A3), QPageLayout.Landscape,
                                  QMarginsF(0, 0, 0, 0)))
    name = os.path.basename(model.path)
    pdf.setTitle(f"조립도 - {name}")
    pdf.setCreator(f"한일 판금 전개 {version_text()}")
    k = 300 / 25.4
    page_w, page_h, margin = 420.0, 297.0, 12.0

    def R(x, y, w, h):
        return QRectF(x * k, y * k, w * k, h * k)

    p = QPainter(pdf)
    gray = QColor(110, 110, 110)

    def header(y, sub=""):
        p.setPen(QColor(20, 20, 20))
        p.setFont(_font(16, True))
        p.drawText(R(margin, y, page_w - 2 * margin, 9), Qt.AlignLeft | Qt.AlignVCenter,
                   f"조립도  {name}" + (f"  ({sub})" if sub else ""))
        p.setFont(_font(8))
        p.setPen(gray)
        bw = _brand(p, R(margin, y + 1.5, page_w - 2 * margin, 6)) / k
        p.setFont(_font(8))
        p.setPen(gray)
        p.drawText(R(margin, y, page_w - 2 * margin - bw - 4, 9), Qt.AlignRight | Qt.AlignVCenter,
                   f"부품 {len(model.solids)}개  ·  {_dt.datetime.now():%Y-%m-%d %H:%M}  ·  "
                   f"한일 판금 전개 {version_text()}")
        p.setPen(QPen(QColor(20, 20, 20), 0.4 * k))
        p.drawLine(QPointF(margin * k, (y + 10) * k), QPointF((page_w - margin) * k, (y + 10) * k))
        return y + 13

    y = header(margin)
    box = R(margin, y, page_w - 2 * margin, page_h - y - margin)
    p.setPen(QPen(QColor(180, 180, 180), 0.2 * k))
    p.drawRect(box)
    inner = R(margin + 4, y + 4, page_w - 2 * margin - 8, page_h - y - margin - 8)
    font_px = int(3.2 * k) if len(parts) <= 60 else int(2.6 * k)
    paint_assembly(p, parts, inner, LIGHT, highlight=highlight, font_px=font_px,
                   line_width=0.18 * k, **kw)

    # 부품표: 쪽마다 3 단
    rows = []
    for i in range(len(model.solids)):
        sx, sy, sz = _bbox_size(model.raw_solids[i])
        rows.append((str(i + 1), model.names[i] if model.names else f"부품 {i + 1}",
                     f"{sx:.0f} × {sy:.0f} × {sz:.0f}"))
    row_h, cols = 6.0, 3
    col_w = (page_w - 2 * margin - (cols - 1) * 6) / cols
    widths = [12, col_w - 12 - 40, 40]
    start = 0
    page = 1
    while start < len(rows):
        pdf.newPage()
        page += 1
        y0 = header(margin, f"부품표 {page - 1}")
        per_col = int((page_h - y0 - margin - row_h) / row_h)
        for c in range(cols):
            x0 = margin + c * (col_w + 6)
            chunk = rows[start:start + per_col]
            if not chunk:
                break
            p.fillRect(R(x0, y0, col_w, row_h), QColor(40, 60, 90))
            p.setPen(QColor(255, 255, 255))
            p.setFont(_font(8, True))
            x = x0
            for h, w in zip(("번호", "부품명", "크기 (mm)"), widths):
                p.drawText(R(x + 1, y0, w - 2, row_h), Qt.AlignCenter, h)
                x += w
            p.setFont(_font(8))
            for i, row in enumerate(chunk):
                ry = y0 + row_h * (i + 1)
                if i % 2:
                    p.fillRect(R(x0, ry, col_w, row_h), QColor(244, 246, 249))
                p.setPen(QColor(20, 20, 20))
                x = x0
                for j, (v, w) in enumerate(zip(row, widths)):
                    align = (Qt.AlignLeft if j == 1 else Qt.AlignCenter) | Qt.AlignVCenter
                    p.drawText(R(x + 1.5, ry, w - 3, row_h), align,
                               p.fontMetrics().elidedText(v, Qt.ElideRight, int((w - 3) * k)))
                    x += w
            start += len(chunk)
    p.end()


def save_nesting_pdf(results, path: str, source_name: str = "", density=None):
    """철판 배치: 첫 쪽에 요약·부품별 수량, 다음 쪽부터 철판 한 장에 한 쪽 (A3 가로)."""
    from PySide6.QtCore import QMarginsF, QPointF, QRectF, Qt
    from PySide6.QtGui import QColor, QPageLayout, QPageSize, QPainter, QPdfWriter, QPen

    from .drawing import LIGHT, paint_sheet

    ensure_gui()
    pdf = QPdfWriter(path)
    pdf.setResolution(300)
    pdf.setPageLayout(QPageLayout(QPageSize(QPageSize.A3), QPageLayout.Landscape,
                                  QMarginsF(0, 0, 0, 0)))
    pdf.setTitle(f"철판 배치 - {source_name}")
    pdf.setCreator(f"한일 판금 전개 {version_text()}")
    k = 300 / 25.4
    page_w, page_h, margin = 420.0, 297.0, 12.0
    total = sum(len(r.sheets) for r in results)

    def R(x, y, w, h):
        return QRectF(x * k, y * k, w * k, h * k)

    p = QPainter(pdf)
    gray = QColor(110, 110, 110)

    def header(sub=""):
        y = margin
        p.setPen(QColor(20, 20, 20))
        p.setFont(_font(16, True))
        p.drawText(R(margin, y, page_w - 2 * margin, 9), Qt.AlignLeft | Qt.AlignVCenter,
                   "철판 배치" + (f"  {source_name}" if source_name else "") + (f"  ({sub})" if sub else ""))
        p.setFont(_font(8))
        p.setPen(gray)
        bw = _brand(p, R(margin, y + 1.5, page_w - 2 * margin, 6)) / k
        p.setFont(_font(8))
        p.setPen(gray)
        p.drawText(R(margin, y, page_w - 2 * margin - bw - 4, 9), Qt.AlignRight | Qt.AlignVCenter,
                   f"철판 {total}장  ·  {_dt.datetime.now():%Y-%m-%d %H:%M}  ·  한일 판금 전개 {version_text()}")
        p.setPen(QPen(QColor(20, 20, 20), 0.4 * k))
        p.drawLine(QPointF(margin * k, (y + 10) * k), QPointF((page_w - margin) * k, (y + 10) * k))
        return y + 14

    # ---- 첫 쪽: 요약 + 부품별 배치 수
    y = header("요약")
    p.setFont(_font(10))
    for r in results:
        lines = r.summary_lines(density)
        p.setPen(QColor(20, 20, 20))
        p.setFont(_font(11, True))
        p.drawText(R(margin, y, 250, 6), Qt.AlignLeft | Qt.AlignVCenter, lines[0])
        y += 6.5
        p.setFont(_font(9))
        for ln in lines[1:]:
            p.setPen(QColor(170, 40, 40) if ln.startswith("못 놓음") else QColor(40, 40, 40))
            p.drawText(R(margin + 4, y, 260, 5), Qt.AlignLeft | Qt.AlignVCenter, ln)
            y += 5
        y += 3
        if y > page_h - 40:
            pdf.newPage()
            y = header("요약")

    rows = []
    for r in results:
        cnt = {}
        for si, s in enumerate(r.sheets):
            for pl in s.placements:
                c = cnt.setdefault(pl.part, [0, set()])
                c[0] += 1
                c[1].add(si + 1)
        for i, part in enumerate(r.parts):
            n, sheets = cnt.get(i, [0, set()])
            rows.append((part.label, part.name, f"t{r.thickness:g}", f"{part.qty}", f"{n}",
                         ", ".join(str(x) for x in sorted(sheets))))
    heads = ("번호", "부품명", "두께", "수량", "배치", "철판")
    widths = [16, 150, 22, 18, 18, 172]
    row_h = 5.5

    def table_head(y):
        p.fillRect(R(margin, y, sum(widths), row_h), QColor(40, 60, 90))
        p.setPen(QColor(255, 255, 255))
        p.setFont(_font(8, True))
        x = margin
        for h, w in zip(heads, widths):
            p.drawText(R(x + 1, y, w - 2, row_h), Qt.AlignCenter, h)
            x += w
        return y + row_h

    y = table_head(y + 2)
    p.setFont(_font(8))
    for i, row in enumerate(rows):
        if y > page_h - margin - row_h:
            pdf.newPage()
            y = table_head(header("부품별 수량"))
            p.setFont(_font(8))
        if i % 2:
            p.fillRect(R(margin, y, sum(widths), row_h), QColor(244, 246, 249))
        x = margin
        for j, (v, w) in enumerate(zip(row, widths)):
            short = row[3] != row[4] and j == 4
            p.setPen(QColor(190, 30, 30) if short else QColor(20, 20, 20))
            align = (Qt.AlignLeft if j in (1, 5) else Qt.AlignCenter) | Qt.AlignVCenter
            p.drawText(R(x + 1.5, y, w - 3, row_h), align,
                       p.fontMetrics().elidedText(v, Qt.ElideRight, int((w - 3) * k)))
            x += w
        y += row_h

    # ---- 철판마다 한 쪽
    for r in results:
        for si, s in enumerate(r.sheets):
            pdf.newPage()
            y = header(f"t{r.thickness:g}  {si + 1} / {len(r.sheets)}장")
            p.setFont(_font(9))
            p.setPen(QColor(40, 40, 40))
            p.drawText(R(margin, y, page_w - 2 * margin, 5), Qt.AlignLeft | Qt.AlignVCenter,
                       f"철판 {r.sheet_h:g} × {r.sheet_w:g}  t{r.thickness:g}  ·  부품 {len(s.placements)}개  ·  "
                       f"사용률 {s.utilization * 100:.1f}%  ·  사용 길이 {s.used_length:.0f} mm  ·  "
                       f"간격 {r.gap:g} / 여유 {r.margin:g} mm")
            box = R(margin, y + 7, page_w - 2 * margin, page_h - y - 7 - margin)
            paint_sheet(p, r, s, box, LIGHT, line_scale=0.25 * k, font_px=int(3.0 * k))
    p.end()


def render_sheet_image(r, sheet, w=1600, h=820):
    """철판 한 장 그림 (흰 바탕, 인쇄·엑셀용)."""
    from PySide6.QtCore import QRectF
    from PySide6.QtGui import QColor, QImage, QPainter

    from .drawing import LIGHT, paint_sheet
    ensure_gui()
    img = QImage(w, h, QImage.Format_RGB32)
    img.fill(QColor(255, 255, 255))
    p = QPainter(img)
    paint_sheet(p, r, sheet, QRectF(6, 6, w - 12, h - 12), LIGHT, font_px=14)
    p.end()
    return img


def render_part_thumb(part, w=208, h=108):
    """펼친 부품 하나 작은 그림 (엑셀 부품별 표용, 흰 바탕)."""
    from PySide6.QtCore import QRectF
    from PySide6.QtGui import QColor, QImage, QPainter

    from .drawing import LIGHT, FlatTransform, _Flat, paint_flat
    ensure_gui()
    img = QImage(w, h, QImage.Format_RGB32)
    img.fill(QColor(255, 255, 255))
    P = np.vstack([np.asarray(c, float) for c in part.cut])
    flat = _Flat(part.entities)
    flat.bbox = (*P.min(0), *P.max(0))
    p = QPainter(img)
    paint_flat(p, flat, FlatTransform.fit(flat, QRectF(0, 0, w, h), margin=0.88), LIGHT,
               strain=False, labels=False, line_scale=0.9)
    p.end()
    return img


def nesting_part_rows(results, density):
    """부품별: (번호, 부품명, 두께, 크기, 면적 m², 수량, 배치 수, 철판, 개당 kg, 합계 kg)."""
    rows = []
    for r in results:
        cnt = {}
        for si, s in enumerate(r.sheets):
            for pl in s.placements:
                c = cnt.setdefault(pl.part, [0, set()])
                c[0] += 1
                c[1].add(si + 1)
        for i, part in enumerate(r.parts):
            n, sheets = cnt.get(i, [0, set()])
            P = np.vstack([np.asarray(c, float) for c in part.cut])
            w, h = np.ptp(P, axis=0)
            kg = part.area * r.thickness / 1e6 * density
            rows.append((part.label, part.name, r.thickness, f"{w:.0f} × {h:.0f}", part.area / 1e6,
                         part.qty, n, ", ".join(str(x) for x in sorted(sheets)), kg, kg * n))
    return rows


def save_nesting_xlsx(results, path: str, source_name: str = "", density=7.85):
    """철판 배치 보고서 (엑셀): 요약 · 부품별 · 철판별(그림)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    st = _xl_styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "요약"

    def head(sheet, row, names, widths=None):
        for j, h in enumerate(names, 1):
            c = sheet.cell(row, j, h)
            c.font, c.fill, c.alignment, c.border = st["head"], st["head_fill"], st["center"], st["border"]
            if widths:
                sheet.column_dimensions[c.column_letter].width = widths[j - 1]

    def put(sheet, row, vals, zebra=False, fmt=None):
        for j, v in enumerate(vals, 1):
            c = sheet.cell(row, j, v)
            c.font, c.border = st["val"], st["border"]
            if zebra:
                c.fill = st["zebra"]
            if fmt and fmt[j - 1]:
                c.number_format = fmt[j - 1]

    from . import company
    ws.cell(1, 1, "철판 배치 보고서").font = st["title"]
    ws.row_dimensions[1].height = 34
    logo_tmp = _add_image(ws, company.logo_image(96), "D1", 100)
    ws.cell(1, 5, company.MAKER).font = _xl_font(st["val"], bold=True)
    ws.cell(1, 5).hyperlink = company.HOMEPAGE
    info = [("제작사", company.NAME), ("도면", source_name), ("만든 때", f"{_dt.datetime.now():%Y-%m-%d %H:%M}"),
            ("프로그램", f"한일 판금 전개 {version_text(long=True)}")]
    if results:
        r0 = results[0]
        info += [("철판 크기", f"{r0.sheet_h:g} × {r0.sheet_w:g} mm"),
                 ("부품 간격 / 가장자리 여유", f"{r0.gap:g} / {r0.margin:g} mm"),
                 ("격자", f"{r0.resolution:g} mm")]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(i, 1, k).font = st["key"]
        ws.cell(i, 2, v).font = st["val"]

    row = 3 + len(info) + 1
    names = ["두께 (mm)", "재질", "철판 (장)", "배치 / 수량", "사용률", "철판 무게 (kg)", "부품 무게 (kg)",
             "자투리 (kg)", "마지막 장 사용 길이 (mm)", "남는 철판 (mm)", "못 놓은 부품"]
    head(ws, row, names, [24, 12, 10, 12, 10, 14, 14, 12, 22, 18, 40])
    fmt = [None, None, "0", None, "0.0%", "0.0", "0.0", "0.0", "0", None, None]
    tot = [0, 0.0, 0.0]
    for i, r in enumerate(results):
        row += 1
        sw = r.sheet_weight_kg(density) * len(r.sheets)
        pw = sum(s.used_area for s in r.sheets) * r.thickness / 1e6 * density
        last = r.sheets[-1] if r.sheets else None
        left = ", ".join(f"{r.parts[pi].label}×{n}" for pi, n, _ in r.unplaced)
        put(ws, row, [r.thickness, r.material, len(r.sheets), f"{r.placed_count} / {r.total_qty}",
                      r.utilization, sw, pw, sw - pw, last.used_length if last else 0,
                      f"{r.sheet_h:g} × {max(0.0, r.sheet_w - last.used_length):.0f}" if last else "",
                      left], i % 2 == 1, fmt)
        tot[0] += len(r.sheets)
        tot[1] += sw
        tot[2] += pw
    row += 1
    put(ws, row, ["합계", "", tot[0], "", (tot[2] / tot[1]) if tot[1] else 0, tot[1], tot[2],
                  tot[1] - tot[2], None, "", ""], False, fmt)
    for j in range(1, len(names) + 1):
        ws.cell(row, j).font = _xl_font(st["head"], color="000000")
    ws.cell(row + 2, 1, f"무게는 밀도 {density:g} g/cm³ 기준. 사용률 = 부품 넓이 합 / 철판 넓이 합.").font = st["key"]
    ws.freeze_panes = None

    # 부품별 (부품명 옆에 펼친 모양 작은 그림)
    wp = wb.create_sheet("부품별")
    names = ["번호", "부품명", "모양", "두께", "크기 (mm)", "면적 (m²)", "수량", "배치", "철판",
             "개당 무게 (kg)", "합계 (kg)"]
    head(wp, 1, names, [8, 30, 16, 8, 16, 12, 8, 8, 16, 14, 12])
    fmt = [None, None, None, "0.0##", None, "0.0000", "0", "0", None, "0.000", "0.000"]
    tmps = []
    parts = [p for r in results for p in r.parts]
    for i, (vals, part) in enumerate(zip(nesting_part_rows(results, density), parts), 2):
        vals = list(vals)
        put(wp, i, vals[:2] + [None] + vals[2:], i % 2 == 1, fmt)
        if vals[5] != vals[6]:
            wp.cell(i, 8).font = _xl_font(st["val"], color="C00000", bold=True)
        wp.row_dimensions[i].height = 42
        for j in range(1, len(names) + 1):
            wp.cell(i, j).alignment = Alignment(vertical="center", horizontal="left" if j == 2 else None)
        tmps.append(_add_image(wp, render_part_thumb(part), f"C{i}", 104))
    wp.freeze_panes = "A2"

    # 철판별 + 그림
    wsh = wb.create_sheet("철판별")
    names = ["두께", "장", "부품 수", "사용률", "사용 길이 (mm)", "부품 번호"]
    head(wsh, 1, names, [8, 8, 10, 10, 14, 90])
    row = 1
    for r in results:
        for k, s in enumerate(r.sheets, 1):
            row += 1
            labels = {}
            for pl in s.placements:
                lb = r.parts[pl.part].label
                labels[lb] = labels.get(lb, 0) + 1
            put(wsh, row, [r.thickness, f"{k} / {len(r.sheets)}", len(s.placements), s.utilization,
                           s.used_length, ", ".join(f"{a}×{n}" if n > 1 else a for a, n in labels.items())],
                row % 2 == 1, [None, None, "0", "0.0%", "0", None])
            wsh.cell(row, 6).alignment = st["wrap"]
    pic = wb.create_sheet("그림")
    prow = 1
    for r in results:
        for k, s in enumerate(r.sheets, 1):
            pic.cell(prow, 1, f"t{r.thickness:g}  {k} / {len(r.sheets)}장 · 부품 {len(s.placements)}개 · "
                              f"사용률 {s.utilization * 100:.1f}% · 사용 길이 {s.used_length:.0f} mm").font = st["val"]
            tmps.append(_add_image(pic, render_sheet_image(r, s), f"A{prow + 1}", 1000))
            prow += 28
    from .company import stamp_workbook
    stamp_workbook(wb)
    wb.save(path)
    for t in tmps + [logo_tmp]:
        try:
            os.remove(t)
        except OSError:
            pass
