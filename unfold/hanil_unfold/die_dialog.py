"""금형 검토 창: 하중·클리어런스·스프링백, 설계 규칙, 스트립 배치, 원통 드로잉 계산."""
from __future__ import annotations

import math
import os

import numpy as np
from PySide6.QtCore import QPointF, QRectF, Qt, QUrl
from PySide6.QtGui import QColor, QDesktopServices, QImage, QPainter, QPen, QPolygonF  # noqa: F401
from PySide6.QtWidgets import (QComboBox, QDialog, QDoubleSpinBox, QFileDialog, QFormLayout, QHBoxLayout,
                               QHeaderView, QLabel, QMessageBox, QPushButton, QTableWidget,
                               QTableWidgetItem, QTabWidget, QVBoxLayout, QWidget)

from . import bendcalc as bc
from . import diecalc as dc

VERDICT_COLOR = {dc.OK: QColor(60, 160, 80), dc.WARN: QColor(220, 150, 0), dc.NG: QColor(220, 50, 50)}


def _table(head, rows, colors=None):
    t = QTableWidget(len(rows), len(head))
    t.setHorizontalHeaderLabels(head)
    t.verticalHeader().setVisible(False)
    t.setEditTriggers(QTableWidget.NoEditTriggers)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            it = QTableWidgetItem(v if isinstance(v, str) else f"{v:,.2f}")
            if not isinstance(v, str):
                it.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
            if colors and colors[i] is not None and j == len(r) - 1:
                it.setForeground(colors[i])
            t.setItem(i, j, it)
    t.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
    t.horizontalHeader().setStretchLastSection(True)
    return t


def load_rows(rep):
    rows = [[L.name, L.length if L.length else "", L.force_n / 1000, L.tonf, L.note] for L in rep.loads]
    rows.append(["합계 (전단 + 스트리퍼)", "", rep.total_n / 1000, rep.total_n / dc.TON, ""])
    rows.append(["권장 프레스 (가장 큰 하중 × 1.3)", "", rep.press_tonf * dc.TON / 1000, rep.press_tonf,
                 "여유 30 % (관례)"])
    return rows


def bend_rows(rep):
    out = []
    for no, length, f, after, die, ks in rep.bends:
        die_txt = f"{die:.2f}" if math.isfinite(die) else "-"
        out.append([f"B{no}", length, f / 1000, f / dc.TON, f"{after:.2f}", die_txt, f"{ks:.4f}"])
    return out


def check_rows(rep):
    return [[c.rule, c.target, c.value, c.limit, c.note, c.verdict] for c in rep.checks]


def render_strip(result, strip, w=1100, h=360, copies=4):
    """스트립 그림: 부품 copies 개를 피치 간격으로, 잔폭·피치 치수."""
    from .nesting import NestPart, _rot_poly
    part = NestPart.from_result("p", result)
    polys = [_rot_poly(c, strip.angle) for c in part.cut]
    P = np.vstack(polys)
    x0, y0 = P.min(0)
    x1, y1 = P.max(0)
    total_w = (copies - 1) * strip.pitch + (x1 - x0) + 2 * strip.bridge
    sw = strip.width
    sc = min((w - 40) / total_w, (h - 70) / sw)
    ox = 20 - (x0 - strip.bridge) * sc
    oy = h - 40 - (strip.side - y0) * sc

    def pt(x, y):
        return QPointF(ox + x * sc, oy - y * sc)

    img = QImage(w, h, QImage.Format_ARGB32)
    img.fill(QColor(255, 255, 255))
    p = QPainter(img)
    p.setRenderHint(QPainter.Antialiasing)
    p.setPen(QPen(QColor(120, 120, 120), 1.5))
    p.setBrush(QColor(235, 240, 248))
    ys0, ys1 = y0 - strip.side, y1 + strip.side
    p.drawRect(QRectF(pt(x0 - strip.bridge, ys1), pt(x0 - strip.bridge + total_w, ys0)))
    p.setPen(QPen(QColor(30, 80, 160), 1.6))
    p.setBrush(QColor(200, 215, 240))
    for k in range(copies):
        for poly in polys:
            q = QPolygonF([pt(a + k * strip.pitch, b) for a, b in poly])
            p.drawPolygon(q)
    f = p.font()
    f.setPixelSize(14)
    p.setFont(f)
    p.setPen(QColor(200, 30, 30))
    ya = ys1 + 3 / sc
    a, b = pt(x0, ya), pt(x0 + strip.pitch, ya)
    p.drawLine(QPointF(a.x(), a.y() - 14), QPointF(b.x(), b.y() - 14))
    p.drawText(QRectF(a.x(), a.y() - 34, b.x() - a.x(), 18), Qt.AlignCenter, f"피치 {strip.pitch:.1f}")
    c1 = pt(x0 - strip.bridge + total_w, ys1)
    p.drawText(QRectF(c1.x() - 200, c1.y() - 22, 200, 18), Qt.AlignRight | Qt.AlignVCenter,
               f"폭 {strip.width:.1f}")
    p.setPen(QColor(40, 40, 40))
    p.drawText(QRectF(10, h - 30, w - 20, 22), Qt.AlignLeft | Qt.AlignVCenter,
               f"부품 {strip.angle:g}° 돌림 · 피치 {strip.pitch:.2f} · 폭 {strip.width:.2f} · "
               f"잔폭 {strip.bridge:.2f} · 옆 잔폭 {strip.side:.2f} · 재료 이용률 {strip.utilization:.1%} · "
               f"1 m 에 {strip.per_meter:.2f} 개")
    p.end()
    return img


class DieDialog(QDialog):
    def __init__(self, result, settings: bc.BendSettings, name="", parent=None, history=None, source_path=""):
        super().__init__(parent)
        self.setWindowTitle(f"금형 검토 — {name}" if name else "금형 검토")
        self.resize(1150, 720)
        self.result, self.settings, self.name = result, settings, name
        self.history, self.source_path = history, source_path
        self.rep = dc.review(result, settings)
        rep = self.rep
        tabs = QTabWidget()

        # 1) 하중
        w1 = QWidget()
        l1 = QVBoxLayout(w1)
        l1.addWidget(QLabel(
            f"<b>{rep.material}</b>  두께 {rep.thickness:g} mm · 전단 저항 τ = 0.8 × 인장강도 = {rep.shear_mpa:.0f} MPa · "
            f"<b>한쪽 클리어런스 {rep.clearance_pct:g} % = {rep.clearance_mm:.3f} mm</b> "
            f"(다이 = 펀치 + {2 * rep.clearance_mm:.3f} mm)<br>"
            f"<span style='color:#555'>독일 VDI 3368: u = 0.01 · s · √(0.1·τ) = {rep.vdi_clearance_mm:.3f} mm (깨끗한 절단, "
            f"버 허용이면 최대 ×3.5) · 다이 두께 어림 H ≈ ∛(0.1·F) = {rep.die_plate_mm:.0f} mm · "
            f"보통 공차 JIS B 0408-B / ISO 2768-m<br>"
            f"전단력이 프레스에 버거우면 펀치·다이 날에 시어 각 (경사) 을 주세요 — 시험에서 펀치 각 16° 로 전단력 최대 80 % 감소 "
            f"(DC01, Trans. FAMENA 2016). 드로잉은 하사점 위에서 힘이 커서 크랭크 프레스 톤수 곡선으로 고릅니다.</span>"))
        l1.addWidget(_table(["하중", "길이 mm", "힘 kN", "톤 (tf)", "비고"], load_rows(rep)))
        if rep.bends:
            l1.addWidget(QLabel("절곡별 굽힘력 · 스프링백 (추정 — 시험 절곡으로 확정)"))
            l1.addWidget(_table(["절곡", "길이 mm", "힘 kN", "톤 (tf)", "보정 없이 접으면 (°)",
                                 "금형 보정 각 (°)", "Ks"], bend_rows(rep)))
        if rep.blank_d:
            l1.addWidget(QLabel(f"성형품: 같은 넓이 원판 지름 Ø{rep.blank_d:.1f} (드로잉 탭에서 공정 계산)"))
        tabs.addTab(w1, "하중 · 클리어런스")

        # 2) 설계 규칙
        w2 = QWidget()
        l2 = QVBoxLayout(w2)
        bad = sum(c.verdict == dc.NG for c in rep.checks)
        warn = sum(c.verdict == dc.WARN for c in rep.checks)
        l2.addWidget(QLabel(f"검사 {len(rep.checks)}건 · <span style='color:#d33'>부적합 {bad}</span> · "
                            f"<span style='color:#d90'>주의 {warn}</span>"))
        l2.addWidget(_table(["규칙", "대상", "값 mm", "기준 mm", "설명", "판정"], check_rows(rep),
                            [VERDICT_COLOR[c.verdict] for c in rep.checks]))
        tabs.addTab(w2, f"설계 규칙 ({bad + warn})" if bad + warn else "설계 규칙")

        # 3) 스트립
        w3 = QWidget()
        l3 = QVBoxLayout(w3)
        if rep.strip is not None:
            from PySide6.QtGui import QPixmap
            self.strip_img = render_strip(result, rep.strip)
            lab = QLabel()
            lab.setPixmap(QPixmap.fromImage(self.strip_img))
            l3.addWidget(lab)
            l3.addWidget(QLabel("순차(프로그레시브) 금형 1열 배치: 부품 사이 잔폭 max(1.5t, 1 mm), "
                                "옆 잔폭(캐리어) max(2t, 1.5 mm). 각도를 바꿔 가며 재료 이용률이 가장 큰 쪽."))
        else:
            self.strip_img = None
            l3.addWidget(QLabel("스트립 배치를 계산하지 못했습니다."))
        l3.addStretch(1)
        tabs.addTab(w3, "스트립 배치")

        # 4) 드로잉
        tabs.addTab(self._draw_tab(), "원통 드로잉 계산")

        # 5) 자동 공정 설계 (문헌 규칙 + AI 도면 검토)
        self.route = None
        tabs.addTab(self._route_tab(), "자동 공정 설계")
        # 6) 금형 공정도
        tabs.addTab(self._process_tab(), "금형 공정도")
        self.tabs = tabs

        row = QHBoxLayout()
        d3 = QPushButton("금형 3D · 도면집 · 가공도 · 파트리스트…")
        d3.setToolTip("블랭킹·피어싱 금형을 판 쌓기·가이드·펀치·스트리퍼·스프링·볼트로 자동 조립해 STEP (3D), "
                      "조립도 PDF (평면도 + 정면도, 풍선 번호), 파트리스트 엑셀로 저장 — 견적·검토용 첫 안")
        d3.clicked.connect(self.save_die3d)
        row.addWidget(d3)
        row.addWidget(QLabel("모든 탭 (하중·설계 규칙·스트립·드로잉·자동 공정 설계·공정도) 을 한 파일로"))
        row.addStretch(1)
        for text_, kind in (("엑셀…", "xlsx"), ("PDF…", "pdf"), ("엑셀 + PDF 모두…", "both")):
            b = QPushButton(text_)
            b.clicked.connect(lambda _=False, k_=kind: self.save(k_))
            row.addWidget(b)
        lay = QVBoxLayout(self)
        lay.addWidget(tabs)
        lay.addLayout(row)

    def save_die3d(self):
        """금형 3D 조립 (STEP) · 조립도 PDF · 파트리스트 엑셀."""
        from . import die3d
        base = os.path.splitext(self.name or "part")[0] + "_금형"
        path, _ = QFileDialog.getSaveFileName(self, "금형 3D · 도면집 · 가공도 · 파트리스트 저장 (이름만 고르면 여섯 파일)",
                                              base + ".step", "STEP (*.step)")
        if not path:
            return
        stem = os.path.splitext(path)[0]
        try:
            plan = None
            try:
                pp = self.process_plan()
                plan = pp if pp is not None and getattr(pp, "kind", "") == "progressive" else None
            except Exception:
                plan = None
            from PySide6.QtCore import QSettings
            from PySide6.QtWidgets import QInputDialog

            from . import quote as qt_
            tols = ["m (보통)", "f (정밀)", "c (거침)"]
            pick, ok = QInputDialog.getItem(self, "일반 공차", "도면에 공차가 없는 치수의 일반 공차 (ISO 2768) — "
                                            "날 치수 마모 여유에 씀", tols, 0, False)
            if not ok:
                return
            machines = qt_.load_rates(QSettings("Hanil", "SheetUnfold").value("quote_rates", ""))["costing"]["machines"]
            assy = die3d.build(self.result, self.settings, plan=plan, tol=pick[0], machines=machines)
            out = [die3d.save_step(assy, stem + ".step"), die3d.save_pdf(assy, stem + "_조립도.pdf", self.name),
                   die3d.save_xlsx(assy, stem + "_파트리스트.xlsx", self.name),
                   die3d.save_dxf(assy, stem + "_가공도.dxf")]
            from . import diedraw
            sheets = diedraw.build(assy, {"product": os.path.splitext(os.path.basename(self.name or "part"))[0],
                                          "material": f"{self.settings.material} t{self.result.thickness:g}"})
            out += [diedraw.save_pdf(sheets, stem + "_도면집.pdf"), diedraw.save_dxf(sheets, stem + "_도면집.dxf")]
        except Exception as e:
            QMessageBox.warning(self, "금형 3D 저장 실패", str(e))
            return
        self._log(out)
        bad = [c for c in assy.checks if c[3] != "적합"]
        msg = (f"{assy.kind} · {assy.L:.0f} × {assy.W:.0f} · 닫힌 높이 {assy.shut:.1f} · 부품 {len(assy.parts)} 종\n"
               f"점검 {len(assy.checks)} 가지 중 검토 {len(bad)}" + "".join(f"\n  · {c[0]}: {c[1]} ({c[2]})" for c in bad)
               + "\n\n" + "\n".join(out))
        QMessageBox.information(self, "금형 3D", msg)
        QDesktopServices.openUrl(QUrl.fromLocalFile(out[4]))

    def _draw_tab(self):
        w = QWidget()
        lay = QHBoxLayout(w)
        form = QFormLayout()

        def spin(v, lo=0.0, hi=5000.0, dec=2):
            s = QDoubleSpinBox()
            s.setRange(lo, hi)
            s.setDecimals(dec)
            s.setValue(v)
            s.valueChanged.connect(self._calc_draw)
            return s

        t = self.result.thickness
        guess_d = 100.0
        if self.rep.blank_d:
            guess_d = round(self.rep.blank_d * 0.6, 1)
        self.d_d = spin(guess_d, 1)
        self.d_h = spin(round(guess_d * 0.5, 1))
        self.d_r = spin(max(4 * t, 1.0))
        self.d_fd = spin(0.0)
        self.d_rf = spin(max(4 * t, 1.0))
        self.d_t = spin(t, 0.05, 50, 3)
        self.d_trim = spin(max(1.0, 0.03 * guess_d * 0.5))
        self.d_mat = QComboBox()
        for m in bc.MATERIALS:
            self.d_mat.addItem(f"{m.code}  {m.name}", m.code)
        self.d_mat.setCurrentIndex(max(0, self.d_mat.findData(self.settings.material)))
        self.d_mat.currentIndexChanged.connect(self._calc_draw)
        form.addRow("컵 지름 d (두께 가운데)", self.d_d)
        form.addRow("높이 h (바닥 ~ 위 끝)", self.d_h)
        form.addRow("바닥 모서리 R", self.d_r)
        form.addRow("플랜지 바깥 지름 (0 = 없음)", self.d_fd)
        form.addRow("플랜지 R", self.d_rf)
        form.addRow("두께 t", self.d_t)
        form.addRow("트림 여유 (높이·플랜지)", self.d_trim)
        form.addRow("재질", self.d_mat)
        note = QLabel("블랭크 지름은 두께 가운데 면 넓이가 같다는 가정 (D = √(d² + 4dh − 1.72dr − 0.56r²) 와 같음). "
                      "드로잉률은 주름 누르개를 쓴 플랜지 없는 원통의 대표 표.")
        note.setWordWrap(True)
        note.setStyleSheet("color:#777")
        left = QVBoxLayout()
        left.addLayout(form)
        left.addWidget(note)
        left.addStretch(1)
        lay.addLayout(left, 1)
        self.draw_out = QLabel()
        self.draw_out.setTextFormat(Qt.RichText)
        self.draw_out.setAlignment(Qt.AlignTop | Qt.AlignLeft)
        self.draw_out.setWordWrap(True)
        lay.addWidget(self.draw_out, 2)
        self._calc_draw()
        return w

    def _route_tab(self):
        from PySide6.QtWidgets import QSpinBox
        w = QWidget()
        lay = QVBoxLayout(w)
        row = QHBoxLayout()
        self.r_qty = QSpinBox()
        self.r_qty.setRange(1, 100_000_000)
        self.r_qty.setSingleStep(10_000)
        self.r_qty.setGroupSeparatorShown(True)
        self.r_qty.setValue(100_000)
        self.r_qty.setSuffix(" 개/년")
        self.r_qty.valueChanged.connect(self._calc_route)
        row.addWidget(QLabel("연간 생산량"))
        row.addWidget(self.r_qty)
        row.addStretch(1)
        ai = QPushButton("AI 도면 검토…")
        ai.setToolTip("도면 PDF(있으면)의 공차·주기를 Claude 가 읽고 공정 초안을 바로잡습니다")
        ai.clicked.connect(self._ai_review)
        go = QPushButton("공정도 보기 →")
        go.clicked.connect(lambda: self.tabs.setCurrentIndex(5))
        row.addWidget(ai)
        row.addWidget(go)
        lay.addLayout(row)
        self.r_head = QLabel()
        self.r_head.setWordWrap(True)
        lay.addWidget(self.r_head)
        self.r_steps = QTableWidget(0, 4)
        self.r_steps.setHorizontalHeaderLabels(["순서", "공정", "내용", "근거"])
        self.r_holes = QTableWidget(0, 7)
        self.r_holes.setHorizontalHeaderLabels(["구멍", "크기", "있는 판", "가까운 절곡 (mm)", "동축 짝",
                                                "뚫는 때", "근거"])
        for t_ in (self.r_steps, self.r_holes):
            t_.verticalHeader().setVisible(False)
            t_.setEditTriggers(QTableWidget.NoEditTriggers)
            t_.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
            t_.horizontalHeader().setStretchLastSection(True)
        lay.addWidget(self.r_steps, 2)
        lay.addWidget(QLabel("구멍별 판단"))
        lay.addWidget(self.r_holes, 1)
        self.r_ai = QLabel()
        self.r_ai.setWordWrap(True)
        self.r_ai.setTextFormat(Qt.RichText)
        lay.addWidget(self.r_ai)
        from . import procplan as pp
        rules = QLabel("근거 규칙: " + " · ".join(f"<b>{k}</b> {v}" for k, v in pp.RULES.items()))
        rules.setWordWrap(True)
        rules.setStyleSheet("color:#777")
        lay.addWidget(rules)
        self._calc_route()
        return w

    def _fill(self, table, rows):
        table.setRowCount(len(rows))
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                table.setItem(i, j, QTableWidgetItem(str(v)))

    def _calc_route(self, *_):
        from . import procplan as pp
        try:
            self.route = pp.route(self.result, self.settings, self.r_qty.value())
        except Exception as e:
            self.route = None
            self.r_head.setText(f"공정을 정하지 못했습니다: {e}")
            return
        rt = self.route
        f = rt.features
        self.r_head.setText(
            f"<h3>추천 금형: {rt.die_type}</h3>{rt.die_reason}<br>"
            f"{f.material} t{f.thickness:g} · 전개 {f.size[0]:.0f} × {f.size[1]:.0f} · 구멍 {len(f.holes)} · "
            f"절곡 {f.bends}" + (" (나란함)" if f.parallel else "") + (" · 성형품" if f.forming else "")
            + "".join(f"<br>⚠ {n}" for n in rt.notes))
        self._fill(self.r_steps, pp.route_rows(rt))
        self._fill(self.r_holes, pp.hole_rows(rt))

    def _ai_review(self):
        from PySide6.QtCore import QSettings
        from PySide6.QtWidgets import QInputDialog, QLineEdit

        from . import procplan as pp
        if self.route is None:
            return
        st = QSettings("Hanil", "SheetUnfold")
        start = st.value("last_dir", "") or ""
        pdf, _ = QFileDialog.getOpenFileName(self, "도면 PDF (없으면 취소 → 3D 특징만)", start, "PDF (*.pdf)")
        key = st.value("anthropic_api_key", "") or os.environ.get("ANTHROPIC_API_KEY", "")
        if not key:
            key, ok = QInputDialog.getText(self, "AI 도면 검토", "Anthropic API 키", QLineEdit.Password)
            if not ok or not key:
                return
        from PySide6.QtWidgets import QApplication
        QApplication.setOverrideCursor(Qt.WaitCursor)
        try:
            ai = pp.ai_review(self.route, pdf or None, self.r_qty.value(), api_key=key)
        except Exception as e:
            from .ai_bom_dialog import _explain
            QMessageBox.warning(self, "AI 도면 검토", _explain(e))
            return
        finally:
            QApplication.restoreOverrideCursor()
        from html import escape
        steps = "".join(f"<tr><td>{i}</td><td>{escape(s_['op'])}</td><td>{escape(s_['detail'])}</td>"
                        f"<td>{escape(s_['reason'])}</td></tr>" for i, s_ in enumerate(ai["steps"], 1))
        self.r_ai.setText(
            f"<h3>AI 검토 ({'도면 ' + escape(os.path.basename(pdf)) if pdf else '도면 없음'}): "
            f"{escape(ai['die_type'])}</h3>{escape(ai['die_reason'])}<br>"
            f"<table border=1 cellspacing=0 cellpadding=3>{steps}</table>"
            + "".join(f"<br>· 도면: {escape(n)}" for n in ai["drawing_notes"])
            + "".join(f"<br>⚠ {escape(n)}" for n in ai["warnings"]))

    def _process_tab(self):
        from PySide6.QtWidgets import QCheckBox, QScrollArea
        w = QWidget()
        lay = QVBoxLayout(w)
        row = QHBoxLayout()
        self.p_kind = QComboBox()
        self.p_kind.addItem("순차 금형 스트립 (판금)", "progressive")
        self.p_kind.addItem("드로잉 공정 ([원통 드로잉 계산] 값)", "draw")
        if self.result.forming is not None or not self.result.bends and not self.rep.checks:
            self.p_kind.setCurrentIndex(1)
        self.p_idle = QCheckBox("절곡 앞에 빈 공정")
        pdf = QPushButton("PDF 저장…")
        dxf = QPushButton("DXF 저장…")
        row.addWidget(QLabel("공정"))
        row.addWidget(self.p_kind)
        row.addWidget(self.p_idle)
        row.addStretch(1)
        row.addWidget(pdf)
        row.addWidget(dxf)
        lay.addLayout(row)
        self.p_img = QLabel()
        self.p_img.setAlignment(Qt.AlignTop | Qt.AlignLeft)
        sc = QScrollArea()
        sc.setWidget(self.p_img)
        sc.setWidgetResizable(True)
        lay.addWidget(sc, 1)
        self.p_table = QTableWidget(0, 5)
        self.p_table.setHorizontalHeaderLabels(["공정", "작업", "내용", "하중 kN", "tf"])
        self.p_table.verticalHeader().setVisible(False)
        self.p_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        self.p_table.horizontalHeader().setStretchLastSection(True)
        self.p_table.setMaximumHeight(200)
        lay.addWidget(self.p_table)
        self.p_sum = QLabel()
        self.p_sum.setWordWrap(True)
        lay.addWidget(self.p_sum)
        self.p_kind.currentIndexChanged.connect(self._calc_process)
        self.p_idle.toggled.connect(self._calc_process)
        pdf.clicked.connect(lambda: self._save_process("pdf"))
        dxf.clicked.connect(lambda: self._save_process("dxf"))
        self._calc_process()
        return w

    def process_plan(self):
        from . import process as pr
        if self.p_kind.currentData() == "draw":
            return pr.drawing(self.result, self.settings, self.draw_plan(), self.d_r.value(), self.d_h.value())
        return pr.progressive(self.result, self.settings, idle=self.p_idle.isChecked())

    def _calc_process(self, *_):
        from PySide6.QtGui import QPixmap

        from . import process as pr
        self.p_idle.setEnabled(self.p_kind.currentData() == "progressive")
        try:
            plan = self.process_plan()
        except Exception as e:
            self.plan = None
            self.p_img.setPixmap(QPixmap())
            self.p_img.setText(f"공정도를 만들 수 없습니다: {e}")
            self.p_table.setRowCount(0)
            self.p_sum.setText("")
            return
        self.plan = plan
        self.p_img.setPixmap(QPixmap.fromImage(pr.render(plan, 1080, 420)))
        rows = pr.station_rows(plan)
        self.p_table.setRowCount(len(rows))
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                it = QTableWidgetItem(v if isinstance(v, str) else f"{v:,.1f}")
                if not isinstance(v, str):
                    it.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                self.p_table.setItem(i, j, it)
        self.p_sum.setText("<br>".join(pr.summary_lines(plan)))

    def _save_process(self, kind):
        from . import process as pr
        if self.plan is None:
            return
        if kind == "dxf" and self.plan.kind != "progressive":
            QMessageBox.information(self, "금형 공정도", "DXF 는 순차 금형 스트립 공정도만 됩니다. PDF 로 저장하세요.")
            return
        base = os.path.splitext(self.name or "part")[0] + f"_공정도.{kind}"
        filt = "PDF (*.pdf)" if kind == "pdf" else "DXF (*.dxf)"
        path, _ = QFileDialog.getSaveFileName(self, "금형 공정도 저장", base, filt)
        if not path:
            return
        if kind == "pdf":
            pr.save_pdf(self.plan, path, self.name, route=self.route)
        else:
            pr.save_dxf(self.plan, path, self.name)
        QDesktopServices.openUrl(QUrl.fromLocalFile(path))

    def _dmat(self):
        code = self.d_mat.currentData()
        return self.settings.mat if code == self.settings.material else bc.MATERIAL_BY_CODE[code]   # 밀시트 값

    def draw_plan(self):
        mat = self._dmat()
        return dc.draw_plan(self.d_d.value(), self.d_h.value(), self.d_r.value(), self.d_t.value(), mat,
                            self.d_fd.value(), self.d_rf.value(), self.d_trim.value())

    def _calc_draw(self, *_):
        if getattr(self, "p_kind", None) is not None and self.p_kind.currentData() == "draw":
            self._calc_process()
        mat = self._dmat()
        try:
            pl = self.draw_plan()
        except (ValueError, ZeroDivisionError) as e:
            self.draw_out.setText(f"계산할 수 없습니다: {e}")
            return
        rows = "".join(f"<tr><td>{s.no}</td><td align=right>{s.m:.3f}</td><td align=right>Ø{s.d:.1f}</td>"
                       f"<td align=right>{s.table_n / 1000:,.1f}</td><td align=right>{s.siebel_n / 1000:,.1f}</td>"
                       f"<td align=right><b>{s.force_n / 1000:,.1f}</b></td><td align=right>{s.force_n / dc.TON:.2f}</td>"
                       f"<td align=right>R{s.die_r:.1f}</td><td align=right>{s.gap:.2f} (최대 {s.gap_max:.2f})</td></tr>"
                       for s in pl.stages)
        from html import escape
        notes = "".join(f"<br>⚠ {escape(n)}" for n in pl.notes)
        self.draw_out.setText(
            f"<h3>블랭크 Ø{pl.D:.2f} mm</h3>"
            f"상대 두께 t/D = {pl.rel_thickness:.2f} % · 전체 드로잉률 d/D = {pl.ratio_total:.3f} · "
            f"<b>드로잉 {len(pl.stages)}회</b><br>"
            f"주름 누르개: {escape(pl.holder_needed)} · MISUMI 판정 {escape(pl.holder_misumi)}<br>"
            f"누르개 힘 {pl.holder_n / 1000:,.1f} kN ({pl.holder_n / dc.TON:.2f} tf, Siebel 면압 {pl.holder_p:.2f} MPa) · "
            f"다이 쿠션 고를 때 {pl.holder_tab_n / 1000:,.1f} kN (교재 표 면압){notes}<br><br>"
            f"한계 드로잉비 LDR = {dc.ldr(mat):.2f} (Whiteley·Hosford, r = {mat.r:g})<br><br>"
            "<table border=1 cellspacing=0 cellpadding=4><tr><th>공정</th><th>드로잉률 m</th><th>지름</th>"
            f"<th>교재 식 kN</th><th>Siebel kN</th><th>드로잉력 kN (큰 값)</th><th>tf</th><th>다이 R</th>"
            f"<th>한쪽 틈새 mm</th></tr>{rows}</table>"
            "<br><span style='color:#777'>교재 식 F = π·d·t·σb·n (n: 드로잉률에 따른 계수 0.4~1.0) · "
            "Siebel F = π·d·t·1.1·k_fm·ln(D/d) ÷ 0.65 (k_fm 평균 유동 응력, 홀로몬) — 큰 값으로 프레스를 고릅니다.<br>"
            "다이 R = 0.035·[50 + (D − d)]·√t (Oehler·Kaiser), 틈새 권장 1.0~1.1 t (MISUMI) · 최대 = 위 끝 두께 "
            "t·β^(1/(1+r)), 누르개 면압 p = [(β − 1)² 또는 ³ + d/(200 t)]·Rm/400 (Siebel)</span>")

    def _log(self, files):
        if self.history is None:
            return
        rep, rt = self.rep, self.route
        plan = getattr(self, "plan", None)
        ops = len(plan.stations) if plan is not None else (len(rt.steps) if rt else 0)
        bad = sum(c.verdict == dc.NG for c in rep.checks)
        warn = sum(c.verdict == dc.WARN for c in rep.checks)
        v = dict(title=self.name, material=rep.material, annual=self.r_qty.value() if rt else None,
                 max_tonf=(plan.press_tonf if plan is not None else rep.press_tonf), total_ops=ops, parts=1,
                 machines="", summary=(f"{rt.die_type} · " if rt else "") + f"공정 {ops} · 권장 프레스 "
                 f"{(plan.press_tonf if plan is not None else rep.press_tonf):.0f} tf · 클리어런스 {rep.clearance_mm:.3f} · "
                 f"설계 규칙 부적합 {bad}·주의 {warn}"
                 + (f" · 스트립 이용률 {rep.strip.utilization:.0%}" if rep.strip else ""))
        detail = {"loads": [[L.name, round(L.force_n / 1000, 2)] for L in rep.loads],
                  "checks": [[c.rule, c.target, round(c.value, 2), c.verdict] for c in rep.checks],
                  "route": [[s_.op, s_.detail, s_.rules] for s_ in rt.steps] if rt else [],
                  "stations": [[st_.no, st_.title, st_.detail] for st_ in plan.stations] if plan is not None else []}
        try:
            self.history.add_doc("금형 검토", files, self.source_path, v, detail)
        except Exception:
            import traceback
            traceback.print_exc()

    def save(self, kind="xlsx"):
        base = os.path.splitext(self.name or "part")[0] + "_금형검토"
        filt = {"xlsx": "엑셀 (*.xlsx)", "pdf": "PDF (*.pdf)", "both": "엑셀 (*.xlsx)"}[kind]
        path, _ = QFileDialog.getSaveFileName(self, "금형 검토 저장", base + (".pdf" if kind == "pdf" else ".xlsx"), filt)
        if not path:
            return
        stem = os.path.splitext(path)[0]
        args = (self.rep, None, self.name, self.strip_img, self.draw_plan(), self.route,
                getattr(self, "plan", None))
        out = []
        try:
            if kind in ("xlsx", "both"):
                save_xlsx(args[0], stem + ".xlsx", *args[2:])
                out.append(stem + ".xlsx")
            if kind in ("pdf", "both"):
                save_pdf(args[0], stem + ".pdf", *args[2:])
                out.append(stem + ".pdf")
        except Exception as e:
            QMessageBox.warning(self, "금형 검토 저장 실패", str(e))
            return
        self._log(out)
        QMessageBox.information(self, "금형 검토", "저장했습니다.\n" + "\n".join(out))
        QDesktopServices.openUrl(QUrl.fromLocalFile(out[-1]))


def save_xlsx(rep, path, name="", strip_img=None, plan=None, route=None, process_plan=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    from . import company
    from .report import _add_image

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E8C")

    def sheet(ws, title, head, rows, start=3):
        ws["A1"] = title
        ws["A1"].font = Font(bold=True, size=13)
        ws.cell(1, len(head) + 1, company.MAKER)
        for j, h in enumerate(head, 1):
            c = ws.cell(start, j, h)
            c.font, c.fill = head_font, head_fill
            ws.column_dimensions[c.column_letter].width = max(12, len(h) * 2 + 2)
        for i, r in enumerate(rows, start + 1):
            for j, v in enumerate(r, 1):
                c = ws.cell(i, j, v)
                if isinstance(v, float):
                    c.number_format = "#,##0.00"
        return start + len(rows) + 2

    ws = wb.active
    ws.title = "하중"
    nxt = sheet(ws, f"금형 검토 — {name}  ({rep.material}, t {rep.thickness:g})",
                ["하중", "길이 mm", "힘 kN", "톤 (tf)", "비고"], load_rows(rep))
    ws.cell(nxt, 1, f"한쪽 클리어런스 {rep.clearance_pct:g} % = {rep.clearance_mm:.3f} mm · "
                    f"전단 저항 {rep.shear_mpa:.0f} MPa")
    ws.cell(nxt + 1, 1, f"VDI 3368 클리어런스 u = 0.01·s·√(0.1τ) = {rep.vdi_clearance_mm:.3f} mm · "
                        f"다이 두께 어림 ∛(0.1F) = {rep.die_plate_mm:.0f} mm · 보통 공차 JIS B 0408-B / ISO 2768-m")
    nxt += 1
    if rep.bends:
        sheet(ws, "", ["절곡", "길이 mm", "힘 kN", "톤 (tf)", "보정 없이 (°)", "금형 보정 각 (°)", "Ks"],
              bend_rows(rep), start=nxt + 2)
    sheet(wb.create_sheet("설계 규칙"), "설계 규칙 검사", ["규칙", "대상", "값 mm", "기준 mm", "설명", "판정"],
          check_rows(rep))
    if rep.strip is not None:
        s = rep.strip
        ws3 = wb.create_sheet("스트립")
        sheet(ws3, "스트립 배치 (1열)", ["항목", "값"],
              [["부품 각도 (°)", float(s.angle)], ["피치 mm", s.pitch], ["폭 mm", s.width],
               ["잔폭 mm", s.bridge], ["옆 잔폭 mm", s.side], ["재료 이용률", s.utilization],
               ["1 m 당 개수", s.per_meter]])
        ws3.cell(9, 2).number_format = "0.0%"
        tmp = _add_image(ws3, strip_img, "D3", 800) if strip_img is not None else None
    else:
        tmp = None
    if plan is not None:
        ws4 = wb.create_sheet("원통 드로잉")
        nxt = sheet(ws4, f"원통 드로잉 — 블랭크 Ø{plan.D:.2f} mm, {len(plan.stages)}회",
                    ["공정", "드로잉률 m", "지름 mm", "교재 식 kN", "Siebel kN", "드로잉력 kN", "tf", "다이 R mm",
                     "한쪽 틈새 mm", "틈새 최대 mm"],
                    [[s.no, s.m, s.d, s.table_n / 1000, s.siebel_n / 1000, s.force_n / 1000, s.force_n / dc.TON,
                      s.die_r, s.gap, s.gap_max] for s in plan.stages])
        ws4.cell(nxt, 1, f"t/D {plan.rel_thickness:.2f} % · 주름 누르개 {plan.holder_needed} (MISUMI {plan.holder_misumi}) · "
                         f"누르개 힘 {plan.holder_n / 1000:,.1f} kN (Siebel {plan.holder_p:.2f} MPa) · "
                         f"다이 쿠션 {plan.holder_tab_n / 1000:,.1f} kN")
        for k, n in enumerate(plan.notes, 1):
            ws4.cell(nxt + k, 1, n)
    tmps = [tmp] if tmp else []
    if route is not None:
        from . import procplan as pp
        ws5 = wb.create_sheet("자동 공정 설계")
        nxt = sheet(ws5, f"추천 금형: {route.die_type} — {route.die_reason}", ["순서", "공정", "내용", "근거"],
                    pp.route_rows(route))
        if route.features.holes:
            nxt = sheet(ws5, "", ["구멍", "크기", "있는 판", "가까운 절곡", "동축 짝", "뚫는 때", "근거"],
                        pp.hole_rows(route), start=nxt)
        for i, (k_, v) in enumerate(pp.RULES.items()):
            ws5.cell(nxt + i, 1, k_)
            ws5.cell(nxt + i, 2, v)
        if route.ai:
            r0 = nxt + len(pp.RULES) + 1
            ws5.cell(r0, 1, f"AI 도면 검토: {route.ai['die_type']} — {route.ai['die_reason']}")
            for i, st_ in enumerate(route.ai["steps"], 1):
                ws5.cell(r0 + i, 1, i)
                ws5.cell(r0 + i, 2, st_["op"])
                ws5.cell(r0 + i, 3, st_["detail"])
                ws5.cell(r0 + i, 4, st_["reason"])
    if process_plan is not None:
        from . import process as pr
        ws6 = wb.create_sheet("금형 공정도")
        nxt = sheet(ws6, "금형 공정도 (초안)", ["공정", "작업", "내용", "하중 kN", "tf"], pr.station_rows(process_plan))
        for i, line in enumerate(pr.summary_lines(process_plan)):
            ws6.cell(nxt + i, 1, line)
        try:
            tmps.append(_add_image(ws6, pr.render(process_plan, 1400, 520), f"A{nxt + len(pr.summary_lines(process_plan)) + 1}",
                                   1000))
        except Exception:
            pass
    from .company import stamp_workbook
    stamp_workbook(wb)
    wb.save(path)
    for t_ in tmps:
        try:
            os.remove(t_)
        except OSError:
            pass


def save_pdf(rep, path, name="", strip_img=None, plan=None, route=None, process_plan=None):
    """금형 검토 PDF (A3 가로): 하중·클리어런스·스프링백 / 설계 규칙 / 스트립·드로잉 / 공정도 / 공정 설계."""
    from PySide6.QtCore import QMarginsF, QPointF, QRectF
    from PySide6.QtGui import QPageLayout, QPageSize, QPdfWriter

    from . import process as pr
    from .report import _brand, _font, ensure_gui, version_text
    ensure_gui()
    pdf = QPdfWriter(path)
    pdf.setResolution(300)
    pdf.setPageLayout(QPageLayout(QPageSize(QPageSize.A3), QPageLayout.Landscape, QMarginsF(0, 0, 0, 0)))
    pdf.setTitle(f"금형 검토 - {name}")
    pdf.setCreator(f"한일 판금 전개 {version_text()}")
    k = 300 / 25.4
    W, H, m = 420.0, 297.0, 12.0

    def R(x, y, w, h):
        return QRectF(x * k, y * k, w * k, h * k)

    p = QPainter(pdf)
    p.setRenderHint(QPainter.Antialiasing)

    def header(title):
        p.setPen(QColor(20, 20, 20))
        p.setFont(_font(16, True))
        p.drawText(R(m, m, W - 2 * m, 9), Qt.AlignLeft | Qt.AlignVCenter, title)
        p.setFont(_font(8))
        _brand(p, R(m, m + 1.5, W - 2 * m, 6))
        p.setPen(QPen(QColor(20, 20, 20), 0.4 * k))
        p.drawLine(QPointF(m * k, (m + 10) * k), QPointF((W - m) * k, (m + 10) * k))
        return m + 15

    def fmt(rows):
        return [[v if isinstance(v, str) else f"{v:,.2f}" for v in r] for r in rows]

    # 1 쪽: 하중 · 절곡 · 설계 규칙
    y = header(f"금형 검토 — {name}  ({rep.material}, t {rep.thickness:g})")
    p.setFont(_font(10))
    p.setPen(QColor(30, 30, 30))
    p.drawText(R(m, y, W - 2 * m, 6), Qt.AlignLeft | Qt.AlignVCenter,
               f"전단 저항 τ = {rep.shear_mpa:.0f} MPa · 한쪽 클리어런스 {rep.clearance_pct:g} % = {rep.clearance_mm:.3f} mm "
               f"(다이 = 펀치 + {2 * rep.clearance_mm:.3f}) · VDI 3368 {rep.vdi_clearance_mm:.3f} mm · "
               f"다이 두께 ≈ {rep.die_plate_mm:.0f} mm · 권장 프레스 {rep.press_tonf:.0f} tf")
    y += 9
    y = pr._table(p, R, k, m, y, [("하중", 80), ("길이 mm", 26), ("힘 kN", 26), ("tf", 22), ("비고", 50)],
                  fmt(load_rows(rep)))
    if rep.bends:
        y = pr._table(p, R, k, m, y + 5, [("절곡", 16), ("길이 mm", 24), ("힘 kN", 24), ("tf", 20),
                                          ("보정 없이 (°)", 30), ("금형 보정 각 (°)", 32), ("Ks", 22)],
                      fmt(bend_rows(rep)), maxy=H - m - 10)
    x2 = m + 216
    pr._table(p, R, k, x2, m + 24, [("규칙", 46), ("대상", 22), ("값", 18), ("기준", 18), ("판정", 16)],
              [[c.rule, c.target, f"{c.value:.2f}", f"{c.limit:.2f}", c.verdict] for c in rep.checks],
              8, maxy=H - m - 10)
    # 2 쪽: 스트립 · 드로잉
    if strip_img is not None or plan is not None:
        pdf.newPage()
        y = header(f"스트립 배치 · 원통 드로잉 — {name}")
        if strip_img is not None:
            iw = W - 2 * m
            ih = iw * strip_img.height() / strip_img.width()
            if ih > 130:                      # 높이를 줄이면 너비도 같은 비율로
                iw, ih = iw * 130 / ih, 130
            p.drawImage(R(m, y, iw, ih), strip_img)
            y += ih + 6
        if plan is not None:
            p.setFont(_font(11, True))
            p.setPen(QColor(31, 78, 140))
            p.drawText(R(m, y, W - 2 * m, 7), Qt.AlignLeft | Qt.AlignVCenter,
                       f"원통 드로잉: 블랭크 Ø{plan.D:.2f} mm · {len(plan.stages)}회 · t/D {plan.rel_thickness:.2f} % · "
                       f"주름 누르개 {plan.holder_needed}")
            y += 9
            pr._table(p, R, k, m, y, [("공정", 20), ("드로잉률 m", 30), ("지름 mm", 30), ("드로잉력 kN", 34), ("tf", 24),
                                      ("다이 R", 24), ("한쪽 틈새 (최대)", 40)],
                      [[str(s_.no), f"{s_.m:.3f}", f"{s_.d:.1f}", f"{s_.force_n / 1000:,.1f}", f"{s_.force_n / dc.TON:.2f}",
                        f"R{s_.die_r:.1f}", f"{s_.gap:.2f} ({s_.gap_max:.2f})"] for s_ in plan.stages])
    # 3 쪽: 공정도, 4 쪽: 자동 공정 설계
    if process_plan is not None:
        pdf.newPage()
        y = header(f"금형 공정도 (초안) — {name}")
        box = R(m, y, W - 2 * m, 160)
        p.setPen(QPen(QColor(180, 180, 180), 0.2 * k))
        p.setBrush(Qt.NoBrush)
        p.drawRect(box)
        p.setFont(_font(9))
        pr.paint(p, box.adjusted(3 * k, 2 * k, -3 * k, -2 * k), process_plan, 3.6 * k)
        yy = pr._table(p, R, k, m, y + 166, [("공정", 16), ("작업", 40), ("내용", 190), ("하중 kN", 30), ("tf", 24)],
                       fmt(pr.station_rows(process_plan)), maxy=H - m - 8)
        del yy
    if route is not None:
        pdf.newPage()
        pr._route_page(p, R, k, W, H, m, route, name)
    p.end()
