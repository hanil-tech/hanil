"""원가계산서 (자동차 협력사 양식): 재료비 · 가공비(노무비 · 경비) · 불량 로스 · 일반관리비 · 이윤 ·
포장·운반비 · 금형 상각 → 판매 단가. 그리고 수량별 단가표, 연차별 단가 인하, 철강 시세 민감도.

  재료비   = 원재료 (소요 kg × 원/kg) − 스크랩 회수 (+ 부재료)
  노무비   = Σ 공정 사이클 × 인원 × 노무 임률            (생산 견적의 인원)
  경비     = 가공비 − 노무비                          (기계 상각·전력·소모품 — 기계 임률에 들어 있던 몫)
  불량 로스 = (재료비 + 가공비) × 불량률
  일반관리비 = (재료비 + 가공비 + 불량 로스) × 관리비율
  이윤     = (가공비 + 일반관리비) × 이윤율             (자동차 협력사 관례: 재료비에는 이윤을 붙이지 않음)
             [이윤 기준 = 1 이면 (제조원가 + 관리비) × 이윤율 — 생산 견적서와 같은 방식]
  포장·운반 = 개당 포장비 + 제품 kg × 운반 원/kg
  금형 상각 = 금형비 ÷ (연간 수량 × 상각 연수)            ([금형 상각 포함] 일 때)
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field

from . import production as prod

DEFAULTS = {
    "labor_rate": 30_000,     # 노무 임률 원/시간 (기계 임률 중 사람 몫)
    "profit_base": 0,         # 0 = (가공비 + 관리비) × 이윤율, 1 = (제조원가 + 관리비) × 이윤율
    "packing": 50,            # 포장비 원/개
    "freight_kg": 80,         # 운반비 원/kg
    "include_die": 1,         # 금형 상각을 단가에 (1) / 따로 (0)
    "die_years": 3,           # 금형 상각 연수
    "cr_pct": 3.0,            # 연차별 단가 인하 %
    "years": 4,               # 연차표 연수
    "qty_list": "10000,50000,100000,300000",
}

NAMES = {
    "labor_rate": "노무 임률 (원/시간)", "profit_base": "이윤 기준 (0 가공비+관리비 · 1 제조원가+관리비)",
    "packing": "포장비 (원/개)", "freight_kg": "운반비 (원/kg)", "include_die": "금형 상각 포함 (1/0)",
    "die_years": "금형 상각 연수", "cr_pct": "연차 단가 인하 (%)", "years": "연차표 연수",
    "qty_list": "수량별 단가표 (개/년, 쉼표)",
}


def params(rates):
    p = dict(DEFAULTS)
    p.update(rates.get("costsheet") or {})
    return p


@dataclass
class PartSheet:
    no: int
    name: str
    qty: int
    spec: str
    raw: float = 0.0          # 원재료비
    scrap: float = 0.0        # 스크랩 회수 (−)
    material: float = 0.0     # 재료비 계
    labor: float = 0.0
    expense: float = 0.0
    process: float = 0.0      # 가공비 계
    defect: float = 0.0
    overhead: float = 0.0
    profit: float = 0.0
    logistics: float = 0.0
    die: float = 0.0          # 금형 상각
    weight: float = 0.0

    @property
    def manufacturing(self):
        return self.material + self.process + self.defect

    @property
    def price(self):
        return self.manufacturing + self.overhead + self.profit + self.logistics + self.die


@dataclass
class CostSheet:
    title: str
    annual: int
    parts: list
    subs: float               # 부재료 (세트당)
    extra: float              # 조립·용접·도장 (세트당)
    p: dict
    rates: dict
    sub_overhead: float = 0.0
    sub_profit: float = 0.0
    notes: list = field(default_factory=list)

    @property
    def set_price(self):
        return math.ceil((sum(ps.price * ps.qty for ps in self.parts) + self.subs + self.extra
                          + self.sub_overhead + self.sub_profit) / 10) * 10

    def total(self, key):
        return sum(getattr(ps, key) * ps.qty for ps in self.parts)


def compute(q, mat_factor=1.0) -> CostSheet:
    """견적(quote.Quote) → 원가계산서. mat_factor: 원재료 단가 배수 (철강 시세 민감도)."""
    rates = q.rates
    p = params(rates)
    pp = prod.params(rates)
    oh, pf = rates["overhead_pct"] / 100, rates["profit_pct"] / 100
    defect = pp["defect_pct"] / 100
    mat_price = rates["material"]
    rows = []
    for pc in q.parts:
        code = pc.spec.split()[0] if pc.spec else ""
        price = float(mat_price.get(code, 0))
        ps = PartSheet(pc.no, pc.name, pc.qty, pc.spec, weight=pc.weight)
        if pc.material_override is None and price:
            raw0 = pc.gross_kg * price
            ps.scrap = min(pc.material_auto - raw0, 0.0)     # 견적의 재료비와 같게 (스크랩 = 차이)
            ps.raw = raw0 * mat_factor
            ps.material = ps.raw + ps.scrap
        else:                                              # 직접 넣은 재료비 · 기계 가공 부품
            ps.material = ps.raw = pc.material * (mat_factor if pc.material_override is None else 1.0)
        ps.process = pc.process_cost
        crew_s = 0.0
        for o in pc.opc:
            crew_s += o.seconds * prod._crew(pp, o.kind, pc.method, pc.n_ops)
        ps.labor = min(crew_s / 3600 * p["labor_rate"], ps.process)
        ps.expense = ps.process - ps.labor
        ps.defect = (ps.material + ps.process) * defect
        ps.overhead = ps.manufacturing * oh
        base = (ps.process + ps.overhead) if int(p["profit_base"]) == 0 else (ps.manufacturing + ps.overhead)
        ps.profit = base * pf
        ps.logistics = p["packing"] + pc.weight * p["freight_kg"]
        if int(p["include_die"]) and pc.die_cost:
            ps.die = pc.die_cost / max(q.annual * pc.qty * max(p["die_years"], 1), 1)
        rows.append(ps)
    cs = CostSheet(q.title, q.annual, rows, q.sub_total, q.extra, p, rates)
    cs.sub_overhead = (q.sub_total + q.extra) * oh
    cs.sub_profit = (q.extra + cs.sub_overhead) * pf if int(p["profit_base"]) == 0 else \
        (q.sub_total + q.extra + cs.sub_overhead) * pf
    if not int(p["include_die"]) and q.die_total:
        cs.notes.append(f"금형비 {q.die_total:,.0f} 원은 단가에 넣지 않고 따로 청구 (일시불)")
    elif q.die_total:
        cs.notes.append(f"금형비 {q.die_total:,.0f} 원을 {p['die_years']:g} 년 × 연 {q.annual:,} 세트로 나눠 단가에 넣음")
    return cs


ROWS = [("원재료비", "raw"), ("스크랩 회수", "scrap"), ("재료비 계", "material"), ("노무비", "labor"),
        ("경비 (기계·전력·소모품)", "expense"), ("가공비 계", "process"), ("불량 로스", "defect"),
        ("제조원가", "manufacturing"), ("일반관리비", "overhead"), ("이윤", "profit"), ("포장·운반비", "logistics"),
        ("금형 상각비", "die"), ("판매 단가 (개당)", "price")]


def sheet_rows(cs):
    """세로 = 원가 항목, 가로 = 부품 (개당) + 세트 합계."""
    out = []
    for label, key in ROWS:
        row = [label] + [float(getattr(ps, key)) for ps in cs.parts] + [float(cs.total(key))]
        out.append(row)
    out.append(["부재료 (세트당)"] + [""] * len(cs.parts) + [float(cs.subs)])
    out.append(["조립·용접·도장 (세트당)"] + [""] * len(cs.parts) + [float(cs.extra)])
    out.append(["부재료·조립 관리비·이윤"] + [""] * len(cs.parts) + [float(cs.sub_overhead + cs.sub_profit)])
    out.append(["세트 판매 단가"] + [""] * len(cs.parts) + [float(cs.set_price)])
    return out


def sheet_head(cs):
    return ["항목"] + [f"{ps.no}. {ps.name} (×{ps.qty})" for ps in cs.parts] + ["세트 합계"]


def qty_table(builder, cs_p, base_q):
    """수량별 단가: builder(annual) → Quote (공정·금형 형식이 수량에 따라 바뀜)."""
    out = []
    for a in [int(float(x)) for x in str(cs_p["qty_list"]).replace(" ", "").split(",") if x]:
        try:
            q = builder(a)
        except Exception:
            continue
        c = compute(q)
        methods = ", ".join(sorted({pc.method.split(" (")[0] for pc in q.parts}))
        out.append([a, float(c.set_price), float(q.die_total), float(c.set_price * a), methods])
    return out


QTY_HEAD = ["연간 수량", "세트 단가", "금형비", "연 매출", "공정 방식"]


def year_table(cs):
    p = cs.p
    out = []
    for y in range(1, int(p["years"]) + 1):
        price = cs.set_price * (1 - p["cr_pct"] / 100) ** (y - 1)
        cum = 1 - (1 - p["cr_pct"] / 100) ** (y - 1)
        out.append([f"{y} 년차", float(price), float(price * cs.annual), f"−{cum:.1%} (처음 단가 대비)"
                    if y > 1 else "기준"])
    return out


YEAR_HEAD = ["연차", "세트 단가", "연 매출", "인하"]


def sensitivity(q):
    base = compute(q).set_price
    out = []
    for f in (0.8, 0.9, 1.0, 1.1, 1.2):
        v = compute(q, f).set_price
        out.append([f"{(f - 1) * 100:+.0f} %", float(v), float(v - base), (v - base) / base if base else 0.0])
    return out


SENS_HEAD = ["철강 (원재료) 단가", "세트 단가", "차이", "차이 %"]


# ---------------------------------------------------------------- 저장

def add_sheets(wb, cs, q, qty=None, title="", first=False):
    from openpyxl.styles import Font, PatternFill
    hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E8C")
    bold = Font(bold=True)

    def table(ws, r0, head, rows, pct=(), bold_rows=()):
        for j, h in enumerate(head, 1):
            c = ws.cell(r0, j, h)
            c.font, c.fill = hf, hfill
            ws.column_dimensions[c.column_letter].width = max(ws.column_dimensions[c.column_letter].width or 0,
                                                              14, min(len(str(h)) * 2 + 2, 40))
        for i, row in enumerate(rows, r0 + 1):
            for j, v in enumerate(row, 1):
                c = ws.cell(i, j, v)
                if isinstance(v, float):
                    c.number_format = "0.0%" if j - 1 in pct else "#,##0"
                if row[0] in bold_rows:
                    c.font = bold
        return r0 + len(rows) + 2

    ws = wb.active if first else wb.create_sheet("원가계산서")
    ws.title = "원가계산서"
    ws["A1"] = f"원가계산서 — {title or cs.title}  (연 {cs.annual:,} 세트, 단위 원)"
    ws["A1"].font = Font(bold=True, size=13)
    r = table(ws, 3, sheet_head(cs), sheet_rows(cs),
              bold_rows=("재료비 계", "가공비 계", "제조원가", "판매 단가 (개당)", "세트 판매 단가"))
    for n in cs.notes + [f"일반관리비 {cs.rates['overhead_pct']:g} % · 이윤 {cs.rates['profit_pct']:g} % "
                         f"({'가공비+관리비' if int(cs.p['profit_base']) == 0 else '제조원가+관리비'} 기준) · "
                         f"불량률 {prod.params(cs.rates)['defect_pct']:g} %"]:
        ws.cell(r, 1, n)
        r += 1
    ws2 = wb.create_sheet("수량·연차·시세")
    r2 = 1
    if qty:
        ws2.cell(r2, 1, "수량별 단가표").font = bold
        r2 = table(ws2, r2 + 1, QTY_HEAD, qty)
    ws2.cell(r2, 1, f"연차별 단가 (연 {cs.p['cr_pct']:g} % 인하)").font = bold
    r2 = table(ws2, r2 + 1, YEAR_HEAD, year_table(cs))
    ws2.cell(r2, 1, "철강 시세 민감도").font = bold
    table(ws2, r2 + 1, SENS_HEAD, sensitivity(q), pct=(3,))
    ws3 = wb.create_sheet("원가 조건")
    for k, v in cs.p.items():
        ws3.append([NAMES.get(k, k), v])
    ws3.column_dimensions["A"].width = 40
    return []


def save_xlsx(cs, q, path, qty=None, title=""):
    from openpyxl import Workbook

    from .company import stamp_workbook
    wb = Workbook()
    add_sheets(wb, cs, q, qty, title, first=True)
    stamp_workbook(wb)
    wb.save(path)


def save_pdf(cs, q, path, qty=None, title=""):
    from PySide6.QtCore import QMarginsF, QPointF, QRectF, Qt
    from PySide6.QtGui import QColor, QPageLayout, QPageSize, QPainter, QPdfWriter, QPen

    from . import process as pr
    from .report import _brand, _font, ensure_gui, version_text
    ensure_gui()
    pdf = QPdfWriter(path)
    pdf.setResolution(300)
    pdf.setPageLayout(QPageLayout(QPageSize(QPageSize.A3), QPageLayout.Landscape, QMarginsF(0, 0, 0, 0)))
    pdf.setTitle(f"원가계산서 - {title or cs.title}")
    pdf.setCreator(f"한일 판금 전개 {version_text()}")
    k = 300 / 25.4
    W, H, m = 420.0, 297.0, 12.0

    def R(x, y, w, h):
        return QRectF(x * k, y * k, w * k, h * k)

    p = QPainter(pdf)
    p.setRenderHint(QPainter.Antialiasing)
    p.setPen(QColor(20, 20, 20))
    p.setFont(_font(16, True))
    p.drawText(R(m, m, W - 2 * m, 9), Qt.AlignLeft | Qt.AlignVCenter,
               f"원가계산서 — {title or cs.title}  (연 {cs.annual:,} 세트, 단위 원)")
    p.setFont(_font(8))
    _brand(p, R(m, m + 1.5, W - 2 * m, 6))
    p.setPen(QPen(QColor(20, 20, 20), 0.4 * k))
    p.drawLine(QPointF(m * k, (m + 10) * k), QPointF((W - m) * k, (m + 10) * k))
    y = m + 15
    head = sheet_head(cs)
    n_parts = min(len(head) - 2, 8)                   # 부품이 많으면 앞 8 개 + 합계
    idx = list(range(1, 1 + n_parts)) + [len(head) - 1]
    shown = [head[0]] + [head[i] for i in idx]
    cw = (W - 2 * m - 58) / max(len(shown) - 1, 1)
    rows = []
    for r in sheet_rows(cs):
        rows.append([r[0]] + [(f"{r[i]:,.0f}" if isinstance(r[i], float) else str(r[i])) for i in idx])
    y = pr._table(p, R, k, m, y, [(shown[0], 58)] + [(h_[:18], cw) for h_ in shown[1:]], rows, 8, maxy=H - 80)
    p.setFont(_font(9))
    p.setPen(QColor(30, 30, 30))
    for n in cs.notes:
        p.drawText(R(m, y, W - 2 * m, 5), Qt.AlignLeft | Qt.AlignVCenter, n)
        y += 5
    y += 3
    x = m
    if qty:
        pr._table(p, R, k, x, y, [(h_, w_) for h_, w_ in zip(QTY_HEAD, (24, 26, 28, 32, 50))],
                  [[f"{a:,}", f"{b:,.0f}", f"{c:,.0f}", f"{d:,.0f}", e] for a, b, c, d, e in qty], 8, maxy=H - m - 3)
        x += 166
    pr._table(p, R, k, x, y, [(h_, w_) for h_, w_ in zip(YEAR_HEAD, (20, 26, 32, 28))],
              [[a, f"{b:,.0f}", f"{c:,.0f}", d] for a, b, c, d in year_table(cs)], 8, maxy=H - m - 3)
    x += 112
    pr._table(p, R, k, x, y, [(h_, w_) for h_, w_ in zip(SENS_HEAD, (34, 24, 22, 18))],
              [[a, f"{b:,.0f}", f"{c:+,.0f}", f"{d:+.1%}"] for a, b, c, d in sensitivity(q)], 8, maxy=H - m - 3)
    p.end()
