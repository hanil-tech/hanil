"""아이템 (견적 요청) 대장: 받은 아이템마다 고객·차종·품번·도면 리비전·연간 수량·양산 시기·견적 마감·상태,
그리고 견적 리비전 (언제 얼마로 냈는지) 을 쌓는다. 기록과 같은 history.sqlite (PC 안 파일) 의 items 표.

상태: 접수 → 견적 중 → 제출 → 수주 / 실주 / 보류
"""
from __future__ import annotations

import datetime as _dt
import json

STATUSES = ["접수", "견적 중", "제출", "수주", "실주", "보류"]

ITEM_COLUMNS = [
    ("id", "INTEGER PRIMARY KEY AUTOINCREMENT"),
    ("created", "TEXT"),
    ("updated", "TEXT"),
    ("customer", "TEXT"),         # 고객
    ("project", "TEXT"),          # 차종·프로젝트
    ("part_no", "TEXT"),          # 고객 품번
    ("part_name", "TEXT"),
    ("drawing_rev", "TEXT"),      # 도면 리비전
    ("annual", "INTEGER"),        # 연간 수량
    ("sop", "TEXT"),              # 양산 시작
    ("eop", "TEXT"),              # 양산 끝
    ("due", "TEXT"),              # 견적 마감
    ("status", "TEXT"),
    ("owner", "TEXT"),            # 담당
    ("source_path", "TEXT"),      # 3D 도면
    ("unit_price", "REAL"),       # 마지막 제출 세트 단가
    ("die_total", "REAL"),        # 마지막 제출 금형비
    ("revisions", "TEXT"),        # JSON [{rev, date, unit_price, die_total, annual, note, files}]
    ("note", "TEXT"),
]
ITEM_LABELS = {
    "id": "번호", "created": "접수일", "updated": "고친 날", "customer": "고객", "project": "차종·프로젝트",
    "part_no": "품번", "part_name": "품명", "drawing_rev": "도면 Rev", "annual": "연간 수량", "sop": "양산 시작",
    "eop": "양산 끝", "due": "견적 마감", "status": "상태", "owner": "담당", "source_path": "3D 도면",
    "unit_price": "제출 단가", "die_total": "제출 금형비", "note": "메모",
}
EDIT_FIELDS = ["customer", "project", "part_no", "part_name", "drawing_rev", "annual", "sop", "eop", "due",
               "status", "owner", "source_path", "note"]


def ensure_table(db):
    cols = ", ".join(f"{n} {t}" for n, t in ITEM_COLUMNS)
    db.execute(f"CREATE TABLE IF NOT EXISTS items ({cols})")
    have = {r[1] for r in db.execute("PRAGMA table_info(items)")}
    for name, typ in ITEM_COLUMNS:
        if name not in have:
            db.execute(f"ALTER TABLE items ADD COLUMN {name} {typ}")
    db.commit()


class _ItemsMixin:
    """History 가 물려받는다."""

    def add_item(self, values) -> int:
        now = _dt.datetime.now().isoformat(timespec="seconds")
        v = {k: values.get(k) for k in EDIT_FIELDS if k in values}
        v.setdefault("status", "접수")
        v.update(created=now, updated=now, revisions="[]")
        cur = self.db.execute(f"INSERT INTO items ({', '.join(v)}) VALUES ({', '.join('?' for _ in v)})",
                              tuple(v.values()))
        self.db.commit()
        return cur.lastrowid

    def update_item(self, rid, values):
        v = {k: values[k] for k in EDIT_FIELDS if k in values}
        v["updated"] = _dt.datetime.now().isoformat(timespec="seconds")
        self.db.execute(f"UPDATE items SET {', '.join(k + '=?' for k in v)} WHERE id=?", (*v.values(), rid))
        self.db.commit()

    def get_item(self, rid):
        from .history import Record
        r = self.db.execute("SELECT * FROM items WHERE id=?", (rid,)).fetchone()
        return Record(r["id"], dict(r)) if r else None

    def search_items(self, text="", status="", limit=5000):
        from .history import Record
        sql = "SELECT * FROM items WHERE 1=1"
        args = []
        for word in text.split():
            sql += (" AND (customer LIKE ? OR project LIKE ? OR part_no LIKE ? OR part_name LIKE ? OR note LIKE ?"
                    " OR owner LIKE ?)")
            args += [f"%{word}%"] * 6
        if status:
            sql += " AND status = ?"
            args.append(status)
        sql += " ORDER BY updated DESC, id DESC LIMIT ?"
        args.append(limit)
        return [Record(r["id"], dict(r)) for r in self.db.execute(sql, args)]

    def delete_items(self, ids):
        for rid in ids:
            self.db.execute("DELETE FROM items WHERE id=?", (rid,))
        self.db.commit()

    def add_revision(self, rid, unit_price, die_total, annual=None, note="", files=(), status="제출"):
        """견적을 낼 때마다 리비전 한 줄 (R0, R1 …) 을 쌓고, 마지막 제출 단가·금형비·상태를 바꾼다."""
        it = self.get_item(rid)
        if it is None:
            return None
        revs = json.loads(it.data.get("revisions") or "[]")
        rev = f"R{len(revs)}"
        revs.append({"rev": rev, "date": _dt.date.today().isoformat(), "unit_price": round(unit_price or 0),
                     "die_total": round(die_total or 0), "annual": annual, "note": note, "files": list(files)})
        st = status if it.data.get("status") not in ("수주", "실주") else it.data.get("status")
        self.db.execute("UPDATE items SET revisions=?, unit_price=?, die_total=?, status=?, updated=? WHERE id=?",
                        (json.dumps(revs, ensure_ascii=False), unit_price, die_total, st,
                         _dt.datetime.now().isoformat(timespec="seconds"), rid))
        if annual:
            self.db.execute("UPDATE items SET annual=? WHERE id=?", (annual, rid))
        self.db.commit()
        return rev

    def item_stats(self):
        rows = self.db.execute("SELECT status, COUNT(*) n, SUM(unit_price * annual) amt FROM items GROUP BY status")
        out = {r["status"]: (r["n"], r["amt"] or 0) for r in rows}
        won, lost = out.get("수주", (0, 0))[0], out.get("실주", (0, 0))[0]
        return out, (won / (won + lost) if won + lost else None)

    def export_items_xlsx(self, path, records=None):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        from .company import stamp_workbook
        recs = records if records is not None else self.search_items()
        wb = Workbook()
        ws = wb.active
        ws.title = "아이템 대장"
        cols = [c for c, _ in ITEM_COLUMNS if c != "revisions"]
        for j, c in enumerate(cols, 1):
            cell = ws.cell(1, j, ITEM_LABELS.get(c, c))
            cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E8C")
            ws.column_dimensions[cell.column_letter].width = 14
        for i, r in enumerate(recs, 2):
            for j, c in enumerate(cols, 1):
                cell = ws.cell(i, j, r.data.get(c))
                if c in ("unit_price", "die_total"):
                    cell.number_format = "#,##0"
        w2 = wb.create_sheet("견적 리비전")
        w2.append(["번호", "고객", "품번", "품명", "Rev", "날짜", "세트 단가", "금형비", "연간 수량", "메모", "앞 Rev 대비"])
        for r in recs:
            prev = None
            for rv in json.loads(r.data.get("revisions") or "[]"):
                diff = (rv["unit_price"] / prev - 1) if prev else None
                w2.append([r.id, r.data.get("customer"), r.data.get("part_no"), r.data.get("part_name"), rv["rev"],
                           rv["date"], rv["unit_price"], rv["die_total"], rv.get("annual"), rv.get("note"),
                           f"{diff:+.1%}" if diff is not None else ""])
                prev = rv["unit_price"] or prev
        stamp_workbook(wb)
        wb.save(path)
