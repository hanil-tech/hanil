"""수주 후 품질 문서 초안: 공정 흐름도 (PFD) · PFMEA · 관리계획서 (Control Plan).

자동차 부품 양산 승인 (PPAP, AIAG) 때 내는 문서의 초안을 견적의 공정 (procplan·costing) 으로 만든다.
  · 공정 흐름도: 원자재 입고 → 수입검사 → 보관 → 공정마다 (기호 ○ 가공 · □ 검사 · ▽ 보관 · ⇨ 운반) → 최종 검사 → 포장 → 출하
  · PFMEA: 공정 종류 (절단·피어싱·절곡·드로잉·가공) 마다 잠재 고장 모드·영향·원인·예방·검출과 S·O·D,
           조치 우선순위 AP (AIAG-VDA FMEA 핸드북 2019 의 H·M·L 표를 줄인 규칙) 와 RPN (= S×O×D, 참고)
           성형 해석 (formsim) 결과가 있으면 파단·주름·절곡 균열의 발생도 O 를 그 판정으로 올린다.
  · 관리계획서: 공정마다 제품 특성·공정 특성·규격·측정 방법·시료·주기·관리 방법·대응 계획
문서는 초안이다 — 품질 담당이 고객 요구·실제 검사구에 맞춰 고쳐 쓴다.
"""
from __future__ import annotations

import datetime as _dt
from dataclasses import dataclass, field

from . import costing as cs
from . import tolerances as tl

SYM = {"가공": "○", "검사": "□", "보관": "▽", "운반": "⇨"}


def op_type(op: str, kind: str = "") -> str:
    o = op.upper()
    if kind == cs.MACHINING or "가공" in op:
        return "machining"
    if o.startswith(("DR", "FL", "RST")) or "성형" in op or "드로" in op:
        return "draw"
    if "PI" in o or "피어싱" in op or "피어스" in op:
        return "pierce" if not ("TR" in o or "트림" in op) else "cut"
    if "절곡" in op or "굽힘" in op:
        return "bend"
    return "cut"            # 블랭킹·트림·노칭·외곽·레이저·분리


@dataclass
class FM:                    # PFMEA 한 줄
    no: int
    process: str
    function: str
    mode: str
    effect: str
    s: int
    cause: str
    o: int
    prevent: str
    detect: str
    d: int

    @property
    def rpn(self):
        return self.s * self.o * self.d

    @property
    def ap(self):
        return action_priority(self.s, self.o, self.d)


def action_priority(s, o, d):
    """AIAG-VDA FMEA (2019) 조치 우선순위 — 표를 줄인 규칙 (H 높음 · M 중간 · L 낮음)."""
    if s >= 9:
        if o >= 4 or (o >= 2 and d >= 5):
            return "H"
        return "M" if o >= 2 or d >= 7 else "L"
    if s >= 7:
        if o >= 8 or (o >= 6 and d >= 2) or (o >= 4 and d >= 7):
            return "H"
        if o >= 4 or (o >= 2 and d >= 5):
            return "M"
        return "L"
    if s >= 4:
        if o >= 8 and d >= 5:
            return "H"
        if o >= 6 or (o >= 4 and d >= 7):
            return "M"
        return "L"
    return "M" if o >= 8 and d >= 5 else "L"


# 공정 종류별 고장 모드: (기능, 고장 모드, 영향, S, 원인, O, 예방 관리, 검출 관리, D)
LIBRARY = {
    "incoming": [("재질·두께가 도면대로인 소재 받기", "재질·두께 다름", "강도 부족·성형 불량", 7,
                  "공급사 혼입·성적서 누락", 2, "승인된 공급사·밀시트 확인", "수입검사 두께 측정·성적서 대조", 4)],
    "cut": [("외곽을 도면 치수로 자르기", "버(burr) 과다", "조립 시 상처·손 베임", 5,
             "클리어런스 과다·칼날 마모", 4, "클리어런스 관리, 칼날 연삭 주기 (타수)", "초·중·종물 버 높이 확인", 5),
            ("외곽을 도면 치수로 자르기", "외곽 치수 불량", "조립 불가", 6,
             "소재 이송·위치 불량, 금형 마모", 3, "파일럿·미스피드 감지, 금형 일상 점검", "초·중·종물 검사구", 4)],
    "pierce": [("구멍을 도면 위치·지름으로 뚫기", "구멍 누락", "조립 불가 (체결 안 됨)", 7,
                "펀치 파손·슬러그 올라옴", 3, "펀치 수명 관리, 슬러그 흡착 다이", "구멍 검사구 (Go/No-go) 초·중·종물", 4),
               ("구멍을 도면 위치·지름으로 뚫기", "구멍 위치 불량", "조립 불량", 6,
                "이송 피치 불량·파일럿 마모", 3, "파일럿 핀 관리, 이송 장치 점검", "위치 검사구", 4)],
    "bend": [("도면 각도·플랜지 길이로 굽히기", "각도 불량 (스프링백)", "조립 틈·간섭", 6,
              "소재 강도 편차, 보정각 부족", 4, "스프링백 보정각 (성형 해석), 소재 성적서 강도 확인", "각도기·검사구 초·중·종물", 5),
             ("도면 각도·플랜지 길이로 굽히기", "굽힘부 바깥 면 균열", "강도 부족·파손", 8,
              "굽힘 R 이 재질 최소 R 보다 작음, 압연 방향과 나란한 굽힘", 3, "금형 R ≥ 최소 R/t, 소재 방향 지정",
              "외관 육안 전수", 6)],
    "draw": [("블랭크를 제품 형상으로 성형", "파단 (터짐)", "제품 기능 상실", 8,
              "주름 누르개 힘 과다·윤활 부족·소재 n 값 낮음", 3, "누르개 압력·윤활 관리, 소재 n·r 성적서",
              "외관 육안 전수, 터짐 감지 센서", 5),
             ("블랭크를 제품 형상으로 성형", "주름", "외관·조립 불량", 6,
              "주름 누르개 힘 부족·비드 마모", 3, "누르개 압력 관리, 비드 점검", "외관 육안 전수", 5),
             ("블랭크를 제품 형상으로 성형", "두께 감소 과다", "강도 부족", 7,
              "금형 R 마모·윤활 부족", 3, "금형 R 점검, 윤활 관리", "초·중·종물 두께 측정", 6)],
    "machining": [("도면 치수로 가공", "치수 불량", "조립 불가", 6, "공구 마모·셋업 오류", 3,
                   "공구 수명 관리, 셋업 확인서", "초·중·종물 측정", 4)],
    "final": [("출하 전 제품 확인", "찍힘·녹", "외관 불량·고객 불만", 4, "적재·보관 불량", 3,
               "적재 기준·방청유", "최종 외관 검사", 5)],
}


@dataclass
class PFDStep:
    no: int
    name: str
    kind: str                # 가공·검사·보관·운반
    machine: str = ""
    items: str = ""          # 관리 항목
    note: str = ""


@dataclass
class CPRow:
    no: int
    process: str
    machine: str
    product: str
    process_char: str
    special: str
    spec: str
    method: str
    sample: str
    freq: str
    control: str
    reaction: str


@dataclass
class PartDocs:
    part_no: str
    part_name: str
    customer: str
    material: str
    pfd: list = field(default_factory=list)
    fmea: list = field(default_factory=list)
    cp: list = field(default_factory=list)
    date: _dt.date = field(default_factory=_dt.date.today)


def _occurrence_from_forming(fa):
    """성형 해석 판정 → (파단 O, 주름 O, 절곡 균열 O)."""
    if fa is None:
        return {}
    out = {}
    if fa.mode == "forming" and fa.zone_frac:
        fr = fa.zone_frac
        out["파단 (터짐)"] = 7 if fr[0] > 0.001 else 5 if fr[1] > 0.001 else 2
        out["주름"] = 5 if fr[6] > 0.05 else 3 if fr[6] > 0.01 else 2
        out["두께 감소 과다"] = 6 if fr[2] > 0.001 else 2
    elif fa.bends:
        worst = max(c.risk for c in fa.bends)
        out["굽힘부 바깥 면 균열"] = 7 if worst > 1 else 4 if worst >= 0.8 else 2
    return out


def build(pc, result=None, settings=None, customer="", part_no="", fa=None) -> PartDocs:
    """견적 부품 (quote.PartCost) → 품질 문서 초안."""
    if fa is None and result is not None and settings is not None:
        try:
            from . import formsim
            fa = formsim.analyze(result, settings)
        except Exception:
            fa = None
    occ = _occurrence_from_forming(fa)
    mat = pc.spec
    t = result.thickness if result is not None else 0.0
    docs = PartDocs(part_no or f"P{pc.no:03d}", pc.name, customer, mat)
    pfd = docs.pfd
    no = 10
    pfd.append(PFDStep(no, "원자재 입고", "운반", "", f"{mat} 코일·시트", ""))
    no += 10
    pfd.append(PFDStep(no, "수입검사", "검사", "두께 게이지", "재질 (밀시트) · 두께 · 외관", ""))
    no += 10
    pfd.append(PFDStep(no, "자재 보관", "보관", "자재 창고", "선입선출", ""))
    proc = [(no + 10 * (i + 1), o) for i, o in enumerate(pc.opc)]
    no = proc[-1][0] if proc else no
    seen_press = set()
    for n, o in proc:
        if o.kind == cs.PRESS and o.seconds == 0 and o.machine in seen_press:
            mach = f"{o.machine} (같은 금형)"
        else:
            mach = o.machine
            seen_press.add(o.machine)
        ty = op_type(o.op, o.kind)
        items = {"cut": "외곽 치수 · 버", "pierce": "구멍 지름 · 위치", "bend": "각도 · 플랜지 길이",
                 "draw": "파단 · 주름 · 두께", "machining": "치수"}[ty]
        pfd.append(PFDStep(n, o.op, "가공", mach, items, o.note[:40]))
    no += 10
    pfd.append(PFDStep(no, "최종 검사", "검사", "검사구·3 차원 측정", "전체 치수 · 외관", ""))
    no += 10
    pfd.append(PFDStep(no, "포장", "가공", "", "포장 사양 · 수량", ""))
    no += 10
    pfd.append(PFDStep(no, "출하", "운반", "", "출하 검사 성적서", ""))

    # PFMEA
    def add(step_no, process, ty, extra_note=""):
        for fn, mode, eff, s, cause, o, prev, det, d in LIBRARY[ty]:
            o2 = occ.get(mode, o)
            docs.fmea.append(FM(step_no, process, fn, mode, eff, s, cause + extra_note, o2, prev, det, d))

    add(20, "수입검사", "incoming")
    for n, o in proc:
        add(n, o.op, op_type(o.op, o.kind))
    add(pfd[-3].no, "최종 검사", "final")

    # 관리계획서
    cp = docs.cp
    cp.append(CPRow(20, "수입검사", "두께 게이지", "재질·두께", "-", "", f"{mat} · t {t:g} ± KS 공차" if t else mat,
                    "밀시트 대조·마이크로미터", "3 개", "입고 로트마다", "수입검사 성적서", "불합격 반품·격리"))
    bends = list(result.bends) if result is not None else []
    w, h = (result.size if result is not None else (0, 0))
    for n, o in proc:
        ty = op_type(o.op, o.kind)
        if ty == "cut":
            cp.append(CPRow(n, o.op, o.machine, "외곽 치수", "클리어런스 · 칼날 상태", "",
                            f"전개 {w:,.1f} {tl.text(w)} × {h:,.1f} {tl.text(h)}" if w else "도면",
                            "검사구 · 버니어", "5 개", "초·중·종물", "자주검사 체크시트", "금형 점검·연삭 후 재가동"))
            cp.append(CPRow(n, o.op, o.machine, "버 높이", "칼날 연삭 주기", "",
                            f"≤ {max(0.1 * t, 0.05):.2f} mm" if t else "≤ 10 % t", "버 게이지", "5 개", "초·중·종물",
                            "자주검사 체크시트", "칼날 연삭"))
        elif ty == "pierce":
            cp.append(CPRow(n, o.op, o.machine, "구멍 지름·위치", "펀치 수명 · 파일럿", "◇",
                            "도면 (구멍 공차) — 없으면 구멍 위치 " + tl.text(max(w, h) / 2 if w else 50),
                            "Go/No-go 핀 게이지 · 위치 검사구", "5 개", "초·중·종물",
                            "자주검사 체크시트", "펀치 교체·전수 선별"))
        elif ty == "bend":
            angs = sorted({round(b.angle_deg) for b in bends}) or [90]
            cp.append(CPRow(n, o.op, o.machine, "굽힘 각도", "보정각 · 금형 R", "◇",
                            " / ".join(f"{a}° ± 1°" for a in angs[:4]) + " · 굽힘 치수 "
                            + tl.text(max(w, h) / 2 if w else 50, "bend"), "각도기 · 검사구", "5 개", "초·중·종물",
                            "X-R 관리도 (각도)", "보정각 조정·선별"))
            cp.append(CPRow(n, o.op, o.machine, "굽힘부 균열", "금형 R ≥ 최소 R", "", "균열 없음",
                            "육안", "전수", "전수", "자주검사", "격리·원인 분석"))
        elif ty == "draw":
            tmin = fa.t_min if fa is not None and fa.mode == "forming" else t * 0.75
            cp.append(CPRow(n, o.op, o.machine, "파단·주름", "누르개 압력 · 윤활", "◇", "파단·주름 없음",
                            "육안 (조명)", "전수", "전수", "자주검사 · 한도 견본", "격리·금형 조정"))
            cp.append(CPRow(n, o.op, o.machine, "두께", "금형 R · 윤활", "",
                            f"≥ {min(tmin, t * 0.75):.2f} mm (얇아짐 25 % 이내)" if t else "얇아짐 25 % 이내",
                            "초음파 두께계·마이크로미터", "3 개", "초·중·종물", "두께 측정표", "금형·윤활 조정"))
        else:
            cp.append(CPRow(n, o.op, o.machine, "가공 치수", "공구 수명", "", "도면", "측정기", "5 개", "초·중·종물",
                            "측정 성적서", "공구 교체·재가공"))
    cp.append(CPRow(pfd[-3].no, "최종 검사", "검사구·3 차원", "전체 치수·외관", "-", "◇",
                    "도면 전 치수 (공차 없는 치수는 " + tl.NOTE + ")",
                    "검사구·3 차원 측정기", "5 개", "로트마다", "출하 검사 성적서", "출하 보류·선별"))
    return docs


PFD_HEAD = ["번호", "공정", "가공", "검사", "보관", "운반", "설비·금형", "관리 항목", "비고"]
FMEA_HEAD = ["번호", "공정", "기능", "잠재 고장 모드", "고장 영향", "S", "고장 원인", "O", "예방 관리", "검출 관리",
             "D", "AP", "RPN"]
CP_HEAD = ["번호", "공정", "설비·금형", "제품 특성", "공정 특성", "특별 특성", "규격·공차", "측정 방법", "시료",
           "주기", "관리 방법", "대응 계획"]


def pfd_rows(docs):
    return [[s.no, s.name] + [SYM[k] if s.kind == k else "" for k in ("가공", "검사", "보관", "운반")]
            + [s.machine, s.items, s.note] for s in docs.pfd]


def fmea_rows(docs):
    return [[f.no, f.process, f.function, f.mode, f.effect, f.s, f.cause, f.o, f.prevent, f.detect, f.d, f.ap, f.rpn]
            for f in docs.fmea]


def cp_rows(docs):
    return [[c.no, c.process, c.machine, c.product, c.process_char, c.special, c.spec, c.method, c.sample, c.freq,
             c.control, c.reaction] for c in docs.cp]


# ---------------------------------------------------------------- 저장

def save_xlsx(all_docs, path, title=""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from . import company
    from .company import stamp_workbook
    wb = Workbook()
    hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E8C")
    ap_fill = {"H": PatternFill("solid", fgColor="F4B6B6"), "M": PatternFill("solid", fgColor="FBE3A6"),
               "L": PatternFill("solid", fgColor="CFE8CF")}
    first = True
    for sheet, head, rows_fn, widths in (
            ("공정 흐름도", PFD_HEAD, pfd_rows, [7, 26, 6, 6, 6, 6, 26, 30, 30]),
            ("PFMEA", FMEA_HEAD, fmea_rows, [7, 22, 26, 20, 20, 4, 30, 4, 30, 30, 4, 5, 6]),
            ("관리계획서", CP_HEAD, cp_rows, [7, 22, 18, 16, 20, 8, 30, 24, 8, 12, 18, 20])):
        ws = wb.active if first else wb.create_sheet(sheet)
        ws.title = sheet
        first = False
        r = 1
        for docs in all_docs:
            ws.cell(r, 1, f"{sheet} (초안) — {docs.part_no} {docs.part_name} · {docs.material} · 고객 {docs.customer or '-'} · "
                          f"작성 {docs.date.isoformat()} · {company.MAKER}").font = Font(bold=True, size=12)
            r += 1
            for j, h in enumerate(head, 1):
                c = ws.cell(r, j, h)
                c.font, c.fill = hf, hfill
                c.alignment = Alignment(horizontal="center", wrap_text=True)
            for j, wdt in enumerate(widths, 1):
                ws.column_dimensions[ws.cell(1, j).column_letter].width = wdt
            for row in rows_fn(docs):
                r += 1
                for j, v in enumerate(row, 1):
                    c = ws.cell(r, j, v)
                    c.alignment = Alignment(wrap_text=True, vertical="top",
                                            horizontal="center" if isinstance(v, int) or len(str(v)) <= 2 else "left")
                    if sheet == "PFMEA" and head[j - 1] == "AP":
                        c.fill = ap_fill.get(v, PatternFill())
            r += 3
        if sheet == "PFMEA":
            ws.cell(r, 1, "S 심각도 · O 발생도 · D 검출도 (1~10), AP 조치 우선순위 (AIAG-VDA FMEA 2019, H 는 조치 필수), "
                          "RPN = S×O×D (참고)")
    stamp_workbook(wb)
    wb.save(path)


def save_pdf(all_docs, path, title=""):
    from PySide6.QtCore import QMarginsF, QPointF, QRectF, Qt
    from PySide6.QtGui import QColor, QPageLayout, QPageSize, QPainter, QPdfWriter, QPen

    from . import process as pr
    from .report import _brand, _font, ensure_gui, version_text
    ensure_gui()
    pdf = QPdfWriter(path)
    pdf.setResolution(300)
    pdf.setPageLayout(QPageLayout(QPageSize(QPageSize.A3), QPageLayout.Landscape, QMarginsF(0, 0, 0, 0)))
    pdf.setTitle(f"품질 문서 - {title}")
    pdf.setCreator(f"한일 판금 전개 {version_text()}")
    k = 300 / 25.4
    W, H, m = 420.0, 297.0, 12.0

    def R(x, y, w, h):
        return QRectF(x * k, y * k, w * k, h * k)

    p = QPainter(pdf)
    p.setRenderHint(QPainter.Antialiasing)
    page = [True]

    def header(t):
        if not page[0]:
            pdf.newPage()
        page[0] = False
        p.setPen(QColor(20, 20, 20))
        p.setFont(_font(15, True))
        p.drawText(R(m, m, W - 2 * m, 9), Qt.AlignLeft | Qt.AlignVCenter, t)
        p.setFont(_font(8))
        _brand(p, R(m, m + 1.5, W - 2 * m, 6))
        p.setPen(QPen(QColor(20, 20, 20), 0.4 * k))
        p.drawLine(QPointF(m * k, (m + 10) * k), QPointF((W - m) * k, (m + 10) * k))
        return m + 14

    def cut(rows, n):
        return [[str(v)[:n] for v in r] for r in rows]

    for docs in all_docs:
        tag = f"{docs.part_no} {docs.part_name} · {docs.material} · 고객 {docs.customer or '-'} · {docs.date.isoformat()}"
        y = header(f"공정 흐름도 (초안) — {tag}")
        pr._table(p, R, k, m, y, [(h_, w_) for h_, w_ in zip(PFD_HEAD, (14, 60, 14, 14, 14, 14, 70, 90, 106))],
                  cut(pfd_rows(docs), 60), 9, maxy=H - m - 5)
        y = header(f"PFMEA (초안) — {tag}")
        pr._table(p, R, k, m, y, [(h_, w_) for h_, w_ in zip(FMEA_HEAD, (12, 34, 40, 32, 30, 7, 52, 7, 52, 52, 7, 9, 12))],
                  cut(fmea_rows(docs), 34), 7, maxy=H - m - 5)
        y = header(f"관리계획서 (초안) — {tag}")
        pr._table(p, R, k, m, y, [(h_, w_) for h_, w_ in zip(CP_HEAD, (12, 34, 30, 26, 30, 14, 56, 44, 14, 22, 32, 42))],
                  cut(cp_rows(docs), 32), 7, maxy=H - m - 5)
    p.end()
