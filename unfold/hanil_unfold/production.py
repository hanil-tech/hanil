"""생산 견적: 받은 수량을 우리 설비로 만들 수 있는가, 얼마나 걸리는가.

  · 설비 부하: 기계마다 1 년 필요 시간 = Σ (개수 × 사이클 ÷ 종합효율) + 준비 횟수 × 준비 시간
                1 년 쓸 수 있는 시간 = 가동일 × 교대 × 교대 시간 → 부하율 %, 모자라면 필요 대수
  · 인원: 기계 필요 시간 × 공정 인원 (프레스 1 · 탠덤 라인 공정마다 1 · 절곡 1 · 레이저 0.5 …)
          ÷ 1 인 1 년 근무 시간
  · 사이클타임: 부품마다 가장 느린 공정 (병목) → 하루 생산량
  · 소재: 원자재 소요 kg × 연간 개수 × (1 + 불량률) → 월 톤, 코일 (순차 금형은 스트립 폭) 개수
  · 금형 제작 일정: 설계 → 소재 (강판 1 주 / 주물 6 주) → 가공 → 조립·트라이 → 초도품 → 양산 (주 단위 간트)

설비 부하 계산은 OEE(종합 설비 효율 = 가용률 × 성능 × 양품률) 를 쓰는 일반적인 방법이다.
값은 모두 [생산 견적] 탭에서 회사 값으로 바꾼다.
"""
from __future__ import annotations

import datetime as _dt
import math
from dataclasses import dataclass, field

from . import costing as cs

DEFAULTS = {
    "days": 240,             # 1 년 가동일
    "shifts": 2,             # 교대
    "shift_hours": 8.0,      # 교대 시간
    "oee": 0.75,             # 종합 설비 효율 (가용률 × 성능 × 양품률)
    "defect_pct": 2.0,       # 불량률 % (소재·시간에 더함)
    "lots_per_year": 12,     # 1 년 생산 로트 수 (준비 횟수)
    "crew_press": 1.0,       # 프레스 1 대 인원 (순차·트랜스퍼)
    "crew_tandem": 1.0,      # 탠덤 라인은 공정마다
    "crew_brake": 1.0,
    "crew_laser": 0.5,
    "crew_machining": 1.0,
    "person_hours": 2000.0,  # 1 인 1 년 근무 시간
    "coil_ton": 8.0,         # 코일 1 개 무게 (톤)
    "start_weeks": 1,        # 수주 → 금형 설계 시작 (주)
}

NAMES = {
    "days": "1 년 가동일", "shifts": "교대", "shift_hours": "교대 시간 (h)", "oee": "종합 설비 효율 OEE",
    "defect_pct": "불량률 (%)", "lots_per_year": "1 년 로트 수 (준비 횟수)", "crew_press": "프레스 인원 (대당)",
    "crew_tandem": "탠덤 라인 인원 (공정당)", "crew_brake": "절곡기 인원", "crew_laser": "레이저 인원",
    "crew_machining": "가공기 인원", "person_hours": "1 인 1 년 근무 (h)", "coil_ton": "코일 1 개 (톤)",
    "start_weeks": "수주 → 설계 시작 (주)",
}


def params(rates):
    p = dict(DEFAULTS)
    p.update(rates.get("production") or {})
    return p


@dataclass
class MachineLoad:
    name: str
    kind: str
    hours: float             # 1 년 필요 시간 (가동)
    setup_hours: float
    available: float         # 1 대 1 년 쓸 수 있는 시간
    crew: float              # 필요 인시 / 가동 시간
    parts: list = field(default_factory=list)

    @property
    def total(self):
        return self.hours + self.setup_hours

    @property
    def load(self):
        return self.total / self.available if self.available else 0.0

    @property
    def units(self):
        return max(1, math.ceil(self.load - 1e-9))


@dataclass
class PartFlow:
    no: int
    name: str
    pieces: int              # 1 년 개수 (불량 포함)
    cycle: float             # 병목 공정 초/개
    bottleneck: str
    per_day: float           # 하루 생산량 (가동 시간 × OEE ÷ 사이클)
    gross_kg: float          # 개당 원자재
    material_t: float        # 1 년 원자재 톤
    spec: str
    coil_width: float = 0.0  # 순차 금형 스트립 폭 (mm)
    coils: float = 0.0       # 1 년 코일 개수


@dataclass
class Phase:
    die: str
    name: str
    start: float             # 주 (수주 = 0)
    weeks: float

    @property
    def end(self):
        return self.start + self.weeks


@dataclass
class ProductionPlan:
    annual: int
    p: dict
    machines: list
    flows: list
    phases: list
    notes: list = field(default_factory=list)
    start: _dt.date = field(default_factory=_dt.date.today)

    @property
    def headcount(self):
        return sum(m.total * m.crew for m in self.machines) / max(self.p["person_hours"], 1)

    @property
    def sop_week(self):
        return max((ph.end for ph in self.phases), default=0.0)

    def week_date(self, w):
        return self.start + _dt.timedelta(days=7 * w)

    @property
    def overloaded(self):
        return [m for m in self.machines if m.load > 1.0]


def _crew(p, kind, method, n_ops):
    if kind == cs.PRESS:                     # 탠덤 라인은 공정마다 프레스 한 대 (기계마다 crew_tandem)
        return p["crew_tandem"] if "탠덤" in method or "단발" in method else p["crew_press"]
    return {cs.BRAKE: p["crew_brake"], cs.LASER: p["crew_laser"]}.get(kind, p["crew_machining"])


def _die_phases(d, p, k):
    """금형 하나의 제작 일정 (주). 주물은 모형·주조 6 주, 형면 NC 시간을 주 80 시간 (2 교대) 으로 나눔."""
    n = max(d.stations, 1)
    cast = getattr(d, "structure", "강판") == "주물"
    forming = any(o.startswith(("DR", "FL", "RST")) for o in d.ops)
    start = p["start_weeks"]
    design = (1.0 + 0.4 * n) * (1.5 if cast else 1.0)
    material = 6.0 if cast else 1.0
    nc_h = 0.0
    for name, formula, amt in d.lines:
        if name.startswith("형면 NC"):
            try:
                nc_h = float(formula.split("=")[-1].split("h")[0].replace(",", ""))
            except ValueError:
                nc_h = 0.0
    machining = 1.0 + 0.5 * n + nc_h / 80.0
    tryout = 1.5 + (1.0 if forming else 0.0) + (1.0 if cast else 0.0)
    tag = f"{d.part_name} 금형"
    out = []
    t = start + 0.25 * k                     # 금형이 여럿이면 설계를 조금씩 엇갈려 시작
    for name, w in (("설계", design), ("소재 (주물)" if cast else "소재 (강판·SKD11)", material)):
        out.append(Phase(tag, name, t, w))
        t += w * (0.6 if name == "설계" else 1.0)    # 설계 60 % 에 소재 발주 (겹침)
    out.append(Phase(tag, "가공 (와이어·MCT·NC)", t, machining))
    t += machining
    out.append(Phase(tag, "조립·트라이", t, tryout))
    return out


def plan(q, rates=None) -> ProductionPlan:
    """견적(quote.Quote) → 생산 견적."""
    rates = rates or q.rates
    p = params(rates)
    defect = 1 + p["defect_pct"] / 100
    avail = p["days"] * p["shifts"] * p["shift_hours"]
    machines = {m["name"]: m for m in rates["costing"]["machines"]}
    loads = {}
    flows = []
    for pc in q.parts:
        pieces = int(math.ceil(q.annual * pc.qty * defect))
        slow, slow_name = 0.0, ""
        for o in pc.opc:
            if not o.machine or o.machine == "-":
                continue
            m = machines.get(o.machine, {})
            ml = loads.get(o.machine)
            if ml is None:
                ml = loads[o.machine] = MachineLoad(o.machine, o.kind, 0.0, 0.0, avail, 0.0)
            ml.hours += pieces * o.seconds / 3600 / max(p["oee"], 0.05)
            if pc.name not in ml.parts:          # 준비는 부품마다 로트당 한 번
                ml.parts.append(pc.name)
                ml.setup_hours += p["lots_per_year"] * m.get("setup_min", 0) / 60
            crew = _crew(p, o.kind, pc.method, pc.n_ops)
            ml.crew = max(ml.crew, crew)
            if o.seconds > slow:
                slow, slow_name = o.seconds, f"{o.op} ({o.machine})"
        per_day = (p["shifts"] * p["shift_hours"] * 3600 * p["oee"] / slow) if slow else 0.0
        mt = pc.gross_kg * pieces / 1000
        cw, coils = 0.0, 0.0
        if pc.die is not None and getattr(pc.die, "plan", None) is not None:
            cw = float(pc.die.plan.strip_width)
            coils = mt / max(p["coil_ton"], 0.1)
        flows.append(PartFlow(pc.no, pc.name, pieces, slow, slow_name, per_day, pc.gross_kg, mt, pc.spec, cw, coils))
    phases = []
    for k, d in enumerate(q.dies):
        phases += _die_phases(d, p, k)
    if phases:
        end = max(ph.end for ph in phases)
        phases += [Phase("제품", "초도품 (P1) · 치수 검사", end, 1.0),
                   Phase("제품", "양산 승인 (PPAP) · 공정 검증", end + 1.0, 2.0),
                   Phase("제품", "양산 (SOP)", end + 3.0, 1.0)]
    pl = ProductionPlan(q.annual, p, sorted(loads.values(), key=lambda m: -m.load), flows, phases,
                        start=q.date if isinstance(q.date, _dt.date) else _dt.date.today())
    for m in pl.overloaded:
        pl.notes.append(f"{m.name}: 부하 {m.load:.0%} — {m.units} 대 필요 (교대를 늘리거나 외주·설비 추가)")
    for f in flows:
        if f.material_t > 0 and f.coils:
            pl.notes.append(f"{f.name}: 코일 폭 {f.coil_width:.0f} mm, 1 년 {f.material_t:,.1f} t "
                            f"(코일 {f.coils:,.1f} 개, 월 {f.material_t / 12:,.2f} t)")
    if not pl.overloaded and pl.machines:
        pl.notes.append(f"모든 설비 부하 100 % 이하 — 가장 높은 곳 {pl.machines[0].name} {pl.machines[0].load:.0%}")
    return pl


def machine_rows(pl):
    return [[m.name, m.kind, m.hours, m.setup_hours, m.total, m.available, m.load, m.units, m.crew,
             ", ".join(m.parts)] for m in pl.machines]


MACHINE_HEAD = ["기계", "종류", "가동 h/년", "준비 h/년", "합계 h/년", "1 대 가능 h/년", "부하율", "필요 대수",
                "인원/대", "부품"]


def flow_rows(pl):
    return [[f.no, f.name, f.pieces, f.cycle, f.bottleneck, f.per_day, f.gross_kg, f.material_t, f.spec,
             f.coil_width or "", f.coils or ""] for f in pl.flows]


FLOW_HEAD = ["품번", "품명", "1 년 개수 (불량 포함)", "사이클 초/개", "병목 공정", "하루 생산량", "원자재 kg/개",
             "원자재 t/년", "재질", "코일 폭 mm", "코일 개수/년"]


def phase_rows(pl):
    return [[ph.die, ph.name, f"{ph.start:.1f}", f"{ph.weeks:.1f}", f"{ph.end:.1f}",
             pl.week_date(ph.start).isoformat(), pl.week_date(ph.end).isoformat()] for ph in pl.phases]


PHASE_HEAD = ["금형·제품", "단계", "시작 (주)", "기간 (주)", "끝 (주)", "시작일", "끝날"]


def summary_lines(pl):
    p = pl.p
    out = [f"연간 {pl.annual:,} 세트 · 가동 {p['days']:g} 일 × {p['shifts']:g} 교대 × {p['shift_hours']:g} h · "
           f"OEE {p['oee']:.0%} · 불량 {p['defect_pct']:g} %",
           f"필요 인원 {pl.headcount:,.1f} 명 · 설비 {len(pl.machines)} 대 (부하 넘침 {len(pl.overloaded)} 대)"]
    if pl.phases:
        out.append(f"금형 {len({ph.die for ph in pl.phases if ph.die != '제품'})} 벌 제작 → 양산 {pl.sop_week:.0f} 주 "
                   f"({pl.week_date(pl.sop_week).isoformat()} 무렵)")
    return out + pl.notes


# ---------------------------------------------------------------- 그림 (간트)

def paint_gantt(p, rect, pl, fp=12.0):
    from PySide6.QtCore import QPointF, QRectF, Qt
    from PySide6.QtGui import QColor, QPen

    from .report import _font
    if not pl.phases:
        return
    pt = fp * 72.0 / max(p.device().logicalDpiY(), 1)     # 글씨는 포인트 (PDF 300 dpi · 화면 96 dpi 모두)
    label_w = fp * 22
    top = rect.top() + fp * 2.2
    rows = pl.phases
    rh = min(fp * 2.0, (rect.height() - fp * 3.5) / max(len(rows), 1))
    weeks = max(math.ceil(pl.sop_week), 1)
    x0, x1 = rect.left() + label_w, rect.right() - fp
    sx = (x1 - x0) / weeks
    p.save()
    p.setFont(_font(pt * 0.75))
    for w in range(weeks + 1):
        x = x0 + w * sx
        p.setPen(QPen(QColor(220, 220, 220), 0.6))
        p.drawLine(QPointF(x, top), QPointF(x, top + rh * len(rows)))
        if w % max(1, weeks // 16) == 0:
            p.setPen(QColor(90, 90, 90))
            p.drawText(QRectF(x - fp * 2, top - fp * 1.6, fp * 4, fp * 1.4), Qt.AlignCenter, f"{w}주")
    colors = {"설계": QColor(80, 130, 200), "소재": QColor(160, 120, 70), "가공": QColor(90, 160, 110),
              "조립": QColor(230, 150, 40), "초도": QColor(150, 90, 190), "양산 승인": QColor(200, 80, 120),
              "양산 (": QColor(200, 40, 40)}
    for i, ph in enumerate(rows):
        y = top + i * rh
        if i % 2:
            p.fillRect(QRectF(rect.left(), y, rect.width(), rh), QColor(246, 248, 252))
        p.setPen(QColor(30, 30, 30))
        p.drawText(QRectF(rect.left() + fp * 0.3, y, label_w - fp * 0.6, rh), Qt.AlignLeft | Qt.AlignVCenter,
                   f"{ph.die} · {ph.name}")
        col = next((c for k, c in colors.items() if ph.name.startswith(k)), QColor(120, 120, 120))
        p.setPen(Qt.NoPen)
        p.setBrush(col)
        p.drawRoundedRect(QRectF(x0 + ph.start * sx, y + rh * 0.2, max(ph.weeks * sx, 2), rh * 0.6), 2, 2)
    p.setPen(QPen(QColor(200, 40, 40), 1.4, Qt.DashLine))
    xs = x0 + pl.sop_week * sx
    p.drawLine(QPointF(xs, top - fp * 0.3), QPointF(xs, top + rh * len(rows)))
    p.setFont(_font(pt, True))
    p.setPen(QColor(20, 20, 20))
    p.drawText(QRectF(rect.left(), rect.top(), rect.width(), fp * 1.6), Qt.AlignLeft | Qt.AlignVCenter,
               f"금형 제작 · 양산 일정 (수주 {pl.start.isoformat()} → 양산 {pl.week_date(pl.sop_week).isoformat()} 무렵)")
    p.restore()


def render_gantt(pl, w=1400, h=None):
    from PySide6.QtCore import QRectF
    from PySide6.QtGui import QColor, QImage, QPainter

    from .report import ensure_gui
    ensure_gui()
    h = h or int(60 + 26 * max(len(pl.phases), 1))
    img = QImage(w, h, QImage.Format_ARGB32)
    img.fill(QColor(255, 255, 255))
    p = QPainter(img)
    p.setRenderHint(QPainter.Antialiasing)
    paint_gantt(p, QRectF(8, 8, w - 16, h - 16), pl, 13.0)
    p.end()
    return img


# ---------------------------------------------------------------- 저장

def add_sheets(wb, pl, title="", first=False):
    from openpyxl.styles import Font, PatternFill

    from .report import _add_image
    hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E8C")

    def table(ws, r0, head, rows, pct=()):
        for j, h in enumerate(head, 1):
            c = ws.cell(r0, j, h)
            c.font, c.fill = hf, hfill
            ws.column_dimensions[c.column_letter].width = max(ws.column_dimensions[c.column_letter].width or 0,
                                                              12, len(h) * 2 + 2)
        for i, row in enumerate(rows, r0 + 1):
            for j, v in enumerate(row, 1):
                c = ws.cell(i, j, v)
                if isinstance(v, float):
                    c.number_format = "0.0%" if j - 1 in pct else "#,##0.0"
        return r0 + len(rows) + 2

    ws = wb.active if first else wb.create_sheet("생산 견적")
    ws.title = "생산 견적"
    ws["A1"] = f"생산 견적 — {title}"
    ws["A1"].font = Font(bold=True, size=13)
    r = 2
    for line in summary_lines(pl):
        ws.cell(r, 1, line)
        r += 1
    r += 1
    r = table(ws, r, MACHINE_HEAD, machine_rows(pl), pct=(6,))
    r = table(ws, r, FLOW_HEAD, flow_rows(pl))
    ws2 = wb.create_sheet("금형·양산 일정")
    r2 = table(ws2, 1, PHASE_HEAD, phase_rows(pl))
    tmps = []
    if pl.phases:
        tmps.append(_add_image(ws2, render_gantt(pl), f"A{r2 + 1}", 1100))
    ws3 = wb.create_sheet("생산 조건")
    for k, v in pl.p.items():
        ws3.append([NAMES.get(k, k), v])
    ws3.column_dimensions["A"].width = 30
    return tmps


def save_xlsx(pl, path, title=""):
    import os

    from openpyxl import Workbook
    wb = Workbook()
    tmps = add_sheets(wb, pl, title, first=True)
    from .company import stamp_workbook
    stamp_workbook(wb)
    wb.save(path)
    for t_ in tmps:
        try:
            os.remove(t_)
        except OSError:
            pass


def save_pdf(pl, path, title=""):
    from PySide6.QtCore import QMarginsF, QPointF, QRectF, Qt
    from PySide6.QtGui import QColor, QPageLayout, QPageSize, QPainter, QPdfWriter, QPen

    from . import process as pr
    from .report import _brand, _font, ensure_gui, version_text
    ensure_gui()
    pdf = QPdfWriter(path)
    pdf.setResolution(300)
    pdf.setPageLayout(QPageLayout(QPageSize(QPageSize.A3), QPageLayout.Landscape, QMarginsF(0, 0, 0, 0)))
    pdf.setTitle(f"생산 견적 - {title}")
    pdf.setCreator(f"한일 판금 전개 {version_text()}")
    k = 300 / 25.4
    W, H, m = 420.0, 297.0, 12.0

    def R(x, y, w, h):
        return QRectF(x * k, y * k, w * k, h * k)

    p = QPainter(pdf)
    p.setRenderHint(QPainter.Antialiasing)

    def header(t):
        p.setPen(QColor(20, 20, 20))
        p.setFont(_font(16, True))
        p.drawText(R(m, m, W - 2 * m, 9), Qt.AlignLeft | Qt.AlignVCenter, t)
        p.setFont(_font(8))
        _brand(p, R(m, m + 1.5, W - 2 * m, 6))
        p.setPen(QPen(QColor(20, 20, 20), 0.4 * k))
        p.drawLine(QPointF(m * k, (m + 10) * k), QPointF((W - m) * k, (m + 10) * k))
        return m + 15

    def fmt(rows, pct=()):
        return [[v if isinstance(v, str) else f"{v:.0%}" if j in pct else f"{v:,.1f}" if isinstance(v, float)
                 else f"{v:,}" for j, v in enumerate(r)] for r in rows]

    y = header(f"생산 견적 — {title}")
    p.setFont(_font(10))
    p.setPen(QColor(30, 30, 30))
    for line in summary_lines(pl):
        p.drawText(R(m, y, W - 2 * m, 5.5), Qt.AlignLeft | Qt.AlignVCenter, line)
        y += 5.5
    y += 3
    y = pr._table(p, R, k, m, y, [(h_, w_) for h_, w_ in zip(MACHINE_HEAD, (44, 18, 24, 22, 24, 26, 18, 18, 16, 186))],
                  fmt(machine_rows(pl), (6,)), 8, maxy=H - 60)
    y = pr._table(p, R, k, m, y + 4, [(h_, w_) for h_, w_ in zip(FLOW_HEAD, (14, 50, 34, 24, 70, 26, 26, 24, 34, 24, 24))],
                  fmt(flow_rows(pl)), 8, maxy=H - m - 5)
    if pl.phases:
        pdf.newPage()
        y = header(f"금형 제작 · 양산 일정 — {title}")
        hh = min(H - y - m - 70, 8 + 6 * len(pl.phases))
        pr_rect = R(m, y, W - 2 * m, hh)
        paint_gantt(p, pr_rect, pl, 3.2 * k)
        pr._table(p, R, k, m, y + hh + 4, [(h_, w_) for h_, w_ in zip(PHASE_HEAD, (90, 70, 26, 26, 26, 40, 40))],
                  phase_rows(pl), 8, maxy=H - m - 5)
    p.end()
