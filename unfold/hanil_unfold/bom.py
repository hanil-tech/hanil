"""BOM (부품표): 조립품의 부품을 같은 것끼리 묶고, 생산용 BOM(MBOM) 처럼 계층·분류·원자재를 붙인다.

같은 부품 찾기: 부피·겉넓이·관성 모멘트(돌려도 변하지 않는 값)가 같으면 같은 부품으로 본다.
조립품 안에서 옮기거나 돌려 놓은 같은 부품도 하나로 묶인다.

계층 (다단계 BOM, 정전개)
  0  완제품      조립품 파일 전체
  1+ 반제품      STEP 의 하위 조립품 (용접·조립 단위). 같은 하위 조립품은 한 번만 펼치고 수량으로
     부품        자작 부품·표준품·구매품. 수량은 바로 위(모품목) 1 개당
     원자재      자작 판금·성형품·가공품 밑에 소재 (판재 재질·두께·크기·kg, 환봉·각재)

분류 (품목 유형 · 조달)
  완제품 / 반제품(조립)          자작   공정: 조립·용접
  자작 부품 (판금·성형·가공)     자작   공정: 레이저 절단 → 절곡, 블랭킹 → 프레스 성형, 절단 → 선반·가공
  표준품 (볼트·너트·와셔·핀…)    구매   이름 또는 작은 축대칭 모양으로 판단
  구매품 (베어링·힌지·실린더…)   구매   이름으로 판단
  원자재 (판재·환봉·각재)        구매   단위 kg

구분 (전개 방식, kind)
  판금      정확 전개된 부품
  성형품    곡면이 있어 근사 전개한 부품
  기타      펼칠 수 없거나 두께가 일정하지 않은 부품
"""
from __future__ import annotations

import datetime as _dt
import os
import re
from dataclasses import dataclass

import numpy as np

from . import version_text

KINDS = ("판금", "성형품", "기타")


@dataclass
class BomItem:
    no: int                          # 품번 (1부터)
    name: str
    indices: list                    # 조립품 안 부품 번호들 (0부터)
    kind: str
    thickness: float | None = None
    flat_w: float | None = None
    flat_h: float | None = None
    blank_area: float | None = None  # mm² (개당)
    weight: float = 0.0              # kg (개당, 3D 부피 × 밀도)
    bends: int = 0
    note: str = ""
    result: object = None
    solid: object = None
    excluded: bool = False
    # 분류 (classify)
    category: str = ""               # 자작 부품(판금) / 표준품 / 구매품 …
    procure: str = ""                # 자작 / 구매
    process: str = ""                # 공정 요약
    code: str = ""                   # 품번 (P001 …)
    rm: dict | None = None           # 원자재: spec, code, kg (개당 소요), net_kg, size
    size_text: str = ""              # 판금이 아닌 부품의 크기 (Ø13 × 6.5, 40 × 20 × 10)

    @property
    def qty(self):
        return len(self.indices)

    @property
    def total_weight(self):
        return self.weight * self.qty

    @property
    def numbers(self):
        """조립도 번호 (1부터) 를 줄여 쓴 글: 3, 6, 12~15"""
        return compress([i + 1 for i in self.indices])


def compress(nums):
    nums = sorted(nums)
    out, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(str(nums[i]) if j == i else f"{nums[i]}~{nums[j]}")
        i = j + 1
    return ", ".join(out)


def signature(solid, rel=2e-4):
    """돌리고 옮겨도 같은 값: (부피, 겉넓이, 주 관성 모멘트 3개) 를 반올림한 것."""
    from OCP.BRepGProp import BRepGProp
    from OCP.GProp import GProp_GProps
    g = GProp_GProps()
    BRepGProp.VolumeProperties_s(solid, g)
    vol = abs(g.Mass())
    m = g.MatrixOfInertia()
    M = np.array([[m.Value(i, j) for j in (1, 2, 3)] for i in (1, 2, 3)])
    eig = np.sort(np.abs(np.linalg.eigvalsh(M)))
    gs = GProp_GProps()
    BRepGProp.SurfaceProperties_s(solid, gs)
    area = gs.Mass()

    def q(x, scale):
        step = max(abs(scale) * rel, 1e-9)
        return int(round(x / step))
    return (q(vol, vol), q(area, area)) + tuple(q(e, eig[-1]) for e in eig)


STD_WORDS = re.compile(r"bolt|nut|washer|screw|rivet|\bpin\b|spring ?washer|nipple|stud|"
                       r"볼트|너트|와셔|나사|스크류|리벳|핀|스터드|니플|\bM\d{1,2}\b", re.I)
BUY_WORDS = re.compile(r"bearing|bush|hinge|gas ?spring|damper|cylinder|motor|lamp|light|rubber|seal|"
                       r"grommet|clip|handle|lock|latch|caster|wheel|cap\b|plug|"
                       r"베어링|부싱|부시|힌지|가스 ?스프링|댐퍼|실린더|모터|램프|등|고무|씰|그로멧|클립|"
                       r"손잡이|핸들|잠금|래치|캐스터|바퀴|캡|플러그", re.I)
SMALL_STD_MM = 50.0     # 축대칭이고 가장 긴 쪽이 이보다 작으면 표준품(너트·와셔·핀) 으로 본다


def _axisymmetric(solid):
    """원통·원뿔·원환·구 면이 넓이의 반 넘고 그 축이 모두 나란하면 True (봉·부싱·너트·핀 모양).
    원판·와셔처럼 얇으면 평면이 모두 축에 수직이고 둥근 면의 축이 나란할 때 True."""
    from OCP.BRepAdaptor import BRepAdaptor_Surface
    from OCP.BRepGProp import BRepGProp
    from OCP.GeomAbs import GeomAbs_Cone, GeomAbs_Cylinder, GeomAbs_Plane, GeomAbs_Sphere, GeomAbs_Torus
    from OCP.GProp import GProp_GProps
    from OCP.TopAbs import TopAbs_FACE
    from OCP.TopExp import TopExp_Explorer
    from OCP.TopoDS import TopoDS
    total = rev = 0.0
    axes = []
    normals = []
    e = TopExp_Explorer(solid, TopAbs_FACE)
    while e.More():
        f = TopoDS.Face(e.Current())
        g = GProp_GProps()
        BRepGProp.SurfaceProperties_s(f, g)
        a = g.Mass()
        total += a
        surf = BRepAdaptor_Surface(f)
        t = surf.GetType()
        ax = None
        if t == GeomAbs_Cylinder:
            ax = surf.Cylinder().Axis().Direction()
        elif t == GeomAbs_Cone:
            ax = surf.Cone().Axis().Direction()
        elif t == GeomAbs_Torus:
            ax = surf.Torus().Axis().Direction()
        if t == GeomAbs_Plane:
            nd = surf.Plane().Axis().Direction()
            normals.append(np.array([nd.X(), nd.Y(), nd.Z()]))
        if ax is not None or t == GeomAbs_Sphere:
            rev += a
            if ax is not None:
                axes.append((np.array([ax.X(), ax.Y(), ax.Z()]), a))
        e.Next()
    if total <= 0 or not axes:
        return False
    main = max(axes, key=lambda x: x[1])[0]
    big = [d for d, a in axes if a > 0.05 * rev]
    if not all(abs(float(d @ main)) > 0.99 for d in big):
        return False
    if rev / total >= 0.5:
        return True
    # 와셔·원판처럼 얇은 것: 평면이 모두 축에 수직이면 돌려 만든 모양
    return bool(normals) and all(abs(float(n @ main)) > 0.99 for n in normals)


def _bbox_dims(solid):
    from OCP.Bnd import Bnd_Box
    from OCP.BRepBndLib import BRepBndLib
    b = Bnd_Box()
    BRepBndLib.AddOptimal_s(solid, b, False, False)
    mn, mx = b.CornerMin(), b.CornerMax()
    return sorted([mx.X() - mn.X(), mx.Y() - mn.Y(), mx.Z() - mn.Z()])


def classify(it, material="", density=7.85):
    """품목 유형·조달·공정·원자재를 정한다 (BOM 분류 규칙, 위 머리말 참고)."""
    name = it.name or ""
    try:
        dims = _bbox_dims(it.solid)
        round_ = _axisymmetric(it.solid)
    except Exception:
        dims, round_ = [0.0, 0.0, 0.0], False
    a, b, c = dims
    if round_:
        d, L = (c, a) if abs(c - b) < 0.05 * max(c, 1e-9) else (b, c)
        it.size_text = f"Ø{d:.1f} × {L:.1f}"
    else:
        it.size_text = f"{c:.1f} × {b:.1f} × {a:.1f}"
    if (it.kind == "판금" and not it.bends and round_ and c <= SMALL_STD_MM
            and not BUY_WORDS.search(name)):
        it.category, it.procure, it.process = "표준품 (추정: 와셔)", "구매", "구매 (규격품)"
        return it
    if it.kind in ("판금", "성형품") and not BUY_WORDS.search(name) and not STD_WORDS.search(name):
        it.procure = "자작"
        t = it.thickness or 0.0
        if it.kind == "성형품":
            it.category = "자작 부품 (프레스 성형)"
            it.process = "블랭킹 → 프레스 성형"
        elif it.bends:
            it.category = "자작 부품 (판금)"
            it.process = f"레이저 절단 → 절곡 {it.bends}회"
        else:
            it.category = "자작 부품 (판재)"
            it.process = "레이저 절단"
        if it.flat_w and it.flat_h and t:
            rect_kg = it.flat_w * it.flat_h * t * density / 1e6
            net_kg = (it.blank_area or it.flat_w * it.flat_h) * t * density / 1e6
            it.rm = {"code": f"RM-{material or '판재'}-T{t:g}", "spec": f"{material} 판재 t{t:g}".strip(),
                     "size": f"{it.flat_w:.0f} × {it.flat_h:.0f}", "kg": rect_kg, "net_kg": net_kg,
                     "unit": "kg"}
        return it
    if STD_WORDS.search(name):
        it.category, it.procure, it.process = "표준품", "구매", "구매 (규격품)"
        return it
    if BUY_WORDS.search(name):
        it.category, it.procure, it.process = "구매품", "구매", "구매"
        return it
    if round_ and dims[2] <= SMALL_STD_MM:
        it.category, it.procure, it.process = "표준품 (추정: 너트·와셔·핀)", "구매", "구매 (규격품)"
        return it
    it.procure = "자작"
    if round_:
        it.category = "자작 부품 (선반 가공)"
        it.process = "환봉 절단 → 선반 가공"
        a, b, c = dims                       # 축대칭: 같은 두 값이 지름
        d, L = (c, a) if abs(c - b) < 0.05 * c else (b, c)
        kg = np.pi * (d / 2) ** 2 * L * density / 1e6
        it.rm = {"code": f"RM-{material or '봉'}-RB{d:.0f}", "spec": f"환봉 Ø{d:.0f}",
                 "size": f"Ø{d:.0f} × {L:.0f}", "kg": kg, "net_kg": it.weight, "unit": "kg"}
    else:
        it.category = "자작 부품 (가공)"
        it.process = "절단 → 기계 가공"
        kg = dims[0] * dims[1] * dims[2] * density / 1e6
        it.rm = {"code": f"RM-{material or '강재'}-T{dims[0]:.0f}", "spec": f"강판·각재 t{dims[0]:.0f}",
                 "size": f"{dims[0]:.0f} × {dims[1]:.0f} × {dims[2]:.0f}", "kg": kg, "net_kg": it.weight,
                 "unit": "kg"}
    return it


def _kind(part):
    if not part.ok:
        return "기타"
    r = part.result
    if abs(getattr(r, "coverage", 1.0) - 1.0) > 0.15:
        return "기타"
    return "성형품" if r.mode == "forming" else "판금"


def build(model, parts, density=7.85, excluded=(), material=""):
    """parts: assembly_calc.compute_all 결과. 품번 순서는 조립품에서 처음 나오는 부품 순서.
    같은 부품끼리 묶은 한 단계 목록 (수량 = 완제품 1 대당 전체 수량)."""
    from OCP.BRepGProp import BRepGProp
    from OCP.GProp import GProp_GProps
    groups = {}
    order = []
    for p in parts:
        solid = model.raw_solids[p.index]
        try:
            key = signature(solid)
        except Exception:
            key = ("?", p.index)
        if key not in groups:
            groups[key] = []
            order.append(key)
        groups[key].append(p)
    items = []
    for no, key in enumerate(order, 1):
        ps = groups[key]
        first = ps[0]
        g = GProp_GProps()
        BRepGProp.VolumeProperties_s(model.raw_solids[first.index], g)
        it = BomItem(no=no, name=re.sub(r" #\d+$", "", first.name or ""), indices=[p.index for p in ps], kind=_kind(first),
                     weight=abs(g.Mass()) * density / 1e6, solid=model.raw_solids[first.index],
                     excluded=all(p.index in excluded for p in ps))
        if first.ok:
            r, m = first.result, first.metrics
            it.result = r
            it.thickness = r.thickness
            it.bends = len(r.bends)
            if it.kind != "기타":
                it.flat_w, it.flat_h, it.blank_area = m.width, m.height, m.blank_area
            else:
                it.note = "판금 아님 (가공·구매품)"
        else:
            it.note = first.error
        names = {re.sub(r" #\d+$", "", p.name or "") for p in ps}   # 같은 부품 #2, #3 은 같은 이름
        generic = all(re.fullmatch(r"부품 \d+", n or "") for n in names)
        if len(names) > 1 and not generic:
            it.note = (it.note + " / " if it.note else "") + "이름 다름: " + ", ".join(sorted(names))[:80]
        if it.excluded:
            it.note = ("제외 / " + it.note) if it.note else "제외"
        it.code = f"P{no:03d}"
        classify(it, material, density)
        items.append(it)
    return items


# ---------------------------------------------------------------- 그림

def thumb(item, w=180, h=110):
    """부품 3D 모양 작은 그림 (흰 바탕)."""
    from PySide6.QtCore import QRectF
    from PySide6.QtGui import QColor, QImage, QPainter

    from .drawing import LIGHT, paint_model, solid_polylines
    from .report import ensure_gui
    ensure_gui()
    img = QImage(w, h, QImage.Format_RGB32)
    img.fill(QColor(255, 255, 255))
    try:
        lines = solid_polylines(item.solid, samples=16, iso=False)
    except Exception:
        lines = []
    p = QPainter(img)
    paint_model(p, lines, QRectF(2, 2, w - 4, h - 4), LIGHT, line_width=0.9)
    p.end()
    return img


# ---------------------------------------------------------------- 저장

# ---------------------------------------------------------------- 계층 BOM (MBOM)

@dataclass
class BomRow:
    level: int
    tree_no: str                 # 1, 1.2, 1.2.3 (원자재는 …-R)
    kind: str                    # "assembly" / "part" / "rm"
    name: str
    code: str
    qty: float                   # 바로 위(모품목) 1 개당
    total: float                 # 완제품 1 대당
    unit: str = "EA"
    category: str = ""
    procure: str = ""
    process: str = ""
    spec: str = ""
    item: object = None          # BomItem (부품·원자재 줄)
    note: str = ""


def build_tree(model, parts, density=7.85, excluded=(), material="", items=None):
    """다단계 BOM 줄 목록 (정전개 순서). model.paths 가 있으면 STEP 조립 구조대로 반제품을 만든다."""
    items = items if items is not None else build(model, parts, density, excluded, material)
    by_index = {i: it for it in items for i in it.indices}
    n = len(model.raw_solids)
    paths = getattr(model, "paths", None)
    root_name = os.path.splitext(os.path.basename(getattr(model, "path", "") or "조립품"))[0]
    if not paths or len(paths) != n:
        paths = [[] for _ in range(n)]
    paths = [[tuple(x) for x in p] for p in paths]
    tops = {p[0] for p in paths if p}
    if len(tops) == 1 and all(p for p in paths):
        top = next(iter(tops))               # 맨 위 조립품 하나 = 완제품
        root_name = top[0] or root_name
        paths = [p[1:] for p in paths]

    # 놓인 자리(occurrence) 나무: 자리 → (하위 자리들, 바로 밑 부품 번호들)
    kids, direct = {}, {}
    for i, p in enumerate(paths):
        for d in range(len(p)):
            kids.setdefault(tuple(p[:d]), [])
            if tuple(p[:d + 1]) not in kids[tuple(p[:d])]:
                kids[tuple(p[:d])].append(tuple(p[:d + 1]))
        direct.setdefault(tuple(p), []).append(i)
    sa_codes = {}
    rows = [BomRow(0, "0", "assembly", root_name, "A00", 1, 1, category="완제품", procure="자작",
                   process="조립·용접")]

    def expand(occ, level, prefix, mult):
        k = 0
        groups = {}
        for child in kids.get(occ, []):
            groups.setdefault(child[-1][0], []).append(child)
        for name, occs in groups.items():          # 반제품 (같은 이름 하위 조립품은 한 줄 + 수량)
            k += 1
            no = f"{prefix}{k}" if prefix else str(k)
            code = sa_codes.setdefault(name, f"SA{len(sa_codes) + 1:02d}")
            q = len(occs)
            rows.append(BomRow(level, no, "assembly", name, code, q, q * mult, category="반제품 (조립품)",
                               procure="자작", process="조립·용접"))
            expand(occs[0], level + 1, no + ".", q * mult)
        cnt = {}
        for i in direct.get(occ, []):
            it = by_index.get(i)
            if it is not None:
                cnt.setdefault(it.no, [it, 0])[1] += 1
        for no_, (it, q) in sorted(cnt.items()):
            k += 1
            no = f"{prefix}{k}" if prefix else str(k)
            spec = (f"t{it.thickness:g} · 전개 {it.flat_w:.0f}×{it.flat_h:.0f}"
                    if it.flat_w and it.procure == "자작" and it.kind != "기타" else it.size_text)
            rows.append(BomRow(level, no, "part", it.name, it.code, q, q * mult, "EA", it.category,
                               it.procure, it.process, spec, it, it.note))
            if it.rm and it.procure == "자작":
                rm = it.rm
                rows.append(BomRow(level + 1, no + "-R", "rm", rm["spec"], rm["code"], round(rm["kg"], 3),
                                   round(rm["kg"] * q * mult, 3), rm.get("unit", "kg"), "원자재", "구매",
                                   "", f"{rm['size']} (순중량 {rm['net_kg']:.3f} kg)", it))

    expand((), 1, "", 1)
    return rows


def material_summary(rows):
    """원자재 소요량 합계: (코드, 규격, 부품 종류 수, 합계 kg, 순중량 합계 kg)."""
    out = {}
    for r in rows:
        if r.kind != "rm" or (r.item is not None and r.item.excluded):
            continue
        e = out.setdefault(r.code, [r.name, 0, 0.0, 0.0])
        e[1] += 1
        e[2] += r.total
        e[3] += r.item.rm["net_kg"] * (r.total / r.qty if r.qty else 0)
    return [(c, v[0], v[1], v[2], v[3]) for c, v in sorted(out.items())]


def flat_totals(rows):
    """부품별 완제품 1 대당 합계 수량 (같은 부품이 여러 반제품에 들어가도 합친다)."""
    out = {}
    for r in rows:
        if r.kind == "part":
            e = out.setdefault(r.item.no, [r.item, 0.0])
            e[1] += r.total
    return [v for _, v in sorted(out.items())]


CLASS_RULES = [
    ("레벨 0 완제품", "조립품 파일 전체 (STEP 맨 위 조립품)"),
    ("반제품 (조립품)", "STEP 의 하위 조립품. 같은 이름의 하위 조립품은 한 번만 펼치고 수량으로 표시"),
    ("자작 부품 (판금)", "두께가 일정하고 정확 전개된 부품, 절곡 있음 → 레이저 절단 → 절곡"),
    ("자작 부품 (판재)", "절곡 없는 평판 → 레이저 절단"),
    ("자작 부품 (프레스 성형)", "곡면이 있어 근사 전개한 부품 → 블랭킹 → 프레스 성형"),
    ("자작 부품 (선반 가공)", "판금이 아니고 축대칭(원통 면이 대부분)이며 50 mm 보다 큼 → 환봉 → 선반"),
    ("자작 부품 (가공)", "판금도 축대칭도 아닌 덩어리 → 강판·각재 → 기계 가공"),
    ("표준품", "이름에 볼트·너트·와셔·나사·핀·리벳·M8 등, 또는 축대칭이고 50 mm 이하 (추정). "
               "절곡 없는 작은 둥근 판(링)은 와셔로 봄"),
    ("구매품", "이름에 베어링·부싱·힌지·가스스프링·실린더·램프·고무·클립·손잡이 등"),
    ("원자재", "자작 부품 밑 소재. 판재는 재질·두께별, 소요량 = 전개 사각 크기 × 두께 × 밀도 (kg), "
               "순중량 = 블랭크 면적 기준"),
    ("수량", "모품목 1 개당 수량. '합계' 는 완제품 1 대당 (위 수량을 곱한 값, 정전개)"),
    ("조달", "자작 = 사내 제작 (Make), 구매 = 외부 구매 (Buy)"),
]


def save_tree_xlsx(rows, path, source_name="", material="", density=7.85):
    """계층 BOM 엑셀: BOM(계층) · 부품 합계 · 원자재 소요량 · 구매·표준품 · 분류 기준."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, PatternFill

    from . import company
    from .report import _add_image, _xl_font, _xl_styles

    st = _xl_styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM (계층)"
    ws.cell(1, 1, "BOM (다단계, 생산용)").font = st["title"]
    ws.row_dimensions[1].height = 34
    tmps = [_add_image(ws, company.logo_image(96), "L1", 100)]
    ws.cell(1, 14, company.MAKER).font = _xl_font(st["val"], bold=True)
    ws.cell(1, 14).hyperlink = company.HOMEPAGE
    parts = [r for r in rows if r.kind == "part"]
    info = [("도면", source_name), ("만든 때", f"{_dt.datetime.now():%Y-%m-%d %H:%M}"),
            ("재질 · 밀도", f"{material} · {density:g} g/cm³"),
            ("반제품 / 부품 종류 / 전체 수량",
             f"{sum(1 for r in rows if r.kind == 'assembly') - 1} / {len({r.item.no for r in parts})} / "
             f"{sum(r.total for r in parts):g} 개"),
            ("프로그램", f"한일 판금 전개 {version_text(long=True)}")]
    for i, (k, v) in enumerate(info, 2):
        ws.cell(i, 1, k).font = st["key"]
        ws.cell(i, 3, v).font = st["val"]
    head = ["레벨", "번호", "품번", "품명", "모양", "수량", "단위", "합계", "품목 유형", "조달", "공정",
            "규격", "재질", "개당 무게 (kg)", "조립도 번호", "비고"]
    widths = [6, 10, 16, 34, 16, 7, 6, 8, 22, 7, 22, 26, 10, 12, 16, 30]
    r0 = 2 + len(info) + 1
    for j, (h, w) in enumerate(zip(head, widths), 1):
        c = ws.cell(r0, j, h)
        c.font, c.fill, c.alignment, c.border = st["head"], st["head_fill"], st["center"], st["border"]
        ws.column_dimensions[c.column_letter].width = w
    fills = {"assembly": PatternFill("solid", fgColor="DCE6F2"), "rm": PatternFill("solid", fgColor="F3F3F3")}
    for k, r in enumerate(rows):
        rr = r0 + 1 + k
        it = r.item
        vals = [r.level, r.tree_no, r.code, ("    " * r.level) + r.name, None,
                r.qty, r.unit, r.total, r.category, r.procure, r.process, r.spec,
                (material if r.kind == "rm" or (it is not None and it.kind != "기타") else ""),
                (it.weight if r.kind == "part" else None),
                (it.numbers if r.kind == "part" else ""), r.note]
        for j, v in enumerate(vals, 1):
            c = ws.cell(rr, j, v)
            c.font, c.border = st["val"], st["border"]
            c.alignment = Alignment(vertical="center", wrap_text=j in (4, 11, 12, 15, 16))
            if r.kind in fills:
                c.fill = fills[r.kind]
            if j in (6, 8):
                c.number_format = "0.###"
            if j == 14:
                c.number_format = "0.000"
        if r.kind == "assembly":
            for j in (1, 2, 3, 4):
                ws.cell(rr, j).font = _xl_font(st["val"], bold=True)
        if r.kind == "part":
            ws.row_dimensions[rr].height = 46
            tmps.append(_add_image(ws, thumb(it, 160, 96), f"E{rr}", 100))
            if it.excluded:
                for j in range(1, len(head) + 1):
                    ws.cell(rr, j).font = _xl_font(st["val"], color="999999")
        if r.level > 0:          # 엑셀 왼쪽 + / − 로 반제품을 접고 펼친다
            ws.row_dimensions[rr].outline_level = min(r.level, 7)
    ws.freeze_panes = ws.cell(r0 + 1, 5)
    ws.sheet_properties.outlinePr.summaryBelow = False
    # 인쇄: 가로, 폭을 한 쪽에 맞추고 머리줄 반복
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{r0}:{r0}"

    # 부품 합계 (완제품 1 대당)
    wf = wb.create_sheet("부품 합계")
    head2 = ["품번", "품명", "품목 유형", "조달", "합계 수량", "개당 무게 (kg)", "합계 무게 (kg)", "공정", "조립도 번호"]
    for j, h in enumerate(head2, 1):
        c = wf.cell(1, j, h)
        c.font, c.fill, c.alignment, c.border = st["head"], st["head_fill"], st["center"], st["border"]
        wf.column_dimensions[c.column_letter].width = [9, 34, 22, 7, 10, 12, 12, 22, 18][j - 1]
    tot = 0.0
    for k, (it, q) in enumerate(flat_totals(rows), 2):
        vals = [it.code, it.name, it.category, it.procure, q, it.weight, it.weight * q, it.process, it.numbers]
        for j, v in enumerate(vals, 1):
            c = wf.cell(k, j, v)
            c.font, c.border = st["val"], st["border"]
            if j in (6, 7):
                c.number_format = "0.000"
        if not it.excluded:
            tot += it.weight * q
    last = wf.max_row + 1
    wf.cell(last, 2, "합계 무게 (제외 빼고)").font = _xl_font(st["val"], bold=True)
    wf.cell(last, 7, round(tot, 3)).font = _xl_font(st["val"], bold=True)

    # 원자재 소요량
    wm = wb.create_sheet("원자재 소요량")
    head3 = ["원자재 코드", "규격", "쓰는 부품 종류", "소요량 합계 (kg)", "순중량 합계 (kg)", "수율"]
    for j, h in enumerate(head3, 1):
        c = wm.cell(1, j, h)
        c.font, c.fill, c.alignment, c.border = st["head"], st["head_fill"], st["center"], st["border"]
        wm.column_dimensions[c.column_letter].width = [20, 22, 14, 16, 16, 9][j - 1]
    for k, (code, spec, n, kg, net) in enumerate(material_summary(rows), 2):
        for j, v in enumerate([code, spec, n, kg, net, net / kg if kg else None], 1):
            c = wm.cell(k, j, v)
            c.font, c.border = st["val"], st["border"]
            c.number_format = ["@", "@", "0", "0.000", "0.000", "0.0%"][j - 1]
    wm.cell(wm.max_row + 2, 1, "소요량 = 전개 사각 크기(가로×세로) × 두께 × 밀도. 철판 배치를 하면 실제 소요가 더 정확합니다.").font = st["key"]

    # 구매·표준품
    wbuy = wb.create_sheet("구매·표준품")
    head4 = ["품번", "품명", "품목 유형", "합계 수량", "추정 규격 (바깥 크기 mm)", "조립도 번호"]
    for j, h in enumerate(head4, 1):
        c = wbuy.cell(1, j, h)
        c.font, c.fill, c.alignment, c.border = st["head"], st["head_fill"], st["center"], st["border"]
        wbuy.column_dimensions[c.column_letter].width = [9, 34, 26, 10, 24, 18][j - 1]
    k = 2
    for it, q in flat_totals(rows):
        if it.procure != "구매":
            continue
        try:
            dims = " × ".join(f"{d:.1f}" for d in _bbox_dims(it.solid))
        except Exception:
            dims = ""
        for j, v in enumerate([it.code, it.name, it.category, q, dims, it.numbers], 1):
            c = wbuy.cell(k, j, v)
            c.font, c.border = st["val"], st["border"]
        k += 1

    # 분류 기준
    wr = wb.create_sheet("분류 기준")
    wr.column_dimensions["A"].width = 26
    wr.column_dimensions["B"].width = 100
    for k, (a, b) in enumerate(CLASS_RULES, 1):
        wr.cell(k, 1, a).font = _xl_font(st["val"], bold=True)
        wr.cell(k, 2, b).alignment = Alignment(wrap_text=True)
    from .company import stamp_workbook
    stamp_workbook(wb)
    wb.save(path)
    for t in tmps:
        try:
            os.remove(t)
        except OSError:
            pass


def save_tree_pdf(rows, path, source_name="", material="", density=7.85):
    """계층 BOM PDF (A4 가로). 반제품 줄은 파란 바탕, 원자재 줄은 회색."""
    from PySide6.QtCore import QMarginsF, QPointF, QRectF, Qt
    from PySide6.QtGui import QColor, QPageLayout, QPageSize, QPainter, QPdfWriter, QPen

    from .report import _brand, _font, ensure_gui

    ensure_gui()
    pdf = QPdfWriter(path)
    pdf.setResolution(300)
    pdf.setPageLayout(QPageLayout(QPageSize(QPageSize.A4), QPageLayout.Landscape, QMarginsF(0, 0, 0, 0)))
    pdf.setTitle(f"BOM - {source_name}")
    pdf.setCreator(f"한일 판금 전개 {version_text()}")
    k = 300 / 25.4
    page_w, page_h, margin = 297.0, 210.0, 10.0

    def R(x, y, w, h):
        return QRectF(x * k, y * k, w * k, h * k)

    p = QPainter(pdf)
    gray = QColor(110, 110, 110)
    cols = [("번호", 15), ("품번", 25), ("모양", 18), ("품명", 58), ("수량", 12), ("단위", 9), ("합계", 13),
            ("품목 유형", 33), ("조달", 10), ("공정", 34), ("규격", 50)]
    part_h, row_h, head_h = 11.0, 6.0, 6.0

    def header(page):
        y = margin
        p.setPen(QColor(20, 20, 20))
        p.setFont(_font(15, True))
        p.drawText(R(margin, y, 160, 9), Qt.AlignLeft | Qt.AlignVCenter, f"BOM (다단계)  {source_name}")
        bw = _brand(p, R(margin, y + 1.5, page_w - 2 * margin, 6)) / k
        p.setFont(_font(7.5))
        p.setPen(gray)
        p.drawText(R(margin, y, page_w - 2 * margin - bw - 4, 9), Qt.AlignRight | Qt.AlignVCenter,
                   f"재질 {material} · {_dt.datetime.now():%Y-%m-%d} · {page} 쪽")
        p.setPen(QPen(QColor(20, 20, 20), 0.4 * k))
        p.drawLine(QPointF(margin * k, (y + 10) * k), QPointF((page_w - margin) * k, (y + 10) * k))
        y += 13
        p.fillRect(R(margin, y, page_w - 2 * margin, head_h), QColor(40, 60, 90))
        p.setPen(QColor(255, 255, 255))
        p.setFont(_font(7.5, True))
        x = margin
        for name, w in cols:
            p.drawText(R(x, y, w, head_h), Qt.AlignCenter, name)
            x += w
        return y + head_h

    page = 1
    y = header(page)
    for r in rows:
        h = part_h if r.kind == "part" else row_h
        if y + h > page_h - margin:
            pdf.newPage()
            page += 1
            y = header(page)
        if r.kind == "assembly":
            p.fillRect(R(margin, y, page_w - 2 * margin, h), QColor(220, 230, 242))
        elif r.kind == "rm":
            p.fillRect(R(margin, y, page_w - 2 * margin, h), QColor(244, 244, 244))
        vals = [r.tree_no, r.code, None, ("   " * r.level) + r.name, f"{r.qty:g}", r.unit, f"{r.total:g}",
                r.category, r.procure, r.process, r.spec]
        x = margin
        for (name, w), v in zip(cols, vals):
            if name == "모양":
                if r.kind == "part":
                    p.drawImage(R(x + 0.5, y + 0.5, w - 1, h - 1), thumb(r.item, 200, 120))
            else:
                excluded = r.item is not None and r.item.excluded
                p.setPen(QColor(150, 150, 150) if excluded else QColor(20, 20, 20))
                p.setFont(_font(7, r.kind == "assembly"))
                flags = (Qt.AlignLeft if name in ("품명", "공정", "규격", "품목 유형") else Qt.AlignHCenter) \
                    | Qt.AlignVCenter | Qt.TextWordWrap
                p.drawText(R(x + 1, y, w - 2, h), flags, v)
            x += w
        p.setPen(QPen(QColor(215, 215, 215), 0.12 * k))
        p.drawLine(QPointF(margin * k, (y + h) * k), QPointF((page_w - margin) * k, (y + h) * k))
        y += h
    p.end()


HEADER = ["품번", "부품명", "모양", "수량", "구분", "재질", "두께 (mm)", "전개 크기 (mm)",
          "블랭크 면적 (m²)", "개당 무게 (kg)", "합계 무게 (kg)", "절곡", "조립도 번호", "비고"]


def _row(it, material):
    size = f"{it.flat_w:.1f} × {it.flat_h:.1f}" if it.flat_w is not None else ""
    return [it.no, it.name, None, it.qty, it.kind, material if it.kind != "기타" else "",
            round(it.thickness, 2) if it.thickness else None, size,
            it.blank_area / 1e6 if it.blank_area is not None else None,
            it.weight, it.total_weight, it.bends or None, it.numbers, it.note]


def save_xlsx(items, path, source_name="", material="", density=7.85):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment

    from . import company
    from .report import _add_image, _xl_font, _xl_styles

    st = _xl_styles()
    wb = Workbook()
    ws = wb.active
    ws.title = "BOM"
    ws.cell(1, 1, "BOM (부품표)").font = st["title"]
    ws.row_dimensions[1].height = 34
    tmps = [_add_image(ws, company.logo_image(96), "K1", 100)]
    ws.cell(1, 12, company.MAKER).font = _xl_font(st["val"], bold=True)
    ws.cell(1, 12).hyperlink = company.HOMEPAGE
    info = [("도면", source_name), ("만든 때", f"{_dt.datetime.now():%Y-%m-%d %H:%M}"),
            ("재질 · 밀도", f"{material} · {density:g} g/cm³"),
            ("품목 / 전체 수량", f"{len(items)} 품목 / {sum(i.qty for i in items)} 개"),
            ("프로그램", f"한일 판금 전개 {version_text(long=True)}")]
    for i, (k, v) in enumerate(info, 2):
        ws.cell(i, 1, k).font = st["key"]
        ws.cell(i, 2, v).font = st["val"]
    row0 = 2 + len(info) + 1
    widths = [7, 30, 16, 7, 9, 10, 10, 18, 13, 12, 12, 7, 22, 40]
    for j, (hname, wdt) in enumerate(zip(HEADER, widths), 1):
        c = ws.cell(row0, j, hname)
        c.font, c.fill, c.alignment, c.border = st["head"], st["head_fill"], st["center"], st["border"]
        ws.column_dimensions[c.column_letter].width = wdt
    fmt = [None, None, None, "0", None, None, "0.0#", None, "0.0000", "0.000", "0.000", "0", None, None]
    tot_qty, tot_w = 0, 0.0
    for k, it in enumerate(items):
        r = row0 + 1 + k
        for j, v in enumerate(_row(it, material), 1):
            c = ws.cell(r, j, v)
            c.font, c.border = st["val"], st["border"]
            c.alignment = Alignment(vertical="center", wrap_text=j in (2, 13, 14))
            if k % 2:
                c.fill = st["zebra"]
            if fmt[j - 1]:
                c.number_format = fmt[j - 1]
            if it.excluded:
                c.font = _xl_font(st["val"], color="999999")
        ws.row_dimensions[r].height = 60
        tmps.append(_add_image(ws, thumb(it), f"C{r}", 112))
        if not it.excluded:
            tot_qty += it.qty
            tot_w += it.total_weight
    r = row0 + 1 + len(items)
    ws.cell(r, 2, "합계 (제외 빼고)").font = _xl_font(st["val"], bold=True)
    ws.cell(r, 4, tot_qty).font = _xl_font(st["val"], bold=True)
    ws.cell(r, 11, round(tot_w, 3)).font = _xl_font(st["val"], bold=True)
    ws.cell(r, 11).number_format = "0.000"
    ws.freeze_panes = ws.cell(row0 + 1, 3)
    ws.print_title_rows = f"{row0}:{row0}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # 두께별 소요량
    wm = wb.create_sheet("두께별 소요량")
    names = ["재질", "두께 (mm)", "수량", "블랭크 면적 합계 (m²)", "블랭크 무게 합계 (kg)"]
    for j, hname in enumerate(names, 1):
        c = wm.cell(1, j, hname)
        c.font, c.fill, c.alignment, c.border = st["head"], st["head_fill"], st["center"], st["border"]
        wm.column_dimensions[c.column_letter].width = 22
    rows = {}
    for it in items:
        if it.excluded or it.kind == "기타" or it.blank_area is None:
            continue
        e = rows.setdefault(round(it.thickness, 2), [0, 0.0])
        e[0] += it.qty
        e[1] += it.blank_area * it.qty
    for i, (t, (n, a)) in enumerate(sorted(rows.items()), 2):
        vals = [material, t, n, a / 1e6, a * t * density / 1e6]
        for j, v in enumerate(vals, 1):
            c = wm.cell(i, j, v)
            c.font, c.border = st["val"], st["border"]
            c.number_format = ["@", "0.0#", "0", "0.0000", "0.000"][j - 1]
    wm.cell(len(rows) + 3, 1, "판금·성형품만 (기타·제외 빼고). 블랭크 무게 = 블랭크 면적 × 두께 × 밀도.").font = st["key"]
    from .company import stamp_workbook
    stamp_workbook(wb)
    wb.save(path)
    for t in tmps:
        try:
            os.remove(t)
        except OSError:
            pass


def save_pdf(items, path, source_name="", material="", density=7.85):
    """A4 가로, 부품 그림이 든 BOM 표 (쪽마다 머리글)."""
    from PySide6.QtCore import QMarginsF, QPointF, QRectF, Qt
    from PySide6.QtGui import QColor, QPageLayout, QPageSize, QPainter, QPdfWriter, QPen

    from .report import _brand, _font, ensure_gui

    ensure_gui()
    pdf = QPdfWriter(path)
    pdf.setResolution(300)
    pdf.setPageLayout(QPageLayout(QPageSize(QPageSize.A4), QPageLayout.Landscape, QMarginsF(0, 0, 0, 0)))
    pdf.setTitle(f"BOM - {source_name}")
    pdf.setCreator(f"한일 판금 전개 {version_text()}")
    k = 300 / 25.4
    page_w, page_h, margin = 297.0, 210.0, 10.0

    def R(x, y, w, h):
        return QRectF(x * k, y * k, w * k, h * k)

    p = QPainter(pdf)
    gray = QColor(110, 110, 110)
    cols = [("품번", 10), ("모양", 26), ("부품명", 52), ("수량", 11), ("구분", 14), ("두께", 12),
            ("전개 크기 (mm)", 32), ("블랭크 (m²)", 20), ("개당 kg", 17), ("합계 kg", 17),
            ("조립도 번호", 33), ("비고", 33)]
    row_h, head_h = 15.0, 6.0
    total_qty = sum(i.qty for i in items if not i.excluded)
    total_w = sum(i.total_weight for i in items if not i.excluded)

    def header(page):
        y = margin
        p.setPen(QColor(20, 20, 20))
        p.setFont(_font(15, True))
        p.drawText(R(margin, y, 150, 9), Qt.AlignLeft | Qt.AlignVCenter, f"BOM  {source_name}")
        bw = _brand(p, R(margin, y + 1.5, page_w - 2 * margin, 6)) / k
        p.setFont(_font(7.5))
        p.setPen(gray)
        p.drawText(R(margin, y, page_w - 2 * margin - bw - 4, 9), Qt.AlignRight | Qt.AlignVCenter,
                   f"{len(items)} 품목 · {total_qty} 개 · {total_w:.2f} kg · 재질 {material}  ·  "
                   f"{_dt.datetime.now():%Y-%m-%d}  ·  {page} 쪽")
        p.setPen(QPen(QColor(20, 20, 20), 0.4 * k))
        p.drawLine(QPointF(margin * k, (y + 10) * k), QPointF((page_w - margin) * k, (y + 10) * k))
        y += 13
        p.fillRect(R(margin, y, page_w - 2 * margin, head_h), QColor(40, 60, 90))
        p.setPen(QColor(255, 255, 255))
        p.setFont(_font(7.5, True))
        x = margin
        for name, w in cols:
            p.drawText(R(x, y, w, head_h), Qt.AlignCenter, name)
            x += w
        return y + head_h

    page = 1
    y = header(page)
    p.setFont(_font(7.5))
    for n, it in enumerate(items):
        if y + row_h > page_h - margin:
            pdf.newPage()
            page += 1
            y = header(page)
            p.setFont(_font(7.5))
        if n % 2:
            p.fillRect(R(margin, y, page_w - 2 * margin, row_h), QColor(244, 246, 249))
        size = f"{it.flat_w:.1f} × {it.flat_h:.1f}" if it.flat_w is not None else ""
        vals = [str(it.no), None, it.name, str(it.qty), it.kind,
                f"{it.thickness:.2f}" if it.thickness else "", size,
                f"{it.blank_area / 1e6:.4f}" if it.blank_area is not None else "",
                f"{it.weight:.3f}", f"{it.total_weight:.3f}", it.numbers, it.note]
        x = margin
        for (name, w), v in zip(cols, vals):
            if name == "모양":
                img = thumb(it, 260, 150)
                p.drawImage(R(x + 0.5, y + 0.5, w - 1, row_h - 1), img)
            else:
                p.setPen(QColor(150, 150, 150) if it.excluded else QColor(20, 20, 20))
                flags = (Qt.AlignLeft if name in ("부품명", "조립도 번호", "비고") else Qt.AlignHCenter) \
                    | Qt.AlignVCenter | Qt.TextWordWrap
                p.drawText(R(x + 1, y, w - 2, row_h), flags, v)
            x += w
        p.setPen(QPen(QColor(210, 210, 210), 0.15 * k))
        p.drawLine(QPointF(margin * k, (y + row_h) * k), QPointF((page_w - margin) * k, (y + row_h) * k))
        y += row_h
    p.end()
