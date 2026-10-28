"""공정별 금형 (탠덤 · 단발 여러 대): 1차 블랭킹 → 2차 피어싱 → 3차~ 절곡 (공정마다 금형 한 벌) → 후 피어싱 · 캠 피어싱.

1) 접기 (Folder): 전개도를 절곡 띠로 나눈 판 (패널) 들의 나무 — 가장 큰 판이 바탕, 절곡마다 바탕에서 먼 쪽이 돈다.
   공정마다 접은 각도 {절곡: 각} 로 두께 가운데 면의 3D 모양 (외곽 · 구멍 · 절곡선) 을 낸다.
   절곡 띠 (펼친 폭 BA) 는 가운데 면 반지름 BA / θ 의 호로, 그 너머는 강체로 돌린다.
2) 공정 짜기 (plan): procplan 규칙 (R1~R8) 을 공정마다 금형 한 벌로 —
   바깥 (잎) 플랜지부터 접고, 같은 판 양쪽의 나란한 같은 방향 플랜지는 U 굽힘 한 공정,
   95° 넘으면 90° 와이핑 + 마무리 (기울인 펀치), 150° 넘으면 90° → 예비 헤밍 (135°) → 헤밍 (눌러 접기).
3) 금형 (design): 블랭킹 · 피어싱은 die3d (피어싱은 블랭크 외곽 네스트 핀으로 자리), 절곡은 bendie.
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field

import numpy as np

from . import bendcalc as bc
from . import diecalc as dc

BLANK, PIERCE, BEND, UBEND, ANGLE, FINISH, PREHEM, HEM, POST, CAM = (
    "블랭킹", "피어싱", "절곡 (와이핑)", "U 굽힘", "각도 굽힘", "마무리 굽힘", "예비 헤밍", "헤밍", "후 피어싱", "캠 피어싱")
BEND_KINDS = (BEND, UBEND, ANGLE, FINISH, PREHEM, HEM)


# ---------------------------------------------------------------- 접기

@dataclass
class FBend:
    number: int
    a: np.ndarray            # 절곡선 한 끝 (전개 좌표)
    l: np.ndarray            # 선 방향 단위
    n: np.ndarray            # 도는 쪽 (바탕에서 먼 판 쪽) 단위 법선
    length: float
    half: float              # 펼친 폭의 반 (BA / 2)
    angle: float             # 최종 각도 °
    up: bool                 # UP: 전개도를 보는 쪽 (+z) 으로 세움
    r: float                 # 안쪽 R
    near: int = -1           # 바탕 쪽 판
    far: int = -1            # 도는 쪽 판
    depth: int = 0           # 바탕에서 몇 번째 절곡 (1 = 바탕에 붙음)

    def u(self, P):
        return (np.asarray(P, float) - self.a) @ self.n

    def v(self, P):
        return (np.asarray(P, float) - self.a) @ self.l


def _rigid(b: FBend, theta_deg):
    """도는 쪽 판의 강체 변환 (4×4, 전개 좌표 → 3D). 절곡 띠 시작 (u = −h) 에서 반지름 BA / θ최종 의 호를 θ 만큼,
    남은 띠는 곧게 (덜 접은 공정: 띠 일부만 굽음)."""
    th = math.radians(theta_deg)
    M = np.eye(4)
    if abs(th) < 1e-9:
        return M
    s = 1.0 if b.up else -1.0
    h = b.half
    rm = 2 * h / math.radians(max(b.angle, theta_deg, 1e-6))
    rest = max(2 * h - rm * th, 0.0)
    en, ez = np.array([b.n[0], b.n[1], 0.0]), np.array([0.0, 0.0, 1.0])
    el = np.array([b.l[0], b.l[1], 0.0])
    a3 = np.array([b.a[0], b.a[1], 0.0])
    start = a3 - h * en
    B = a3 + h * en
    A = start + rm * (math.sin(th) * en + s * (1 - math.cos(th)) * ez) + rest * (math.cos(th) * en + s * math.sin(th) * ez)
    # 회전: en → cos en + s sin ez,  ez → −s sin en + cos ez,  el 그대로
    R_loc = np.array([[math.cos(th), -s * math.sin(th), 0.0], [s * math.sin(th), math.cos(th), 0.0], [0.0, 0.0, 1.0]])
    Bm = np.column_stack([en, ez, el])                   # 국소 (n, z, l) → 전역
    R = Bm @ R_loc @ Bm.T
    M[:3, :3] = R
    M[:3, 3] = A - R @ B
    return M


def _zone(b: FBend, theta_deg, P2):
    """절곡 띠 안 점 (전개 좌표, z = 0) → 3D (호 위)."""
    P2 = np.asarray(P2, float)
    th = math.radians(theta_deg)
    s = 1.0 if b.up else -1.0
    h = b.half
    u = b.u(P2)
    v = b.v(P2)
    out = np.zeros((len(P2), 3))
    base = b.a[None, :] + v[:, None] * b.l[None, :]
    if abs(th) < 1e-9:
        out[:, :2] = P2
        return out
    rm = 2 * h / math.radians(max(b.angle, theta_deg, 1e-6))
    sv = np.clip(u + h, 0.0, 2 * h)
    phi = np.minimum(sv / rm, th)
    extra = np.maximum(sv - rm * th, 0.0)                  # 덜 접은 공정: 호 끝에서 곧게
    d_n = -h + rm * np.sin(phi) + extra * math.cos(th)
    d_z = s * (rm * (1 - np.cos(phi)) + extra * math.sin(th))
    out[:, :2] = base + d_n[:, None] * b.n[None, :]
    out[:, 2] = d_z
    return out


@dataclass
class Panel:
    index: int
    pts: np.ndarray          # 판 안 격자 점 (전개 좌표)
    area: float
    chain: list = field(default_factory=list)   # 바탕에서 이 판까지 절곡 번호 (가까운 것부터)


class Folder:
    """전개도 → 판 나무 · 단계별 3D."""

    def __init__(self, result):
        from scipy import ndimage

        from .nesting import NestPart, _rasterize
        part = NestPart.from_result("p", result)
        if part.outer is None:
            raise ValueError("전개 외곽이 닫혀 있지 않습니다")
        self.result = result
        self.t = float(result.thickness)
        self.outer = np.asarray(part.outer, float)
        self.holes = [np.asarray(h, float) for h in part.holes]
        w, h = np.ptp(self.outer, 0)
        res = max(max(w, h) / 500.0, 0.2)
        self.res = res
        m, (ox, oy) = _rasterize(part, 0.0, res, conservative=False)
        iy, ix = np.nonzero(m)
        P = np.column_stack([ox + (ix + 0.5) * res, oy + (iy + 0.5) * res])
        self.bends = []
        for b in result.bends:
            a, c = np.array(b.line[0], float), np.array(b.line[1], float)
            d = c - a
            L = float(np.linalg.norm(d)) or 1.0
            l_ = d / L
            n_ = np.array([-l_[1], l_[0]])
            self.bends.append(FBend(b.number, a, l_, n_, L, b.allowance / 2, float(b.angle_deg),
                                    b.direction == "UP", float(b.inner_radius)))
        # 절곡 띠 지움 → 판 (이어진 칸)
        cut = np.zeros(m.shape, bool)
        for fb in self.bends:
            u, v = fb.u(P), fb.v(P)
            sel = (np.abs(u) <= fb.half + 1.5 * res) & (v >= -2 * res) & (v <= fb.length + 2 * res)
            cut[iy[sel], ix[sel]] = True
        lab, nlab = ndimage.label(m & ~cut)
        self._lab, self._org = lab, (ox, oy)
        panels = []
        for k in range(1, nlab + 1):
            yy, xx = np.nonzero(lab == k)
            if len(yy) < 6:
                continue
            pts = np.column_stack([ox + (xx + 0.5) * res, oy + (yy + 0.5) * res])
            panels.append(Panel(k, pts, len(yy) * res * res))
        if not panels:
            raise ValueError("판을 나누지 못했습니다")
        self.panels = {p.index: p for p in panels}
        # 절곡마다 양쪽 판
        for fb in self.bends:
            mid = fb.a + 0.5 * fb.length * fb.l
            sides = []
            for sg in (-1, 1):
                sides.append(self._label_near(mid + sg * (fb.half + 3 * res) * fb.n, fb.n * sg))
            fb._sides = sides                                         # [−n 쪽, +n 쪽]
        base = max(panels, key=lambda p: p.area).index
        self.base = base
        # 나무 (너비 우선)
        seen = {base}
        frontier = [base]
        self.panels[base].chain = []
        while frontier:
            nxt = []
            for pi in frontier:
                for fb in self.bends:
                    if fb.far != -1:
                        continue
                    a_, b_ = fb._sides
                    if pi in (a_, b_):
                        other = b_ if pi == a_ else a_
                        if other in seen or other < 0:
                            continue
                        if other == b_:
                            pass                                   # +n 쪽이 돈다
                        else:
                            fb.n = -fb.n                           # 도는 쪽이 −n → 법선 뒤집기
                            fb.a = fb.a + fb.length * fb.l         # 선 방향도 뒤집어 오른손 유지 (l × n 같게)
                            fb.l = -fb.l
                        fb.near, fb.far = pi, other
                        fb.depth = len(self.panels[pi].chain) + 1
                        self.panels[other].chain = self.panels[pi].chain + [fb.number]
                        seen.add(other)
                        nxt.append(other)
            frontier = nxt
        self.bmap = {fb.number: fb for fb in self.bends}
        self.loose = [fb.number for fb in self.bends if fb.far == -1]   # 판을 못 가른 절곡 (부분 절곡 등)

    def _label_near(self, q, direction):
        ox, oy = self._org
        res = self.res
        for k in range(0, 12):
            p = q + direction * k * res
            ix, iy = int((p[0] - ox) / res), int((p[1] - oy) / res)
            if 0 <= iy < self._lab.shape[0] and 0 <= ix < self._lab.shape[1] and self._lab[iy, ix]:
                return int(self._lab[iy, ix])
        return -1

    # ---- 점이 어느 판 / 띠
    def locate(self, P):
        """[(종류, 번호)] — ("panel", 판) 또는 ("zone", 절곡 번호)."""
        P = np.asarray(P, float)
        out = []
        ox, oy = self._org
        res = self.res
        for p in P:
            z = None
            for fb in self.bends:
                if fb.far == -1:
                    continue
                u, v = fb.u(p), fb.v(p)
                if abs(u) < fb.half and -res <= v <= fb.length + res:
                    z = fb.number
                    break
            if z is not None:
                out.append(("zone", z))
                continue
            ix, iy = int((p[0] - ox) / res), int((p[1] - oy) / res)
            lab = 0
            for r in range(0, 6):
                ys, xs = slice(max(iy - r, 0), iy + r + 1), slice(max(ix - r, 0), ix + r + 1)
                win = self._lab[ys, xs]
                vals = win[win > 0]
                if vals.size:
                    lab = int(np.bincount(vals).argmax())
                    break
            out.append(("panel", lab if lab in self.panels else self.base))
        return out

    def hole_panel(self, i):
        """구멍 i 가 있는 판 (구멍 둘레 점들의 많은 쪽, 절곡 띠 안이면 그 띠의 바탕 쪽 판)."""
        votes = {}
        for kind, k in self.locate(self.holes[i]):
            pi = k if kind == "panel" else self.bmap[k].near
            votes[pi] = votes.get(pi, 0) + 1
        return max(votes, key=votes.get) if votes else self.base

    def panel_matrix(self, pi, done):
        M = np.eye(4)
        for num in self.panels[pi].chain:                         # 바탕 쪽부터 곱함 (멀리 있는 것이 먼저 적용)
            M = M @ _rigid(self.bmap[num], done.get(num, 0.0))
        return M

    def fold(self, P2, done, where=None):
        """전개 좌표 점들 → 접은 3D (done = {절곡 번호: 각도})."""
        P2 = np.asarray(P2, float)
        where = where or self.locate(P2)
        out = np.zeros((len(P2), 3))
        cache = {}
        for i, (kind, k) in enumerate(where):
            if kind == "panel":
                if k not in cache:
                    cache[k] = self.panel_matrix(k, done)
                M = cache[k]
                out[i] = M[:3, :3] @ np.array([P2[i, 0], P2[i, 1], 0.0]) + M[:3, 3]
            else:
                fb = self.bmap[k]
                key = ("z", fb.near)
                if key not in cache:
                    cache[key] = self.panel_matrix(fb.near, done)
                M = cache[key]
                q = _zone(fb, done.get(k, 0.0), P2[i:i + 1])[0]
                out[i] = M[:3, :3] @ q + M[:3, 3]
        return out

    def dense(self, P, step=None):
        """닫힌 고리를 촘촘히 (띠를 지나는 변이 호가 되게)."""
        step = step or max(self.res, 0.5)
        P = np.asarray(P, float)
        Q = np.vstack([P, P[:1]])
        out = []
        for a, b in zip(Q[:-1], Q[1:]):
            k = max(1, int(np.linalg.norm(b - a) / step))
            out.append(a + np.outer(np.arange(k) / k, b - a))
        return np.vstack(out)

    def wire(self, done, holes=None):
        """접은 모양 선 (3D 꺾은선 목록): 외곽 · 구멍 (holes: 번호 목록, None = 모두) · 절곡 띠 양 끝선."""
        lines = []
        hs = self.holes if holes is None else [self.holes[i] for i in holes]
        for loop in [self.outer] + hs:
            D = self.dense(loop)
            F = self.fold(D, done)
            lines.append(np.vstack([F, F[:1]]))
        for fb in self.bends:
            if fb.far == -1:
                continue
            for uu in (-fb.half, fb.half):
                seg = np.array([fb.a + uu * fb.n + s_ * fb.length * fb.l for s_ in (0.0, 1.0)])
                seg = seg + np.sign(uu) * 1e-6 * fb.n
                lines.append(self.fold(seg, done, where=[("zone", fb.number)] * 2 if abs(uu) < fb.half
                                       else self.locate(seg)))
        return lines

    def panel_poly(self, pi):
        """판 외곽 (전개 좌표, 격자 점의 볼록 껍질 — 다이 블록 자리 어림)."""
        from scipy.spatial import ConvexHull
        pts = self.panels[pi].pts
        try:
            hull = ConvexHull(pts)
            return pts[hull.vertices]
        except Exception:  # noqa: BLE001
            lo, hi = pts.min(0), pts.max(0)
            return np.array([lo, [hi[0], lo[1]], hi, [lo[0], hi[1]]])

    def subtree(self, num):
        """절곡 num 너머 (도는 쪽) 판들."""
        return [pi for pi, p in self.panels.items() if num in p.chain]


# ---------------------------------------------------------------- 공정 짜기

@dataclass
class Op:
    no: int
    kind: str
    title: str
    detail: str
    bends: list = field(default_factory=list)     # 이 공정에서 접는 절곡 번호
    before: dict = field(default_factory=dict)    # 공정 전 {절곡: 각}
    after: dict = field(default_factory=dict)     # 공정 뒤
    holes: list = field(default_factory=list)     # 뚫는 구멍 (Folder.holes 번호)
    panel: int = -1                               # 금형에 눕히는 판 (절곡: 바탕 쪽 판, 피어싱: 구멍 판)
    rules: list = field(default_factory=list)
    force_n: float = 0.0
    notes: list = field(default_factory=list)
    die: object = None
    error: str = ""
    layout: object = None                         # 1차 블랭킹 고른 배치 (Layout)
    layouts: list = field(default_factory=list)   # 배치 후보 (재료 적은 순)

    @property
    def label(self):
        return f"{self.no}차 {self.title}"


WIPE_MIN = 75.0           # 이 각도부터 와이핑 (수직 펀치), 아래는 각도 굽힘 (기울인 다이 면에 눌러 붙임)
WIPE_MAX = 95.0           # 넘으면 90° 와이핑 + 마무리
HEM_FROM = 150.0          # 넘으면 헤밍


def plan(result, settings: bc.BendSettings, folder: Folder | None = None):
    """공정 목록 [Op] — 공정마다 금형 한 벌."""
    from . import procplan
    if getattr(result, "forming", None) is not None or getattr(result, "rolled", None):
        raise ValueError("성형품 (원스텝) · 말아 만든 부품은 절곡 공정별 금형 대상이 아닙니다 — 드로잉 금형은 [공정별] 성형 해석으로")
    F = folder or Folder(result)
    t = F.t
    mat = settings.mat
    tau = dc.shear_strength(mat)
    feat = procplan.features(result, settings)
    info = {h.index: h for h in feat.holes}
    ops = []

    def per(P):
        P = np.asarray(P, float)
        return float(np.sum(np.linalg.norm(np.roll(P, -1, 0) - P, axis=1)))
    ops.append(Op(0, BLANK, "블랭킹", f"외곽 {per(F.outer):,.0f} mm 따내기 (코일 · 스트립에서)",
                  force_n=per(F.outer) * t * tau, rules=["R6"]))
    # 구멍이 있는 판 (Folder 기준)
    hole_panel = {i: F.hole_panel(i) for i in range(len(F.holes))}
    pre = [i for i in range(len(F.holes)) if not info[i].after_bend]
    groups = sorted({info[i].group for i in pre})
    for g in groups:
        ids = [i for i in pre if info[i].group == g]
        ops.append(Op(0, PIERCE, "피어싱" if len(groups) == 1 else f"피어싱 {g + 1}",
                      ", ".join(info[i].label for i in ids) + " — 블랭크 외곽 네스트로 자리"
                      + (" (가까운 구멍 나눔)" if g else ""),
                      holes=ids, force_n=sum(per(F.holes[i]) for i in ids) * t * tau,
                      rules=["R2"] + (["R5"] if g else [])))
    # 절곡: 잎 (바깥) 부터 — 아직 안 접은 자손이 없는 절곡만 고를 수 있음
    done = {}
    todo = [fb for fb in F.bends if fb.far != -1]
    area = {pi: p.area for pi, p in F.panels.items()}

    def sub_area(fb):
        return sum(area[pi] for pi in F.subtree(fb.number))

    def ready(fb):
        return all(o.number in done for o in todo if o is not fb and fb.number in F.panels[o.far].chain)

    def step(fb, ang):
        return {fb.number: ang}
    while len([fb for fb in todo if done.get(fb.number) != fb.angle]) > 0:
        cand = [fb for fb in todo if fb.number not in done and ready(fb)]
        if not cand:
            break
        cand.sort(key=lambda fb: (-fb.depth, sub_area(fb)))
        fb = cand[0]
        grp = [fb]
        if fb.angle <= WIPE_MAX:
            for o in cand[1:]:
                if (o.near == fb.near and abs(fb.l[0] * o.l[1] - fb.l[1] * o.l[0]) < 0.02 and fb.n @ o.n < -0.9
                        and o.up == fb.up and abs(o.angle - fb.angle) <= 5 and o.angle <= WIPE_MAX
                        and WIPE_MIN <= o.angle and WIPE_MIN <= fb.angle):
                    grp.append(o)
                    break
        names = " · ".join(f"B{g.number} {g.angle:.0f}° {'UP' if g.up else 'DOWN'}" for g in grp)
        before = dict(done)
        if len(grp) == 2:
            for g in grp:
                done[g.number] = g.angle
            ops.append(Op(0, UBEND, "U 굽힘", f"{names} — 양쪽 와이핑 (패드로 바닥 누름)", [g.number for g in grp],
                          before, dict(done), panel=fb.near, rules=["R7"]))
            continue
        if fb.angle < WIPE_MIN:
            done[fb.number] = fb.angle
            ops.append(Op(0, ANGLE, "각도 굽힘", f"{names} — 기울인 다이 면에 눌러 붙임 (바텀)", [fb.number], before,
                          dict(done), panel=fb.near, rules=["R7"]))
            continue
        done[fb.number] = min(fb.angle, 90.0)
        ops.append(Op(0, BEND, "절곡", f"{names}" + (" — 1단 90°" if fb.angle > WIPE_MAX else "") + " — 와이핑 (패드 + 펀치)",
                      [fb.number], before, dict(done), panel=fb.near, rules=["R7"] + (["R8"] if fb.angle > WIPE_MAX else [])))
        if fb.angle > HEM_FROM:
            b2 = dict(done)
            done[fb.number] = 135.0
            ops.append(Op(0, PREHEM, "예비 헤밍", f"B{fb.number} 90° → 135° (기울인 펀치)", [fb.number], b2, dict(done),
                          panel=fb.near, rules=["R8"]))
            b3 = dict(done)
            done[fb.number] = fb.angle
            ops.append(Op(0, HEM, "헤밍", f"B{fb.number} 135° → {fb.angle:.0f}° (눌러 접기)", [fb.number], b3, dict(done),
                          panel=fb.near, rules=["R8"]))
        elif fb.angle > WIPE_MAX:
            b2 = dict(done)
            done[fb.number] = fb.angle
            ops.append(Op(0, FINISH, "마무리 굽힘", f"B{fb.number} 90° → {fb.angle:.0f}° (기울인 펀치로 넘김)", [fb.number], b2,
                          dict(done), panel=fb.near, rules=["R8"]))
    final = dict(done)
    post = [i for i in range(len(F.holes)) if info[i].after_bend]
    flat = [i for i in post if hole_panel[i] == F.base]
    side = [i for i in post if hole_panel[i] != F.base]
    if flat:
        ops.append(Op(0, POST, "후 피어싱", ", ".join(info[i].label for i in flat) + " (절곡선 가까움 — 절곡 뒤)",
                      before=final, after=final, holes=flat, panel=F.base,
                      force_n=sum(per(F.holes[i]) for i in flat) * t * tau, rules=["R3"]))
    by_panel = {}
    for i in side:
        by_panel.setdefault(hole_panel[i], []).append(i)
    # 플랜지마다 (그 판을 세운 절곡), 같은 판 양쪽 맞은편 플랜지 (동축 구멍 R4) 는 한 금형에 캠 둘
    groups_c = []
    for pi, ids in by_panel.items():
        bn = F.panels[pi].chain[-1] if F.panels[pi].chain else None
        if bn is None:
            continue
        fb = F.bmap[bn]
        for g in groups_c:
            o = F.bmap[g[0][0]]
            if len(g) == 1 and o.near == fb.near and o.n @ fb.n < -0.9:
                g.append((bn, ids))
                break
        else:
            groups_c.append([(bn, ids)])
    for g in groups_c:
        ids = [i for _, hs in g for i in hs]
        ops.append(Op(0, CAM, "캠 피어싱", procplan.pair_names([info[i] for i in ids]) + " (세운 플랜지 — 옆으로"
                      + (", 맞은편 둘 한 번에)" if len(g) == 2 else ")"),
                      bends=[bn for bn, _ in g], before=final, after=final, holes=ids, panel=F.bmap[g[0][0]].near,
                      force_n=sum(per(F.holes[i]) for i in ids) * t * tau,
                      rules=sorted({info[i].reason.split(",")[0] for i in ids})))
    for i, op in enumerate(ops, 1):
        op.no = i
    for op in ops:
        if op.kind in (BLANK, PIERCE):
            op.before = {}
            op.after = {}
    return ops


# ---------------------------------------------------------------- 금형 (공정마다 한 벌)

@dataclass
class Layout:
    """1차 블랭킹 스트립 배치 후보: 돌린 각도 · 열 수 · 피치 · 스트립 폭 · 재료 이용률."""
    name: str
    angle: float
    rows: int
    pitch: float
    width: float
    util: float
    blanks: list = field(default_factory=list)     # 한 피치 안 블랭크들 (돌린 좌표, 스트립 가운데 y = 0)
    kg_per_part: float = 0.0

    @property
    def strip_per_part(self):
        return self.pitch * self.width / self.rows


TWO_ROW_GAIN = 0.03     # 2열 배치를 고르려면 1열보다 재료가 이만큼 (3 %) 넘게 줄어야 (금형 값 · 프레스 힘 두 배)


def _clear(P, Q, gap):
    """두 닫힌 외곽이 겹치지 않고 (엇갈림 · 안에 들어감 없음) gap 이상 떨어짐."""
    from . import diecalc as dc
    from .die3d import _inside
    from .unfoldcheck import crossings
    if crossings([P, Q]) > 0 or _inside(P, Q[0]) or _inside(Q, P[0]):
        return False
    return dc._poly_dist(P, Q) >= gap


def blank_layouts(result, t, bridge=None, side=None):
    """1차 블랭킹 배치 후보 — 가로 (0°) · 세로 (90°) · 가장 좋은 각 (5° 씩) 을 1열로, 그리고 가로 · 세로를 2열 (둘째 줄은
    180° 뒤집어 반 피치 어긋나게) 로. 재료비 = 한 개가 차지하는 스트립 넓이 × t × 밀도 — 가장 적은 것이 앞."""
    from . import diecalc as dc
    from .die3d import DENSITY
    from .nesting import NestPart, _rot_poly
    part = NestPart.from_result("p", result)
    bridge = max(1.5 * t, 1.0) if bridge is None else bridge
    side = max(2.0 * t, 1.5) if side is None else side
    area = part.area

    def rot(angle):
        O = np.asarray(_rot_poly(part.outer, angle), float)
        if np.linalg.norm(O[0] - O[-1]) < 1e-9:
            O = O[:-1]
        return O - 0.5 * (O.min(0) + O.max(0))                  # 상자 가운데 (뒤집어도 같은 자리)

    def one_row(angle, name):
        O = rot(angle)
        wd = float(np.ptp(O[:, 1])) + 2 * side
        lo, hi = float(np.ptp(O[:, 0])) * 0.3, float(np.ptp(O[:, 0])) + bridge   # 피치: 옆 부품과 bridge 띄움
        for _ in range(18):
            mid = 0.5 * (lo + hi)
            if dc._poly_dist(O, O + (mid, 0)) >= bridge:
                hi = mid
            else:
                lo = mid
        p = math.ceil(hi * 2) / 2
        return Layout(name, angle, 1, p, wd, area / (p * wd), [O])

    def two_row(angle, name):
        O = rot(angle)
        p = one_row(angle, name).pitch
        Q = -O + (p / 2, 0.0)                                   # 180° 뒤집어 반 피치 어긋나게
        gap2 = bridge + 1.0                                     # 줄 사이: 다이 날 사이 벽 (클리어런스 몫 더함)
        # 둘째 줄 y 띄움: gap2 이상 (이웃 피치의 것과도)
        dy_lo, dy_hi = 0.0, float(np.ptp(O[:, 1])) + gap2

        def ok(dy):
            Q2 = Q + (0.0, dy)
            return all(_clear(O + (k * p, 0.0), Q2, gap2) for k in (-1, 0, 1))
        for _ in range(18):
            mid = 0.5 * (dy_lo + dy_hi)
            if ok(mid):
                dy_hi = mid
            else:
                dy_lo = mid
        Q = Q + (0.0, dy_hi)
        allp = np.vstack([O, Q])
        yc = 0.5 * (allp[:, 1].min() + allp[:, 1].max())
        wd = float(np.ptp(allp[:, 1])) + 2 * side
        return Layout(name + " 2열", angle, 2, p, wd, 2 * area / (p * wd), [O - (0, yc), Q - (0, yc)])
    strip = dc.strip_layout(result, part, bridge, side)
    cands = [one_row(0.0, "가로"), one_row(90.0, "세로"), two_row(0.0, "가로"), two_row(90.0, "세로")]
    if min(abs(strip.angle - a_) for a_ in (0.0, 90.0, 180.0)) > 2.5:
        cands.append(one_row(strip.angle, f"기울임 {strip.angle:.0f}°"))
    for c in cands:
        c.kg_per_part = c.strip_per_part * t * DENSITY
    cands.sort(key=lambda c: (c.strip_per_part, c.rows))
    # 2열은 금형이 크고 (펀치 둘 · 힘 두 배) 비싸니 재료가 TWO_ROW_GAIN 넘게 줄 때만 앞에 — 아니면 1열 중 가장 좋은 것
    best1 = min((c for c in cands if c.rows == 1), key=lambda c: c.strip_per_part)
    best2 = min((c for c in cands if c.rows == 2), key=lambda c: c.strip_per_part)
    if best2.strip_per_part > (1 - TWO_ROW_GAIN) * best1.strip_per_part:
        cands = [c for c in cands if c.rows == 1] + [c for c in cands if c.rows == 2]
    return cands


def _cut_die(F, op, result, settings, tol, machines, layout=None):
    """1차 블랭킹 (스트립 — layout 배치) · 피어싱 (블랭크 네스트) — die3d 전단 금형."""
    from . import die3d as d3
    t = F.t
    lo, hi = F.outer.min(0), F.outer.max(0)
    c0 = 0.5 * (lo + hi)
    blank = F.outer - c0
    holes = [h - c0 for h in F.holes]
    L, W = hi - lo
    if op.kind == BLANK:
        if layout is None:
            layout = blank_layouts(result, t)[0]
        ops = [(B, "blank") for B in layout.blanks]
        blank = layout.blanks[0]
        holes = []
        xs = np.vstack(layout.blanks)[:, 0]
        work, strip, nest = (float(np.ptp(xs)), layout.width), (layout.pitch, layout.width), None
    else:
        ops = [(holes[i], "pierce") for i in op.holes]
        work, strip, nest = (L, W), (0.0, W), blank
    cop = d3._cop([P for P, _ in ops])
    sh = -np.asarray(cop, float)
    ops = [(P + sh, k) for P, k in ops]
    blank_s = blank + sh
    half = max(float(np.abs(np.vstack([P for P, _ in ops] + [blank_s])[:, 0]).max()), work[0] / 2)
    halfy = max(float(np.abs(np.vstack([P for P, _ in ops] + [blank_s])[:, 1]).max()), work[1] / 2)
    work = (2 * half, 2 * halfy)
    for k in range(8):
        try:
            a = d3._assemble(result, settings, op.label, None, t, ops, [], work, strip, blank_s,
                             [h + sh for h in holes], extra=10.0 * k, tol=tol, machines=machines,
                             nest=nest if op.kind == PIERCE else None)
            if op.kind != BLANK:
                a.stations = [(op.label, 0.0, blank_s)]
            else:                                            # 배치도에 이웃 피치 (앞 · 뒤) 까지 — 배치가 보이게
                a.stations = []
                for k, (B, _) in enumerate(ops):
                    for j in (-1, 0, 1):
                        nm = (f"{op.label}" + (f" ({k + 1}열)" if len(ops) > 1 else "")) if j == 0 else ""
                        a.stations.append((nm, float(B[:, 0].mean()) + j * layout.pitch, B + (j * layout.pitch, 0.0)))
                a.layout = layout
            return a
        except d3._NoRoom:
            continue
    raise ValueError("표준 부품 자리를 못 찾았습니다")


def _post_die(F, op, result, settings, tol, machines):
    """후 피어싱: 다 접은 제품을 바탕 판 수평 (플랜지가 아래로 많은 쪽) 으로 다이에 씌움 — 다이 쌓기 판을 제품 안쪽에 맞춤."""
    from . import bendie
    from . import die3d as d3
    t = F.t
    done = op.after
    # 바탕 판 좌표 (바탕은 접지 않으므로 전개 좌표 그대로, z 위 +)
    P3 = np.vstack([F.fold(F.dense(lp), done) for lp in [F.outer] + F.holes])
    flip = (P3[:, 2] < -0.3).sum() < (P3[:, 2] > 0.3).sum()
    if flip:
        P3 = P3 * np.array([1, -1, -1])
    base = F.panels[F.base].pts
    if flip:
        base = base * np.array([1, -1])
    holes = [h * (np.array([1, -1]) if flip else 1) for h in F.holes]
    if flip:
        holes = [h[::-1] for h in holes]
    # 다이 자리: 바탕 판 안, 아래로 내려간 점에서 t + 2 띄움
    rect = [float(base[:, 0].min()), float(base[:, 0].max()), float(base[:, 1].min()), float(base[:, 1].max())]
    hs = np.vstack([holes[i] for i in op.holes])
    keep = [float(hs[:, 0].min()) - 5, float(hs[:, 0].max()) + 5, float(hs[:, 1].min()) - 5, float(hs[:, 1].max()) + 5]
    obst = [tuple(p[:2]) for p in P3 if p[2] < -0.3]
    rect, left = bendie._rect_fit(rect, obst, keep, t + 2.0)
    c0 = np.array([0.5 * (rect[0] + rect[1]), 0.5 * (rect[2] + rect[3])])
    ops = [(holes[i] - c0, "pierce") for i in op.holes]
    depth = float(max(-P3[:, 2].min(), 0.0)) + t
    up_h = float(max(P3[:, 2].max() - t / 2, 0.0))
    ext = np.ptp(P3[:, :2], 0)
    a = None
    for k in range(6):
        try:
            a = d3._assemble(result, settings, op.label, None, t, ops, [], (rect[1] - rect[0], rect[3] - rect[2]),
                             (0.0, rect[3] - rect[2]), F.outer - c0, [h - c0 for h in holes], extra=0.0, tol=tol,
                             machines=machines, nest=None, fit=(rect[1] - rect[0], rect[3] - rect[2]),
                             outer=(float(ext[0]) + 10 * k, float(ext[1]) + 10 * k), depth_below=depth, up_h=up_h)
            break
        except d3._NoRoom:
            continue
    if a is None:
        raise ValueError("후 피어싱 금형: 제품 안쪽에 다이 자리가 모자랍니다")
    a.checks.append(["다이 ↔ 아래로 접힌 플랜지", f"{left} 점", "0", "적합" if left == 0 else "검토"])
    # 제품 선 (다이 윗면 = 제품 안쪽 면)
    try:
        die_top = a.plate("다이 플레이트").prims[0]
        z_top = die_top.p[2] + die_top.p[5]
    except StopIteration:
        z_top = 0.0
    work = []
    for lp in [F.outer] + F.holes:
        Q = F.fold(F.dense(lp), done)
        if flip:
            Q = Q * np.array([1, -1, -1])
        Q = Q - np.array([c0[0], c0[1], 0.0]) + np.array([0, 0, z_top + t / 2])
        work.append((Q, "before"))
    a.work = work
    a.notes.insert(0, f"{op.label}: {op.detail} — 제품을 다이에 씌워 자리 (플랜지가 다이 옆을 감쌈)")
    return a


def design(result, settings: bc.BendSettings, tol="m", machines=None, folder=None, ops=None):
    """공정마다 금형 한 벌 — [Op] (op.die = DieAssy, 못 만들면 op.error)."""
    from . import bendie
    F = folder or Folder(result)
    ops = ops if ops is not None else plan(result, settings, F)
    for op in ops:
        try:
            if op.kind == BLANK:
                lays = blank_layouts(result, F.t)
                op.layouts = lays
                op.die = None
                for lay in lays:                              # 재료 적게 드는 배치부터, 금형이 되는 것
                    try:
                        op.die = _cut_die(F, op, result, settings, tol, machines, layout=lay)
                        op.layout = lay
                        break
                    except Exception as e:  # noqa: BLE001
                        op.error = str(e)
                if op.die is not None:
                    op.error = ""
                    lay = op.layout
                    op.detail = (f"{lay.name} 배치 ({lay.angle:.0f}°, {lay.rows}열) — 피치 {lay.pitch:.1f} × 폭 {lay.width:.1f}, "
                                 f"이용률 {lay.util * 100:.1f} %, 한 개 {lay.kg_per_part * 1000:.1f} g")
                    op.die.layout_note = ("배치 비교 (한 개 재료): " + " · ".join(
                        f"{c.name} {c.util * 100:.1f} % {c.kg_per_part * 1000:.0f} g" for c in lays)
                        + f" → {lay.name} (2열은 재료가 {TWO_ROW_GAIN * 100:.0f} % 넘게 줄 때만 — 금형 · 프레스 힘 두 배)")
                    op.die.notes.insert(1, op.die.layout_note)
            elif op.kind == PIERCE:
                op.die = _cut_die(F, op, result, settings, tol, machines)
            elif op.kind == POST:
                op.die = _post_die(F, op, result, settings, tol, machines)
            else:
                op.die = bendie.design(F, op, result, settings, machines)
            if op.die is not None:
                op.die.kind = op.label
                if op.kind in BEND_KINDS + (CAM,):
                    op.force_n = op.die.force_n
        except Exception as e:  # noqa: BLE001 — 한 공정이 안 되어도 나머지는 계속
            op.error = str(e)
    return ops


# ---------------------------------------------------------------- 공정표 · 저장

def stage_holes(ops, k):
    """k 번째 공정 (0 부터) 까지 뚫은 구멍."""
    out = []
    for op in ops[:k + 1]:
        if op.kind in (PIERCE, POST, CAM):
            out += op.holes
    return out


def line_sheets(F, ops, info, prefix, rates=None):
    """공정표: 공정마다 칸 (공정 뒤 제품 입체 · 금형 크기 · 하중 · 프레스 · 점검 · 금형비) — A3 한 장에 6 칸, 아래 표."""
    from . import diecost
    from . import diedraw as dd
    from . import dieview
    per = 6
    sheets = []
    pages = [ops[i:i + per] for i in range(0, len(ops), per)] or [[]]
    for pg, chunk in enumerate(pages):
        sh = dd.Sheet("공정별 금형 공정표" + (f" ({pg + 1})" if len(pages) > 1 else ""), f"{prefix}-L{pg + 1}", dd.A3,
                      material=info.get("material", ""), qty=1)
        sh.scale = "NTS"
        W, H = sh.size
        x0, y_top = dd.MARGIN + 6, H - dd.MARGIN - 6
        tot = diecost.total(ops, rates)
        sh.text(x0, y_top - 6, f"{info.get('product', '')} — 공정별 금형 {len(ops)} 벌 (1차 ~ {len(ops)}차) · 금형비 어림 합 "
                f"{tot / 1e4:,.0f} 만원 (단가는 예시 — [견적 단가] 창에서 고침)", 4.2, "L", bold=True)
        cw, ch = (W - 2 * dd.MARGIN - 20) / 3, (92.0 if len(chunk) <= 3 else 78.0)
        for j, op in enumerate(chunk):
            k = ops.index(op)
            cx = x0 + (j % 3) * (cw + 3)
            cy = y_top - 14 - (j // 3 + 1) * (ch + 3)
            sh.rect(cx, cy, cw, ch, "THIN", 0.25)
            sh.text(cx + 2, cy + ch - 5.5, op.label, 3.4, "L", bold=True)
            sh.text(cx + 2, cy + ch - 10, op.detail[:60], 2.0, "L")
            try:
                lines = F.wire(op.after if op.kind not in (BLANK, PIERCE) else {}, stage_holes(ops, k))
                L2 = [np.array([dieview.project_point(q) for q in P]) for P in lines]
                dd._iso(sh, L2, cx + 2, cy + 22, cw - 4, ch - 36, "", "PART", frame=False)
            except Exception:  # noqa: BLE001
                pass
            a = op.die
            if a is not None:
                bad = sum(1 for c in a.checks if c[3] != "적합")
                f_kn = op.force_n / 1000 if op.force_n else (a.press or {}).get("cut_kn", 0.0)
                mach = (a.press or {}).get("machine", "") or "-"
                lines_ = [f"금형 {a.L:.0f} × {a.W:.0f} · 닫힌 {a.shut:.0f} / 열린 {a.open_h:.0f}",
                          f"하중 {f_kn:,.1f} kN · {mach}",
                          f"점검 {len(a.checks)} 중 검토 {bad}" + (" ✔" if not bad else "")
                          + f" · 금형비 어림 {diecost.estimate(a, rates)[1] / 1e4:,.0f} 만원"]
            else:
                lines_ = ["금형을 못 만듦: " + (op.error or "")[:70]]
            for m, ln in enumerate(lines_):
                sh.text(cx + 2, cy + 15 - 5 * m, ln, 2.3, "L")
        if pg == len(pages) - 1:
            lay_op = next((o for o in ops if o.kind == BLANK and o.layouts), None)
            yt = y_top - 14 - ((len(chunk) + 2) // 3) * (ch + 3) - 4
            room = max(int((yt - dd.MARGIN - dd.TB_H - 6) / 4.2) - 1, 0)
            if lay_op is not None and room >= 2:
                rows_l = [[c.name, f"{c.angle:.0f}°", c.rows, f"{c.pitch:.1f}", f"{c.width:.1f}", f"{c.util * 100:.1f} %",
                           f"{c.kg_per_part * 1000:.1f}", "◀ 고름" if c is lay_op.layout else ""] for c in lay_op.layouts]
                dd._table(sh, x0 + 226, yt, [26, 12, 8, 18, 18, 18, 20, 16],
                          ["1차 배치", "각도", "열", "피치", "폭", "이용률", "g/개", ""], rows_l, h=4.2, max_rows=room)
            rows = []
            for op in ops:
                a = op.die
                rows.append([op.no, op.title, op.detail[:36],
                             f"{a.L:.0f}×{a.W:.0f}×{a.shut:.0f}" if a else "-",
                             f"{(op.force_n or 0) / 1000:,.1f}", ((a.press or {}).get("machine", "") if a else "") or "-",
                             (f"검토 {sum(1 for c in a.checks if c[3] != '적합')}" if a else "못 만듦"),
                             f"{diecost.estimate(a, rates)[1] / 1e4:,.0f}" if a else "-"])
            if room >= 2:
                dd._table(sh, x0, yt, [8, 24, 70, 32, 16, 32, 16, 20], ["차", "공정", "내용", "금형 L×W×닫힌", "하중 kN",
                                                                     "프레스", "점검", "금형비 만원"], rows, h=4.2, max_rows=room)
        sheets.append(sh)
    return sheets


def save_all(result, settings, ops, stem, info=None, dwg_converter="", folder=None, rates=None):
    """공정마다 STEP · 도면집 (한 PDF · DXF) · 부품표 엑셀 (공정표 + 공정마다 부품표 · 점검). [저장한 파일], [알림]."""
    from . import die3d as d3
    from . import diedraw as dd
    from . import drawfmt
    F = folder or Folder(result)
    info = dict(info or {})
    pre = info.get("prefix") or "HD-" + __import__("datetime").date.today().strftime("%y%m%d")
    out, notes, sheets = [], [], []
    cover = line_sheets(F, ops, info, pre, rates)
    for op in ops:
        a = op.die
        if a is None:
            notes.append(f"{op.label}: 금형을 못 만듦 — {op.error}")
            continue
        out.append(d3.save_step(a, f"{stem}_{op.no:02d}차_{op.title}.step"))
        inf = dict(info, prefix=f"{pre}-{op.no:02d}", die=op.label, product_sheet=(op.no == 1))
        if getattr(a, "y_cut", None) is not None:
            inf["y_cut"] = a.y_cut
        try:
            sheets += dd.build(a, inf)
        except Exception as e:  # noqa: BLE001
            notes.append(f"{op.label}: 도면을 못 만듦 — {e}")
    for i, sh in enumerate(cover, 1):
        dd._frame(sh, dict(info, die=f"공정별 금형 {len(ops)} 벌"), i, len(cover))
    allsh = cover + sheets
    out.append(dd.save_pdf(allsh, stem + "_공정별금형_도면집.pdf", title="공정별 금형 도면집"))
    dxf = dd.save_dxf(allsh, stem + "_공정별금형_도면집.dxf")
    out.append(dxf)
    try:
        out.append(drawfmt.dxf_to_dwg(dxf, stem + "_공정별금형_도면집.dwg", dwg_converter))
    except Exception:  # noqa: BLE001 — 변환기 없으면 DXF 만
        notes.append("DWG: 변환기 (ODA File Converter) 가 없어 DXF 만 저장")
    out.append(save_xlsx(ops, stem + "_공정별금형_부품표.xlsx", info.get("product", ""), rates))
    return out, notes


def save_xlsx(ops, path, name="", rates=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from . import diecost
    from . import die3d as d3
    wb = Workbook()
    ws = wb.active
    ws.title = "공정표"
    ws["A1"] = f"공정별 금형 공정표 — {name}"
    ws["A1"].font = Font(bold=True, size=13)
    head = ["차", "공정", "내용", "금형 L", "금형 W", "닫힌 높이", "열린 높이", "하중 kN", "프레스", "점검 검토", "근거 규칙",
            "금형비 어림 (원)"]
    hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E8C")
    for j, h in enumerate(head, 1):
        c = ws.cell(3, j, h)
        c.font, c.fill = hf, hfill
    for i, op in enumerate(ops, 4):
        a = op.die
        row = [op.no, op.title, op.detail, round(a.L) if a else "", round(a.W) if a else "",
               round(a.shut, 1) if a else "", round(a.open_h, 1) if a else "", round((op.force_n or 0) / 1000, 1),
               ((a.press or {}).get("machine", "") if a else "") or "",
               (sum(1 for c in a.checks if c[3] != "적합") if a else op.error), " ".join(op.rules),
               round(diecost.estimate(a, rates)[1]) if a else ""]
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
    for col, wd in zip("ABCDEFGHIJKL", (5, 14, 50, 8, 8, 10, 10, 9, 16, 10, 12, 16)):
        ws.column_dimensions[col].width = wd
    r0 = 5 + len(ops)
    ws.cell(r0, 11, "합계").font = Font(bold=True)
    ws.cell(r0, 12, round(diecost.total(ops, rates))).font = Font(bold=True)
    ws.cell(r0 + 1, 1, "금형비는 설계한 부품표 (무게 · 표준품 개수) 와 예시 단가 (costing.DEFAULT_DIE, [견적 단가] 창) 로 어림 — 업체 견적으로 맞출 것")
    wc = wb.create_sheet("금형비")
    wc["A1"] = f"금형비 어림 — {name}"
    wc["A1"].font = Font(bold=True, size=12)
    rr = 3
    for op in ops:
        if op.die is None:
            continue
        wc.cell(rr, 1, op.label).font = Font(bold=True)
        rr += 1
        for j, h in enumerate(("항목", "근거", "원"), 1):
            c = wc.cell(rr, j, h)
            c.font, c.fill = hf, hfill
        rr += 1
        lines_, tot_ = diecost.estimate(op.die, rates)
        for nm_, why_, v_ in lines_:
            wc.cell(rr, 1, nm_)
            wc.cell(rr, 2, why_)
            wc.cell(rr, 3, round(v_))
            rr += 1
        wc.cell(rr, 1, "합계").font = Font(bold=True)
        wc.cell(rr, 3, round(tot_)).font = Font(bold=True)
        rr += 2
    for col, wd in zip("ABC", (30, 70, 16)):
        wc.column_dimensions[col].width = wd
    for op in ops:
        a = op.die
        if a is None:
            continue
        w2 = wb.create_sheet(f"{op.no}차 부품표"[:31])
        w2["A1"] = f"{op.label} — 부품표"
        w2["A1"].font = Font(bold=True, size=12)
        for j, h in enumerate(d3.HEAD, 1):
            c = w2.cell(3, j, h)
            c.font, c.fill = hf, hfill
        for i, row in enumerate(a.rows(), 4):
            for j, v in enumerate(row, 1):
                w2.cell(i, j, v).alignment = Alignment(wrap_text=True, vertical="top")
        r = 5 + len(a.parts)
        w2.cell(r, 1, "점검").font = Font(bold=True)
        for c_ in a.checks:
            r += 1
            for j, v in enumerate(c_, 1):
                w2.cell(r, j, v)
        r += 2
        for n in a.notes:
            w2.cell(r, 1, "• " + n)
            r += 1
        for col, wd in zip("ABCDEFGH", (6, 30, 6, 22, 30, 9, 10, 56)):
            w2.column_dimensions[col].width = wd
    wb.save(path)
    return path
