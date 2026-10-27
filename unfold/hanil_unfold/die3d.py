"""금형 3D 조립 (블랭킹·피어싱 단발 / 순차 금형): 판 쌓기 · 가이드 포스트·부시 · 펀치 · 스트리퍼 · 스프링 · 볼트 ·
다우얼 핀 · 생크 를 솔리드로 만들고, STEP (부품 이름 붙임) · 2D 조립도 (평면도 + 정면도, 풍선 번호) · 파트리스트 엑셀로.

VISI·PressCAD 의 3D 금형 설계 순서 (레이아웃 → 상·하 스틸 → 스트리퍼·펀치 → 홀더·표준 부품 → 조립도·부품도 → 파트리스트)
를 견적·검토용 첫 안으로 자동으로 한다. 크기는 어림 규칙:
  다이 플레이트 두께  H ≈ ∛(0.1 F) (F 절단력, 독일 교재 · diecalc.die_plate_thickness), 5 mm 단위 올림, 최소 20
  다이 플레이트 여유  max(1.5 H, 30) mm (날 끝에서 판 끝까지)
  다이 세트 (상·하형) 두께 max(40, 0.08 × 긴 변), 가이드 포스트는 다이 플레이트 바깥 네 귀
  펀치 플레이트 max(20, 0.8 H) · 스트리퍼 max(15, 0.6 H) · 받침판 10 · 스트리퍼 아래 열림 max(10, 4 t)
  스트리퍼 힘 = 절단력의 7 % (다이 스프링 개수), 클리어런스 = diecalc.clearance_pct
설계 확정은 금형 설계자가 한다 (표준품 규격은 업체 표로 확인).
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field

import numpy as np

from . import bendcalc as bc
from . import diecalc as dc

BOX, CYL, PRISM = "box", "cyl", "prism"
# STEP 안의 부품 이름 (영문 — 한글은 CAD 마다 STEP 문자 인코딩이 달라 깨짐)
EN = {"하형 다이 세트": "LOWER DIE SET", "다이 받침판": "DIE BACKING PLATE", "다이 플레이트": "DIE PLATE",
      "스트리퍼 플레이트": "STRIPPER PLATE", "펀치 플레이트": "PUNCH PLATE", "펀치 받침판": "PUNCH BACKING PLATE",
      "상형 다이 세트": "UPPER DIE SET", "가이드 포스트": "GUIDE POST", "가이드 부시": "GUIDE BUSH",
      "블랭킹 펀치": "BLANKING PUNCH", "피어싱 펀치": "PIERCE PUNCH", "이형 피어싱 펀치": "SHAPED PIERCE PUNCH",
      "스트리퍼 볼트": "STRIPPER BOLT", "다이 스프링": "DIE SPRING", "육각 구멍 붙이 볼트": "SOCKET HEAD BOLT",
      "다우얼 핀": "DOWEL PIN", "생크": "SHANK", "(하형)": "(LOWER)", "(상형)": "(UPPER)"}


def en_name(name):
    out = name
    for k in sorted(EN, key=len, reverse=True):
        out = out.replace(k, EN[k])
    return out


@dataclass
class Prim:
    kind: str
    p: tuple                      # box: (x0, y0, z0, dx, dy, dz) · cyl: (x, y, z0, r, h) · prism: (z0, h)
    poly: np.ndarray | None = None   # prism 의 2D 다각형 (x, y)
    cuts: list = field(default_factory=list)   # 이 판에서 빼낼 프리즘 [(다각형, 넓힘)]


@dataclass
class DiePart:
    no: int
    name: str
    qty: int
    material: str
    size: str
    std: bool = False
    note: str = ""
    prims: list = field(default_factory=list)
    upper: bool = False           # 상형 쪽 (정면도에서 위)


@dataclass
class DieAssy:
    kind: str
    parts: list
    L: float
    W: float
    shut: float                   # 닫힌 높이 (어림)
    force_n: float
    notes: list = field(default_factory=list)
    blank: np.ndarray | None = None
    holes: list = field(default_factory=list)

    def rows(self):
        return [[p.no, p.name, p.qty, p.material, p.size, "표준품" if p.std else "가공품", p.note] for p in self.parts]


HEAD = ["번호", "품명", "수량", "재질", "크기 (mm)", "구분", "비고"]


def _up5(x):
    return float(math.ceil(x / 5.0) * 5)


def _outline(result):
    from .nesting import NestPart
    np_ = NestPart.from_result("p", result)
    if np_.outer is None:
        raise ValueError("전개 외곽이 닫혀 있지 않아 금형을 만들 수 없습니다")
    P = np.asarray(np_.outer, float)
    c = 0.5 * (P.min(0) + P.max(0))
    return P - c, [np.asarray(h, float) - c for h in np_.holes]


def _offset(poly, d):
    """다각형을 바깥으로 d 만큼 (꼭짓점 법선, 어림)."""
    P = np.asarray(poly, float)
    area = 0.5 * np.sum(P[:, 0] * np.roll(P[:, 1], -1) - np.roll(P[:, 0], -1) * P[:, 1])
    t = np.roll(P, -1, axis=0) - np.roll(P, 1, axis=0)
    n = np.column_stack([t[:, 1], -t[:, 0]]) * (1 if area > 0 else -1)
    n /= np.maximum(np.linalg.norm(n, axis=1), 1e-30)[:, None]
    return P + d * n


def _circle_d(h):
    c = h.mean(0)
    return 2 * float(np.mean(np.linalg.norm(h - c, axis=1))), c


def build(result, settings: bc.BendSettings, kind="단발 (블랭킹 · 피어싱)", plan=None):
    """전개 결과 하나의 블랭킹·피어싱 금형 조립. plan: 순차 금형 공정도 (process.ProcessPlan) — 있으면 스트립 길이로."""
    mat = settings.mat
    t = result.thickness
    blank, holes = _outline(result)
    tau = dc.shear_strength(mat)
    per = lambda P: float(np.linalg.norm(np.roll(P, -1, axis=0) - P, axis=1).sum())  # noqa: E731
    L_cut = per(blank) + sum(per(h) for h in holes)
    F = L_cut * t * tau
    H = max(20.0, _up5(dc.die_plate_thickness(F)))
    clr = dc.clearance_pct(mat, t) / 100 * t
    bx0, by0 = blank.min(0)
    bx1, by1 = blank.max(0)
    work_L, work_W = bx1 - bx0, by1 - by0
    progressive = plan is not None and getattr(plan, "pitch", 0)
    if progressive:
        n_st = len(plan.stations)
        work_L = n_st * plan.pitch
        work_W = plan.strip_width
        kind = f"순차 ({n_st} 공정)"
    m = max(1.5 * H, 30.0)
    Ld, Wd = _up5(work_L + 2 * m), _up5(work_W + 2 * m)
    post_d = 25.0 if max(Ld, Wd) < 250 else 32.0 if max(Ld, Wd) < 450 else 38.0
    Ls, Ws = _up5(Ld + 2 * (post_d + 35)), _up5(Wd + 20)
    ts = max(40.0, _up5(0.08 * max(Ls, Ws)))
    tb, tpp, tst = 10.0, max(20.0, _up5(0.8 * H)), max(15.0, _up5(0.6 * H))
    gap = max(10.0, 4 * t)                       # 다이 윗면 ↔ 스트리퍼 아랫면 (열림)
    sgap = 20.0                                  # 스트리퍼 윗면 ↔ 펀치 플레이트 아랫면 (스프링 자리)
    z = 0.0
    parts = []

    def add(name, qty, material, size, prims, std=False, note="", upper=False):
        parts.append(DiePart(len(parts) + 1, name, qty, material, size, std, note, prims, upper))
        return parts[-1]

    # 아래에서부터: 하형 다이 세트 · 다이 받침판 · 다이 플레이트
    open_blank = _offset(blank, clr) if not progressive else None
    z_ls = z
    lower = Prim(BOX, (-Ls / 2, -Ws / 2, z, Ls, Ws, ts))
    z += ts
    back = Prim(BOX, (-Ld / 2, -Wd / 2, z, Ld, Wd, tb))
    z += tb
    die = Prim(BOX, (-Ld / 2, -Wd / 2, z, Ld, Wd, H))
    # 다이 구멍: 블랭크 외곽 (+클리어런스) 과 구멍들 (+클리어런스), 아래 판에는 떨어질 구멍 (+1 mm)
    for poly in ([open_blank] if open_blank is not None else []) + [_offset(h, clr) for h in holes]:
        die.cuts.append((poly, 0.0))
        back.cuts.append((poly, 1.0))
        lower.cuts.append((poly, 2.0))
    z_die_top = z + H
    add("하형 다이 세트", 1, "SS400 (FC250)", f"{Ls:.0f} × {Ws:.0f} × {ts:.0f}", [lower], note="떨어질 구멍")
    add("다이 받침판", 1, "S50C (열처리)", f"{Ld:.0f} × {Wd:.0f} × {tb:.0f}", [back])
    add("다이 플레이트", 1, "SKD11 HRC 58~62", f"{Ld:.0f} × {Wd:.0f} × {H:.0f}", [die],
        note=f"한쪽 클리어런스 {clr:.3f} mm ({dc.clearance_pct(mat, t):g} % t)")
    z = z_die_top + gap
    z_st = z
    stripper = Prim(BOX, (-Ld / 2, -Wd / 2, z, Ld, Wd, tst))
    for poly in ([blank] if not progressive else []) + list(holes):
        stripper.cuts.append((poly, 0.05))
    add("스트리퍼 플레이트", 1, "S50C (SKD11 인서트)", f"{Ld:.0f} × {Wd:.0f} × {tst:.0f}", [stripper],
        upper=True, note=f"스트리퍼 힘 ≈ 절단력 7 % = {0.07 * F / 1000:,.1f} kN")
    z += tst + sgap
    z_pp = z
    pp = Prim(BOX, (-Ld / 2, -Wd / 2, z, Ld, Wd, tpp))
    add("펀치 플레이트", 1, "S50C", f"{Ld:.0f} × {Wd:.0f} × {tpp:.0f}", [pp], upper=True)
    z += tpp
    pb = Prim(BOX, (-Ld / 2, -Wd / 2, z, Ld, Wd, tb))
    add("펀치 받침판", 1, "SKD11 / S50C (열처리)", f"{Ld:.0f} × {Wd:.0f} × {tb:.0f}", [pb], upper=True)
    z += tb
    z_us = z
    upper_shoe = Prim(BOX, (-Ls / 2, -Ws / 2, z, Ls, Ws, ts - 5))
    add("상형 다이 세트", 1, "SS400 (FC250)", f"{Ls:.0f} × {Ws:.0f} × {ts - 5:.0f}", [upper_shoe], upper=True)
    z_top = z + ts - 5
    # 가이드 포스트 · 부시 (다이 플레이트 바깥 네 귀)
    px, py = Ld / 2 + post_d / 2 + 15, Wd / 2 - post_d
    corners = [(sx * px, sy * py) for sx in (-1, 1) for sy in (-1, 1)]
    post_L = _up5(z_us - z_ls + 20)
    add("가이드 포스트", 4, "SUJ2 HRC 60", f"Ø{post_d:g} × {post_L:.0f}",
        [Prim(CYL, (x, y, z_ls + 5, post_d / 2, post_L)) for x, y in corners], std=True, note="표준품 (압입)")
    add("가이드 부시", 4, "SUJ2 / 동합금", f"Ø{post_d:g} 안지름 × {ts - 5:.0f}",
        [Prim(CYL, (x, y, z_us, post_d / 2 + 6, ts - 5)) for x, y in corners], std=True, upper=True,
        note="표준품 (미끄럼 또는 볼 부시)")
    # 펀치: 블랭킹 (외곽) · 피어싱 (구멍마다) — 펀치 플레이트 윗면에서 스트리퍼 아랫면 1 mm 위까지
    p_len = (z_pp + tpp) - (z_st + 1.0)
    if not progressive:
        add("블랭킹 펀치", 1, "SKD11 HRC 58~62", f"{work_L:.0f} × {work_W:.0f} × {p_len:.0f}",
            [Prim(PRISM, (z_st + 1.0, p_len), poly=blank, cuts=[(h, 0.0) for h in holes])],
            upper=True, note="외곽 = 블랭크 (다이 쪽에 클리어런스)")
    groups = {}
    for h in holes:
        d, c = _circle_d(h)
        round_ = np.std(np.linalg.norm(h - c, axis=1)) < 0.02 * d
        key = (round(d, 1) if round_ else None)
        groups.setdefault(key, []).append((h, c, d))
    for key, hs in groups.items():
        if key is not None:
            add(f"피어싱 펀치 Ø{key:g}", len(hs), "SKH51 HRC 62~64", f"Ø{key:g} × {p_len:.0f}",
                [Prim(CYL, (c[0], c[1], z_st + 1.0, d / 2, p_len)) for _, c, d in hs], std=key <= 32, upper=True,
                note="표준품 (머리 붙임)" if key <= 32 else "")
        else:
            add("이형 피어싱 펀치", len(hs), "SKD11 HRC 58~62", "구멍 모양", [Prim(PRISM, (z_st + 1.0, p_len), poly=h)
                                                                       for h, _, _ in hs], upper=True)
    # 스트리퍼 볼트 · 다이 스프링
    sb_pos = [(sx * (Ld / 2 - 18), sy * (Wd / 2 - 18)) for sx in (-1, 1) for sy in (-1, 1)]
    sb_L = _up5(z_top - z_st)
    add("스트리퍼 볼트", 4, "SCM435", f"M8 (어깨 Ø10) × {sb_L:.0f}",
        [Prim(CYL, (x, y, z_st + 3, 5.0, sb_L - 3)) for x, y in sb_pos], std=True, upper=True)
    f_strip = 0.07 * F
    n_sp = max(4, min(12, 2 * math.ceil(f_strip / 2 / 1500.0)))   # 한 개 약 1.5 kN (Ø25 중하중, 20 % 휨) 어림
    sp = []
    for i in range(n_sp):
        a = 2 * math.pi * (i + 0.5) / n_sp
        sp.append(Prim(CYL, (0.42 * Ld * math.cos(a), 0.36 * Wd * math.sin(a), z_st + tst, 12.5, sgap + tpp * 0.6)))
    add("다이 스프링", n_sp, "SWOSC-V", f"Ø25 × {_up5(sgap + tpp * 0.6):.0f}", sp, std=True, upper=True,
        note=f"스트리퍼 힘 {f_strip / 1000:,.1f} kN ÷ {n_sp} — 하중 등급은 업체 표로")
    # 볼트 · 다우얼 핀 (하형: 하형 세트 → 다이 플레이트, 상형: 상형 세트 → 펀치 플레이트)
    bolt_d = 8.0 if max(Ld, Wd) < 300 else 10.0
    bl = [(sx * (Ld / 2 - 12), sy * (Wd / 2 - 30)) for sx in (-1, 1) for sy in (-1, 1)]
    add(f"육각 구멍 붙이 볼트 M{bolt_d:g} (하형)", 4, "SCM435", f"M{bolt_d:g} × {_up5(ts * 0.5 + tb + H * 0.7):.0f}",
        [Prim(CYL, (x, y, z_ls + ts * 0.5, bolt_d / 2, ts * 0.5 + tb + H * 0.7)) for x, y in bl], std=True)
    add(f"육각 구멍 붙이 볼트 M{bolt_d:g} (상형)", 4, "SCM435", f"M{bolt_d:g} × {_up5(ts * 0.5 + tb + tpp * 0.7):.0f}",
        [Prim(CYL, (x, y, z_pp + tpp * 0.3, bolt_d / 2, ts * 0.5 + tb + tpp * 0.7)) for x, y in bl], std=True,
        upper=True)
    dw = [(-(Ld / 2 - 25), 0.0), (Ld / 2 - 25, 0.0)]
    add("다우얼 핀 (하형)", 2, "SUJ2", f"Ø{bolt_d:g} × {_up5(tb + H + 10):.0f}",
        [Prim(CYL, (x, y, z_ls + ts - 10, bolt_d / 2, tb + H + 10)) for x, y in dw], std=True)
    add("다우얼 핀 (상형)", 2, "SUJ2", f"Ø{bolt_d:g} × {_up5(tb + tpp + 10):.0f}",
        [Prim(CYL, (x, y, z_pp, bolt_d / 2, tb + tpp + 10)) for x, y in dw], std=True, upper=True)
    if max(Ls, Ws) < 450:
        add("생크", 1, "S45C", "Ø38 × 60", [Prim(CYL, (0, 0, z_top, 19.0, 60.0))], upper=True,
            note="작은 금형 — 큰 금형은 상형 세트를 슬라이드에 클램프")
    notes = [f"절단 길이 {L_cut:,.0f} mm × t {t:g} × τ {tau:.0f} MPa = 절단력 {F / 1000:,.1f} kN "
             f"(권장 프레스 {F * 1.3 / dc.TON:,.0f} tf)",
             f"다이 플레이트 두께 ∛(0.1 F) = {dc.die_plate_thickness(F):.1f} → {H:.0f} mm, 날 끝 여유 {m:.0f} mm",
             f"열린 상태로 그림 (다이 ↔ 스트리퍼 {gap:.0f} mm) · 닫힌 높이 어림 {z_top - gap:.0f} mm"]
    return DieAssy(kind, parts, Ls, Ws, z_top - gap, F, notes, blank, holes)


# ---------------------------------------------------------------- 솔리드 · STEP

def _solid(pr: Prim):
    from OCP.BRepAlgoAPI import BRepAlgoAPI_Cut
    from OCP.BRepBuilderAPI import BRepBuilderAPI_MakeFace, BRepBuilderAPI_MakePolygon
    from OCP.BRepPrimAPI import BRepPrimAPI_MakeBox, BRepPrimAPI_MakeCylinder, BRepPrimAPI_MakePrism
    from OCP.gp import gp_Ax2, gp_Dir, gp_Pnt, gp_Vec

    def prism(poly, z0, h, grow=0.0):
        P = _offset(poly, grow) if grow else np.asarray(poly, float)
        mk = BRepBuilderAPI_MakePolygon()
        for x, y in P:
            mk.Add(gp_Pnt(float(x), float(y), float(z0)))
        mk.Close()
        face = BRepBuilderAPI_MakeFace(mk.Wire(), True).Face()
        return BRepPrimAPI_MakePrism(face, gp_Vec(0, 0, float(h))).Shape()
    if pr.kind == BOX:
        x0, y0, z0, dx, dy, dz = pr.p
        s = BRepPrimAPI_MakeBox(gp_Pnt(x0, y0, z0), dx, dy, dz).Shape()
        z_lo, hh = z0, dz
    elif pr.kind == CYL:
        x, y, z0, r, h = pr.p
        return BRepPrimAPI_MakeCylinder(gp_Ax2(gp_Pnt(x, y, z0), gp_Dir(0, 0, 1)), r, h).Shape()
    else:
        z0, h = pr.p
        s = prism(pr.poly, z0, h)
        z_lo, hh = z0, h
    for poly, grow in pr.cuts:
        try:
            s = BRepAlgoAPI_Cut(s, prism(poly, z_lo - 1, hh + 2, grow)).Shape()
        except Exception:
            pass
    return s


def save_step(assy: DieAssy, path):
    """부품마다 이름 붙인 STEP (번호 · 품명)."""
    from OCP.IFSelect import IFSelect_RetDone
    from OCP.STEPCAFControl import STEPCAFControl_Writer
    from OCP.STEPControl import STEPControl_AsIs
    from OCP.TCollection import TCollection_ExtendedString
    from OCP.TDataStd import TDataStd_Name
    from OCP.TDocStd import TDocStd_Document
    from OCP.XCAFDoc import XCAFDoc_DocumentTool
    doc = TDocStd_Document(TCollection_ExtendedString("MDTV-XCAF"))
    st = XCAFDoc_DocumentTool.ShapeTool_s(doc.Main())
    for p in assy.parts:
        for k, pr in enumerate(p.prims, 1):
            lab = st.AddShape(_solid(pr), False)
            nm = f"{p.no:02d} {en_name(p.name)}" + (f" #{k}" if len(p.prims) > 1 else "")
            TDataStd_Name.Set_s(lab, TCollection_ExtendedString(nm))
    st.UpdateAssemblies()
    w = STEPCAFControl_Writer()
    w.SetNameMode(True)
    w.Transfer(doc, STEPControl_AsIs)
    if w.Write(path) != IFSelect_RetDone:
        raise ValueError("STEP 을 쓰지 못했습니다")
    return path


# ---------------------------------------------------------------- 2D 조립도 (평면도 + 정면도, 풍선)

COLORS = {"다이": (90, 140, 200), "펀치": (210, 120, 60), "스트리퍼": (120, 170, 110), "세트": (170, 170, 175),
          "표준": (230, 190, 70)}


def _color(p: DiePart):
    if p.std:
        return COLORS["표준"]
    for k, c in COLORS.items():
        if k in p.name:
            return c
    return (150, 160, 175)


def paint(p, rect, assy: DieAssy, title=True):
    """왼쪽 평면도 (하형을 위에서), 오른쪽 정면도 (판 쌓기, 열린 상태). 풍선 = 파트리스트 번호."""
    from PySide6.QtCore import QPointF, QRectF, Qt
    from PySide6.QtGui import QBrush, QColor, QFont, QPen, QPolygonF
    x0, y0, W, H = rect.x(), rect.y(), rect.width(), rect.height()
    top_h = 34 if title else 0
    if title:
        f = QFont()
        f.setPixelSize(int(max(12, H * 0.028)))
        f.setBold(True)
        p.setFont(f)
        p.drawText(QRectF(x0, y0, W, top_h), Qt.AlignLeft | Qt.AlignVCenter,
                   f"금형 조립도 — {assy.kind} · {assy.L:.0f} × {assy.W:.0f} · 닫힌 높이 약 {assy.shut:.0f} mm")
    half = W * 0.5
    zmax = max(pr.p[2] + pr.p[5] if pr.kind == BOX else (pr.p[2] + pr.p[4] if pr.kind == CYL else pr.p[0] + pr.p[1])
               for part in assy.parts for pr in part.prims)
    balloons = []
    f = QFont()
    f.setPixelSize(int(max(9, H * 0.018)))
    p.setFont(f)
    # 평면도: 하형 (upper 아닌 것) + 펀치 위치 점선
    sx = (half - 40) / assy.L
    sy = (H - top_h - 60) / assy.W
    s = min(sx, sy)
    cx, cy = x0 + half / 2, y0 + top_h + (H - top_h) / 2

    def T(x, y):
        return QPointF(cx + x * s, cy - y * s)
    for part in assy.parts:
        col = QColor(*_color(part))
        for pr in part.prims:
            if part.upper and "펀치" not in part.name:
                continue
            pen = QPen(col.darker(150), 1.2, Qt.DashLine if part.upper else Qt.SolidLine)
            p.setPen(pen)
            p.setBrush(QBrush(QColor(col.red(), col.green(), col.blue(), 40 if part.upper else 90)))
            if pr.kind == BOX:
                bx, by, _, dx, dy, _ = pr.p
                p.drawRect(QRectF(T(bx, by + dy), T(bx + dx, by)))
                for poly, grow in pr.cuts:
                    P = _offset(poly, grow) if grow else poly
                    p.setBrush(QBrush(QColor(255, 255, 255)))
                    p.drawPolygon(QPolygonF([T(*q) for q in P]))
            elif pr.kind == CYL:
                x, y, _, r, _ = pr.p
                c = T(x, y)
                p.drawEllipse(c, r * s, r * s)
            else:
                p.drawPolygon(QPolygonF([T(*q) for q in pr.poly]))
        pr = part.prims[0]
        if pr.kind == CYL:
            at = T(pr.p[0], pr.p[1])
        elif pr.kind == BOX:
            if part.name not in ("하형 다이 세트", "다이 플레이트"):
                continue                                   # 받침판은 다이 플레이트 밑에 가려 보이지 않음
            at = T(pr.p[0] + 0.03 * pr.p[3], pr.p[1] + 0.85 * pr.p[4])
        else:
            at = T(*pr.poly.mean(0))
        if not part.upper or "펀치" in part.name:
            balloons.append((part.no, at, "top"))
    # 정면도: x-z (앞에서), 판은 사각형, 원기둥은 세로 사각형
    fx0 = x0 + half + 20
    sz = min((half - 150) / assy.L, (H - top_h - 60) / max(zmax, 1.0))
    bz = y0 + H - 30

    def Tf(x, zz):
        return QPointF(fx0 + (x + assy.L / 2) * sz, bz - zz * sz)
    for part in assy.parts:
        col = QColor(*_color(part))
        p.setPen(QPen(col.darker(160), 1.0))
        p.setBrush(QBrush(QColor(col.red(), col.green(), col.blue(), 140)))
        for pr in part.prims:
            if pr.kind == BOX:
                bx, _, z0, dx, _, dz = pr.p
                p.drawRect(QRectF(Tf(bx, z0 + dz), Tf(bx + dx, z0)))
            elif pr.kind == CYL:
                x, _, z0, r, h = pr.p
                p.drawRect(QRectF(Tf(x - r, z0 + h), Tf(x + r, z0)))
            else:
                z0, h = pr.p
                xs = pr.poly[:, 0]
                p.drawRect(QRectF(Tf(xs.min(), z0 + h), Tf(xs.max(), z0)))
        # 지시선이 그림을 가로지르지 않게 오른쪽 끝 (x 가 가장 큰) 프리미티브에 붙임
        def right(pr):
            return pr.p[0] + pr.p[3] if pr.kind in (BOX, CYL) else float(pr.poly[:, 0].max())
        pr = max(part.prims, key=right)
        if pr.kind == BOX:
            at = Tf(pr.p[0] + pr.p[3], pr.p[2] + pr.p[5] / 2)
        elif pr.kind == CYL:
            at = Tf(pr.p[0] + pr.p[3], pr.p[2] + pr.p[4] / 2)
        else:
            at = Tf(pr.poly[:, 0].max(), pr.p[0] + pr.p[1] / 2)
        balloons.append((part.no, at, "front"))
    # 풍선: 정면도는 오른쪽에 세로로, 평면도는 위쪽 가장자리로 끌어냄
    r = max(9.0, H * 0.017)
    fr = [b for b in balloons if b[2] == "front"]
    fr.sort(key=lambda b: b[1].y())                     # 위에서부터 — 지시선이 엇갈리지 않게
    xs_ = min(fx0 + assy.L * sz + 30, x0 + W - 3 * r)
    step = (H - top_h - 40) / max(len(fr), 1)
    for k, (no, at, _) in enumerate(fr):
        bpos = QPointF(xs_ + (k % 2) * 2.4 * r, y0 + top_h + 20 + k * step)
        _balloon(p, at, bpos, no, r)
    tp = [b for b in balloons if b[2] == "top"]
    tp.sort(key=lambda b: b[1].x())
    for k, (no, at, _) in enumerate(tp):
        bpos = QPointF(x0 + 30 + k * (half - 60) / max(len(tp) - 1, 1), y0 + top_h + 14 + (k % 2) * 2.4 * r)
        _balloon(p, at, bpos, no, r)


def _balloon(p, at, pos, no, r):
    from PySide6.QtCore import QRectF, Qt
    from PySide6.QtGui import QBrush, QColor, QPen
    p.setPen(QPen(QColor(40, 40, 40), 0.8))
    p.drawLine(at, pos)
    p.setBrush(QBrush(QColor(255, 255, 255)))
    p.drawEllipse(pos, r, r)
    p.drawText(QRectF(pos.x() - r, pos.y() - r, 2 * r, 2 * r), Qt.AlignCenter, str(no))


def render(assy: DieAssy, w=1800, h=900):
    from PySide6.QtCore import QRectF
    from PySide6.QtGui import QColor, QImage, QPainter

    from .report import ensure_gui
    ensure_gui()
    img = QImage(w, h, QImage.Format_ARGB32)
    img.fill(QColor(255, 255, 255))
    p = QPainter(img)
    p.setRenderHint(QPainter.Antialiasing)
    paint(p, QRectF(10, 10, w - 20, h - 20), assy)
    p.end()
    return img


def save_xlsx(assy: DieAssy, path, name=""):
    import os

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from .report import _add_image
    wb = Workbook()
    ws = wb.active
    ws.title = "파트리스트"
    ws["A1"] = f"금형 파트리스트 — {name} · {assy.kind}"
    ws["A1"].font = Font(bold=True, size=13)
    hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E8C")
    for j, hd in enumerate(HEAD, 1):
        c = ws.cell(3, j, hd)
        c.font, c.fill = hf, hfill
    for i, row in enumerate(assy.rows(), 4):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, wd in zip("ABCDEFG", (6, 30, 6, 22, 26, 9, 50)):
        ws.column_dimensions[col].width = wd
    r = 5 + len(assy.parts)
    for n in assy.notes:
        ws.cell(r, 1, "• " + n)
        r += 1
    ws2 = wb.create_sheet("조립도")
    tmp = _add_image(ws2, render(assy, 1800, 900), "A1", 1100)
    from .company import stamp_workbook
    stamp_workbook(wb)
    wb.save(path)
    try:
        os.remove(tmp)
    except OSError:
        pass
    return path


def save_pdf(assy: DieAssy, path, name=""):
    from PySide6.QtCore import QRectF, Qt
    from PySide6.QtGui import QColor, QFont

    from .process import _open_pdf
    pdf, p, R, k, W, H, m = _open_pdf(path, f"금형 조립도 {name}")
    paint(p, R(m, m, W - 2 * m, H * 0.62), assy)
    f = QFont()
    f.setPointSizeF(8.5)
    p.setFont(f)
    y = m + H * 0.64
    cols = [12, 70, 12, 50, 60, 18, 150]
    x = m
    for wd, hd in zip(cols, HEAD):
        p.fillRect(R(x, y, wd, 6), QColor(31, 78, 140))
        p.setPen(QColor(255, 255, 255))
        p.drawText(R(x + 1, y, wd - 2, 6), Qt.AlignVCenter | Qt.AlignLeft, hd)
        x += wd
    p.setPen(QColor(20, 20, 20))
    y += 6
    for row in assy.rows():
        if y > H - m - 5:
            pdf.newPage()
            y = m
        x = m
        for wd, v in zip(cols, row):
            p.drawText(R(x + 1, y, wd - 2, 5), Qt.AlignVCenter | Qt.AlignLeft, str(v))
            x += wd
        y += 5
    for n in assy.notes:
        if y > H - m - 5:
            pdf.newPage()
            y = m
        p.drawText(R(m, y + 1, W - 2 * m, 5), Qt.AlignVCenter | Qt.AlignLeft, "• " + n)
        y += 5
    p.end()
    return path
