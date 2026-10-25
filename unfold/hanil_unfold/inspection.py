"""검사 성적서: 도면 (전개 결과·3D) 에서 검사 항목과 공차를 뽑고, 측정값을 넣어 합격·불합격 · 공정 능력 (Cp·Cpk).

검사 항목 (자동)
  · 재료   — 소재 두께 (JIS G 3141 냉연 강판 두께 허용차, 폭 630 mm 미만 대표값), 밀시트 확인
  · 블랭크 — 전개 가로·세로, 구멍 지름·중심 위치 (JIS B 0408 打抜き 보통 공차)
  · 완성품 — 제품 전체 길이·폭·높이 (JIS B 0408 曲げ・絞り), 절곡 각도 (ISO 2768-1 각도, 짧은 변 길이별)
  · 성형   — 성형 해석의 문제 부위 (파단·위험·얇아짐) 최소 두께 ≥ t × (1 − 25 %), 성형 깊이
  · 외관   — 균열·주름·버·긁힘 (합/부)
공정 능력: Cp = (USL − LSL) / 6σ, Cpk = min(USL − μ, μ − LSL) / 3σ (σ = 시료 표준편차).
  AIAG PPAP 4 판 초기 공정 연구: Ppk ≥ 1.67 승인, 1.33 ~ 1.67 고객과 협의, < 1.33 개선 필요.
  시료가 적으면 (25 개 미만) 참고값 — AIAG SPC 매뉴얼은 25 군 이상 (100 개 이상) 을 권한다.

엑셀 성적서는 측정값 칸만 채우면 엑셀이 평균·Cpk·판정을 바로 계산하는 식이 들어 있고,
같은 파일에 다시 저장하면 검사 회차마다 시트가 쌓이고 [이력] 시트에 한 줄씩 남는다.
"""
from __future__ import annotations

import datetime as _dt
import math
import os
from dataclasses import dataclass, field

from . import tolerances as tl

VAR, ATTR, MIN, MAX = "var", "attr", "min", "max"      # 계량 (양쪽 공차) · 합부 · 하한만 · 상한만
KINDS = ["초도품 검사 (ISIR)", "정기 검사", "출하 검사", "수입 검사", "금형 수리 후 검사"]
GROUPS = ["재료", "블랭크", "완성품", "성형", "외관"]
DEFAULT_N = 5
MAX_N = 30
MAX_HOLES = 12

# JIS G 3141 (냉간 압연 강판) 두께 허용차 — 폭 630 mm 미만 (두께 위 끝 mm, ± mm)
JIS_G3141 = [(0.25, 0.03), (0.40, 0.04), (0.60, 0.05), (0.80, 0.06), (1.00, 0.06), (1.25, 0.07), (1.60, 0.08),
             (2.00, 0.10), (2.50, 0.12), (3.15, 0.14), (99.0, 0.16)]
# ISO 2768-1 각도 보통 공차 (짧은 변 길이 위 끝 mm, (m, c) ± 도)
ISO_ANGLE = [(10, (1.0, 1.5)), (50, (0.5, 1.0)), (120, (1 / 3, 0.5)), (400, (1 / 6, 0.25)), (1e9, (1 / 12, 1 / 6))]

GAUGES = {"len": "버니어 캘리퍼스", "big": "줄자·하이트 게이지·3차원 측정기", "hole": "핀 게이지·캘리퍼스",
          "pos": "3차원 측정기·검사 지그", "angle": "각도기·3차원 측정기", "thick": "마이크로미터",
          "thick_formed": "초음파 두께계 (또는 단면 절단)", "eye": "육안 (검사 표준 견본)"}


def thickness_tol(t):
    for top, v in JIS_G3141:
        if t <= top + 1e-9:
            return v
    return JIS_G3141[-1][1]


def angle_tol(short_leg, cls="c"):
    i = 0 if cls == "m" else 1
    for top, v in ISO_ANGLE:
        if short_leg <= top:
            return v[i]
    return ISO_ANGLE[-1][1][i]


def deg_text(x):
    """0.5 → 0°30′"""
    d = int(x)
    m = round((x - d) * 60)
    return f"{d}°{m:02d}′" if m else f"{d}°"


@dataclass
class Char:
    no: int
    group: str
    name: str
    nominal: float | None = None
    upper: float = 0.0              # + 공차
    lower: float = 0.0              # − 공차 (음수 또는 0)
    kind: str = VAR
    gauge: str = ""
    source: str = ""                # 공차 근거
    special: str = ""               # 특별 특성 (◇ 등)
    unit: str = "mm"
    values: list = field(default_factory=list)   # float · None · "OK"/"NG"

    @property
    def usl(self):
        if self.kind in (ATTR, MIN) or self.nominal is None:
            return None
        return self.nominal + self.upper

    @property
    def lsl(self):
        if self.kind in (ATTR, MAX) or self.nominal is None:
            return None
        return self.nominal + self.lower

    def spec_text(self):
        if self.kind == ATTR:
            return "이상 없음"
        if self.kind == MIN:
            return f"≥ {self.lsl:g} {self.unit}".strip()
        if self.kind == MAX:
            return f"≤ {self.usl:g} {self.unit}".strip()
        u, l_ = self.upper, self.lower
        tol = f"±{u:g}" if abs(u + l_) < 1e-12 else f"+{u:g} / {l_:+g}"
        return f"{self.nominal:g} {tol}"

    def judge(self, v):
        """한 값 → 'OK' / 'NG' / '' (비어 있음)."""
        if v is None or v == "":
            return ""
        if self.kind == ATTR:
            s = str(v).strip().upper()
            return "NG" if s in ("NG", "X", "불", "불량", "0") else "OK"
        try:
            x = float(v)
        except (TypeError, ValueError):
            return ""
        lo, hi = self.lsl, self.usl
        if lo is not None and x < lo - 1e-9:
            return "NG"
        if hi is not None and x > hi + 1e-9:
            return "NG"
        return "OK"

    def numbers(self):
        out = []
        for v in self.values:
            if self.kind == ATTR or v is None or v == "":
                continue
            try:
                out.append(float(v))
            except (TypeError, ValueError):
                pass
        return out


@dataclass
class Stats:
    n: int
    mean: float | None
    lo: float | None
    hi: float | None
    sd: float | None
    cp: float | None
    cpk: float | None
    ng: int
    verdict: str          # "합격" / "불합격" / "" (값 없음)

    attr: bool = False

    @property
    def capability(self):
        return "" if self.attr else capability_text(self.cpk, self.n)


def stats(c: Char) -> Stats:
    js = [c.judge(v) for v in c.values]
    ng = js.count("NG")
    filled = sum(1 for j in js if j)
    verdict = "" if not filled else "불합격" if ng else "합격"
    if c.kind == ATTR:
        return Stats(filled, None, None, None, None, None, None, ng, verdict, True)
    x = c.numbers()
    n = len(x)
    if not n:
        return Stats(0, None, None, None, None, None, None, ng, verdict)
    mu = sum(x) / n
    sd = math.sqrt(sum((v - mu) ** 2 for v in x) / (n - 1)) if n > 1 else None
    cp = cpk = None
    if sd is not None and sd > 1e-12:
        lo, hi = c.lsl, c.usl
        if lo is not None and hi is not None:
            cp = (hi - lo) / (6 * sd)
        sides = [v for v in ((hi - mu) / (3 * sd) if hi is not None else None,
                             (mu - lo) / (3 * sd) if lo is not None else None) if v is not None]
        cpk = min(sides) if sides else None
    return Stats(n, mu, min(x), max(x), sd, cp, cpk, ng, verdict)


def capability_text(cpk, n):
    if cpk is None:
        return "" if n < 2 else "편차 없음"
    tag = "승인" if cpk >= 1.67 else "고객 협의" if cpk >= 1.33 else "개선 필요"
    return tag + (" (참고)" if n < 25 else "")      # 시료 25 개 미만은 참고값


@dataclass
class Inspection:
    part_name: str
    part_no: str = ""
    customer: str = ""
    material: str = ""
    thickness: float = 0.0
    kind: str = KINDS[0]
    lot: str = ""
    inspector: str = ""
    n: int = DEFAULT_N
    when: _dt.datetime = field(default_factory=lambda: _dt.datetime.now().replace(microsecond=0))
    chars: list = field(default_factory=list)
    note: str = ""

    @property
    def verdict(self):
        vs = [stats(c).verdict for c in self.chars]
        if "불합격" in vs:
            return "불합격"
        if not any(vs):
            return "측정 전"
        if "" in vs:
            return "측정 중"
        return "합격"

    def ng_chars(self):
        return [c for c in self.chars if stats(c).verdict == "불합격"]

    def cpk_min(self):
        ks = [s.cpk for s in map(stats, self.chars) if s.cpk is not None]
        return min(ks) if ks else None

    def resize(self, n):
        self.n = max(1, min(MAX_N, int(n)))
        for c in self.chars:
            c.values = (list(c.values) + [None] * self.n)[:self.n]

    def renumber(self):
        for i, c in enumerate(self.chars, 1):
            c.no = i


# ---------------------------------------------------------------- 검사 항목 만들기

def solid_size(solid):
    """3D 부품의 (x, y, z) 크기 mm — 축에 나란한 감싸는 상자."""
    if solid is None:
        return None
    try:
        from OCP.Bnd import Bnd_Box
        from OCP.BRepBndLib import BRepBndLib
        b = Bnd_Box()
        BRepBndLib.AddOptimal_s(solid, b, False, False)
        mn, mx = b.CornerMin(), b.CornerMax()
        return (mx.X() - mn.X(), mx.Y() - mn.Y(), mx.Z() - mn.Z())
    except Exception:
        return None


def _short_legs(dims):
    """절곡 번호 → 짧은 변 길이 (전개도 연속 치수에서 절곡선 양쪽 간격)."""
    out = {}
    for ch in dims.chains:
        sp = ch.spans
        for k, lab in enumerate(ch.labels[1:-1], 1):
            out[lab] = min(sp[k - 1], sp[k])
    return out


def build(result, settings, solid=None, fa=None, name="", cls="B", density=None) -> Inspection:
    """전개 결과 (와 3D, 성형 해석) 로 검사 항목."""
    from . import dims as dimsmod
    mat = settings.mat
    t = result.thickness
    ins = Inspection(part_name=name, material=mat.code, thickness=t)
    C = ins.chars

    def add(group, nm, nominal=None, up=0.0, lo=0.0, kind=VAR, gauge="", source="", unit="mm"):
        C.append(Char(len(C) + 1, group, nm, None if nominal is None else round(float(nominal), 3),
                      round(up, 4), round(lo, 4), kind, gauge, source, unit=unit))

    # 재료
    tt = thickness_tol(t)
    add("재료", "소재 두께", t, tt, -tt, gauge=GAUGES["thick"], source="JIS G 3141 (폭 630 미만)")
    add("재료", f"재질 {mat.code} — 밀시트 (재질·두께·기계적 성질)", kind=ATTR, gauge="밀시트 대조", source="소재 규격")
    # 블랭크
    d = dimsmod.compute(result, density or mat.density)
    pk = "punch"
    for nm, v in (("전개 가로", d.width), ("전개 세로", d.height)):
        tol = tl.jis(v, pk, cls)
        add("블랭크", nm, v, tol, -tol, gauge=GAUGES["big" if v > 300 else "len"], source=f"JIS B 0408-{cls} 打抜き")
    for h in d.holes[:MAX_HOLES]:
        if h.diameter:
            tol = tl.jis(h.diameter, pk, cls)
            add("블랭크", f"{h.label} 구멍 지름 Ø", h.diameter, tol, -tol, gauge=GAUGES["hole"],
                source=f"JIS B 0408-{cls} 打抜き")
        for ax, v in (("X", h.center[0]), ("Y", h.center[1])):
            tol = tl.jis(v, pk, cls)
            add("블랭크", f"{h.label} 구멍 위치 {ax} (전개도 원점)", v, tol, -tol, gauge=GAUGES["pos"],
                source=f"JIS B 0408-{cls} 打抜き")
    if len(d.holes) > MAX_HOLES:
        ins.note = f"구멍 {len(d.holes)} 개 중 {MAX_HOLES} 개만 넣음 — 필요하면 [항목 추가]"
    # 완성품
    size = solid_size(solid)
    if size is not None and (result.bends or result.mode == "forming"):
        for nm, v in zip(("제품 길이 (X)", "제품 폭 (Y)", "제품 높이 (Z)"), size):
            if v < 2 * t:
                continue
            tol = tl.jis(v, "bend", cls)
            add("완성품", nm + " 바깥", v, tol, -tol, gauge=GAUGES["big" if v > 300 else "len"], source=f"JIS B 0408-{cls} 曲げ・絞り")
    legs = _short_legs(d)
    for b in result.bends:
        leg = legs.get(f"B{b.number}", 30.0)
        tol = angle_tol(leg, "c")
        add("완성품", f"B{b.number} 절곡 각도 (짧은 변 {leg:.0f})", b.angle_deg, tol, -tol, gauge=GAUGES["angle"],
            source=f"ISO 2768-c 각도 (±{deg_text(tol)})", unit="°")
    # 성형
    if fa is not None and fa.mode == "forming":
        from . import formsim as fs
        t_lim = round(t * (1 - fs.THIN_MAX), 3)
        ex = fa.explain
        spots = [i for i in (ex.issues if ex else []) if i.zone in (0, 1, 2)][:3]
        for i in spots:
            add("성형", f"성형 부위 {i.mark} 두께 — {i.short or i.region} (예측 {t * (1 - i.thin_max):.2f})", t_lim,
                kind=MIN, gauge=GAUGES["thick_formed"], source=f"얇아짐 ≤ {fs.THIN_MAX:.0%} (성형 해석)")
        if not spots:
            add("성형", f"가장 얇은 곳 두께 (예측 {fa.t_min:.2f})", t_lim, kind=MIN, gauge=GAUGES["thick_formed"],
                source=f"얇아짐 ≤ {fs.THIN_MAX:.0%} (성형 해석)")
    # 외관
    for nm in ("균열·터짐 (특히 R 부·귀·구멍 가장자리)", "주름·꺼짐 (플랜지·넓은 면)", "버 (날카로운 버·말림 없음)",
               "긁힘·찍힘·녹·이물"):
        add("외관", nm, kind=ATTR, gauge=GAUGES["eye"], source="검사 표준 견본")
    ins.resize(DEFAULT_N)
    return ins


# ---------------------------------------------------------------- 엑셀

HEAD = ["번호", "구분", "검사 항목", "기준", "+공차", "−공차", "하한", "상한", "단위", "측정기", "공차 근거"]
TAIL = ["평균", "최소", "최대", "σ", "Cp", "Cpk", "판정", "공정 능력"]
FIRST = 9          # 첫 항목 줄 (8 = 머리줄)


def _col(i):
    from openpyxl.utils import get_column_letter
    return get_column_letter(i)


def sheet_name(ins, wb=None):
    base = f"{ins.kind.split(' ')[0]} {ins.when:%y%m%d-%H%M}"[:28]
    nm, k = base, 2
    while wb is not None and nm in wb.sheetnames:
        nm, k = f"{base} ({k})", k + 1
    return nm


def write_sheet(ws, ins: Inspection):
    """한 회차 성적서 시트 (엑셀 식으로 평균·Cpk·판정)."""
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    from . import company
    hf, hfill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E8C")
    thin = Side(style="thin", color="B0B8C8")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    n = ins.n
    s0 = len(HEAD) + 1                     # 첫 시료 칸
    s1 = s0 + n - 1
    t0 = s1 + 1
    ws["A1"] = f"검사 성적서 — {ins.kind}"
    ws["A1"].font = Font(bold=True, size=15)
    ws.cell(1, t0, company.MAKER).font = Font(bold=True, color="1F4E8C")
    info = [("품명", ins.part_name), ("품번", ins.part_no), ("고객", ins.customer),
            ("재질·두께", f"{ins.material} t{ins.thickness:g}"), ("로트", ins.lot), ("검사자", ins.inspector),
            ("검사 일시", ins.when.strftime("%Y-%m-%d %H:%M")), ("시료 수", n)]
    for k, (a, b) in enumerate(info):
        r, c = 3 + k // 4, 1 + (k % 4) * 3
        ws.cell(r, c, a).font = Font(bold=True)
        ws.cell(r, c + 1, b)
    last = FIRST + len(ins.chars) - 1
    ws.cell(5, 1, "종합 판정").font = Font(bold=True)
    vcol = _col(t0 + 6)
    ws.cell(5, 2, f'=IF(COUNTIF({vcol}{FIRST}:{vcol}{last},"NG")>0,"불합격",'
                  f'IF(COUNTIF({vcol}{FIRST}:{vcol}{last},"OK")=0,"측정 전",'
                  f'IF(COUNTIF({vcol}{FIRST}:{vcol}{last},"OK")<{len(ins.chars)},"측정 중","합격")))').font = \
        Font(bold=True, size=13)
    ws.cell(5, 4, "불합격 항목 수")
    ws.cell(5, 5, f'=COUNTIF({vcol}{FIRST}:{vcol}{last},"NG")')
    ws.cell(5, 7, "가장 낮은 Cpk")
    kcol = _col(t0 + 5)
    ws.cell(5, 8, f'=IF(COUNT({kcol}{FIRST}:{kcol}{last})=0,"",MIN({kcol}{FIRST}:{kcol}{last}))').number_format = "0.00"
    if ins.note:
        ws.cell(6, 1, ins.note).font = Font(color="888888")
    head = HEAD + [f"S{k}" for k in range(1, n + 1)] + TAIL
    for j, h in enumerate(head, 1):
        c = ws.cell(FIRST - 1, j, h)
        c.font, c.fill, c.border = hf, hfill, box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    widths = [5, 7, 42, 9, 7, 7, 9, 9, 5, 22, 22] + [8] * n + [9, 9, 9, 8, 7, 7, 7, 22]
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[_col(j)].width = w
    for i, ch in enumerate(ins.chars):
        r = FIRST + i
        rng = f"{_col(s0)}{r}:{_col(s1)}{r}"
        row = [ch.no, ch.group, ch.name, "" if ch.nominal is None else ch.nominal,
               "" if ch.kind in (ATTR, MIN) else ch.upper, "" if ch.kind in (ATTR, MAX) else ch.lower]
        for j, v in enumerate(row, 1):
            ws.cell(r, j, v)
        D, E, F, G, H = (f"{_col(k)}{r}" for k in (4, 5, 6, 7, 8))
        if ch.kind == ATTR:
            ws.cell(r, 7, "")
            ws.cell(r, 8, "")
        else:
            ws.cell(r, 7, f"={D}+{F}" if ch.kind != MAX else "")
            ws.cell(r, 8, f"={D}+{E}" if ch.kind != MIN else "")
        ws.cell(r, 9, ch.unit if ch.kind != ATTR else "")
        ws.cell(r, 10, ch.gauge)
        ws.cell(r, 11, ch.source)
        for k in range(n):
            v = ch.values[k] if k < len(ch.values) else None
            ws.cell(r, s0 + k, v if v not in (None, "") else None)
        if ch.kind == ATTR:
            ws.cell(r, t0 + 6, f'=IF(COUNTA({rng})=0,"",IF(COUNTIF({rng},"NG")>0,"NG","OK"))')
        else:
            cnt = f"COUNT({rng})"
            mean = f"{_col(t0)}{r}"
            sd = f"{_col(t0 + 3)}{r}"
            ws.cell(r, t0, f'=IF({cnt}=0,"",AVERAGE({rng}))')
            ws.cell(r, t0 + 1, f'=IF({cnt}=0,"",MIN({rng}))')
            ws.cell(r, t0 + 2, f'=IF({cnt}=0,"",MAX({rng}))')
            ws.cell(r, t0 + 3, f'=IF({cnt}<2,"",STDEV({rng}))')
            ok_sd = f'AND({cnt}>=2,N({sd})>0)'
            if ch.kind == VAR:
                ws.cell(r, t0 + 4, f'=IF({ok_sd},({H}-{G})/(6*{sd}),"")')
                ws.cell(r, t0 + 5, f'=IF({ok_sd},MIN({H}-{mean},{mean}-{G})/(3*{sd}),"")')
                bad = f'OR(MIN({rng})<{G}-1E-9,MAX({rng})>{H}+1E-9)'
            elif ch.kind == MIN:
                ws.cell(r, t0 + 5, f'=IF({ok_sd},({mean}-{G})/(3*{sd}),"")')
                bad = f'MIN({rng})<{G}-1E-9'
            else:
                ws.cell(r, t0 + 5, f'=IF({ok_sd},({H}-{mean})/(3*{sd}),"")')
                bad = f'MAX({rng})>{H}+1E-9'
            ws.cell(r, t0 + 6, f'=IF({cnt}=0,"",IF({bad},"NG","OK"))')
            cpk = f"{_col(t0 + 5)}{r}"
            ws.cell(r, t0 + 7, f'=IF(N({cpk})=0,"",IF({cpk}>=1.67,"승인",IF({cpk}>=1.33,"고객 협의","개선 필요"))'
                               f'&IF({cnt}<25," (참고)",""))')
            for k in list(range(4, 9)) + list(range(s0, t0 + 3)):
                ws.cell(r, k).number_format = "0.000" if ch.unit == "mm" else "0.00"
            for k in (t0 + 3, t0 + 4, t0 + 5):
                ws.cell(r, k).number_format = "0.000" if k == t0 + 3 else "0.00"
        for j in range(1, t0 + 8):
            ws.cell(r, j).border = box
            ws.cell(r, j).alignment = Alignment(vertical="center", wrap_text=j in (3, 10, 11))
        if ch.kind != ATTR:                   # 측정값 칸: 공차 밖이면 빨강 (엑셀에서 넣어도)
            red = PatternFill("solid", fgColor="F8C8C8")
            if ch.kind != MAX:
                ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=[f"$G${r}-1E-9"], fill=red))
            if ch.kind != MIN:
                ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=[f"$H${r}+1E-9"], fill=red))
    vr = f"{vcol}{FIRST}:{vcol}{last}"
    ws.conditional_formatting.add(vr, CellIsRule(operator="equal", formula=['"NG"'],
                                                 fill=PatternFill("solid", fgColor="F4B6B6")))
    ws.conditional_formatting.add(vr, CellIsRule(operator="equal", formula=['"OK"'],
                                                 fill=PatternFill("solid", fgColor="CFE8CF")))
    ws.freeze_panes = ws.cell(FIRST, 4)
    r = last + 2
    for line in ("측정값 칸 (S1 …) 에 숫자를 넣으면 평균·σ·Cp·Cpk·판정이 바로 계산됩니다. 외관·밀시트 칸은 OK / NG.",
                 "Cp = (상한 − 하한) / 6σ, Cpk = min(상한 − 평균, 평균 − 하한) / 3σ. AIAG PPAP: 1.67 이상 승인, "
                 "1.33~1.67 고객 협의, 1.33 미만 개선 필요 (시료 25 개 미만은 참고값).",
                 "보통 공차: JIS B 0408 (打抜き·曲げ・絞り) · ISO 2768-1 각도 · 소재 두께 JIS G 3141 — 도면에 공차가 "
                 "있으면 도면 값으로 바꾸세요."):
        ws.cell(r, 1, line).font = Font(color="666666")
        r += 1
    r += 1
    ws.cell(r, 3, "작성 :                검토 :                승인 :").font = Font(bold=True)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0


HIST = "이력"
HIST_HEAD = ["검사 일시", "검사 종류", "로트", "검사자", "시료 수", "종합 판정", "불합격 항목", "가장 낮은 Cpk", "시트"]


def save_xlsx(ins: Inspection, path, append=True):
    """성적서 저장. append 이고 파일이 있으면 같은 파일에 시트를 하나 더 쌓고 [이력] 에 한 줄."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill

    from .company import stamp_workbook
    if append and os.path.exists(path):
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.active.title = HIST
    if HIST not in wb.sheetnames:
        wb.create_sheet(HIST, 0)
    hs = wb[HIST]
    if hs.max_row < 3 or hs.cell(2, 1).value != HIST_HEAD[0]:
        hs["A1"] = f"검사 이력 — {ins.part_name} {ins.part_no}".strip()
        hs["A1"].font = Font(bold=True, size=13)
        for j, h in enumerate(HIST_HEAD, 1):
            c = hs.cell(2, j, h)
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F4E8C")
        for j, w in enumerate([17, 18, 12, 10, 8, 10, 40, 12, 22], 1):
            hs.column_dimensions[_col(j)].width = w
    name = sheet_name(ins, wb)
    write_sheet(wb.create_sheet(name), ins)
    r = hs.max_row + 1
    ks = ins.cpk_min()
    for j, v in enumerate([ins.when.strftime("%Y-%m-%d %H:%M"), ins.kind, ins.lot, ins.inspector, ins.n, ins.verdict,
                           ", ".join(f"{c.no}. {c.name}" for c in ins.ng_chars()) or "-",
                           round(ks, 2) if ks is not None else "", name], 1):
        hs.cell(r, j, v)
    hs.cell(r, 9).hyperlink = f"#'{name}'!A1"
    hs.cell(r, 6).font = Font(bold=True, color={"합격": "2A8C3C", "불합격": "D22828"}.get(ins.verdict, "666666"))
    stamp_workbook(wb)
    wb.save(path)
    return name


def read_xlsx(path, sheet=None) -> Inspection:
    """우리 성적서 (엑셀에서 측정값을 채운 것) 를 다시 읽는다. sheet 가 없으면 마지막 회차 시트."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=False)
    names = [n for n in wb.sheetnames if n != HIST]
    if not names:
        raise ValueError("검사 시트가 없습니다")
    ws = wb[sheet or names[-1]]
    hdr = [ws.cell(FIRST - 1, j).value for j in range(1, ws.max_column + 1)]
    if not hdr or hdr[0] != "번호":
        raise ValueError(f"'{ws.title}' 은 이 프로그램의 검사 성적서 양식이 아닙니다 (8 째 줄 머리줄)")
    s_cols = [j for j, h in enumerate(hdr, 1) if isinstance(h, str) and h.startswith("S") and h[1:].isdigit()]
    title = str(ws["A1"].value or "")
    kind = title.split("—", 1)[-1].strip() if "—" in title else KINDS[0]
    info = {}
    for r in (3, 4):
        for c in (1, 4, 7, 10):
            k = ws.cell(r, c).value
            if k:
                info[str(k)] = ws.cell(r, c + 1).value
    mat, _, th = str(info.get("재질·두께") or "").partition(" t")
    ins = Inspection(part_name=str(info.get("품명") or ""), part_no=str(info.get("품번") or ""),
                     customer=str(info.get("고객") or ""), material=mat, kind=kind,
                     lot=str(info.get("로트") or ""), inspector=str(info.get("검사자") or ""))
    try:
        ins.thickness = float(th)
    except ValueError:
        pass
    try:
        ins.when = _dt.datetime.strptime(str(info.get("검사 일시")), "%Y-%m-%d %H:%M")
    except ValueError:
        pass
    r = FIRST
    while ws.cell(r, 1).value not in (None, ""):
        if not isinstance(ws.cell(r, 1).value, (int, float)):
            break
        nominal = ws.cell(r, 4).value
        up, lo = ws.cell(r, 5).value, ws.cell(r, 6).value
        unit = ws.cell(r, 9).value or ""
        if nominal in (None, ""):
            kind_ = ATTR
        elif up in (None, "") and lo not in (None, ""):
            kind_ = MIN
        elif lo in (None, "") and up not in (None, ""):
            kind_ = MAX
        else:
            kind_ = VAR
        vals = []
        for j in s_cols:
            v = ws.cell(r, j).value
            vals.append(None if v in (None, "") else v if kind_ == ATTR else _num(v))
        ins.chars.append(Char(int(ws.cell(r, 1).value), str(ws.cell(r, 2).value or ""), str(ws.cell(r, 3).value or ""),
                              None if kind_ == ATTR else _num(nominal),
                              _num(up) or 0.0, _num(lo) or 0.0, kind_, str(ws.cell(r, 10).value or ""),
                              str(ws.cell(r, 11).value or ""), unit=str(unit), values=vals))
        r += 1
    ins.n = len(s_cols)
    return ins


def _num(v):
    if v in (None, ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------- PDF

def save_pdf(ins: Inspection, path):
    """A3 가로 성적서: 머리 정보 · 항목별 기준/공차/측정값/통계/판정 · 서명 칸."""
    from PySide6.QtCore import QMarginsF, QPointF, QRectF, Qt
    from PySide6.QtGui import QColor, QPageLayout, QPageSize, QPainter, QPdfWriter, QPen

    from .report import _brand, _font, ensure_gui, version_text
    ensure_gui()
    pdf = QPdfWriter(path)
    pdf.setResolution(300)
    pdf.setPageLayout(QPageLayout(QPageSize(QPageSize.A3), QPageLayout.Landscape, QMarginsF(0, 0, 0, 0)))
    pdf.setTitle(f"검사 성적서 - {ins.part_name}")
    pdf.setCreator(f"한일 판금 전개 {version_text()}")
    k = 300 / 25.4
    W, H, m = 420.0, 297.0, 12.0

    def R(x, y, w, h):
        return QRectF(x * k, y * k, w * k, h * k)

    p = QPainter(pdf)
    p.setRenderHint(QPainter.Antialiasing)
    vcol = {"합격": QColor(40, 140, 60), "불합격": QColor(210, 40, 40)}
    ns = min(ins.n, 10)                         # PDF 에는 시료 10 개까지 (나머지는 엑셀)
    cols = [("No", 9), ("구분", 13), ("검사 항목", 82), ("규격", 30), ("측정기", 34)] + \
           [(f"S{j}", 15) for j in range(1, ns + 1)] + [("평균", 17), ("σ", 14), ("Cpk", 13), ("판정", 15)]
    scale = (W - 2 * m) / sum(c[1] for c in cols)
    cols = [(a, b * scale) for a, b in cols]

    def header(page):
        p.setPen(QColor(20, 20, 20))
        p.setFont(_font(17, True))
        p.drawText(R(m, m, W - 2 * m, 10), Qt.AlignLeft | Qt.AlignVCenter, f"검사 성적서 — {ins.kind}")
        p.setFont(_font(8))
        _brand(p, R(m, m + 2, W - 2 * m, 6))
        p.setPen(QPen(QColor(20, 20, 20), 0.4 * k))
        p.drawLine(QPointF(m * k, (m + 11) * k), QPointF((W - m) * k, (m + 11) * k))
        y = m + 14
        info = [("품명", ins.part_name), ("품번", ins.part_no or "-"), ("고객", ins.customer or "-"),
                ("재질·두께", f"{ins.material} t{ins.thickness:g}"), ("로트", ins.lot or "-"),
                ("검사자", ins.inspector or "-"), ("검사 일시", ins.when.strftime("%Y-%m-%d %H:%M")),
                ("시료 수", f"{ins.n}" + (" (PDF 는 10 개까지, 전체는 엑셀)" if ins.n > 10 else "")),
                ("쪽", f"{page}")]
        bw = (W - 2 * m - 60) / 3
        for i, (a, b) in enumerate(info):
            x, yy = m + (i % 3) * bw, y + (i // 3) * 6.5
            p.setFont(_font(9, True))
            p.setPen(QColor(31, 78, 140))
            p.drawText(R(x, yy, 22, 6), Qt.AlignLeft | Qt.AlignVCenter, a)
            p.setFont(_font(9))
            p.setPen(QColor(20, 20, 20))
            p.drawText(R(x + 22, yy, bw - 24, 6), Qt.AlignLeft | Qt.AlignVCenter, str(b))
        v = ins.verdict
        p.setPen(QPen(vcol.get(v, QColor(120, 120, 120)), 0.8 * k))
        p.setBrush(Qt.NoBrush)
        p.drawRect(R(W - m - 56, y - 1, 56, 20))
        p.setFont(_font(18, True))
        p.setPen(vcol.get(v, QColor(120, 120, 120)))
        p.drawText(R(W - m - 56, y - 1, 56, 13), Qt.AlignCenter, v)
        p.setFont(_font(8.5))
        ks = ins.cpk_min()
        p.drawText(R(W - m - 56, y + 11, 56, 7), Qt.AlignCenter,
                   f"불합격 {len(ins.ng_chars())} 항목" + (f" · 최저 Cpk {ks:.2f}" if ks is not None else ""))
        y += 22
        p.setPen(Qt.NoPen)
        p.setBrush(QColor(31, 78, 140))
        p.drawRect(R(m, y, W - 2 * m, 7))
        p.setPen(QColor(255, 255, 255))
        p.setFont(_font(8.5, True))
        x = m
        for title, cw in cols:
            p.drawText(R(x + 1, y, cw - 2, 7), Qt.AlignCenter, title)
            x += cw
        return y + 7

    def fmt(v, unit):
        if v is None:
            return ""
        return f"{v:.3f}" if unit == "mm" else f"{v:.2f}"

    page = 1
    y = header(page)
    rh = 6.4
    for i, c in enumerate(ins.chars):
        if y + rh > H - m - 26:
            pdf.newPage()
            page += 1
            y = header(page)
        s = stats(c)
        if i % 2:
            p.setPen(Qt.NoPen)
            p.setBrush(QColor(240, 244, 250))
            p.drawRect(R(m, y, W - 2 * m, rh))
        cells = [str(c.no), c.group, c.name, c.spec_text(), c.gauge]
        vals = []
        for j in range(ns):
            v = c.values[j] if j < len(c.values) else None
            vals.append(("" if v in (None, "") else str(v) if c.kind == ATTR else fmt(float(v), c.unit), c.judge(v)))
        tail = [fmt(s.mean, c.unit), "" if s.sd is None else f"{s.sd:.3f}",
                "" if s.cpk is None else f"{s.cpk:.2f}", s.verdict]
        x = m
        p.setFont(_font(7.8))
        for j, (title, cw) in enumerate(cols):
            if j < 5:
                txt, col = cells[j], QColor(20, 20, 20)
            elif j < 5 + ns:
                txt, jj = vals[j - 5]
                col = QColor(210, 40, 40) if jj == "NG" else QColor(20, 20, 20)
                if jj == "NG":
                    p.setPen(Qt.NoPen)
                    p.setBrush(QColor(248, 205, 205))
                    p.drawRect(R(x, y, cw, rh))
            else:
                txt = tail[j - 5 - ns]
                col = vcol.get(txt, QColor(20, 20, 20)) if j == len(cols) - 1 else \
                    (QColor(210, 40, 40) if j == len(cols) - 2 and s.cpk is not None and s.cpk < 1.33 else QColor(20, 20, 20))
            p.setPen(col)
            p.setFont(_font(7.8, j == len(cols) - 1))
            al = Qt.AlignLeft if j in (2, 4) else Qt.AlignCenter
            p.drawText(R(x + 1, y, cw - 2, rh), al | Qt.AlignVCenter, txt)
            x += cw
        p.setPen(QPen(QColor(210, 214, 222), 0.15 * k))
        p.drawLine(QPointF(m * k, (y + rh) * k), QPointF((W - m) * k, (y + rh) * k))
        y += rh
    # 바닥: 알림 · 서명
    yb = H - m - 24
    p.setFont(_font(8))
    p.setPen(QColor(90, 90, 90))
    for kx, line in enumerate((
            "규격: 도면에 공차가 없는 치수는 보통 공차 JIS B 0408 (打抜き·曲げ・絞り) · ISO 2768-1 각도 · 소재 두께 JIS G 3141.",
            "Cpk = min(상한 − 평균, 평균 − 하한) / 3σ — AIAG PPAP: 1.67 이상 승인, 1.33~1.67 고객 협의, 1.33 미만 개선 "
            "필요 (시료 25 개 미만은 참고값).")):
        p.drawText(R(m, yb + kx * 5, W - 2 * m - 100, 5), Qt.AlignLeft | Qt.AlignVCenter, line)
    p.setPen(QPen(QColor(60, 60, 60), 0.3 * k))
    p.setBrush(Qt.NoBrush)
    for kx, lab in enumerate(("작성", "검토", "승인")):
        x = W - m - 96 + kx * 32
        p.drawRect(R(x, yb - 2, 32, 22))
        p.drawLine(QPointF(x * k, (yb + 4) * k), QPointF((x + 32) * k, (yb + 4) * k))
        p.setFont(_font(8.5, True))
        p.drawText(R(x, yb - 2, 32, 6), Qt.AlignCenter, lab)
    if ins.inspector:
        p.setFont(_font(9))
        p.drawText(R(W - m - 96, yb + 6, 32, 12), Qt.AlignCenter, ins.inspector)
    p.end()


def record(ins: Inspection):
    """[견적·금형 기록] 에 남길 (values, detail)."""
    ks = ins.cpk_min()
    v = dict(title=ins.part_name, customer=ins.customer, material=f"{ins.material} t{ins.thickness:g}", parts=1,
             summary=f"{ins.kind} {ins.verdict} · 항목 {len(ins.chars)} · 시료 {ins.n} · 불합격 {len(ins.ng_chars())}"
                     + (f" · 최저 Cpk {ks:.2f}" if ks is not None else "") + (f" · 로트 {ins.lot}" if ins.lot else ""))
    detail = {"checks": [[f"{c.no}. {c.name}", c.spec_text(), stats(c).verdict or "-"] for c in ins.chars]}
    return v, detail
