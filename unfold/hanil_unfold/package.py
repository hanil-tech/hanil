"""전체 분석: 3D 조립품 → 부품 분리 → 전개 → 자재비 → 부품별 공정도 → 생산 견적 → 금형 견적을 한 폴더에.

  01_BOM.xlsx            다단계 BOM (부품 분리·묶음·분류·원자재)
  02_자재비.xlsx          요약 · 부품별 자재비 · 재질별 소요량 · 부자재
  03_공정도.pdf           차례 + 자작 부품마다 공정 (순차 금형 스트립 공정도·자동 공정 설계 / 레이저·절곡 작업표)
  04_생산견적서.xlsx·pdf
  05_금형견적서.xlsx·pdf  (금형이 필요한 부품이 있을 때)
"""
from __future__ import annotations

import datetime as _dt
import os
import re

from . import bendcalc as bc
from . import quote as qt_


def _safe(name):
    return re.sub(r'[\\/:*?"<>|]+', "_", name).strip() or "제품"


def material_rows(q: qt_.Quote, settings: bc.BendSettings):
    """부품별 자재비 줄."""
    price = q.rates["material"].get(settings.mat.code, 0)
    rows = []
    for p in q.parts:
        scrap = max(p.gross_kg - p.weight, 0)
        rows.append([p.no, p.name, p.spec, p.qty, round(p.gross_kg, 3), round(p.weight, 3), round(scrap, 3),
                     price,
                     round(p.material), round(p.material * p.qty), round(p.gross_kg * p.qty * q.sets, 2),
                     round(p.material * p.qty * q.sets), "직접 입력" if p.material_override is not None else ""])
    return rows


def by_spec(q: qt_.Quote):
    """재질·규격별 합계: (규격, 부품 종류, 세트당 kg, 견적 수량 kg, 세트당 금액, 견적 수량 금액)."""
    out = {}
    for p in q.parts:
        e = out.setdefault(p.spec or "-", [0, 0.0, 0.0])
        e[0] += 1
        e[1] += p.gross_kg * p.qty
        e[2] += p.material * p.qty
    return [(k, v[0], v[1], v[1] * q.sets, v[2], v[2] * q.sets) for k, v in sorted(out.items())]


def save_material_xlsx(q: qt_.Quote, settings, path, files=()):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    from . import company
    wb = Workbook()
    head = PatternFill("solid", fgColor="1F4E8C")
    hf = Font(bold=True, color="FFFFFF")

    def table(ws, r0, heads, rows, money=()):
        for j, h in enumerate(heads, 1):
            c = ws.cell(r0, j, h)
            c.font, c.fill = hf, head
            ws.column_dimensions[c.column_letter].width = max(ws.column_dimensions[c.column_letter].width or 0,
                                                              max(10, len(h) * 2))
        for i, row in enumerate(rows, r0 + 1):
            for j, v in enumerate(row, 1):
                c = ws.cell(i, j, v)
                if j in money:
                    c.number_format = "#,##0"
        return r0 + len(rows) + 2

    ws = wb.active
    ws.title = "요약"
    ws["A1"] = f"전체 분석 — {q.title}"
    ws["A1"].font = Font(bold=True, size=15)
    ws["E1"] = company.MAKER
    total_kg = sum(p.gross_kg * p.qty for p in q.parts)
    lines = [
        ("분석 일시", f"{_dt.datetime.now():%Y-%m-%d %H:%M}"), ("수신", q.customer or "-"),
        ("견적 수량", f"{q.sets:,} 세트"), ("연간 생산량 (금형 형식 기준)", f"{q.annual:,} 개"),
        ("자작 부품", f"{len(q.parts)} 종 · 세트당 {sum(p.qty for p in q.parts)} 개"),
        ("부자재", f"{len(q.subs)} 종"),
        ("원자재 소요 (세트당)", f"{total_kg:,.2f} kg  (견적 수량 {total_kg * q.sets:,.1f} kg)"),
        ("자재비 (세트당)", f"{qt_.won(q.material)} 원"), ("가공비 (세트당)", f"{qt_.won(q.process)} 원"),
        ("부자재비 (세트당)", f"{qt_.won(q.sub_total)} 원"), ("공정 합계", f"{q.total_ops}"),
        ("필요 톤수 최대", f"{q.max_tonf:,.1f} tf"), ("사용 기계", ", ".join(q.machine_list) or "-"),
        ("세트 단가", f"{qt_.won(q.unit_price)} 원"), ("생산 견적 합계 (VAT 포함)", f"{qt_.won(q.total)} 원"),
        ("금형 견적 (VAT 별도)", f"{qt_.won(q.die_total)} 원 · {len(q.dies)} 벌"),
    ]
    for i, (k, v) in enumerate(lines, 3):
        ws.cell(i, 1, k).font = Font(bold=True)
        ws.cell(i, 2, v)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 60
    r = len(lines) + 4
    ws.cell(r, 1, "만든 파일").font = Font(bold=True)
    for i, f in enumerate(files, r + 1):
        ws.cell(i, 1, os.path.basename(f))

    w2 = wb.create_sheet("부품별 자재비")
    table(w2, 1, ["품번", "품명", "재질·규격", "수량/세트", "소요 kg/개", "제품 kg/개", "스크랩 kg/개", "단가 원/kg",
                  "자재비/개", "자재비/세트", f"견적 수량 kg ({q.sets:,} 세트)", "견적 수량 자재비", "비고"],
          material_rows(q, settings), money=(8, 9, 10, 12))
    w2.column_dimensions["B"].width = 26
    w3 = wb.create_sheet("재질별 소요량")
    table(w3, 1, ["재질·규격", "부품 종류", "세트당 kg", f"{q.sets:,} 세트 kg", "세트당 금액", f"{q.sets:,} 세트 금액"],
          [[a, b, round(c, 3), round(d, 2), round(e), round(f)] for a, b, c, d, e, f in by_spec(q)], money=(5, 6))
    w4 = wb.create_sheet("부자재")
    table(w4, 1, ["품명", "규격", "수량/세트", "단위", "단가", "금액/세트", f"{q.sets:,} 세트 수량", "비고"],
          [[s.name, s.spec, s.qty, s.unit, s.unit_price, s.amount, s.qty * q.sets, s.note] for s in q.subs],
          money=(5, 6))
    from .company import stamp_workbook
    stamp_workbook(wb)
    wb.save(path)


def _ops_page(p, R, k, W, H, m, name, part):
    """금형 없이 만드는 부품 (레이저·절곡기·기계 가공) 작업표 한 쪽."""
    from PySide6.QtCore import QPointF, Qt
    from PySide6.QtGui import QColor, QPen

    from .process import _table
    from .report import _brand, _font
    p.setPen(QColor(20, 20, 20))
    p.setFont(_font(16, True))
    p.drawText(R(m, m, W - 2 * m, 9), Qt.AlignLeft | Qt.AlignVCenter, f"작업 공정표 — {name}  ·  {part.method}")
    p.setFont(_font(8))
    _brand(p, R(m, m + 1.5, W - 2 * m, 6))
    p.setPen(QPen(QColor(20, 20, 20), 0.4 * k))
    p.drawLine(QPointF(m * k, (m + 10) * k), QPointF((W - m) * k, (m + 10) * k))
    p.setFont(_font(10))
    p.setPen(QColor(30, 30, 30))
    p.drawText(R(m, m + 14, W - 2 * m, 7), Qt.AlignLeft | Qt.AlignVCenter,
               f"{part.spec} · 공정 {part.n_ops} · 필요 톤수 {part.need_tonf:,.1f} tf · {part.seconds:,.0f} 초/개 · "
               f"자재비 {qt_.won(part.material)} · 가공비 {qt_.won(part.process_cost)} 원/개")
    rows = [[str(i), o.op, o.kind, o.machine, f"{o.tonf:,.2f}", f"{o.seconds:,.1f}", f"{o.cost:,.0f}", o.note]
            for i, o in enumerate(part.opc, 1)]
    _table(p, R, k, m, m + 24, [("순서", 14), ("공정", 40), ("종류", 22), ("기계", 46), ("힘 tf", 22), ("초/개", 22),
                                ("원/개", 24), ("비고", 200)], rows, 9, maxy=H - m - 10)


def save_process_pdf(q: qt_.Quote, items, settings, path, title):
    """자작 부품 전체 공정도 한 권: 차례 → 부품마다 (순차 금형이면 스트립 공정도) + 자동 공정 설계 / 작업표."""
    from PySide6.QtCore import Qt
    from PySide6.QtGui import QColor

    from . import process as pr
    from . import procplan as pp
    from .process import _open_pdf, _plan_page, _route_page, _table
    from .report import _brand, _font
    by_no = {it.no: it for it in items}
    pages = []
    for part in q.parts:
        it = by_no.get(part.no)
        rt = plan = None
        if it is not None and it.result is not None and part.method != "기계 가공" and "레이저" not in part.method:
            try:
                rt = pp.route(it.result, settings, q.annual)
            except Exception:
                rt = None
            if "순차" in part.method:
                try:
                    plan = pr.progressive(it.result, settings)
                except Exception:
                    plan = None
        pages.append((part, plan, rt))
    pdf, p, R, k, W, H, m = _open_pdf(path, f"금형 공정도 - {title}")
    p.setPen(QColor(20, 20, 20))
    p.setFont(_font(18, True))
    p.drawText(R(m, m, W - 2 * m, 10), Qt.AlignLeft | Qt.AlignVCenter, f"공정도 (초안) — {title}")
    p.setFont(_font(8))
    _brand(p, R(m, m + 2, W - 2 * m, 6))
    rows, page = [], 2
    for part, plan, rt in pages:
        n = (1 if plan is not None else 0) + (1 if rt is not None else 0) or 1
        rows.append([str(part.no), part.name[:28], part.spec, str(part.qty), qt_.short_method(part.method),
                     str(part.n_ops), f"{part.need_tonf:,.1f}", ", ".join(part.machines)[:30], str(page)])
        page += n
    _table(p, R, k, m, m + 16, [("품번", 14), ("품명", 70), ("재질", 30), ("수량", 16), ("방식", 40), ("공정", 16),
                                ("톤 tf", 20), ("기계", 70), ("쪽", 14)], rows, 9, maxy=H - m - 10)
    for part, plan, rt in pages:
        name = f"{part.no}. {part.name}"
        if plan is not None:
            pdf.newPage()
            _plan_page(p, R, k, W, H, m, plan, name)
        if rt is not None:
            pdf.newPage()
            _route_page(p, R, k, W, H, m, rt, name)
        if plan is None and rt is None:
            pdf.newPage()
            _ops_page(p, R, k, W, H, m, name, part)
    p.end()
    return len(pages)


def run(q: qt_.Quote, model, parts, settings: bc.BendSettings, folder, items=None, progress=None):
    """전체 분석 파일들을 folder 에 만들고 경로 목록을 돌려준다. progress(단계 글) 는 화면 표시용."""
    from . import bom
    os.makedirs(folder, exist_ok=True)

    def step(t):
        if progress:
            progress(t)

    title = _safe(q.title)
    files = []
    step("BOM (부품 분리·묶음·분류)")
    if items is None:
        items = bom.build(model, parts, settings.mat.density, (), settings.mat.code)
    rows = bom.build_tree(model, parts, settings.mat.density, (), settings.mat.code, items=items)
    f = os.path.join(folder, f"01_{title}_BOM.xlsx")
    bom.save_tree_xlsx(rows, f, os.path.basename(model.path), settings.mat.code, settings.mat.density)
    files.append(f)
    step("공정도 (부품마다)")
    f = os.path.join(folder, f"03_{title}_공정도.pdf")
    save_process_pdf(q, items, settings, f, q.title)
    files.append(f)
    step("생산 견적서")
    for fn, name in ((qt_.save_xlsx, "04_{}_생산견적서.xlsx"), (qt_.save_pdf, "04_{}_생산견적서.pdf")):
        f = os.path.join(folder, name.format(title))
        fn(q, f)
        files.append(f)
    if q.dies:
        step("금형 견적서")
        for fn, name in ((qt_.save_die_xlsx, "05_{}_금형견적서.xlsx"), (qt_.save_die_pdf, "05_{}_금형견적서.pdf")):
            f = os.path.join(folder, name.format(title))
            fn(q, f)
            files.append(f)
    step("자재비")
    f = os.path.join(folder, f"02_{title}_자재비.xlsx")
    save_material_xlsx(q, settings, f, [f] + files)
    files.insert(1, f)
    return files
