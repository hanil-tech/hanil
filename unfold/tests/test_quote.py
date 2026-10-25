import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(__file__))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from models import nested_assembly_step  # noqa: E402

from hanil_unfold import assembly_calc, loader  # noqa: E402
from hanil_unfold import bendcalc as bc  # noqa: E402
from hanil_unfold import quote as q  # noqa: E402
from hanil_unfold.unfolder import unfold  # noqa: E402


@pytest.fixture(scope="module")
def assy(tmp_path_factory):
    path = tmp_path_factory.mktemp("q") / "truck.step"
    nested_assembly_step(str(path))
    m = loader.load(str(path), use_cache=False)
    st = bc.BendSettings(material="SPCC")
    return m, assembly_calc.compute_all(m, st), st


def test_korean_amount():
    assert q.korean_amount(0) == "영"
    assert q.korean_amount(1_038_400) == "일백삼만팔천사백"
    assert q.korean_amount(33_528_000) == "삼천삼백오십이만팔천"
    assert q.korean_amount(1_0000_0000) == "일억"
    assert q.korean_amount(2_0005_0010) == "이억오만십" or q.korean_amount(2_0005_0010) == "이억오만일십"


def test_rates_load_merges_saved():
    r = q.load_rates('{"material": {"SPCC": 999}, "profit_pct": 5, "supplier": {"전화": "02"}}')
    assert r["material"]["SPCC"] == 999 and r["material"]["SUS304"] == q.MATERIAL_PRICE["SUS304"]
    assert r["profit_pct"] == 5 and r["supplier"]["전화"] == "02" and r["supplier"]["주소"] == ""
    assert q.load_rates("깨진 글")["profit_pct"] == q.DEFAULT_RATES["profit_pct"]


def test_quote_low_volume_laser(assy):
    m, parts, st = assy
    rates = q.load_rates(None)
    qt = q.build(m, parts, st, rates, sets=10, annual=1000, title="TRUCK")
    names = {p.name: p for p in qt.parts}
    assert {"SIDE PANEL", "BRACKET", "FLOOR PLATE", "BOSS"} <= set(names)
    assert {s.name for s in qt.subs} >= {"NUT M8"}
    panel = names["SIDE PANEL"]
    assert "레이저" in panel.method and panel.ops[0] == "레이저 절단"
    assert panel.n_ops == 1 + sum(o.startswith("절곡") for o in panel.ops) == 5     # 트레이: 절곡 4
    assert panel.qty == 2 and panel.die_cost == 0
    # 재료비 = 소요 kg × 단가 − 스크랩
    assert panel.material == pytest.approx(panel.gross_kg * 1300 - (panel.gross_kg - panel.weight) * 300, rel=0.05)
    assert names["BOSS"].method == "기계 가공"
    # 합계 관계
    assert qt.manufacturing == pytest.approx(qt.material + qt.process + qt.sub_total)
    assert qt.unit_price >= qt.manufacturing * 1.21 - 10
    assert qt.supply == qt.unit_price * 10
    assert qt.total == qt.supply + qt.vat
    assert qt.total_ops == sum(p.n_ops for p in qt.parts)


def test_quote_high_volume_dies_and_overrides(assy):
    m, parts, st = assy
    rates = q.load_rates(None)
    qt = q.build(m, parts, st, rates, sets=1000, annual=300_000)
    panel = next(p for p in qt.parts if p.name == "SIDE PANEL")
    assert "순차" in panel.method and panel.die is not None and panel.die_cost > 0
    assert qt.supply == qt.unit_price * 1000 + qt.die_total
    before = qt.unit_price
    qt.amortize_die = True
    assert qt.unit_price > before and qt.supply == qt.unit_price * 1000
    panel.process_override = 0.0
    panel.material_override = 123.0
    assert panel.process_cost == 0 and panel.unit == 123.0
    qt.subs[0].unit_price = 100
    assert qt.sub_total == pytest.approx(100 * qt.subs[0].qty)


def test_tonnage_and_machine_choice(assy):
    from hanil_unfold import costing as cs
    from hanil_unfold import diecalc as dc
    m, parts, st = assy
    rates = q.load_rates(None)
    qt = q.build(m, parts, st, rates, sets=1000, annual=300_000)
    panel = next(p for p in qt.parts if p.name == "SIDE PANEL")
    # 순차 금형: 한 프레스 = 공정 힘의 합, 그 × 1.3 이상인 가장 작은 프레스
    total = sum(o.tonf for o in panel.opc)
    assert panel.need_tonf == pytest.approx(total)
    from hanil_unfold import bom
    item = next(it for it in bom.build(m, parts, st.mat.density, (), st.mat.code) if it.name == "SIDE PANEL")
    rep = dc.review(item.result, st, strip=False)
    assert total >= (rep.loads[0].force_n + rep.loads[1].force_n) / dc.TON
    press = [x for x in rates["costing"]["machines"] if x["kind"] == cs.PRESS]
    need = total * 1.3
    expect = min((x for x in press if x["capacity"] >= need), key=lambda x: x["capacity"])
    assert panel.machines == [expect["name"]]
    # 저 생산량: 레이저 + 절곡기, 절곡마다 절곡기 톤수
    low = q.build(m, parts, st, rates, sets=10, annual=1000)
    lp = next(p for p in low.parts if p.name == "SIDE PANEL")
    kinds = [o.kind for o in lp.opc]
    assert kinds[0] == cs.LASER and kinds[1:] == [cs.BRAKE] * 4
    assert all(o.tonf > 0 and o.machine.startswith("절곡기") for o in lp.opc[1:])
    # 임률을 두 배로 → 가공비 두 배 (준비 포함)
    r2 = q.load_rates(None)
    for x in r2["costing"]["machines"]:
        x["rate"] *= 2
    lp2 = next(p for p in q.build(m, parts, st, r2, sets=10, annual=1000).parts if p.name == "SIDE PANEL")
    assert lp2.process == pytest.approx(2 * lp.process)
    # 기계가 모자라면 표시
    r3 = q.load_rates(None)
    r3["costing"]["machines"] = [x for x in r3["costing"]["machines"] if x["kind"] != cs.PRESS] + [
        {"name": "작은 프레스", "kind": cs.PRESS, "capacity": 1, "rate": 10000, "spm": 30, "setup_min": 0}]
    hp = next(p for p in q.build(m, parts, st, r3, sets=10, annual=300_000).parts if p.name == "SIDE PANEL")
    assert any("⚠" in o.note for o in hp.opc)


def test_die_quote_parts(assy):
    m, parts, st = assy
    rates = q.load_rates(None)
    qt = q.build(m, parts, st, rates, sets=1000, annual=300_000)
    d = next(p.die for p in qt.parts if p.name == "SIDE PANEL")
    names = [x[0] for x in d.lines]
    assert names[:7] == ["설계", "다이 세트·플레이트", "인서트 (SKD11)", "열처리", "가공 (와이어·MCT·연삭)",
                         "표준 부품", "조립·트라이"]
    sub = sum(x[2] for x in d.lines[:7])
    assert d.cost == pytest.approx(sub * 1.1 * 1.1)
    rates["costing"]["die"]["machining_per_station"] *= 2
    d2 = next(p.die for p in q.build(m, parts, st, rates, sets=1000, annual=300_000).parts if p.name == "SIDE PANEL")
    assert d2.cost > d.cost
    supply, vat, total = q.die_total_with_vat(qt)
    assert supply == pytest.approx(qt.die_total) and total == supply + vat


def test_outputs(assy, tmp_path):
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    m, parts, st = assy
    qt = q.build(m, parts, st, q.load_rates(None), sets=5, annual=300_000, title="TRUCK")
    qt.customer = "고객사"
    q.save_xlsx(qt, str(tmp_path / "q.xlsx"))
    q.save_pdf(qt, str(tmp_path / "q.pdf"))
    q.save_die_xlsx(qt, str(tmp_path / "d.xlsx"))
    q.save_die_pdf(qt, str(tmp_path / "d.pdf"))
    import openpyxl
    wb = openpyxl.load_workbook(tmp_path / "q.xlsx")
    assert wb.sheetnames == ["견적서", "산출 근거", "공정별", "기계·임률", "단가표"]
    ws = wb["견적서"]
    texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
    assert any("일금" in t and q.korean_amount(qt.total) in t for t in texts)
    assert any("고객사" in t for t in texts)
    assert any("tf" in t for t in texts)
    assert wb["공정별"].max_row > len(qt.parts)
    wd = openpyxl.load_workbook(tmp_path / "d.xlsx")
    assert wd.sheetnames == ["금형 견적서", "금형 산출 근거", "금형 단가", "금형 간략도"]
    assert len(wd["금형 간략도"]._images) == len(qt.dies)
    assert (tmp_path / "q.pdf").stat().st_size > 5000 and (tmp_path / "d.pdf").stat().st_size > 3000


def test_saved_machines_survive_load():
    import json
    r = q.load_rates(None)
    r["costing"]["machines"] = [{"name": "우리 프레스 150", "kind": "프레스", "capacity": 150, "rate": 1, "spm": 20}]
    r["costing"]["die"]["design_per_station"] = 1
    back = q.load_rates(json.dumps(r, ensure_ascii=False))
    assert back["costing"]["machines"][0]["name"] == "우리 프레스 150"
    assert back["costing"]["die"]["design_per_station"] == 1
    assert back["costing"]["die"]["heat_kg"] == q.DEFAULT_RATES["costing"]["die"]["heat_kg"]


def test_dialog_save_all_records_latest_customer(assy, tmp_path, monkeypatch):
    from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox
    QApplication.instance() or QApplication([])
    from hanil_unfold import quote_dialog as qd
    from hanil_unfold.history import History
    m, parts, st = assy
    monkeypatch.setattr(qd.QDesktopServices, "openUrl", staticmethod(lambda *a: None))
    monkeypatch.setattr(QMessageBox, "information", staticmethod(lambda *a, **k: None))
    monkeypatch.setattr(QFileDialog, "getSaveFileName",
                        staticmethod(lambda *a, **k: (str(tmp_path / "T_견적서.xlsx"), "")))
    h = History(str(tmp_path / "기록"))
    dlg = qd.QuoteDialog(m, parts, st, (), "TRUCK", history=h)
    dlg.annual.setValue(300_000)
    dlg.customer.setText("마바상사")          # 고친 뒤 다시 계산 없이 바로 저장
    dlg.save_all()
    rows = h.search_docs()
    assert {r.data["kind"] for r in rows} == {"생산 견적", "금형 견적"}
    assert all(r.data["customer"] == "마바상사" for r in rows)
    assert (tmp_path / "T_견적서.pdf").exists() and (tmp_path / "T_금형견적서.xlsx").exists()


def test_full_analysis_package(assy, tmp_path):
    import openpyxl
    from hanil_unfold import package
    m, parts, st = assy
    qt = q.build(m, parts, st, q.load_rates(None), sets=10, annual=300_000, title="TRUCK/BOX")
    steps = []
    files = package.run(qt, m, parts, st, str(tmp_path / "out"), progress=steps.append)
    names = [os.path.basename(f) for f in files]
    assert names == ["01_TRUCK_BOX_BOM.xlsx", "02_TRUCK_BOX_자재비.xlsx", "03_TRUCK_BOX_공정도.pdf",
                     "04_TRUCK_BOX_생산견적서.xlsx", "04_TRUCK_BOX_생산견적서.pdf",
                     "05_TRUCK_BOX_금형견적서.xlsx", "05_TRUCK_BOX_금형견적서.pdf"]
    assert all(os.path.getsize(f) > 1000 for f in files) and len(steps) == 5
    wb = openpyxl.load_workbook(files[1])
    assert wb.sheetnames == ["요약", "부품별 자재비", "재질별 소요량", "부자재"]
    ws = wb["부품별 자재비"]
    total = sum(ws.cell(r, 10).value for r in range(2, ws.max_row + 1))
    assert total == pytest.approx(qt.material, abs=len(qt.parts))
    pdf = open(files[2], "rb").read()
    assert pdf.count(b"/Type /Page") >= 1 + len(qt.parts)          # 차례 + 부품마다
    # 금형이 필요 없으면 금형 견적서는 안 만든다
    low = q.build(m, parts, st, q.load_rates(None), sets=10, annual=1000, title="T")
    names = [os.path.basename(f) for f in package.run(low, m, parts, st, str(tmp_path / "low"))]
    assert not any("금형견적서" in n for n in names) and len(names) == 5


def test_die_sketch(assy):
    from hanil_unfold import diesketch
    m, parts, st = assy
    qt = q.build(m, parts, st, q.load_rates(None), sets=100, annual=300_000)
    prog = next(d for d in qt.dies if "순차" in d.die_type)
    assert prog.plan is not None and prog.press_tonf > 0
    # 금형 길이 = 공정 수 × 피치 + 여유 (간략도와 금형비가 같은 크기)
    mg = q.load_rates(None)["costing"]["die"]["margin"]
    assert prog.size[0] == pytest.approx(len(prog.plan.stations) * prog.plan.pitch + 2 * mg)
    assert prog.size[1] == pytest.approx(prog.plan.strip_width + 2 * mg)
    img = diesketch.render(prog, 800, 500)
    assert img.width() == 800
    comp = next(d for d in qt.dies if "복합" in d.die_type)
    assert comp.plan is None and comp.outline is not None
    diesketch.render(comp, 800, 500)


def test_large_die_switches_to_casting():
    """대형 금형: 공정 하나 몸체 긴 변이 1000 mm 를 넘으면 주물 (FC300·FCD) 로 계산한다."""
    from models import drawn_cup, tray
    from hanil_unfold import costing as cs, procplan as pp
    from hanil_unfold.die_cost_dialog import part_item
    st = bc.BendSettings(material="SPCC")
    out = {}
    for key, shape in (("small", tray()), ("big", drawn_cup(300, 150, 20, 1.6))):
        r = unfold(loader.list_solids(shape)[0], settings=st)
        it = part_item(r, key, 7.85)
        rt = pp.route(r, st, 100_000)
        out[key] = cs.die_quote(it, rt.die_type, [s.op for s in rt.steps], cs.default_costing(), st)
    small, big = out["small"], out["big"]
    assert small.structure == "강판" and small.cast_kg == 0
    assert big.structure == "주물" and big.cast_kg > 1000 and "1,000" in big.reason
    names = [x[0] for x in big.lines]
    assert any("FC300" in n for n in names) and any("FCD" in n for n in names)
    assert any("모형" in n for n in names) and any("NC" in n for n in names)
    # 무게 × 단가: 주물 줄 금액 = kg × 원/kg
    fc = next(x for x in big.lines if "FC300" in x[0])
    assert fc[2] == pytest.approx(float(fc[1].split("=")[-1].split("kg")[0].replace(",", "")) * 3000, rel=0.01)
    # 직접 고르기: 강판 고정이면 큰 금형도 강판, 주물 고정이면 작은 금형도 주물
    c = cs.default_costing()
    c["die"]["structure"] = 1
    r = unfold(loader.list_solids(drawn_cup(300, 150, 20, 1.6))[0], settings=st)
    rt = pp.route(r, st, 100_000)
    assert cs.die_quote(part_item(r, "b", 7.85), rt.die_type, [s.op for s in rt.steps], c, st).structure == "강판"


def test_die_cost_dialog_and_maker_stamp(tmp_path):
    from openpyxl import load_workbook
    from models import drawn_cup
    from hanil_unfold import company
    from hanil_unfold import quote as q
    from hanil_unfold.die_cost_dialog import DieCostDialog
    st = bc.BendSettings(material="SPCC")
    r = unfold(loader.list_solids(drawn_cup(300, 150, 20, 1.6))[0], settings=st)
    dlg = DieCostDialog(r, st, "big_cup.step")
    assert dlg.q.dies and dlg.q.dies[0].structure == "주물"
    assert "주물" in dlg.text.toPlainText()
    x = str(tmp_path / "d.xlsx")
    q.save_die_xlsx(dlg.q, x)
    wb = load_workbook(x)
    assert wb.properties.creator == company.NAME
    assert company.MAKER in wb.worksheets[0].oddFooter.center.text
    q.save_die_pdf(dlg.q, str(tmp_path / "d.pdf"))
    dlg.close()
    assert not company.splash_pixmap("1.0").isNull()
