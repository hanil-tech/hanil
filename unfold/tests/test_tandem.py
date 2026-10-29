"""공정별 금형 (1차 ~ N차): 접기 · 공정 짜기 · 절곡 · 캠 · 후 피어싱 금형 · 저장."""
import math
import os

import numpy as np
import pytest

import models
from hanil_unfold import bendcalc as bc
from hanil_unfold import bendie, die3d, loader, tandem
from hanil_unfold.unfolder import unfold

ST = bc.BendSettings(material="SPCC")


def channel_with_holes():
    return models.channel_holes()


@pytest.fixture(scope="module")
def chan():
    s = channel_with_holes()
    r = unfold(s, settings=ST)
    F = tandem.Folder(r)
    ops = tandem.design(r, ST, folder=F)
    return s, r, F, ops


def _bbox(s):
    from OCP.Bnd import Bnd_Box
    from OCP.BRepBndLib import BRepBndLib
    b = Bnd_Box()
    BRepBndLib.Add_s(s, b)
    p0, p1 = b.CornerMin(), b.CornerMax()
    return np.array([p1.X() - p0.X(), p1.Y() - p0.Y(), p1.Z() - p0.Z()])


@pytest.mark.parametrize("steps,width", [
    ([("L", 40), ("B", 90, 2.0), ("L", 100), ("B", 90, 2.0), ("L", 40)], 120),
    ([("L", 25), ("B", -90, 2.0), ("L", 40), ("B", 90, 2.0), ("L", 80), ("B", 90, 2.0), ("L", 40), ("B", -90, 2.0),
      ("L", 25)], 150)])
def test_fold_matches_3d(steps, width):
    """다 접은 두께 가운데 면 = 3D 모델 (상자 크기, 두께 가운데 → 바깥 t/2 · K 만큼 차이)."""
    s = models.profile_part(steps, 2.0, width)
    r = unfold(s, settings=ST)
    F = tandem.Folder(r)
    P = np.vstack(F.wire({b.number: b.angle_deg for b in r.bends}))
    got = np.sort(np.ptp(P, 0))
    want = np.sort(_bbox(s))
    # 판 끝이 상자를 정하는 쪽은 그대로, 면이 정하는 쪽은 t (가운데 면) — 0.3 (K 0.44 의 호 차이) 안
    assert np.all(np.abs(got - want) <= 2.0 + 0.3)
    assert not F.loose and len(F.panels) == len(r.bends) + 1


def test_fold_is_isometric():
    """접어도 외곽 길이 그대로 (반만 접은 공정도) — 절곡 띠는 호 + 남은 곧은 부분."""
    s = models.profile_part([("L", 60), ("B", 180, 1.0), ("L", 15)], 1.0, 80)
    r = unfold(s, settings=ST)
    F = tandem.Folder(r)
    D = F.dense(F.outer, 0.2)

    def per(Q):
        return float(np.sum(np.linalg.norm(np.diff(np.vstack([Q, Q[:1]]), axis=0), axis=1)))
    flat = per(np.column_stack([D, np.zeros(len(D))]))
    for ang in (45.0, 90.0, 135.0, 180.0):
        assert per(F.fold(D, {1: ang})) == pytest.approx(flat, rel=2e-3)


def test_plan_one_to_five(chan):
    """구멍 있는 ㄷ 채널: 1차 블랭킹 → 2차 피어싱 (먼 구멍) → 3차 U 굽힘 → 4차 후 피어싱 (R3) → 5차 캠 피어싱 (R4, 맞은편 둘)."""
    _, _, F, ops = chan
    assert [op.kind for op in ops] == [tandem.BLANK, tandem.PIERCE, tandem.UBEND, tandem.POST, tandem.CAM]
    assert [op.label.split()[0] for op in ops] == ["1차", "2차", "3차", "4차", "5차"]
    assert len(ops[2].bends) == 2 and len(ops[4].bends) == 2 and "R4" in ops[4].rules
    assert len(ops[1].holes) == 1 and len(ops[3].holes) == 1 and len(ops[4].holes) == 2
    assert all(op.die is not None and not op.error for op in ops)


@pytest.mark.parametrize("steps,t,kinds", [
    ([("L", 60), ("B", 180, 1.0), ("L", 15)], 1.0, [tandem.BLANK, tandem.BEND, tandem.PREHEM, tandem.HEM]),
    ([("L", 60), ("B", 120, 2.0), ("L", 30)], 2.0, [tandem.BLANK, tandem.BEND, tandem.FINISH]),
    ([("L", 60), ("B", 45, 2.0), ("L", 40)], 2.0, [tandem.BLANK, tandem.ANGLE])])
def test_plan_angles_and_hem(steps, t, kinds):
    """95° 넘으면 90° 와이핑 + 마무리, 150° 넘으면 + 예비 헤밍 + 헤밍, 75° 아래는 각도 굽힘 — 공정마다 금형이 됨."""
    s = models.profile_part(steps, t, 80)
    r = unfold(s, settings=ST)
    ops = tandem.design(r, ST)
    assert [op.kind for op in ops] == kinds
    assert all(op.die is not None for op in ops), [op.error for op in ops]
    for op in ops[1:]:
        assert all(c[3] == "적합" for c in op.die.checks), (op.label, [c for c in op.die.checks if c[3] != "적합"])


def test_u_bend_die(chan):
    """U 굽힘 금형: 다이 모서리 R = 제품 안쪽 R, 펀치 틈 1.05 t, 펀치 바깥에 힐 블록, 굽힘력 = 0.33 UTS L t² / W × 2."""
    _, r, F, ops = chan
    a = ops[2].die
    t = r.thickness
    die = a.plate("U 벤딩 다이").prims[0]
    P = np.asarray(die.poly)
    assert np.ptp(P[:, 0]) == pytest.approx(100 + 2 * 2.0, abs=0.05)          # 채널 안쪽 폭 (바닥 100 + 양쪽 R2)
    punches = [p.prims[0] for p in a.parts if p.name.startswith("와이핑 펀치")]
    assert len(punches) == 2
    right = max(punches, key=lambda q: np.asarray(q.poly)[:, 0].mean())
    gap = np.asarray(right.poly)[:, 0].max() - P[:, 0].max()
    assert gap > 0
    inner = np.asarray(right.poly)[np.asarray(right.poly)[:, 1] > np.asarray(right.poly)[:, 1].min() + 30][:, 0].min()
    assert inner - P[:, 0].max() == pytest.approx(1.05 * t, abs=0.6)          # 틈 (위쪽, 스프링백 기울임 빼고)
    heels = a.plate("힐 블록").prims
    assert len(heels) == 2 and min(h.p[0] for h in heels if h.p[0] > 0) >= np.asarray(right.poly)[:, 0].max() - 1e-6
    L = r.bends[0].length
    W = r.bends[0].inner_radius + 1.05 * t + max(2 * t, 3.0)
    assert a.force_n == pytest.approx(2 * bendie.K_WIPE * ST.mat.tensile_mpa * L * t * t / W, rel=1e-6)
    assert all(c[3] == "적합" for c in a.checks), [c for c in a.checks if c[3] != "적합"]
    # 금형 안 제품: 공정 전 (평판) · 뒤 (U) 선
    tags = {tag for _, tag in a.work}
    assert tags == {"before", "after"}


def test_bend_die_step_valid(chan, tmp_path):
    from OCP.BRepCheck import BRepCheck_Analyzer
    _, _, _, ops = chan
    for op in (ops[2], ops[4]):
        p = die3d.save_step(op.die, str(tmp_path / f"{op.no}.step"))
        m = loader.load(p, use_cache=False)
        assert len(m.solids) == sum(len(x.prims) for x in op.die.parts)
        assert all(BRepCheck_Analyzer(s_).IsValid() for s_ in m.solids)


def test_cam_and_post_dies(chan):
    """캠 피어싱: 앤빌 옆면 다이 구멍 2 (양쪽), 캠 슬라이더 · 드라이버 2 벌, 가로 펀치. 후 피어싱: 다이가 채널 안쪽 폭에 맞음."""
    _, r, F, ops = chan
    cam = ops[4].die
    anvil = cam.plate("캠 피어싱 앤빌 (다이)").prims[0]
    side = [h for h in anvil.holes if h.axis in ("+x", "-x") and h.kind == "다이 구멍"]
    assert len(side) == 2 and {h.axis for h in side} == {"+x", "-x"}
    assert sum(1 for p in cam.parts if p.name.startswith("캠 슬라이더")) == 2
    assert sum(1 for p in cam.parts if p.name == "캠 드라이버") == 2
    pun = [p for p in cam.parts if p.name.startswith("캠 피어싱 펀치")]
    assert sum(len(p.prims) for p in pun) == 2 and all(q.ax == "x" for p in pun for q in p.prims)
    post = ops[3].die
    dp = post.plate("다이 플레이트").prims[0]
    inner = 100.0 - 2 * r.thickness
    assert min(dp.p[3], dp.p[4]) <= inner                     # 채널 안 (플랜지 사이) 에 들어감
    back = post.plate("다이 받침판").prims[0]
    assert dp.p[5] + back.p[5] >= 40.0                       # 플랜지 40 이 하형 세트에 닿지 않게 받침판을 높임


def test_save_all(chan, tmp_path):
    import ezdxf
    from openpyxl import load_workbook
    s, r, F, ops = chan
    out, notes = tandem.save_all(r, ST, ops, str(tmp_path / "ch"), {"product": "ch", "material": "SPCC t2", "solid": s,
                                                                  "result": r}, folder=F)
    names = [os.path.basename(p) for p in out]
    assert sum(n.endswith(".step") for n in names) == 5
    assert "ch_03차_U 굽힘.step" in names
    pdf = next(p for p in out if p.endswith(".pdf"))
    assert open(pdf, "rb").read(4) == b"%PDF"
    dxf = next(p for p in out if p.endswith(".dxf"))
    assert not ezdxf.readfile(dxf).audit().errors
    wb = load_workbook(next(p for p in out if p.endswith(".xlsx")))
    assert wb.sheetnames[0] == "공정표" and "5차 부품표" in wb.sheetnames
    assert wb["공정표"].cell(8, 2).value == "캠 피어싱"
    sh = tandem.line_sheets(F, ops, {"product": "ch"}, "HD")
    texts = [e[3] for e in sh[0].ents if e[0] == "text"]
    assert any(t.startswith("5차 캠 피어싱") for t in texts)


def test_dialog_tandem(tmp_path, monkeypatch):
    from PySide6.QtGui import QDesktopServices
    from PySide6.QtWidgets import QFileDialog, QInputDialog, QMessageBox

    from hanil_unfold import die_dialog, drawfmt
    s = models.profile_part([("L", 30), ("B", 90, 2.0), ("L", 60), ("B", 90, 2.0), ("L", 30)], 2.0, 60)
    r = unfold(s, settings=ST)
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *a, **k: (str(tmp_path / "ㄷ_공정별금형.pdf"), ""))
    monkeypatch.setattr(QInputDialog, "getItem", lambda *a, **k: ("m (보통)", True))
    shown = {}
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: shown.setdefault("msg", a[2]))
    monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: shown.setdefault("err", a[2]))
    monkeypatch.setattr(QDesktopServices, "openUrl", lambda *a, **k: True)
    monkeypatch.setattr(drawfmt, "find_dwg_writer", lambda user_path="": None)
    dlg = die_dialog.DieDialog(r, ST, "ㄷ.step", solid=s)
    dlg.save_tandem()
    assert "err" not in shown, shown.get("err")
    assert "1차 블랭킹" in shown["msg"] and "2차 U 굽힘" in shown["msg"]
    assert (tmp_path / "ㄷ_공정별금형_도면집.pdf").exists() and (tmp_path / "ㄷ_02차_U 굽힘.step").exists()
    dlg.close()


def test_rect_fit():
    r, left = bendie._rect_fit([0, 100, 0, 50], [(50, 48), (95, 25)], [0, 60, 10, 40], 2.0)
    assert left == 0 and r[3] <= 46 and r[1] <= 93
    assert math.isclose(r[0], 0) and math.isclose(r[2], 0)


def test_blank_layouts(chan):
    """1차 블랭킹 배치: 가로 · 세로 · 2열 후보, 이용률 ≤ 100 %, 2열은 두 블랭크가 겹치지 않음, 재료 적은 순 (2열은 3 % 넘게 줄 때만)."""
    _, r, F, ops = chan
    lays = tandem.blank_layouts(r, r.thickness)
    names = [c.name for c in lays]
    assert {"가로", "세로", "가로 2열", "세로 2열"} <= set(names)
    for c in lays:
        assert 0.3 < c.util <= 1.0
        assert c.kg_per_part == pytest.approx(c.pitch * c.width / c.rows * r.thickness * die3d.DENSITY)
        if c.rows == 2:
            A, B = c.blanks
            assert tandem._clear(A, B, 1.0) and tandem._clear(A + (c.pitch, 0), B, 1.0)
            assert np.ptp(np.vstack(c.blanks)[:, 1]) + 2 * max(2 * r.thickness, 1.5) == pytest.approx(c.width)
    one = [c for c in lays if c.rows == 1]
    assert one[0].name == "세로" and one[0].strip_per_part <= one[1].strip_per_part        # ㄷ 채널은 세로 (189 폭 → 피치 123.5)
    assert lays[0].rows == 1                                                              # 2열 이득 < 3 % → 1열
    op = ops[0]
    assert op.layout is lays[0] or op.layout.name == lays[0].name
    assert "세로 배치" in op.detail and "이용률" in op.detail
    assert op.die.strip[0] == pytest.approx(op.layout.pitch) and op.die.strip[1] == pytest.approx(op.layout.width)
    assert len(op.die.stations) == 3                                                      # 배치도에 앞 · 뒤 피치
    assert any("배치 비교" in n for n in op.die.notes)


def test_two_row_when_it_pays(monkeypatch):
    """2열 이득이 문턱을 넘으면 2열 — 블랭킹 금형에 펀치 둘, 스트립 폭이 두 줄."""
    s = models.profile_part([("L", 30), ("B", 90, 2.0), ("L", 60), ("B", 90, 2.0), ("L", 30)], 2.0, 60)
    r = unfold(s, settings=ST)
    monkeypatch.setattr(tandem, "TWO_ROW_GAIN", -1.0)                                     # 늘 2열을 앞에
    lays = tandem.blank_layouts(r, r.thickness)
    assert lays[0].rows == 2
    ops = tandem.design(r, ST)
    op = ops[0]
    assert op.layout.rows == 2 and op.die is not None, op.error
    assert op.die.plate("블랭킹 펀치").qty == 1 and sum(1 for p in op.die.parts if p.name == "블랭킹 펀치") == 2
    assert len(op.die.stations) == 6


def test_long_channel_dies_scale_with_product():
    """큰 제품 (1.2 m 채널): 블랭킹 (스트리퍼 볼트 긴 것 · 주문 길이) 과 U 굽힘 (좁은 패드에도 가스 스프링) 금형이 다 나옴."""
    r = unfold(models.sharp_channel(w=100.0, h=40.0, t=3.2, length=1200.0), settings=ST)
    ops = tandem.design(r, ST)
    for op in ops:
        assert op.die is not None, (op.kind, op.error)
    blank = next(op for op in ops if op.kind == tandem.BLANK).die
    assert max(blank.L, blank.W) > 1200 and any("주문 길이" in p.size for p in blank.parts if p.name.startswith("스트리퍼 볼트"))
    ub = next(op for op in ops if op.kind == tandem.UBEND).die
    assert ub.W > 1200 and any("가스 스프링" in p.name for p in ub.parts)


def test_pad_spring_pack_grid():
    """좁고 긴 패드 (채널 안 92 × 326): 고른 격자로 Ø63 가스 스프링 3 개가 들어감 (전에는 욕심껏 골라 2 개)."""
    from hanil_unfold import bendie
    rect = [-93.1, -0.5, -313.0, 13.0]
    bolts = [(-46.8, -289.0), (-46.8, -11.0)]
    assert len(bendie._pack(rect, 63.0, 8.0, bolts, 12.0)) >= 3
    gas, xy = bendie._pick_gas(25801.0, 46.8, rect, bolts, 12.0)
    assert gas is not None and len(xy) * gas.f_each >= 25801.0
