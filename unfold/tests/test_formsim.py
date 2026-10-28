"""제품 성형 해석 (formsim): 성형성 존, 절곡 균열, 하중, 저장, 창."""
import math
import os

import numpy as np
import pytest

from hanil_unfold import bendcalc as bc
from hanil_unfold import forming
from hanil_unfold import formsim as fs
from hanil_unfold import loader
from hanil_unfold.unfolder import unfold


def solid(shape):
    return loader.list_solids(shape)[0]


@pytest.fixture(scope="module")
def cup():
    from models import drawn_cup
    st = bc.BendSettings(material="SPCC")
    r = unfold(solid(drawn_cup()), settings=st, mode="forming")
    return r, st, fs.analyze(r, st)


def test_cup_zones_and_thickness(cup):
    r, st, fa = cup
    assert fa.mode == "forming"
    assert sum(fa.zone_frac) == pytest.approx(1.0, abs=1e-6)
    assert fa.zone_frac[0] < 0.001                      # 드로잉률 0.59 의 컵은 터지지 않는다
    assert fa.verdict != fs.NG
    # 딥 드로잉 컵: 벽 위쪽이 두꺼워지고 (압축), 바닥은 거의 그대로
    assert -fa.thin_lo > 0.05
    assert fa.zone_frac[5] + fa.zone_frac[6] > 0.15
    names = [n for n, *_ in fa.loads]
    assert "드로잉 (펀치) 힘" in names and "주름 누르개 힘" in names
    # 드로잉력 = 둘레 × t × σb × n(m): Ø100 컵 (가운데 면 Ø99) 둘레 ≈ 311 mm
    f = dict((n, v) for n, v, _ in fa.loads)["드로잉 (펀치) 힘"]
    # 교재 식 (π·d·t·σb·n) 과 Siebel 식 중 큰 값 — 두 식은 같은 크기 (±40 %)
    assert 311 * 1.0 * 414 * 0.6 < f < 311 * 1.0 * 414 * 1.4
    assert fa.ldr and fa.ldr[0] < fa.ldr[1]                 # β ≈ 1.7 < LDR 2.2 (SPCC r 1.3)
    assert "성형 해석" in fs.summary_html(fa)


def test_cup_mesh_has_no_wall_slivers(cup):
    """긴 모서리 나누기: 원통 벽의 벽 높이만 한 바늘 삼각형이 없어지고 변형이 벽 높이에 따라 고르게 바뀐다."""
    r, _, fa = cup
    z = fa.fi.V[fa.fi.F].mean(1) @ np.array([0, 0, 1.0])
    th = fa.fi.thinning
    band = (z > 20) & (z < 35)                          # 벽 가운데 띠: 같은 높이면 두께 변화도 비슷해야
    assert band.sum() > 200
    assert np.std(th[band]) < 0.04
    top = np.median(th[z > z.max() - 8])
    bottom = np.median(th[z < z.min() + 1])
    assert top < -0.05 < bottom                         # 위 끝은 두꺼워지고 바닥은 거의 그대로


def test_one_step_cup_blank_close_to_volume_constant(cup):
    r, _, fa = cup
    d_equal = math.sqrt(4 * fa.fi.product_area / math.pi)
    assert abs(max(r.size) - d_equal) < 0.03 * d_equal
    # 블랭크 넓이 = 제품 넓이 × 평균 (t / t0): 벽이 두꺼워지는 만큼 블랭크가 조금 크다
    assert 1.0 <= fa.fi.blank_area / fa.fi.product_area < 1.06


def test_sheet_bends_crack_check():
    from models import tray
    st = bc.BendSettings(material="SPCC")
    r = unfold(solid(tray()), settings=st)              # R 1.6, t 1.6 (R/t 1)
    fa = fs.analyze(r, st)
    assert fa.mode == "sheet" and fa.bends
    assert fa.verdict == fs.OK
    c = fa.bends[0]
    assert c.strain == pytest.approx(math.log(1 + 1 / (2 * 1.0 + 1)), rel=1e-6)
    assert c.die_angle > c.angle                       # 스프링백만큼 더 접는다
    st6 = bc.BendSettings(material="AL6061")            # 6061-T6 최소 R 2.5 t → R/t 1 은 균열
    fa6 = fs.analyze(unfold(solid(tray()), settings=st6), st6)
    assert fa6.verdict == fs.NG
    assert fa6.bends[0].r_min == pytest.approx(2.5 * 1.6)
    assert any("균열" in m for m in fa6.messages)


def test_save_xlsx_pdf_and_report_sheet(cup, tmp_path):
    from openpyxl import load_workbook

    from hanil_unfold import report
    r, st, fa = cup
    x, p = str(tmp_path / "a.xlsx"), str(tmp_path / "a.pdf")
    fs.save_xlsx(fa, x, "cup")
    fs.save_pdf(fa, p, "cup")
    assert os.path.getsize(p) > 10000
    wb = load_workbook(x)
    assert {"성형 해석", "제품 그림"} <= set(wb.sheetnames)
    d = report.ReportData("cup", "", r, st)
    y = str(tmp_path / "rep.xlsx")
    report.save_xlsx(d, y)
    assert "성형 해석" in load_workbook(y).sheetnames


def test_render_views(cup):
    _, _, fa = cup
    for v in fs.VIEWS:
        img = fs.render_product(fa, v, 400, 300)
        assert img.width() == 400
        # 흰 바탕만이 아니어야 한다
        assert img.pixelColor(200, 150).name() != "#ffffff"


def test_dialog_builds(cup):
    from hanil_unfold.formsim_dialog import FormSimDialog
    r, st, _ = cup
    dlg = FormSimDialog(r, st, "cup")
    assert dlg.fa.verdict in (fs.OK, fs.WARN, fs.NG)
    dlg.view.set_view(fs.VIEWS[1])
    dlg.close()


def test_burring_edge_stretch_vs_hole_expansion_ratio():
    """구멍 버링: 세운 가장자리 늘림 ≈ R/(R−H) − 1 을 구멍 확장률 HER (ISO 16630) 과 견줌."""
    from models import burring
    sh = burring()
    out = {}
    for mat in ("SPCC", "SPFC590"):
        st = bc.BendSettings(material=mat)
        r = unfold(solid(sh), settings=st, mode="forming")
        out[mat] = fs.analyze(r, st)
    e, where, lim = out["SPCC"].edge
    assert where == "구멍 가장자리" and 0.4 < e < 0.7 and lim == 1.0
    assert out["SPCC"].edge_verdict == fs.OK
    assert out["SPFC590"].edge_verdict == fs.NG and out["SPFC590"].verdict == fs.NG
    assert any("HER" in m for m in out["SPFC590"].messages)


def test_press_tonnage_curve():
    """크랭크 프레스: 하사점 앞 30° 까지 정격, 그 위는 sin 30° / sin α 로 줄어듦."""
    import math
    from hanil_unfold import costing as cs
    m = {"capacity": 200, "stroke": 200}
    h30 = 100 * (1 - math.cos(math.radians(30)))
    assert cs.press_factor(m, h30 * 0.99) == 1.0
    assert cs.press_factor(m, 100) == pytest.approx(0.5)               # 90°
    h60 = 100 * (1 - math.cos(math.radians(60)))
    assert cs.press_factor(m, h60) == pytest.approx(0.5 / math.sin(math.radians(60)))
    # 깊은 드로잉은 같은 힘이라도 더 큰 프레스 (또는 긴 행정)
    ms = cs.default_costing()["machines"]
    a, _ = cs.pick(ms, cs.PRESS, 50)
    b, _ = cs.pick(ms, cs.PRESS, 50, height=40, depth=50)
    assert b["capacity"] > a["capacity"] and b["stroke"] >= 100


def test_explanation_locates_problems_and_gives_actions(cup, tmp_path):
    """성형 해석 설명: 문제 부위마다 위치 · 값 · 뜻 · 원인 · 대책, 그림 읽는 법, 엑셀 [해석 설명] 시트 · PDF 설명 쪽."""
    from openpyxl import load_workbook

    from hanil_unfold import formexplain as fx
    r, st, fa = cup
    ex = fa.explain
    assert ex is not None and ex.headline and ex.reasons
    kinds = {row[0] for row in ex.regions}
    assert {"바닥 (펀치 면)", "벽"} <= kinds                      # 컵: 바닥 · 벽 을 나눔
    assert sum(f for _, f in ex.thickness) == pytest.approx(1.0, abs=0.01)
    for i in ex.issues:
        assert i.region and i.value and i.meaning and i.cause and i.actions
    assert set(fx.VIEW_HELP) >= set(fs.VIEWS)
    assert "그림 읽는 법" in fx.html(fa)
    fs.save_xlsx(fa, tmp_path / "e.xlsx", "cup")
    assert "해석 설명" in load_workbook(tmp_path / "e.xlsx").sheetnames
    fs.save_pdf(fa, str(tmp_path / "e.pdf"), "cup")
    assert (tmp_path / "e.pdf").stat().st_size > 10_000


def test_explanation_sheet_bends():
    """절곡 부품: 최소 R 보다 작은 절곡마다 원인 · 대책 설명."""
    from models import tray
    st = bc.BendSettings(material="SPFC980")
    fa = fs.analyze(unfold(solid(tray(r=1.0)), settings=st), st)
    ex = fa.explain
    assert fa.verdict == fs.NG and ex.headline.startswith("부적합")
    assert ex.issues and all(i.region.startswith("절곡 B") and "최소 R" in i.cause for i in ex.issues)
    assert any("스프링백" in row[0] for row in ex.process)


def test_explanation_deep_oil_pan_corner():
    """깊은 오일 팬: 파단을 귀 (모서리) 로 찾고 다단 드로잉 · 귀 R 대책."""
    from models import oil_pan
    st = bc.BendSettings(material="SPCC")
    fa = fs.analyze(unfold(solid(oil_pan(H=160, r_corner=15, flange=0)), settings=st), st)
    ex = fa.explain
    crack = [i for i in ex.issues if i.zone == 0]
    assert crack and crack[0].corner and "귀 R" in " ".join(crack[0].actions)
    assert any("다단" in a for a in ex.advice)
    img = fs.render_product(fa, fs.VIEWS[0], 400, 300)           # 번호 표시 포함
    assert not img.isNull()


def test_measured_flc_curve(cup):
    """실측 FLC (ISO 12004-2 점들): 점 사이 보간·밖은 끝 기울기, 해석이 Keeler 대신 씀, 단위 바꾸기."""
    from hanil_unfold import flcdata, forming
    r, st, fa0 = cup
    pts = flcdata.parse("-0.3 0.65\n-0.1, 0.42\n0\t0.30\n0.1 0.33\n0.3 0.42", "true")
    assert flcdata.check(pts)[0] and pts[2] == (0.0, 0.30)
    assert flcdata.parse("-30 90\n0 35\n30 52", "pct")[1][1] == pytest.approx(math.log(1.35))
    assert not flcdata.check([(0.0, 0.3), (0.1, 0.33)])[0]
    try:
        forming.MEASURED_FLC["SPCC"] = pts
        lim = forming.flc_limit(np.array([-0.2, 0.0, 0.2, 0.5]), 0.29, "SPCC")
        assert lim[0] == pytest.approx((0.65 + 0.42) / 2) and lim[1] == pytest.approx(0.30)
        assert lim[3] == pytest.approx(0.42 + (0.42 - 0.33) / 0.2 * 0.2)              # 오른쪽 끝 기울기로
        fa = fs.analyze(r, st)
        assert fa.flc0 == pytest.approx(0.30) and any("실측 FLC" in m for m in fa.messages)
        assert fa.risk_top != pytest.approx(fa0.risk_top)
    finally:
        forming.MEASURED_FLC.pop("SPCC", None)


def test_section_thickness_graph(cup):
    """가운데 X 단면: 컵 두께 가운데 단면 길이 (벽 2 + 모서리 R 2 + 바닥) 와 맞고, 두께 값은 해석 범위 안."""
    r, st, fa = cup
    sec = fs.section(fa, "X")
    Rm, rr, h = 50 - 0.5, 8 + 0.5, 50 - 0.5
    expect = 2 * (h - rr) + math.pi * rr + 2 * (Rm - rr)
    assert sec.s[-1] == pytest.approx(expect, rel=0.02)
    assert sec.thickness.min() >= fa.t_min - 0.02 and sec.thickness.max() <= fa.t_max + 0.02
    assert sec.pts[0, 1] < sec.pts[-1, 1]                                        # 그림 왼쪽에서 시작
    assert fs.section(fa, "X", 10_000) is None                                    # 제품 밖
    img = fs.render_section(sec, 600, 400)
    assert not img.isNull() and "단면" in fs.section_note(sec)
    assert len(sec.rows()[0]) == len(fs.SECTION_HEAD)


def test_load_stroke_and_press(cup):
    r, st, fa = cup
    ls = fs.load_stroke(fa, st.mat)
    assert ls is not None and ls.h[0] == 0 and ls.h[-1] == pytest.approx(ls.depth)
    assert ls.draw.max() == pytest.approx(fa.draw_geom["f_draw"])                 # 곡선 최대 = 해석 드로잉력
    assert ls.holder[0] >= ls.holder[-1] >= 0                                     # 플랜지가 들어가며 누르개 넓이 줄어듦
    if ls.ok:
        assert np.all(ls.cap >= ls.total) and ls.margin[0] >= 1
    assert "최대 하중" in ls.note and not fs.render_load_stroke(ls, 500, 300).isNull()


def test_compare_materials_and_thickness():
    from models import drawn_cup
    st = bc.BendSettings(material="SPCC")
    sol = solid(drawn_cup())
    rows = fs.compare(sol, st, [("SPCC", None), ("SPFC590", None), ("SPCC", 1.4)])
    assert [r[0] for r in rows] == ["SPCC", "SPFC590", "SPCC"] and rows[2][1] == pytest.approx(1.4)
    assert all(r[2] in (fs.OK, fs.WARN, fs.NG) for r in rows) and len(rows[0]) == len(fs.COMPARE_HEAD)
    assert fs.best_row(rows) in rows


def test_dialog_new_tabs(cup):
    from hanil_unfold.report import ensure_gui
    ensure_gui()
    from PySide6.QtWidgets import QApplication, QTabWidget
    QApplication.instance() or QApplication([])
    from hanil_unfold.formsim_dialog import FormSimDialog
    from models import drawn_cup
    r, st, _ = cup
    d = FormSimDialog(r, st, "cup", solid=solid(drawn_cup()))
    names = [d.findChildren(QTabWidget)[0].tabText(i) for i in range(d.findChildren(QTabWidget)[0].count())]
    for want in ("성형 과정 (3D)", "단면 (두께 그래프)", "하중 - 행정", "FLC 입력 (실측)", "재질·두께 비교"):
        assert want in names


def test_forming_process_onsets_and_frames(tmp_path):
    """성형 과정: 깊은 오일 팬은 행정 도중 귀에서 먼저 터지기 시작 · 주름 시작, 처음 = 평판 (누르개 높이), 끝 = 제품."""
    from models import oil_pan

    from hanil_unfold import formprocess as fp
    st = bc.BendSettings(material="SPCC")
    r = unfold(solid(oil_pan(H=160, r_corner=15, flange=0)), settings=st, mode="forming")
    fa = fs.analyze(r, st)
    pr = fp.build(fa)
    on = {o.kind: o for o in pr.onsets}
    assert "crack" in on and 0 < on["crack"].s < 1 and "귀" in on["crack"].region
    assert on["risk"].s <= on["crack"].s and "wrinkle" in on
    assert on["crack"].depth == pytest.approx(on["crack"].s * pr.depth)
    xx, yy, zz = pr.frame
    V0, V1 = fp.stage_vertices(pr, 0.0), fp.stage_vertices(pr, 1.0)
    assert np.ptp(V0 @ zz) == pytest.approx(0, abs=1e-6)                    # 처음은 평판
    assert np.allclose(V1, pr.Vf)                                            # 끝은 제품
    z0, _ = fp.stage_zones(pr, 0.0)
    assert (z0 == 4).all()                                                   # 처음엔 모두 늘림 부족 (변형 0)
    lines = fp.summary(pr)
    assert len(lines) == 3 and "터지기 시작" in lines[0]
    g = fp.save_gif(pr, str(tmp_path / "p.gif"), frames=4, w=200, h=150)
    from PIL import Image
    assert Image.open(g).n_frames >= 5
    assert any("성형 과정" in t for t in fa.explain.reasons)


def test_forming_process_safe_cup(cup):
    from hanil_unfold import formprocess as fp
    r, st, fa = cup
    pr = fp.build(fa)
    assert "끝까지 터지지 않음" in fp.summary(pr)[0]
    assert not fp.render_stage(pr, 0.5, 300, 220).isNull()


def test_process_play_rotate_controls(cup):
    """성형 과정 탭: ▶ 재생은 행정을 올리고 끝에서 반복, 🔄 회전은 방향을 돌림, 멈추면 원래 메시로 다시 그림.
    제품 3D 화면의 🔄 회전도."""
    from hanil_unfold.report import ensure_gui
    ensure_gui()
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from hanil_unfold.formsim_dialog import ProcessPanel, ProductView
    _, _, fa = cup
    pp = ProcessPanel(fa)
    pp.toggle_play()
    assert pp.playing and pp.slider.value() == 0 and pp.timer.isActive()
    for _ in range(10):
        pp._tick()
    assert pp.slider.value() == 20
    for _ in range(60):                                   # 끝에 닿으면 잠깐 멈췄다가 처음부터 (반복)
        pp._tick()
    assert pp.playing and pp.slider.value() < 100
    pp.toggle_rotate()
    az0 = pp.az
    pp._tick()
    assert pp.rotating and pp.az == pytest.approx((az0 + pp.TURN) % 360)
    pp.loop.setChecked(False)
    pp.toggle_play()
    pp.toggle_rotate()
    assert not pp.busy and not pp.timer.isActive()
    pp._go(100)
    pp._go_crack()                                         # 안전한 컵: 터지는 곳 없음 → 그대로
    assert pp.slider.value() == 100
    v = ProductView(fa)
    v.resize(400, 300)
    v.set_spin(True)
    a0 = v.az
    v._spin_step()
    assert v.az == pytest.approx((a0 + 3.0) % 360) and v._pix is not None
    v.set_spin(False)
    assert not v._spin.isActive()
    # 줄인 메시: 삼각형이 줄고 원래 번호를 가리킴
    from hanil_unfold import formsim as fs_
    inv, Fc, keep = fs_.cluster_mesh(fa.fi.V, fa.fi.F, target=500)
    assert len(Fc) < len(fa.fi.F) and keep.max() < len(fa.fi.F) and len(inv) == len(fa.fi.V)


def test_product_view_zoom_pan(cup):
    """성형 해석 그림: 휠로 커서 자리를 중심으로 확대 · 축소, 오른쪽 단추로 끌어 이동, 두 번 누르면 처음대로."""
    from hanil_unfold.report import ensure_gui
    ensure_gui()
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from hanil_unfold import formprocess as fp
    from hanil_unfold.formsim_dialog import ProductView, zoom_at
    _, _, fa = cup
    # zoom_at: 커서 (x, y) 가 가리키던 그림 점이 확대 뒤에도 같은 자리
    z, pan = zoom_at(1.0, (0.0, 0.0), 120, 300.0, 100.0, 200.0, 150.0)
    assert z == pytest.approx(1.15)
    # 그림 점 q 의 화면 자리: c + pan + q·z  — 커서 자리에 있던 점 q = (300-200, 100-150) = (100, -50)
    assert 200 + pan[0] + 100 * z == pytest.approx(300.0) and 150 + pan[1] - 50 * z == pytest.approx(100.0)
    z2, _ = zoom_at(19.0, (0.0, 0.0), 1200, 0, 0, 0, 0)
    assert z2 == 20.0                                       # 상한
    v = ProductView(fa)
    v.resize(400, 300)
    v._render()
    base = v._pix.toImage()
    v.zoom, v.pan = 3.0, (40.0, -20.0)
    v._render()
    assert v._pix.toImage() != base and v._pix.width() == v.width()
    v.mouseDoubleClickEvent(None)
    assert v.zoom == 1.0 and v.pan == (0.0, 0.0) and v._pix.toImage() == base
    # 성형 과정 그림도 확대 · 이동
    pr = fp.build(fa)
    a = fp.render_stage(pr, 1.0, 400, 300)
    b = fp.render_stage(pr, 1.0, 400, 300, zoom=2.5, pan=(30.0, 10.0))
    assert a != b and b.width() == 400
