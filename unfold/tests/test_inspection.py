"""검사 성적서: 도면에서 검사 항목·보통 공차, 측정값 판정, Cp·Cpk, 엑셀 (식·이력 시트) · PDF · 다시 읽기."""
import os

import pytest

from hanil_unfold import bendcalc as bc
from hanil_unfold import inspection as ip
from hanil_unfold import loader
from hanil_unfold.unfolder import unfold


@pytest.fixture(scope="module")
def tray_ins():
    from models import tray
    st = bc.BendSettings(material="SPCC")
    sol = loader.list_solids(tray())[0]
    r = unfold(sol, settings=st)
    return r, st, sol, ip.build(r, st, sol, None, "트레이")


def test_plan_from_drawing(tray_ins):
    r, st, sol, ins = tray_ins
    groups = {c.group for c in ins.chars}
    assert groups == {"재료", "블랭크", "완성품", "외관"}
    t = next(c for c in ins.chars if c.name == "소재 두께")
    assert t.nominal == pytest.approx(1.6) and t.upper == pytest.approx(0.08)      # JIS G 3141 1.25~1.6
    w = next(c for c in ins.chars if c.name == "전개 가로")
    assert w.upper == pytest.approx(0.5) and "JIS B 0408-B" in w.source             # 120~400 B 급 ±0.5
    angles = [c for c in ins.chars if "절곡 각도" in c.name]
    assert len(angles) == len(r.bends) and all(c.nominal == pytest.approx(90) and c.unit == "°" for c in angles)
    assert any(c.name.startswith("제품 높이") for c in ins.chars)                   # 3D 바깥 치수
    assert all(len(c.values) == ip.DEFAULT_N for c in ins.chars)


def test_judge_and_capability():
    c = ip.Char(1, "블랭크", "길이", 100.0, 0.5, -0.5)
    assert (c.usl, c.lsl) == (100.5, 99.5)
    assert c.judge(100.4) == "OK" and c.judge(100.6) == "NG" and c.judge(None) == ""
    c.values = [100.0, 100.1, 99.9, 100.05, 99.95]
    s = ip.stats(c)
    assert s.verdict == "합격" and s.cp == pytest.approx(1.0 / (6 * s.sd))
    assert s.cpk == pytest.approx(min(100.5 - s.mean, s.mean - 99.5) / (3 * s.sd))
    assert "참고" in s.capability                                                     # 시료 25 개 미만
    m = ip.Char(2, "성형", "두께", 0.75, kind=ip.MIN, values=[0.8, 0.74])
    assert m.usl is None and ip.stats(m).verdict == "불합격"
    a = ip.Char(3, "외관", "균열", kind=ip.ATTR, values=["OK", "ng", None])
    assert ip.stats(a).verdict == "불합격" and ip.stats(a).n == 2
    assert ip.angle_tol(32, "c") == 1.0 and ip.angle_tol(5, "m") == 1.0 and ip.deg_text(0.5) == "0°30′"


def test_xlsx_history_and_read_back(tray_ins, tmp_path):
    from openpyxl import load_workbook
    *_, ins = tray_ins
    import copy
    ins = copy.deepcopy(ins)
    for c in ins.chars:
        c.values = ["OK"] * ins.n if c.kind == ip.ATTR else [c.nominal] * ins.n
    ins.chars[2].values[1] = ins.chars[2].nominal + 2 * ins.chars[2].upper        # 하나 불합격
    assert ins.verdict == "불합격" and ins.ng_chars() == [ins.chars[2]]
    f = tmp_path / "ins.xlsx"
    s1 = ip.save_xlsx(ins, f)
    s2 = ip.save_xlsx(ins, f)                                                       # 같은 파일 → 시트 하나 더
    wb = load_workbook(f)
    assert wb.sheetnames == ["이력", s1, s2] and s1 != s2
    assert wb["이력"].max_row == 4 and wb["이력"].cell(3, 6).value == "불합격"
    ws = wb[s2]
    assert str(ws.cell(ip.FIRST, len(ip.HEAD) + ins.n + 7).value).startswith("=IF(")  # 엑셀 식으로 판정
    back = ip.read_xlsx(f)
    assert len(back.chars) == len(ins.chars) and back.n == ins.n and back.verdict == "불합격"
    assert back.chars[0].kind == ip.VAR and back.chars[1].kind == ip.ATTR
    ip.save_pdf(ins, str(tmp_path / "ins.pdf"))
    assert os.path.getsize(tmp_path / "ins.pdf") > 10_000
    v, detail = ip.record(ins)
    assert "불합격" in v["summary"] and len(detail["checks"]) == len(ins.chars)


def test_forming_spots_from_analysis():
    """성형품: 성형 해석 문제 부위 (깊은 오일 팬 귀) 의 두께 ≥ t × 75 % 를 검사 항목으로."""
    from models import oil_pan

    from hanil_unfold import formsim as fs
    st = bc.BendSettings(material="SPCC")
    sol = loader.list_solids(oil_pan(H=160, r_corner=15, flange=0))[0]
    r = unfold(sol, settings=st)
    ins = ip.build(r, st, sol, fs.analyze(r, st), "오일 팬")
    spots = [c for c in ins.chars if c.group == "성형"]
    assert spots and all(c.kind == ip.MIN and c.lsl == pytest.approx(0.75) for c in spots)
    assert "귀" in spots[0].name


def test_dialog_builds(tray_ins):
    from hanil_unfold.report import ensure_gui
    ensure_gui()
    from PySide6.QtWidgets import QApplication
    app = QApplication.instance() or QApplication([])  # noqa: F841
    from hanil_unfold.inspection_dialog import InspectionDialog
    import copy
    *_, ins = tray_ins
    d = InspectionDialog(copy.deepcopy(ins))
    t = d.table
    t.item(2, 7).setText(                                # 7 = 첫 시료 칸 (S1)
        str(ins.chars[2].nominal + 5))
    assert d.ins.chars[2].values[0] == pytest.approx(ins.chars[2].nominal + 5) and "불합격" in d.verdict.text()
    d.n.setValue(8)
    assert t.columnCount() == 7 + 8 + 5 and len(d.ins.chars[0].values) == 8
