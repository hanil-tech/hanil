import json
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(__file__))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from models import profile_part, rotated, tray  # noqa: E402

from hanil_unfold import ai_bom, assembly_calc, bom, loader  # noqa: E402
from hanil_unfold import bendcalc as bc  # noqa: E402


class FakeModel:
    def __init__(self, shapes):
        self.raw_solids = [loader.list_solids(s)[0] for s in shapes]
        self.solids = self.raw_solids
        self.names = [f"부품 {i + 1}" for i in range(len(shapes))]


@pytest.fixture(scope="module")
def assy():
    t = tray()
    br = profile_part([("L", 50), ("B", 90, 2), ("L", 30)], 2.0, 40)
    m = FakeModel([t, br, rotated(t), br])         # 트레이 2 (하나는 돌려 놓음), 브래킷 2
    parts = assembly_calc.compute_all(m, bc.BendSettings())
    return m, parts


def test_bom_groups_identical_parts_even_when_rotated(assy):
    m, parts = assy
    items = bom.build(m, parts, 7.85)
    assert [(it.no, it.qty, it.indices) for it in items] == [(1, 2, [0, 2]), (2, 2, [1, 3])]
    assert items[0].kind == "판금" and items[0].numbers == "1, 3"
    assert items[0].blank_area == pytest.approx(parts[0].metrics.blank_area)
    assert bom.compress([1, 2, 3, 5, 7, 8]) == "1~3, 5, 7~8"


def test_bom_excel_and_pdf(assy, tmp_path):
    import openpyxl
    m, parts = assy
    items = bom.build(m, parts, 7.85, excluded={1, 3})
    bom.save_xlsx(items, str(tmp_path / "b.xlsx"), "t.stp", "SPCC")
    wb = openpyxl.load_workbook(tmp_path / "b.xlsx")
    ws = wb["BOM"]
    assert len(ws._images) == 3                     # 로고 + 품목 2
    assert "제외" in [ws.cell(r, 14).value for r in range(1, ws.max_row + 1) if ws.cell(r, 14).value][-1]
    assert wb["두께별 소요량"].cell(2, 3).value == 2  # 제외한 브래킷 빼고 트레이 2개
    bom.save_pdf(items, str(tmp_path / "b.pdf"), "t.stp", "SPCC")
    assert (tmp_path / "b.pdf").stat().st_size > 2000


class _Block:
    def __init__(self, text):
        self.type, self.text = "text", text


class _Msg:
    def __init__(self, text, stop="end_turn"):
        self.content, self.stop_reason = [_Block(text)], stop


class FakeClient:
    """client.beta.messages.stream(...) 만 흉내 낸다 (네트워크 없이)."""

    def __init__(self, answer):
        self.answer, self.kw = answer, None
        outer = self

        class _Stream:
            def __enter__(self):
                return self

            def __exit__(self, *a):
                return False

            def get_final_message(self):
                return _Msg(json.dumps(outer.answer, ensure_ascii=False))

        class _Messages:
            def stream(self, **kw):
                outer.kw = kw
                return _Stream()

        class _Beta:
            messages = _Messages()

        self.beta = _Beta()


def test_ai_bom_request_and_merge(assy, tmp_path):
    import openpyxl
    m, parts = assy
    items = bom.build(m, parts, 7.85)
    pdf = tmp_path / "d.pdf"
    pdf.write_bytes(b"%PDF-1.4\n%fake\n")
    answer = {"drawing": {"drawing_no": "A-100", "title": "BRACKET ASSY", "revision": "B", "company": "",
                          "date": "", "scale": "1:5"},
              "items": [{"item_no": "1", "part_no": "A-101", "name": "TRAY", "material": "SPCC", "spec": "t1.6",
                         "qty": 2, "kind": "제작품", "matched_3d": [1], "confidence": "높음", "remark": ""},
                        {"item_no": "2", "part_no": "A-102", "name": "BRACKET", "material": "SPCC",
                         "spec": "t2.0", "qty": 3, "kind": "제작품", "matched_3d": [2], "confidence": "보통",
                         "remark": "수량 다름"},
                        {"item_no": "3", "part_no": "", "name": "BOLT M8", "material": "", "spec": "M8x20",
                         "qty": 4, "kind": "표준품", "matched_3d": [], "confidence": "낮음", "remark": ""}],
              "notes": ["시험"]}
    fc = FakeClient(answer)
    ai = ai_bom.analyze(str(pdf), items, client=fc)
    kw = fc.kw
    assert kw["model"] == "claude-opus-5-5"
    assert kw["output_config"]["format"]["type"] == "json_schema"
    doc = kw["messages"][0]["content"][0]
    assert doc["type"] == "document" and doc["source"]["media_type"] == "application/pdf"
    summary = kw["messages"][0]["content"][1]["text"]
    assert '"no": 1' in summary and '"qty": 2' in summary
    rows, unmatched = ai_bom.merge(ai, items)
    assert [r["qty_ok"] for r in rows] == [True, False, False]
    assert unmatched == []
    out = tmp_path / "ai.xlsx"
    ai_bom.save_xlsx(ai, items, str(out), "t.stp", "d.pdf", "SPCC")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["AI BOM", "3D 에만 있음", "도면 정보"]
    ws = wb["AI BOM"]
    names = [ws.cell(r, 3).value for r in range(1, ws.max_row + 1)]
    assert "TRAY" in names and "BOLT M8" in names
    assert wb["도면 정보"].cell(1, 2).value == "A-100"


def test_ai_bom_refusal_is_reported(assy, tmp_path):
    m, parts = assy
    items = bom.build(m, parts, 7.85)
    pdf = tmp_path / "d.pdf"
    pdf.write_bytes(b"%PDF-1.4\n")
    fc = FakeClient({})
    fc.beta.messages.stream = lambda **kw: type("S", (), {
        "__enter__": lambda s: s, "__exit__": lambda s, *a: False,
        "get_final_message": lambda s: _Msg("", "refusal")})()
    with pytest.raises(RuntimeError):
        ai_bom.analyze(str(pdf), items, client=fc)


def test_ai_bom_without_pdf_sends_part_images(assy, tmp_path):
    import openpyxl
    m, parts = assy
    items = bom.build(m, parts, 7.85)
    answer = {"drawing": {"drawing_no": "", "title": "t", "revision": "", "company": "", "date": "", "scale": ""},
              "items": [{"item_no": str(i), "part_no": "", "name": n, "material": "SPCC", "spec": "PL t2",
                         "qty": q, "kind": "제작품", "matched_3d": [i], "confidence": "보통", "remark": ""}
                        for i, n, q in ((1, "트레이", 2), (2, "브래킷", 2))],
              "notes": []}
    fc = FakeClient(answer)
    ai = ai_bom.analyze(None, items, client=fc, source_name="t.stp", material="SPCC")
    content = fc.kw["messages"][0]["content"]
    assert not any(c["type"] == "document" for c in content)
    imgs = [c for c in content if c["type"] == "image"]
    assert len(imgs) == 2 and imgs[0]["source"]["media_type"] == "image/png"
    assert "도면은 없습니다" in content[-1]["text"]
    assert fc.kw["system"] == ai_bom.SYSTEM_3D
    rows, unmatched = ai_bom.merge(ai, items)
    assert all(r["qty_ok"] for r in rows) and not unmatched
    out = tmp_path / "ai3d.xlsx"
    ai_bom.save_xlsx(ai, items, str(out), "t.stp", "", "SPCC")
    ws = openpyxl.load_workbook(out)["AI BOM"]
    assert ws.cell(1, 1).value == "BOM (3D, AI 분석)"
    fc2 = FakeClient(answer)
    ai_bom.analyze(None, items, client=fc2, images=False)
    assert not any(c["type"] == "image" for c in fc2.kw["messages"][0]["content"])


def test_multilevel_bom_from_step_assembly(tmp_path):
    import openpyxl
    from models import nested_assembly_step
    path = nested_assembly_step(str(tmp_path / "nested.stp"))
    m = loader.load(path, use_cache=False)
    assert m.paths and m.paths[0][0][0] == "TRUCK BOX"
    parts = assembly_calc.compute_all(m, bc.BendSettings())
    rows = bom.build_tree(m, parts, 7.85, material="SPCC")
    by = {r.name: r for r in rows if r.kind != "rm"}
    assert rows[0].level == 0 and rows[0].name == "TRUCK BOX" and rows[0].category == "완제품"
    sa = by["SIDE PANEL ASSY"]
    assert (sa.level, sa.qty, sa.total, sa.category) == (1, 2, 2, "반제품 (조립품)")
    assert (by["SIDE PANEL"].level, by["SIDE PANEL"].qty, by["SIDE PANEL"].total) == (2, 1, 2)
    assert (by["BRACKET"].qty, by["BRACKET"].total) == (2, 4)       # BRACKET #2 도 같은 부품
    assert by["NUT M8"].category == "표준품" and by["NUT M8"].procure == "구매" and by["NUT M8"].total == 8
    assert by["W1"].category.startswith("표준품") and by["W1"].spec == "Ø16.0 × 1.6"
    assert by["BOSS"].category == "자작 부품 (선반 가공)"
    assert by["FLOOR PLATE"].category == "자작 부품 (판재)" and by["FLOOR PLATE"].level == 1
    rms = [r for r in rows if r.kind == "rm"]
    assert {r.code for r in rms} == {"RM-SPCC-T1.6", "RM-SPCC-T2", "RM-SPCC-T3.2", "RM-SPCC-RB30"}
    floor_rm = next(r for r in rms if r.code == "RM-SPCC-T3.2")
    assert floor_rm.total == pytest.approx(400 * 300 * 3.2 * 7.85e-6, rel=1e-3)
    summ = {c: kg for c, _, _, kg, _ in bom.material_summary(rows)}
    side_rm = next(r for r in rms if r.tree_no == "1.1-R")
    assert summ["RM-SPCC-T1.6"] == pytest.approx(side_rm.total, rel=1e-6)
    out = tmp_path / "t.xlsx"
    bom.save_tree_xlsx(rows, str(out), "nested.stp", "SPCC")
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["BOM (계층)", "부품 합계", "원자재 소요량", "구매·표준품", "분류 기준"]
    # 계층은 시트 안의 레벨 칸 (레벨 0 · 1 · 2 …) 에 번호로 — 엑셀 바깥 접기 막대 (아웃라인) 는 쓰지 않음
    wt = wb["BOM (계층)"]
    hr = next(r for r in range(1, 12) if wt.cell(r, 1).value == "레벨 0")
    heads = [wt.cell(hr, j).value for j in range(1, 8)]
    n_lv = max(r.level for r in rows) + 1
    assert heads[:n_lv] == [f"레벨 {i}" for i in range(n_lv)] and heads[n_lv] == "품번"
    body = {wt.cell(r, n_lv + 1).value: [wt.cell(r, j).value for j in range(1, n_lv + 1)] for r in range(hr + 1, wt.max_row + 1)}
    assert body["A00"][0] == "0" and body["A00"][1:] == [None] * (n_lv - 1)
    assert body["SA01"][1] == "1" and body["SA01"][0] is None
    assert all(wt.row_dimensions[r].outline_level in (0, None) for r in range(hr + 1, wt.max_row + 1))
    buy = [wb["구매·표준품"].cell(r, 2).value for r in range(2, wb["구매·표준품"].max_row + 1)]
    assert set(buy) == {"NUT M8", "W1"}
    bom.save_tree_pdf(rows, str(tmp_path / "t.pdf"), "nested.stp", "SPCC")
    assert (tmp_path / "t.pdf").stat().st_size > 2000


def test_flat_file_bom_tree_has_root_and_parts(assy):
    m, parts = assy
    m.path = "flat.stp"
    rows = bom.build_tree(m, parts, 7.85, material="SPCC")
    assert rows[0].name == "flat" and rows[0].level == 0
    assert [r.level for r in rows if r.kind == "part"] == [1, 1]
    assert all(r.level == 2 for r in rows if r.kind == "rm")
