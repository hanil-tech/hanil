import math
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(__file__))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from models import dish, drill, profile_part, tray  # noqa: E402

from hanil_unfold import bendcalc as bc  # noqa: E402
from hanil_unfold import loader, metrics, report  # noqa: E402
from hanil_unfold.unfolder import unfold  # noqa: E402


def solid(shape):
    return loader.list_solids(shape)[0]


def test_metrics_area_with_hole_and_weight():
    t = 2.0
    part = drill(profile_part([("L", 50), ("B", 90, 2), ("L", 30)], t, 40), 20, 0, 15, 8)
    r = unfold(solid(part), k_factor=0.44)
    m = metrics.compute(r, 7.85)
    w = r.size[0]
    assert m.blank_area == pytest.approx(w * 40 - math.pi * 16, rel=1e-4)
    assert m.holes == 1 and m.pierce == 2
    assert m.cut_length == pytest.approx(2 * (w + 40) + math.pi * 8, rel=1e-6)
    assert m.weight_kg == pytest.approx(m.blank_area * t * 7.85e-6, rel=1e-9)


def test_metrics_tray_corner_notches():
    r = unfold(solid(tray(100, 80, 30, 1.6, 1.6, 5)), k_factor=0.44)
    m = metrics.compute(r, 7.85)
    a = math.pi / 2 * (1.6 + 0.44 * 1.6) + 30
    assert m.blank_area == pytest.approx(m.rect_area - 4 * a * (a + 5), rel=1e-6)


@pytest.fixture(scope="module")
def reports(tmp_path_factory):
    out = tmp_path_factory.mktemp("rep")
    s = bc.BendSettings(method=bc.K_MATERIAL, material="SPCC")
    sheet = solid(profile_part([("L", 50), ("B", 90, 2), ("L", 30), ("B", -45, 2), ("L", 20)], 2, 40))
    curved = solid(dish(150, 40, 1.2))
    items = []
    for name, sol in (("sheet", sheet), ("dish", curved)):
        r = unfold(sol, settings=s)
        d = report.ReportData(name, f"{name}.step", r, s, sol)
        report.save_pdf(d, str(out / f"{name}.pdf"))
        report.save_xlsx(d, str(out / f"{name}.xlsx"))
        items.append(d)
    report.save_parts_list(items, str(out / "parts.xlsx"))
    return out, items


def test_pdf_reports_are_written(reports):
    out, _ = reports
    for name in ("sheet", "dish"):
        data = (out / f"{name}.pdf").read_bytes()
        assert data.startswith(b"%PDF") and len(data) > 5000


def test_excel_report_contents(reports):
    import openpyxl
    out, items = reports
    wb = openpyxl.load_workbook(out / "sheet.xlsx")
    assert wb.sheetnames == ["보고서", "절곡표", "전개 검증", "치수", "성형 해석", "해석 설명"]
    assert wb["전개 검증"].cell(2, 6).value == "적합"
    ws = wb["절곡표"]
    assert ws.cell(1, 1).value == "번호"
    assert [ws.cell(r, 3).value for r in (2, 3)] == pytest.approx([90, 45], abs=1e-6) or \
        sorted(ws.cell(r, 3).value for r in (2, 3)) == pytest.approx([45, 90], abs=1e-6)
    assert len(wb["보고서"]._images) == 3          # 전개도, 3D, 로고

    wb = openpyxl.load_workbook(out / "dish.xlsx")
    assert "두께 변화" in wb.sheetnames
    ws = wb["두께 변화"]
    shares = [ws.cell(r, 3).value for r in range(4, 10)]
    assert sum(shares) == pytest.approx(1.0, abs=1e-6)
    assert len(ws._images) == 1                     # 성형 한계도 (FLD)


def test_parts_list(reports):
    import openpyxl
    out, items = reports
    ws = openpyxl.load_workbook(out / "parts.xlsx")["부품 목록"]
    assert ws.cell(2, 1).value == "sheet.step"
    assert ws.cell(3, 3).value == "성형품(근사)"
    assert ws.cell(2, 9).value == pytest.approx(items[0].metrics.weight_kg, rel=1e-3)
    assert ws.cell(4, 9).value == "=SUM(I2:I3)"


def test_assembly_names_and_numbered_pdf(tmp_path):
    """조립 구조가 있는 STEP 은 부품 이름을 읽고, 번호 풍선 조립도와 부품표를 만든다."""
    from OCP.gp import gp_Trsf, gp_Vec
    from OCP.STEPCAFControl import STEPCAFControl_Writer
    from OCP.STEPControl import STEPControl_AsIs
    from OCP.TCollection import TCollection_ExtendedString
    from OCP.TDataStd import TDataStd_Name
    from OCP.TDocStd import TDocStd_Document
    from OCP.TopLoc import TopLoc_Location
    from OCP.XCAFDoc import XCAFDoc_DocumentTool

    doc = TDocStd_Document(TCollection_ExtendedString("MDTV-XCAF"))
    st = XCAFDoc_DocumentTool.ShapeTool_s(doc.Main())
    assy = st.NewShape()
    TDataStd_Name.Set_s(assy, TCollection_ExtendedString("TAILGATE ASSY"))
    a = st.AddShape(profile_part([("L", 50), ("B", 90, 2), ("L", 30)], 2, 40), False)
    TDataStd_Name.Set_s(a, TCollection_ExtendedString("61900-BRACKET LH"))
    b = st.AddShape(tray(), False)
    TDataStd_Name.Set_s(b, TCollection_ExtendedString("61910-TRAY"))
    for lab, x in ((a, 0), (b, 200), (a, 400)):
        t = gp_Trsf()
        t.SetTranslation(gp_Vec(x, 0, 0))
        st.AddComponent(assy, lab, TopLoc_Location(t))
    st.UpdateAssemblies()
    path = str(tmp_path / "assy.step")
    w = STEPCAFControl_Writer()
    w.SetNameMode(True)
    w.Transfer(doc, STEPControl_AsIs)
    w.Write(path)

    model = loader.load(path)
    assert model.names == ["61900-BRACKET LH", "61910-TRAY", "61900-BRACKET LH #2"]
    assert model.label(1) == "2. 61910-TRAY"
    # 같은 부품을 옮겨 놓은 세 번째도 따로 전개된다
    r = unfold(model.solids[2], k_factor=0.44)
    assert len(r.bends) == 1

    out = tmp_path / "assy.pdf"
    report.save_assembly_pdf(model, str(out))
    data = out.read_bytes()
    assert data.startswith(b"%PDF") and data.count(b"/Type /Page") >= 2


def test_single_product_file_gets_numbered_names(tmp_path):
    from OCP.STEPControl import STEPControl_AsIs, STEPControl_Writer
    path = str(tmp_path / "one.step")
    w = STEPControl_Writer()
    w.Transfer(profile_part([("L", 50), ("B", 90, 2), ("L", 30)], 2, 40), STEPControl_AsIs)
    w.Write(path)
    model = loader.load(path)
    assert len(model.solids) == 1 and model.label(0).startswith("1. ")


def test_basic_dimensions_and_area():
    from hanil_unfold import dims

    steps = [("L", 50), ("B", 90, 2), ("L", 30), ("B", -90, 2), ("L", 20)]
    part = drill(drill(profile_part(steps, 2, 40), 20, 0, 15, 8), 20, 0, 30, 5)
    r = unfold(solid(part), k_factor=0.44)
    d = dims.compute(r, 7.85)
    ba = math.pi / 2 * (2 + 0.44 * 2)
    # 절곡선(가운데) 위치: 첫 플랜지 + 절곡 여유의 절반
    chain = d.chains[0]
    assert chain.axis == "x"
    assert chain.positions[1] == pytest.approx(50 + ba / 2, abs=1e-6)
    assert chain.positions[2] == pytest.approx(50 + ba + 30 + ba / 2, abs=1e-6)
    assert sum(chain.spans) == pytest.approx(d.width, abs=1e-9)
    holes = sorted(d.holes, key=lambda h: h.diameter)
    assert [round(h.diameter, 6) for h in holes] == [5, 8]
    assert holes[1].center == pytest.approx((20, 15), abs=1e-6)
    assert d.thickness == pytest.approx(2)
    assert d.metrics.blank_area == pytest.approx(d.width * 40 - math.pi * (16 + 6.25), rel=1e-4)
    keys = [k for k, _ in d.rows()]
    assert "두께" in keys and "블랭크 면적 (구멍 뺌)" in keys and "H1 구멍" in keys


def test_dxf_has_real_dimensions(tmp_path):
    import ezdxf
    from hanil_unfold import dxf_export

    part = drill(profile_part([("L", 50), ("B", 90, 2), ("L", 30)], 2, 40), 20, 0, 15, 8)
    r = unfold(solid(part), k_factor=0.44)
    out = tmp_path / "d.dxf"
    dxf_export.save(r, str(out))
    doc = ezdxf.readfile(out)
    texts = []
    for dim in doc.modelspace().query("DIMENSION"):
        texts += [t.dxf.text for t in doc.blocks.get(dim.dxf.geometry).query("MTEXT TEXT")]
    w = r.size[0]
    assert f"{w:.2f}" in texts            # 전체 길이, 배율 1
    assert "40" in texts
    assert any("%%c8" in t for t in texts)  # 구멍 지름
    notes = [t.dxf.text for t in doc.modelspace().query('TEXT[layer=="NOTE"]')]
    assert any(t.startswith("AREA ") and "WEIGHT" in t for t in notes)


def test_user_dims_in_dxf_and_snap_points(tmp_path):
    import ezdxf
    from hanil_unfold import dxf_export
    from hanil_unfold.drawing import dim_text, snap_points
    r = unfold(solid(tray()), k_factor=0.44)
    snaps = snap_points(r)
    x0, y0, x1, y1 = r.bbox
    assert any(abs(px - x1) < 1e-9 and abs(py - y1) < 1e-9 for px, py in snaps)   # 바깥 귀퉁이
    path = tmp_path / "u.dxf"
    dxf_export.save(r, str(path), user_dims=[((0, 0), (30, 40)), ((0, 0), (60, 0))])
    dims = [d for d in ezdxf.readfile(str(path)).modelspace().query("DIMENSION")]
    meas = sorted(round(d.get_measurement(), 3) for d in dims)
    assert 50.0 in meas and 60.0 in meas
    assert dim_text((0, 0), (30, 40)) == "50.00  (X 30.00 · Y 40.00)"
    assert dim_text((0, 0), (60, 0)) == "60.00"
