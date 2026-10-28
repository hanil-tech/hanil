"""주물 드로잉 금형 (castdie) 과 금형비 어림 (diecost)."""
import numpy as np
import pytest

import models
from hanil_unfold import bendcalc as bc
from hanil_unfold import castdie, diecost, die3d, dieface, loader, tandem
from hanil_unfold.unfolder import unfold

ST = bc.BendSettings(material="SPCC")


@pytest.fixture(scope="module")
def pan():
    s = models.oil_pan()
    r = unfold(s, settings=ST, mode="forming")
    df = dieface.build(r.forming.V, r.forming.F, r.thickness)
    assert df is not None
    return s, r, df, castdie.design(df, r, ST)


def test_cast_die_parts_and_geometry(pan):
    """상형 다이 (FC300) · 펀치 · 누르개 (FCD) · 베이스, 가이드 · 쿠션 핀 · 밸런스 · 마모판 · 볼트 · 키, 펀치는 볼록 (제품 바닥이 꼭대기)."""
    s, r, df, a = pan
    names = [p.name for p in a.parts]
    for nm in ("하형 베이스", "드로우 펀치", "누르개 (블랭크 홀더)", "상형 다이", "가이드 포스트", "쿠션 핀", "밸런스 블록 (스토퍼)",
               "마모판 (누르개 안내)", "위치 키 (다우얼)"):
        assert any(n.startswith(nm) for n in names), nm
    pp = a.plate("드로우 펀치")
    pd = a.plate("상형 다이")
    assert "FCD" in pp.material and "FC300" in pd.material and pd.upper and not pp.prims[0].density == 0
    # 펀치 꼭대기 = 베이스 윗면 + 펀치 높이 = 깊이만큼 바인더 위
    pr = pp.prims[0]
    base = a.plate("하형 베이스").prims[0]
    depth = float(np.max(pr.grid["hz"])) - float(np.min(pr.grid["hz"]))
    assert depth == pytest.approx(110.0, abs=3.0)                            # 오일 팬 깊이 110 (두께 가운데 면)
    assert pr.p[2] == pytest.approx(base.p[2] + base.p[5], abs=0.5)           # 펀치가 베이스 위에
    assert pr.p[2] + pr.p[5] > base.p[2] + base.p[5] + depth                  # 꼭대기가 깊이보다 위 (누르개 두께 몫)
    # 상형: 열린 상태로 펀치 꼭대기보다 위, 형면 = 제품 + t
    assert pd.prims[0].p[2] > pr.p[2] + pr.p[5]
    assert float(np.max(pd.prims[0].grid["hz"])) - float(np.max(pr.grid["hz"])) == pytest.approx(
        r.thickness, abs=0.01)                                                # dz (열림) 는 따로
    assert all(c[3] == "적합" for c in a.checks), [c for c in a.checks if c[3] != "적합"]
    assert a.press["draw_kn"] > 100 and a.press["bhf_kn"] > 20
    assert sum(p.weight * p.qty for p in a.parts if "주물" in p.material) > 800
    assert all(len(_grid_or_none(p)) >= 0 for p in a.parts)


def _grid_or_none(p):
    return getattr(p.prims[0], "grid", None) or {}


def test_cast_die_solids_step(pan, tmp_path):
    from OCP.BRepCheck import BRepCheck_Analyzer
    _, _, _, a = pan
    for nm in ("하형 베이스", "드로우 펀치", "누르개 (블랭크 홀더)", "상형 다이"):
        sol = die3d._solid(a.plate(nm).prims[0])
        assert BRepCheck_Analyzer(sol).IsValid(), nm
    p = die3d.save_step(a, str(tmp_path / "cast.step"))
    m = loader.load(p, use_cache=False)
    assert len(m.solids) == sum(len(x.prims) for x in a.parts)


def test_cast_die_section_and_contours(pan):
    """단면 A-A: 상형 (위) · 누르개 (쿠션 위) · 펀치 · 베이스 네 몸체가 잘림, 펀치 등고선 · 실루엣."""
    from hanil_unfold import diedraw, dieview
    s, r, df, a = pan
    cache = {}
    sh = diedraw.section_sheet(a, {"product": "pan"}, "X", cache, y_cut=0.0)
    hatch = sum(1 for e in sh.ents if e[0] == "line" and e[5] == "HATCH")
    assert hatch > 100
    pr = a.plate("드로우 펀치").prims[0]
    top = castdie.grid_lines(pr, dieview.TOP[0], dieview.TOP[1])
    front = castdie.grid_lines(pr, dieview.FRONT[0], dieview.FRONT[1])
    assert len(top) >= 5 and len(front) >= 1
    sil = front[-1]
    assert np.ptp(sil[:, 1]) == pytest.approx(110.0, abs=5.0)                 # 실루엣 높이 = 깊이
    psh = diedraw.shape_part_sheet(a.plate("드로우 펀치"), a, {"product": "pan"}, "X")
    texts = [e[3] for e in psh.ents if e[0] == "text"]
    assert "우측면도" in texts and any("주물 FCD550" in t for t in texts)


def test_cast_die_cost_and_xlsx(pan, tmp_path):
    from openpyxl import load_workbook
    _, _, _, a = pan
    lines, tot = diecost.estimate(a)
    names = [n for n, _, _ in lines]
    assert "주물 FC300 (윗형 · 아랫형)" in names and "주물 FCD550 (펀치 · 누르개)" in names and "형면 NC 가공" in names
    assert 5e6 < tot < 1e8
    p = castdie.save_xlsx(a, str(tmp_path / "c.xlsx"), "pan")
    wb = load_workbook(p)
    assert "금형비" in wb.sheetnames and wb["금형비"].cell(4, 1).value == "설계 (주물 · 대형)"


def test_tandem_die_cost():
    """공정별 금형 금형비: 공정마다 어림 + 합, 캠 공정은 가공 배수 1.6."""
    r = unfold(models.channel_holes(), settings=ST)
    ops = tandem.design(r, ST)
    tot = diecost.total(ops)
    assert 1e7 < tot < 1e8
    cam = next(op for op in ops if op.kind == tandem.CAM)
    lines, _ = diecost.estimate(cam.die)
    mach = next(x for x in lines if x[0].startswith("가공"))
    assert "× 1.6" in mach[1]
    sh = tandem.line_sheets(tandem.Folder(r), ops, {"product": "ch"}, "HD")
    texts = [e[3] for e in sh[0].ents if e[0] == "text"]
    assert any("금형비 어림 합" in t for t in texts) and "금형비 만원" in texts
