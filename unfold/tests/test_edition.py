"""두 가지 판: 전체판 (기본) 과 전개 전용판 (HANIL_EDITION=unfold, 설치 파일은 _edition.py 로 못 박음).
판은 처음 불러올 때 정해지므로 전개 전용판은 따로 띄운 파이썬에서 시험한다."""
import os
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

SCRIPT = r'''
import os, sys, json, tempfile
sys.path.insert(0, ROOT)
os.environ["QT_QPA_PLATFORM"] = "offscreen"
from PySide6.QtWidgets import QApplication
app = QApplication([])
from hanil_unfold import edition, company
from hanil_unfold.app import MainWindow
w = MainWindow()
tabs = [w.ribbon.tabText(i) for i in range(w.ribbon.count())]
acts = [a.text() for a in w.actions()]
w.open_file(os.path.join(ROOT, "samples", "tray_t1.6.step"))
from hanil_unfold import report
d = w._report_data()
out = os.path.join(tempfile.mkdtemp(), "r.xlsx")
report.save_xlsx(d, out)
from openpyxl import load_workbook
print(json.dumps({"edition": edition.EDITION, "product": company.PRODUCT, "tabs": tabs, "acts": acts,
                  "title": w.windowTitle(), "sheets": load_workbook(out).sheetnames, "result": w.result is not None}))
'''


def run(env_edition):
    env = dict(os.environ, HANIL_EDITION=env_edition, QT_QPA_PLATFORM="offscreen")
    p = subprocess.run([sys.executable, "-c", f"ROOT = {ROOT!r}\n" + SCRIPT], capture_output=True, text=True,
                       env=env, timeout=300, cwd=ROOT)
    assert p.returncode == 0, p.stderr[-2000:]
    import json
    return json.loads(p.stdout.strip().splitlines()[-1])


def test_unfold_only_edition():
    r = run("unfold")
    assert r["edition"] == "unfold" and "전개 전용" in r["product"] and "전개 전용" in r["title"]
    assert r["tabs"] == ["파일", "전개", "기록", "설정·도움말"]
    for gone in ("성형 해석", "금형 검토", "견적", "원가계산서", "검사 성적서", "철판 배치", "BOM", "실적 보정"):
        assert gone not in r["acts"], gone
    for keep in ("열기", "DXF 저장", "보고서 PDF", "보고서 엑셀", "여러 파일 한꺼번에", "전개 기록"):
        assert keep in r["acts"], keep
    assert r["result"] and "성형 해석" not in r["sheets"] and "보고서" in r["sheets"]


def test_full_edition_default():
    r = run("")
    assert r["edition"] == "full" and "전개 전용" not in r["product"]
    assert "견적" in r["tabs"] and "금형" in r["tabs"] and "성형 해석" in r["acts"]


def test_unfold_only_cli_blocks_full_features():
    env = dict(os.environ, HANIL_EDITION="unfold")
    p = subprocess.run([sys.executable, "-m", "hanil_unfold", "nest", "x.step"], capture_output=True, text=True,
                       env=env, cwd=ROOT, timeout=120)
    assert p.returncode == 2 and "전체판" in p.stdout
    p = subprocess.run([sys.executable, "-m", "hanil_unfold", "--version"], capture_output=True, text=True,
                       env=env, cwd=ROOT, timeout=120)
    assert "전개 전용" in p.stdout
