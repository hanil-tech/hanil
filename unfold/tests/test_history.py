import datetime as dt
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(__file__))

from models import profile_part  # noqa: E402

from hanil_unfold import bendcalc as bc  # noqa: E402
from hanil_unfold import loader  # noqa: E402
from hanil_unfold.history import History, bend_summary  # noqa: E402
from hanil_unfold.unfolder import unfold  # noqa: E402


@pytest.fixture
def result():
    s = loader.list_solids(profile_part([("L", 50), ("B", 90, 2), ("L", 30), ("B", 90, 2), ("L", 20)], 2, 40))[0]
    return unfold(s, k_factor=0.44)


def test_record_accumulates_by_day_and_updates_same_day(tmp_path, result):
    h = History(str(tmp_path / "h"))
    s = bc.BendSettings(material="SUS304")
    day1 = dt.datetime(2026, 10, 1, 9, 0)
    a = h.record("C:/parts/bracket.step", result, s, when=day1)
    b = h.record("C:/parts/bracket.step", result, s, when=day1.replace(hour=15))
    assert a == b and h.count() == 1          # 같은 날 같은 부품은 갱신
    c = h.record("C:/parts/bracket.step", result, s, when=dt.datetime(2026, 10, 2, 9, 0))
    d = h.record("C:/parts/bracket.step", result, s, part_index=2, part_count=2, when=day1)
    assert len({a, c, d}) == 3 and h.count() == 3   # 다른 날·다른 부품은 새 줄

    rec = {r.id: r for r in h.search()}[a]
    assert rec["material"] == "SUS304"
    assert rec["bends"] == 2 and rec["bend_summary"] == "90.0°R2.00 ×2"
    assert rec["weight_kg"] > 0
    assert os.path.exists(rec["archive_dxf"])
    assert "2026-10" in rec["archive_dxf"]


def test_search_note_export_delete(tmp_path, result):
    import openpyxl

    h = History(str(tmp_path / "h"))
    s = bc.BendSettings()
    a = h.record("a/door_panel.step", result, s, when=dt.datetime(2026, 9, 1))
    b = h.record("a/bracket.step", result, s, when=dt.datetime(2026, 10, 5))
    h.set_note(b, "현대 주문 50개")
    h.add_export(b, "out/bracket_flat.dxf")
    h.add_export(b, "out/bracket_flat.dxf")

    assert [r.id for r in h.search("door")] == [a]
    assert [r.id for r in h.search("현대")] == [b]
    assert [r.id for r in h.search(date_from="2026-10-01")] == [b]
    assert h.search("bracket")[0]["exported"].count(";") == 0

    out = tmp_path / "기록.xlsx"
    h.export_xlsx(str(out))
    ws = openpyxl.load_workbook(out)["전개 기록"]
    assert ws.cell(1, 1).value == "번호"
    assert ws.cell(4, 1).value == "합계"

    dxf = h.search("door")[0]["archive_dxf"]
    h.delete([a])
    assert h.count() == 1 and not os.path.exists(dxf)


def test_history_survives_reopen(tmp_path, result):
    folder = str(tmp_path / "h")
    h = History(folder)
    h.record("x.step", result, bc.BendSettings())
    h.close()
    assert History(folder).count() == 1


def test_bend_summary_groups(result):
    assert bend_summary(result.bends) == "90.0°R2.00 ×2"


def test_assembly_totals_by_thickness(tmp_path):
    from models import tray
    from hanil_unfold import assembly_calc
    from hanil_unfold.loader import LazySolids, LoadedModel

    a = profile_part([("L", 50), ("B", 90, 2), ("L", 30)], 2, 40)
    b = tray(100, 80, 30, 1.6, 1.6, 5)
    c = profile_part([("L", 20), ("B", 90, 2), ("L", 20)], 2, 30)
    solids = [loader.list_solids(x)[0] for x in (a, b, c)]
    model = LoadedModel("x.step", None, LazySolids(solids), ["A", "B", "C"])
    s = bc.BendSettings(material="SPCC")
    parts = assembly_calc.compute_all(model, s)
    assert all(p.ok for p in parts)
    groups, total = assembly_calc.summarize(parts, "SPCC")
    assert [g.thickness for g in groups] == [1.6, 2.0]
    assert [g.count for g in groups] == [1, 2]
    assert total.weight_kg == pytest.approx(sum(p.metrics.weight_kg for p in parts))
    out = tmp_path / "assy.xlsx"
    assembly_calc.save_xlsx(str(out), model, parts, "SPCC")
    import openpyxl
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["두께별 합계", "부품별"]
    assert wb["두께별 합계"].cell(4, 2).value == "합계"


def test_version_text_and_old_history_file_upgrade(tmp_path, result):
    import sqlite3
    from hanil_unfold import __version__, version_text

    assert version_text().startswith(__version__)
    # 버전 칸이 없던 예전 기록 파일도 열면 칸이 덧붙고 새 기록에 버전이 남는다
    folder = tmp_path / "old"
    folder.mkdir()
    db = sqlite3.connect(folder / "history.sqlite")
    db.execute("CREATE TABLE unfolds (id INTEGER PRIMARY KEY AUTOINCREMENT, created TEXT, updated TEXT, "
               "day TEXT, source_path TEXT, part_index INTEGER)")
    db.commit()
    db.close()
    h = History(str(folder))
    rid = h.record("x.step", result, bc.BendSettings())
    assert h.search()[0]["app_version"] == version_text()
    assert rid == 1


def test_quote_and_die_docs_are_recorded(tmp_path):
    import json
    import os
    import sys
    sys.path.insert(0, os.path.dirname(__file__))
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    QApplication.instance() or QApplication([])
    from models import nested_assembly_step

    from hanil_unfold import assembly_calc, loader
    from hanil_unfold import bendcalc as bc
    from hanil_unfold import quote as q
    from hanil_unfold.history import History
    step = tmp_path / "truck.step"
    nested_assembly_step(str(step))
    m = loader.load(str(step), use_cache=False)
    st = bc.BendSettings()
    qt = q.build(m, assembly_calc.compute_all(m, st), st, q.load_rates(None), sets=20, annual=300_000, title="TRUCK")
    qt.customer = "가나상사"
    out = tmp_path / "견적서.xlsx"
    q.save_xlsx(qt, str(out))
    h = History(str(tmp_path / "기록"))
    rid = h.add_doc("생산 견적", [str(out)], str(step), q.record_values(qt), q.record_detail(qt))
    h.add_doc("금형 견적", [], str(step), q.record_values(qt, die=True), q.record_detail(qt, die=True))
    rows = h.search_docs("가나")
    assert len(rows) == 2 and h.count_docs() == 2
    r = next(x for x in rows if x.id == rid).data
    assert r["total"] == qt.total and r["unit_price"] == qt.unit_price and r["max_tonf"] == pytest.approx(qt.max_tonf)
    assert r["kind"] == "생산 견적" and r["source_name"] == "truck.step"
    assert os.path.exists(r["archived"]) and r["archived"] != str(out)        # 보관 사본
    assert json.loads(r["detail"])[0]["name"]
    assert len(h.search_docs(kind="금형 견적")) == 1
    h.set_doc_note(rid, "수주")
    assert h.search_docs("수주")[0].id == rid
    h.export_docs_xlsx(str(tmp_path / "log.xlsx"))
    import openpyxl
    wb = openpyxl.load_workbook(tmp_path / "log.xlsx")
    assert wb.sheetnames == ["견적·금형 기록", "부품별"] and wb["부품별"].max_row > 2
    h.delete_docs([rid])
    assert h.count_docs() == 1 and not os.path.exists(r["archived"])
